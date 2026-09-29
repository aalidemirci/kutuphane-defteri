"""Uydurma deneme verisi üreticisi — saha kabul protokolü için (`docs/saha-kabulu.md` §1.2).

Saha kabulünde (tasarım §14.1 F12; kullanıcı kararı 27.09.2026) program gerçek
öğrenci ve personel listesi YERİNE bu dosyalarla sınanır. Üretilenler:

    eokul-ogrenci-listesi.xlsx          e-Okul "OOG01001R020 — Sınıf/Şube Öğrenci Listesi"
                                        yerleşimi: şube şube bloklar, sayaç dipnotları
    eokul-ogrenci-listesi-2-donem.xlsx  aynı okulun sonraki listesi: ayrılan, gelen ve şube
                                        değiştiren öğrencilerle (Ayrılış Havuzu sınaması)
    eokul-personel-listesi.xlsx         e-Okul "OOK01001R1 — Personel Listesi" yerleşimi
    katalog-deneme.xlsx                 katalog Excel şablonunun "Katalog" sayfası
                                        (varsayılan 2.000 eser)
    katalog-sorunlu-satirlar.xlsx       önizlemenin uyarı, hata ve karar yollarını gösteren
                                        küçük dosya (satırların adı "Sınama —" ile başlar)
    OZET.txt                            önizlemede görülmesi gereken sayılar

**Hepsi UYDURMADIR** (CLAUDE.md §2-12, sözlük §2.1): kişi adları ad ve soyad
havuzlarının rastgele birleşimidir; okul "Örnek" ön ekiyle kurulur; T.C. kimlik
numarası, telefon, adres, e-posta ve veli bilgisi ÜRETİLMEZ (program bunları
almaz). e-Okul yerleşiminin "Cinsiyeti" ve "Branşı" sütunları gerçek raporda
bulunduğu için uydurma değerle doldurulur; program onları OKUMAZ (tasarım §6.1,
V2-01) ve bu da sınanan davranışlardan biridir. Kitap künyeleri ya bilinen
eserlerindir ya da uydurmadır (künye ne kişisel veri ne telif konusudur; eserlerin
metni yer almaz — listedeki bazı yazarların eserleri FSEK md. 27'nin ölümden sonraki
70 yıllık süresi içindedir, yani "kamu malı" DEĞİLDİR); ISBN'ler sağlaması tutan UYDURMA numaralardır
(gerçek bir kitaba denk gelebilir; künye kişisel veri değildir).

**Veri dosyaları depoya GİRMEZ:** varsayılan çıktı klasörü `deneme-verisi/`
`.gitignore`'dadır ve `*.xlsx` zaten yasaktır. Depoya yalnız bu üretici girer.

**Neden .xlsx?** e-Okul'un kendi dosyası Excel 97-2003 (.XLS) biçimindedir; o
biçimi yazmak ek bir kitaplık isterdi ve proje yeni Python bağımlılığı almaz.
Program iki biçimi de okur ve biçimi uzantıdan değil dosyanın içinden tanır
(`apps.okul.excel_ogrenci.read_sheet`); e-Okul'a özgü YERLEŞİM (şube blokları,
dipnotlar) iki biçimde de aynı önişleyiciden geçer. Gerçek .XLS kabının okunması
depodaki sentetik örnekle ayrıca sınanır (`backend/apps/okul/tests/veri/`).

Yalnız standart kitaplık ve openpyxl kullanılır (openpyxl programın zaten
bağımlılığıdır). Program modülleri İÇE AKTARILMAZ: üretici programdan bağımsız
bir "dışarıdan gelen dosya" yazar; uyumu test sınar
(`backend/apps/kutuphane/tests/test_deneme_verisi.py`). Aynı tohumla her koşu
aynı veriyi üretir.

Kullanım (host'ta Python yok; Docker'da, depo kökünden):

    docker compose run --rm -T -w /repo backend python scripts/deneme_verisi.py

Seçenekler: `--help`.
"""

from __future__ import annotations

import argparse
import random
import sys
import unicodedata
from collections.abc import Callable, Iterable, Sequence
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

# ---------------------------------------------------------------------------
# Sabitler
# ---------------------------------------------------------------------------
VARSAYILAN_CIKTI = "deneme-verisi"
VARSAYILAN_TOHUM = 2026
VARSAYILAN_OGRENCI = 750
VARSAYILAN_PERSONEL = 60
VARSAYILAN_ESER = 2000

#: Komut satırı sınırları (programın hedeflediği ölçek: 1.000-10.000 kitap — S8).
OGRENCI_SINIRI = (50, 5000)
PERSONEL_SINIRI = (10, 500)
ESER_SINIRI = (40, 10000)

DOSYA_OGRENCI = "eokul-ogrenci-listesi.xlsx"
DOSYA_OGRENCI_2 = "eokul-ogrenci-listesi-2-donem.xlsx"
DOSYA_PERSONEL = "eokul-personel-listesi.xlsx"
DOSYA_KATALOG = "katalog-deneme.xlsx"
DOSYA_SORUNLU = "katalog-sorunlu-satirlar.xlsx"
DOSYA_OZET = "OZET.txt"

#: Katalog şablonunun sayfa adı ve sütun başlıkları — `apps/kutuphane/import_schema.py`
#: ile BİREBİR (sıra dahil). Burada elle yazılıdır: üretici programı içe aktarmaz,
#: eşitliği test sınar (şablon değişirse test kırmızıya döner).
KATALOG_SAYFASI = "Katalog"
KATALOG_BASLIKLARI: tuple[str, ...] = (
    "Eser Adı",
    "Yazar",
    "Çevirmen",
    "Yayınevi",
    "Baskı",
    "Yayın Yılı",
    "ISBN",
    "Konu",
    "Sınıflama Kodu",
    "Dil",
    "Kaynak Türü",
    "Nüsha Sayısı",
    "Bölüm",
    "Eski Kayıt No",
    "Ciltli Süreli Yayın",
    "Danışma Kaynağı",
)

EVET = "Evet"
HAYIR = "Hayır"
TUR_KITAP = "Kitap"
TUR_DERS = "Ders kitabı"
TUR_SURELI = "Süreli yayın"
TUR_GORSEL = "Görsel-işitsel materyal"

#: e-Okul sayfa altındaki tarih/saat seri numarası (gerçek raporda da sayıdır).
_EOKUL_TARIH_SERISI = 46292.0

KADEMELER: dict[str, tuple[str, tuple[int, ...], str]] = {
    # anahtar: (blok başlığının öneki, sınıf seviyeleri, okulun uydurma adı)
    "lise": ("AL", (9, 10, 11, 12), "Örnek Anadolu Lisesi"),
    "ortaokul": ("ORTAOKUL", (5, 6, 7, 8), "Örnek Ortaokulu"),
    "ilkokul": ("İLKOKUL", (1, 2, 3, 4), "Örnek İlkokulu"),
}

#: İlk seviyenin şubeleri: Türk alfabesinde I ve İ AYRI iki harftir ve aynı okulda
#: ikisi birden bulunur (`normalize.tr_upper` gerekçesi). Öbür seviyeler beş şubelidir.
SUBELER_ILK = ("A", "B", "C", "D", "E", "F", "G", "H", "I", "İ")
SUBELER_DIGER = ("A", "B", "C", "D", "E")

# ---------------------------------------------------------------------------
# Ad havuzları (UYDURMA birleşimler — gerçek bir kişiyi göstermez)
# ---------------------------------------------------------------------------
KIZ_ADLARI = (
    "Ayşe", "Fatma", "Zeynep", "Elif", "Merve", "Büşra", "Esra", "İrem", "Nazlı",
    "Derya", "Selin", "Ece", "Işıl", "İlayda", "Şevval", "Gizem", "Beyza", "Dilara",
    "Nisa", "Öykü", "Çağla", "Tuğba", "Aslı", "İpek", "Melis", "Defne", "Azra", "Ela",
    "Mine", "Yağmur", "Şimal", "Duru", "Ilgın", "Nehir", "Sıla", "Gülsüm", "Hilal",
    "Eylül", "Ceren", "Özge",
)  # fmt: skip
ERKEK_ADLARI = (
    "Mehmet", "Mustafa", "Emre", "Burak", "Yusuf", "Ömer", "Ali", "Hüseyin", "İbrahim",
    "Kerem", "Berkay", "Çağan", "Onur", "Uğur", "Ege", "Arda", "Tuna", "Barış", "Işık",
    "İlker", "Şahin", "Gökhan", "Serkan", "Emir", "Efe", "Yiğit", "Çınar", "Alperen",
    "Batuhan", "Oğuz", "Doruk", "Ilgaz", "Kaan", "Mert", "Tolga", "Umut", "Volkan",
    "Selim", "Taner", "Rüzgar",
)  # fmt: skip
SOYADLARI = (
    "Yılmaz", "Kaya", "Şahin", "Çelik", "Yıldız", "Yıldırım", "Öztürk", "Aydın",
    "Özdemir", "Arslan", "Doğan", "Kılıç", "Aslan", "Çetin", "Kara", "Koç", "Kurt",
    "Özkan", "Şimşek", "Polat", "Özcan", "Korkmaz", "Çakır", "Yavuz", "Can", "Acar",
    "Aksoy", "Güneş", "Işıkçı", "Uludağlı", "Çınarlı", "Gökmen", "Tunç", "Akgül",
    "İnce", "Ilgaz", "Önal", "Karataş", "Erdem", "Bulut", "Tekin", "Güler", "Sarı",
    "Ünal", "Özer", "Keskin", "Çiftçi", "Taş", "Akın", "Uysal", "Bozkurt", "Oral",
    "İleri", "Ilıcak", "Şen", "Başaran", "Ağaoğlu", "Gündüz", "Duman", "Eren",
)  # fmt: skip

# ---------------------------------------------------------------------------
# Personel görevleri: (e-Okul "GÖREVİ" metni, programın beklenen üye türü, tanınır mı)
# Program görev metnini SAKLAMAZ; yalnız üye türünü seçmek için okur (F1 ekleri 2).
# Tanınmayan görev öğretmen sayılır ve önizlemede "üye türünü denetleyin" uyarısı çıkar.
# ---------------------------------------------------------------------------
OGRETMEN = "öğretmen"
DIGER_PERSONEL = "diğer personel"
OGRETMEN_GOREVLERI: tuple[tuple[str, str, int], ...] = (
    # (görev, kadro durumu, ağırlık)
    ("Öğretmen", "KADROLU", 70),
    ("Sözleşmeli Öğretmen(657 S.K. 4/B)", "SÖZLEŞMELİ", 10),
    ("Ücretli Öğretmen", "ÜCRETLİ", 10),
    ("Rehber Öğretmen", "KADROLU", 5),
    ("Usta Öğretici", "KADROLU", 3),
)
DIGER_GOREVLER: tuple[str, ...] = (
    "Memur",
    "Hizmetli",
    "Teknisyen",
    "Şef",
    "Aşçı",
    "Bekçi",
    "Veri Hazırlama ve Kontrol İşletmeni",
)
#: Programın tanımadığı görev (öğretmen sayılır + uyarı). Tek satır.
TANINMAYAN_GOREV = "Kütüphane Görevlisi"
BRANSLAR = (
    "Türk Dili ve Edebiyatı", "Matematik", "Fizik", "Kimya", "Biyoloji", "Tarih",
    "Coğrafya", "Felsefe", "İngilizce", "Almanca", "Beden Eğitimi ve Spor",
    "Görsel Sanatlar", "Müzik", "Din Kültürü ve Ahlak Bilgisi", "Bilişim Teknolojileri",
    "Rehberlik",
)  # fmt: skip

# ---------------------------------------------------------------------------
# Katalog havuzları
# ---------------------------------------------------------------------------
YAYINEVLERI = (
    "Örnek Yayınevi", "Deneme Kitap", "Kumsal Yayınları", "Çıra Kitap", "Ilgın Yayınevi",
    "Palamut Yayınları", "Işıltı Kitabevi", "Şafak Örnek Yayıncılık", "Ağaçkakan Yayınları",
    "Üveyik Kitap",
)  # fmt: skip

BOLUM_EDEBIYAT = "Edebiyat"
BOLUM_YABANCI = "Yabancı Dil"
BOLUM_COCUK = "Çocuk ve Gençlik"
BOLUM_TARIH = "Tarih ve Coğrafya"
BOLUM_BILIM = "Bilim ve Teknoloji"
BOLUM_TOPLUM = "Felsefe ve Toplum"
BOLUM_SANAT = "Sanat"
BOLUM_DANISMA = "Danışma"
BOLUM_DERS = "Ders Kitapları"
BOLUM_SURELI = "Süreli Yayınlar"
BOLUM_GORSEL = "Görsel-İşitsel"


@dataclass(frozen=True)
class Klasik:
    """Bilinen bir eserin künyesi (künye ne kişisel veri ne telif konusudur; metin yer almaz)."""

    ad: str
    yazar: str
    konu: str
    dos: str
    dil: str = "Türkçe"
    ceviri: bool = False
    bolum: str = BOLUM_EDEBIYAT
    yil: int | None = None
    isbn_yok: bool = False
    danisma: bool = False


#: İlk iki satır `katalog-sorunlu-satirlar.xlsx`'in "mevcut eser" ve "şüpheli" satırlarının
#: dayanağıdır: ISBN'li olmalı ve katalog dosyasında her boyutta yer almalıdır.
KLASIKLER: tuple[Klasik, ...] = (
    Klasik("Çalıkuşu", "Reşat Nuri Güntekin", "Türk edebiyatı, roman", "894.353"),
    Klasik("Suç ve Ceza", "Fyodor Dostoyevski", "Rus edebiyatı, roman", "891.73", ceviri=True),
    Klasik("Mai ve Siyah", "Halit Ziya Uşaklıgil", "Türk edebiyatı, roman", "894.353"),
    Klasik("Aşk-ı Memnu", "Halit Ziya Uşaklıgil", "Türk edebiyatı, roman", "894.353"),
    Klasik("Sinekli Bakkal", "Halide Edip Adıvar", "Türk edebiyatı, roman", "894.353"),
    Klasik("Ateşten Gömlek", "Halide Edip Adıvar", "Türk edebiyatı, roman", "894.353"),
    Klasik("Kürk Mantolu Madonna", "Sabahattin Ali", "Türk edebiyatı, roman", "894.353"),
    Klasik("Kuyucaklı Yusuf", "Sabahattin Ali", "Türk edebiyatı, roman", "894.353"),
    Klasik("İçimizdeki Şeytan", "Sabahattin Ali", "Türk edebiyatı, roman", "894.353"),
    Klasik("Safahat", "Mehmet Akif Ersoy", "Türk edebiyatı, şiir", "894.351"),
    Klasik("Rübab-ı Şikeste", "Tevfik Fikret", "Türk edebiyatı, şiir", "894.351"),
    Klasik("İntibah", "Namık Kemal", "Türk edebiyatı, roman", "894.353"),
    Klasik("Vatan yahut Silistre", "Namık Kemal", "Türk edebiyatı, tiyatro", "894.352"),
    Klasik("Kaşağı", "Ömer Seyfettin", "Türk edebiyatı, öykü", "894.353"),
    Klasik("Yaban", "Yakup Kadri Karaosmanoğlu", "Türk edebiyatı, roman", "894.353"),
    Klasik("Kiralık Konak", "Yakup Kadri Karaosmanoğlu", "Türk edebiyatı, roman", "894.353"),
    Klasik("Araba Sevdası", "Recaizade Mahmut Ekrem", "Türk edebiyatı, roman", "894.353"),
    Klasik("Eylül", "Mehmet Rauf", "Türk edebiyatı, roman", "894.353"),
    Klasik("Dokuzuncu Hariciye Koğuşu", "Peyami Safa", "Türk edebiyatı, roman", "894.353"),
    Klasik("Dede Korkut Hikâyeleri", "", "Türk edebiyatı, destan", "894.3"),
    Klasik(
        "Kamus-ı Türkî",
        "Şemseddin Sami",
        "Türkçe, sözlük",
        "494.353",
        dil="Osmanlı Türkçesi",
        bolum=BOLUM_DANISMA,
        yil=1901,
        isbn_yok=True,
        danisma=True,
    ),
    Klasik("Savaş ve Barış", "Lev Tolstoy", "Rus edebiyatı, roman", "891.73", ceviri=True),
    Klasik("Palto", "Nikolay Gogol", "Rus edebiyatı, öykü", "891.73", ceviri=True),
    Klasik("Sefiller", "Victor Hugo", "Fransız edebiyatı, roman", "843", ceviri=True),
    Klasik("Seksen Günde Devrialem", "Jules Verne", "Fransız edebiyatı, roman", "843", ceviri=True),
    Klasik(
        "Denizler Altında Yirmi Bin Fersah",
        "Jules Verne",
        "Fransız edebiyatı, roman",
        "843",
        ceviri=True,
    ),
    Klasik("Gurur ve Önyargı", "Jane Austen", "İngiliz edebiyatı, roman", "823", ceviri=True),
    Klasik("Oliver Twist", "Charles Dickens", "İngiliz edebiyatı, roman", "823", ceviri=True),
    Klasik("Robinson Crusoe", "Daniel Defoe", "İngiliz edebiyatı, roman", "823", ceviri=True),
    Klasik(
        "Tom Sawyer'ın Maceraları",
        "Mark Twain",
        "Amerikan edebiyatı, roman",
        "813",
        ceviri=True,
    ),
    Klasik("Beyaz Diş", "Jack London", "Amerikan edebiyatı, roman", "813", ceviri=True),
    Klasik("Don Kişot", "Miguel de Cervantes", "İspanyol edebiyatı, roman", "863", ceviri=True),
    Klasik(
        "Genç Werther'in Acıları",
        "Johann Wolfgang von Goethe",
        "Alman edebiyatı, roman",
        "833",
        ceviri=True,
    ),
    Klasik("Hamlet", "William Shakespeare", "İngiliz edebiyatı, tiyatro", "822.33", ceviri=True),
    Klasik("İlyada", "Homeros", "Yunan edebiyatı, destan", "883", ceviri=True),
    Klasik("Quo Vadis", "Henryk Sienkiewicz", "Polonya edebiyatı, roman", "891.85", ceviri=True),
    Klasik(
        "Küçük Prens", "Antoine de Saint-Exupéry", "Fransız edebiyatı, roman", "843", ceviri=True
    ),
    Klasik(
        "Wuthering Heights",
        "Emily Brontë",
        "İngiliz edebiyatı, roman",
        "823",
        dil="İngilizce",
        bolum=BOLUM_YABANCI,
    ),
    Klasik(
        "Great Expectations",
        "Charles Dickens",
        "İngiliz edebiyatı, roman",
        "823",
        dil="İngilizce",
        bolum=BOLUM_YABANCI,
    ),
    Klasik(
        "The Adventures of Sherlock Holmes",
        "Arthur Conan Doyle",
        "İngiliz edebiyatı, öykü",
        "823",
        dil="İngilizce",
        bolum=BOLUM_YABANCI,
    ),
    Klasik(
        "Die Verwandlung",
        "Franz Kafka",
        "Alman edebiyatı, öykü",
        "833",
        dil="Almanca",
        bolum=BOLUM_YABANCI,
    ),
    Klasik(
        "Candide",
        "Voltaire",
        "Fransız edebiyatı, roman",
        "843",
        dil="Fransızca",
        bolum=BOLUM_YABANCI,
    ),
)

# Kurgu başlıkları: sıfat + ad, tamlama (ilgi + iyelik ekli ad), yer + "Bir …".
# Ekler hazır yazılıdır (ünlü uyumunu kurala bağlamak yerine elle doğru biçim).
SIFATLAR = (
    "Sessiz", "Kayıp", "Işıklı", "Ilık", "İnce", "Uzak", "Yakın", "Gizli", "Son", "İlk",
    "Mavi", "Kırmızı", "Yeşil", "Sarı", "Beyaz", "Kara", "Eski", "Yeni", "Küçük", "Büyük",
    "Derin", "Sisli", "Rüzgârlı", "Karlı", "Yağmurlu", "Sıcak", "Soğuk", "Güzel", "Yalnız",
    "Kırık", "Altın", "Gümüş", "Uçan", "Dönen", "Susan", "Bekleyen", "Unutulmuş", "Öksüz",
    "Şen", "Tuzlu",
)  # fmt: skip
ADLAR = (
    "Liman", "Şehir", "Köprü", "Sokak", "Bahçe", "Deniz", "Orman", "Dağ", "Kıyı", "Ada",
    "Ev", "Kapı", "Pencere", "Yol", "Gece", "Sabah", "Akşam", "Mevsim", "Kış", "Yaz", "Güz",
    "Bahar", "Mektup", "Defter", "Şarkı", "Türkü", "Masal", "Rüya", "Saat", "Kuyu", "Çeşme",
    "Ağaç", "Kuş", "Yıldız", "Ay", "Güneş", "Fener", "Tren", "Gemi", "Çocukluk",
)  # fmt: skip
ILGI_ADLARI = (
    "Denizin", "Rüzgârın", "Ormanın", "Şehrin", "Gecenin", "Işığın", "Kışın", "Yazın",
    "Ağacın", "Dağın", "Nehrin", "Köyün", "Sokağın", "Yolun", "Evin", "Zamanın", "Suyun",
    "Taşın", "Kalbin", "Bulutun",
)  # fmt: skip
IYELIK_ADLARI = (
    "Sesi", "Rengi", "Sırrı", "Kıyısı", "Yolu", "Düşü", "Türküsü", "Gölgesi", "Hikâyesi",
    "Şarkısı", "Ötesi", "Kokusu", "Anısı", "Sessizliği", "Çocukları", "Bekçisi", "Kapısı",
    "Haritası", "Mevsimi",
)  # fmt: skip
YERLER = (
    "İstanbul'da", "Ankara'da", "Karadeniz'de", "Kapadokya'da", "Ege'de", "Toroslar'da",
    "Çukurova'da", "Van'da", "Trabzon'da", "Muğla'da", "Erzurum'da", "Bozcaada'da",
)  # fmt: skip
BIR_ADLAR = (
    "Bir Yaz", "Bir Kış", "Bir Sabah", "Bir Akşam", "Bir Yolculuk", "Bir Mektup",
    "Bir Dostluk", "Bir Bayram", "Bir Düğün", "Bir Tren Yolculuğu",
)  # fmt: skip
#: (başlık eki, konu, DOS) — "Konu" edebî türü de taşır (şablonun kuralı).
KURGU_TURLERI: tuple[tuple[str, str, str], ...] = (
    ("", "Türk edebiyatı, roman", "894.353"),
    (": Öyküler", "Türk edebiyatı, öykü", "894.353"),
    (": Şiirler", "Türk edebiyatı, şiir", "894.351"),
    (": Anılar", "Türk edebiyatı, anı", "894.358"),
    (": Denemeler", "Türk edebiyatı, deneme", "894.354"),
    (": Oyun", "Türk edebiyatı, tiyatro", "894.352"),
    (": Masallar", "Türk edebiyatı, masal", "398.2"),
    (": Roman", "Türk edebiyatı, roman", "894.353"),
)
#: Bilgi kitapları: (yönelme ekli ad, konu, DOS, bölüm).
BILIM_DALLARI: tuple[tuple[str, str, str, str], ...] = (
    ("Fiziğe", "Fizik", "530", BOLUM_BILIM),
    ("Kimyaya", "Kimya", "540", BOLUM_BILIM),
    ("Biyolojiye", "Biyoloji", "570", BOLUM_BILIM),
    ("Astronomiye", "Astronomi", "520", BOLUM_BILIM),
    ("Matematiğe", "Matematik", "510", BOLUM_BILIM),
    ("Ekolojiye", "Ekoloji", "577", BOLUM_BILIM),
    ("Bilgisayara", "Bilgisayar", "004", BOLUM_BILIM),
    ("Felsefeye", "Felsefe", "100", BOLUM_TOPLUM),
    ("Psikolojiye", "Psikoloji", "150", BOLUM_TOPLUM),
    ("Sosyolojiye", "Sosyoloji", "301", BOLUM_TOPLUM),
    ("Mantığa", "Mantık", "160", BOLUM_TOPLUM),
    ("Ekonomiye", "Ekonomi", "330", BOLUM_TOPLUM),
    ("Tarihe", "Tarih", "900", BOLUM_TARIH),
    ("Coğrafyaya", "Coğrafya", "910", BOLUM_TARIH),
    ("Müziğe", "Müzik", "780", BOLUM_SANAT),
    ("Resme", "Resim", "750", BOLUM_SANAT),
)
YONELME_EKLERI = ("Giriş", "İlk Adımlar", "Yolculuk", "Yeni Bir Bakış", "Kısa Bir Yol")
UZERINE_EKLERI = ("Denemeler", "Notlar", "Konuşmalar", "Sorular")
DONEMLER = (
    "Osmanlı", "Selçuklu", "Cumhuriyet", "Kurtuluş Savaşı", "Anadolu Uygarlıkları",
    "Orta Asya Türk", "Bizans", "Roma", "Mezopotamya", "Hitit", "Lale Devri", "Tanzimat",
)  # fmt: skip
DONEM_EKLERI = ("Tarihi", "Döneminde Gündelik Hayat", "Üzerine İncelemeler", "Kronolojisi")
DOGA_OZNELERI = (
    "Böceklerin", "Kuşların", "Denizlerin", "Yıldızların", "Bitkilerin", "Dinozorların",
    "Volkanların", "Mağaraların", "Arıların", "Kaplumbağaların",
)  # fmt: skip
DOGA_EKLERI = ("Dünyası", "Gizemi", "Sırları")
COCUK_OZNELERI = (
    "Minik Kaplumbağanın", "Uçan Balonun", "Meraklı Kedinin", "Tembel Ayının",
    "Cesur Tavşanın", "Küçük Bulutun", "Yaramaz Keçinin", "Uykucu Baykuşun",
    "Sevimli Ejderhanın", "Kayıp Oyuncağın", "Neşeli Karıncanın", "Utangaç Zürafanın",
    "Gezgin Kırlangıcın", "Akıllı Robotun", "Çalışkan Arının", "Yalnız Deniz Fenerinin",
    "Konuşan Ağacın", "Şaşkın Penguenin", "Korkak Aslanın", "Renkli Şemsiyenin",
)  # fmt: skip
COCUK_EKLERI = (
    "Maceraları", "Yolculuğu", "Bayram Sabahı", "Doğum Günü", "Dileği", "Kış Uykusu",
    "İlk Okul Günü", "Hazinesi",
)  # fmt: skip
DERSLER = (
    "Türk Dili ve Edebiyatı", "Matematik", "Fizik", "Kimya", "Biyoloji", "Tarih",
    "Coğrafya", "İngilizce",
)  # fmt: skip
#: (başlık, konu, DOS) — "Danışma Kaynağı: Evet".
DANISMA_ESERLERI: tuple[tuple[str, str, str], ...] = (
    *((f"Resimli Bilim Ansiklopedisi — Cilt {i}", "Genel ansiklopedi", "030") for i in range(1, 9)),
    *((f"Ünlüler Ansiklopedisi — Cilt {i}", "Biyografi, ansiklopedi", "920") for i in range(1, 5)),
    ("Temel Türkçe Sözlük", "Türkçe, sözlük", "494.353"),
    ("Yazım Kılavuzu", "Türkçe, yazım", "494.35"),
    ("Deyimler ve Atasözleri Sözlüğü", "Türkçe, sözlük", "398.9"),
    ("Türkçe-İngilizce Sözlük", "İngilizce, sözlük", "423"),
    ("İngilizce-Türkçe Sözlük", "İngilizce, sözlük", "423"),
    ("Almanca-Türkçe Sözlük", "Almanca, sözlük", "433"),
    ("Genel Dünya Atlası", "Coğrafya, atlas", "912"),
    ("Türkiye Fiziki Atlası", "Coğrafya, atlas", "912.561"),
    ("Tarih Atlası", "Tarih, atlas", "911"),
    ("Bilim Terimleri Sözlüğü", "Bilim, sözlük", "503"),
    ("Matematik Formülleri Rehberi", "Matematik, rehber", "510.2"),
    ("Osmanlıca-Türkçe Sözlük", "Osmanlı Türkçesi, sözlük", "494.35"),
)
#: (dergi adı, konu, yayımlayan) — ciltlenmiş yıllık cilt, "Ciltli Süreli Yayın: Evet".
DERGILER: tuple[tuple[str, str, str], ...] = (
    ("Merak Penceresi Dergisi", "Süreli yayın, bilim", "Örnek Yayınevi"),
    ("Genç Kalemler Dergisi", "Süreli yayın, edebiyat", "Kumsal Yayınları"),
    ("Doğa ve İnsan Dergisi", "Süreli yayın, doğa", "Deneme Kitap"),
    ("Tarih Sayfaları Dergisi", "Süreli yayın, tarih", "Çıra Kitap"),
    ("Sanat Köşesi Dergisi", "Süreli yayın, sanat", "Işıltı Kitabevi"),
)
DERGI_YILLARI = tuple(range(2015, 2023))
#: (ad, biçim, konu, DOS) — "Kaynak Türü: Görsel-işitsel materyal".
GORSEL_ESERLER: tuple[tuple[str, str, str, str], ...] = (
    ("Anadolu'nun Kuşları", "Belgesel DVD", "Kuşlar, belgesel", "598"),
    ("Türk Halk Müziği Seçkisi", "Ses kaydı", "Halk müziği", "781.62"),
    ("Güneş Sistemi Belgeseli", "Belgesel DVD", "Astronomi, belgesel", "523.2"),
    ("Kapadokya Belgeseli", "Belgesel DVD", "Coğrafya, belgesel", "915.6"),
    ("Çocuk Şarkıları", "Ses kaydı", "Çocuk müziği", "782.42"),
    ("Klasik Batı Müziği Seçkisi", "Ses kaydı", "Klasik müzik", "781.68"),
    ("Deniz Canlıları Belgeseli", "Belgesel DVD", "Deniz biyolojisi, belgesel", "591.77"),
    ("İstanbul'un Tarihi Yapıları", "Belgesel DVD", "Tarih, belgesel", "956.1"),
    ("Masal Kuşağı Canlandırmaları", "Çizgi film DVD", "Çocuk edebiyatı, canlandırma", "791.43"),
    ("İngilizce Dinleme Alıştırmaları", "Ses kaydı", "İngilizce, dinleme", "428"),
    ("Volkanlar ve Depremler", "Belgesel DVD", "Yer bilimleri, belgesel", "551.2"),
    ("Ebru Sanatı Belgeseli", "Belgesel DVD", "Geleneksel sanatlar, belgesel", "745.5"),
    ("Türk Sanat Müziği Seçkisi", "Ses kaydı", "Türk sanat müziği", "781.62"),
    ("Satranç Dersleri", "Eğitim DVD'si", "Satranç", "794.1"),
    ("Göç Eden Kuşlar", "Belgesel DVD", "Kuşlar, belgesel", "598.156"),
)

#: Protokolün Türkçe arama sınaması (CLAUDE.md §2-8 kapısı): büyük/küçük harf
#: çiftleri aynı sonucu, noktalı/noktasız i farklı sonucu vermelidir.
ARAMA_SINAMALARI: tuple[tuple[str, str], ...] = (
    ("şiir", "ŞİİR"),
    ("ılık", "ILIK"),
    ("ince", "İNCE"),
    ("ilik", "İLİK"),
    ("ınce", "INCE"),
)


# ---------------------------------------------------------------------------
# Türkçe yardımcılar (programın kurallarının küçük eşleri; ölçüt test tarafındadır)
# ---------------------------------------------------------------------------
def tr_upper(metin: str) -> str:
    """Türkçe büyük harf: 'i' → 'İ', 'ı' → 'I' (çıplak `.upper()` 'i'yi 'I' basar)."""
    return metin.replace("i", "İ").replace("ı", "I").upper()


_TURK_HARFLERI = frozenset("ÇĞİÖŞÜ")
#: Düzeltme işaretli i Türkçede NOKTALI i'dir ("millî", "Türkî"): NFD onu noktasız
#: I'ya indirirdi (programdaki `normalize._AKSAN_ISTISNALARI` ile aynı kural).
_NOKTALI_I = frozenset("ÎÍÌÏĨĪĮ")


def arama_katlamasi(metin: str) -> str:
    """Programın arama katlamasının eşi: TR büyük harf, aksan katlaması, harf/rakam dışı → boşluk.

    Türk alfabesinin harfleri (Ç Ğ I İ Ö Ş Ü) korunur, düzeltme işareti düşer
    ('Rüzgâr' → 'RÜZGAR', 'Türkî' → 'TÜRKİ'). OZET'teki arama sayıları bununla
    hesaplanır; programın kendi katlamasıyla aynı sonucu verdiğini test sınar.
    """
    parcalar: list[str] = []
    bosluk = True
    for ch in tr_upper(metin):
        if ch in _NOKTALI_I:
            ch = "İ"
        elif ch not in _TURK_HARFLERI:
            ayrik = unicodedata.normalize("NFD", ch)
            ch = "".join(k for k in ayrik if not unicodedata.combining(k)) or ch
        for alt in ch:
            if alt.isalnum():
                parcalar.append(alt)
                bosluk = False
            elif not bosluk:
                parcalar.append(" ")
                bosluk = True
    return "".join(parcalar).strip()


def kaba_katlama(metin: str) -> str:
    """Eşsizlik denetimi için KABA katlama: Türk harfleri de düzleşir ('İ' ≡ 'I').

    Programın katlamasından daha fazlasını birleştirir; burada eşsiz olan başlık
    programın eşleştirmesinde de eşsizdir (aynı adlı iki eser "şüpheli" olurdu).
    """
    ayrik = unicodedata.normalize("NFKD", tr_upper(metin))
    duz = "".join(k for k in ayrik if not unicodedata.combining(k))
    return " ".join("".join(k if k.isalnum() else " " for k in duz).split())


def isbn13_saglama(govde: str) -> str:
    """12 hanelik gövdenin ISBN-13 sağlama hanesi."""
    toplam = sum(int(r) * (1 if i % 2 == 0 else 3) for i, r in enumerate(govde))
    return str((10 - toplam % 10) % 10)


def isbn10_saglama(govde: str) -> str:
    """9 hanelik gövdenin ISBN-10 sağlama hanesi ('X' = 10)."""
    toplam = sum(int(r) * (10 - i) for i, r in enumerate(govde))
    kalan = (11 - toplam % 11) % 11
    return "X" if kalan == 10 else str(kalan)


# ---------------------------------------------------------------------------
# Veri yapıları
# ---------------------------------------------------------------------------
@dataclass(frozen=True)
class Ogrenci:
    numara: int
    ad: str  # e-Okul gibi BÜYÜK HARF
    soyad: str
    cinsiyet: str  # yalnız e-Okul yerleşimi için; program okumaz


@dataclass(frozen=True)
class Sube:
    seviye: int
    harf: str

    @property
    def etiket(self) -> str:
        return f"{self.seviye}/{self.harf}"


Blok = tuple[Sube, list[Ogrenci]]


@dataclass(frozen=True)
class Personel:
    ad_soyad: str
    gorev: str
    kadro: str
    brans: str
    uye_turu: str  # programın beklenen sınıflaması
    gorev_taninir: bool


@dataclass
class KatalogSatiri:
    eser_adi: str
    yazar: str = ""
    cevirmen: str = ""
    yayinevi: str = ""
    baski: str = ""
    yil: int | None = None
    isbn: str = ""
    konu: str = ""
    sinif_kodu: str = ""
    dil: str = "Türkçe"
    tur: str = TUR_KITAP
    nusha: int | str = 1
    bolum: str = ""
    eski_no: str = ""
    ciltli: str | bool = ""
    danisma: str | bool = ""

    def hucreler(self) -> list[object]:
        """Şablonun sütun sırasıyla hücre değerleri (boş hücre `None`)."""
        degerler: list[object] = [
            self.eser_adi,
            self.yazar,
            self.cevirmen,
            self.yayinevi,
            self.baski,
            self.yil,
            self.isbn,
            self.konu,
            self.sinif_kodu,
            self.dil,
            self.tur,
            self.nusha,
            self.bolum,
            self.eski_no,
            self.ciltli,
            self.danisma,
        ]
        return [None if deger == "" else deger for deger in degerler]


@dataclass(frozen=True)
class SorunluSatir:
    """`katalog-sorunlu-satirlar.xlsx`'in bir satırı ve önizlemede beklenen sonucu."""

    satir_no: int  # Excel satır numarası (önizleme bununla konuşur)
    hucreler: tuple[object, ...]  # dosyadaki sütun sırasıyla
    beklenen: str  # kısa kod: mevcut / supheli / yeni / aktarilmadi / yazilamadi / bolum
    uyari: bool  # yeni satırda uyarı beklenir mi
    aciklama: str  # OZET ve protokol için


@dataclass
class Ozet:
    """Üretilen dosyaların beklenen sayıları (OZET.txt ve testin ölçütü)."""

    tohum: int
    kademe: str
    okul: str
    # öğrenci
    ogrenci: int = 0
    subeler: list[str] = field(default_factory=list)
    ayrilan: int = 0
    gelen: int = 0
    sube_degisen: int = 0
    donem2_toplam: int = 0
    # personel
    personel: int = 0
    ogretmen: int = 0
    diger_personel: int = 0
    taninmayan_gorev_satiri: int = 0
    # katalog
    katalog_satir: int = 0
    eser: int = 0
    nusha: int = 0
    bolunmus_ek_satir: int = 0
    ders_kitabi_satiri: int = 0
    danisma_eseri: int = 0
    sureli_eser: int = 0
    gorsel_eser: int = 0
    isbn13_eser: int = 0
    isbn10_eser: int = 0
    isbnsiz_eser: int = 0
    eski_kayitli_nusha: int = 0
    bolumler: list[str] = field(default_factory=list)
    #: bölüm → nüsha sayısı (protokol §8.3 etiket partisi bunu okur; F12 düzeltme turu)
    bolum_nusha: dict[str, int] = field(default_factory=dict)
    arama: dict[str, int] = field(default_factory=dict)
    # sorunlu satırlar
    sorunlu: list[SorunluSatir] = field(default_factory=list)

    def metin(self) -> str:
        """OZET.txt'nin metni (protokol "ne görülmeli" satırlarını buradan okur)."""
        s: list[str] = []
        s.append("KÜTÜPHANE DEFTERİ — DENEME VERİSİ ÖZETİ")
        s.append("=" * 40)
        s.append("Bu klasördeki bütün dosyalar UYDURMADIR. Gerçek kişi ya da okul verisi")
        s.append("içermez; depoya, e-postaya ve dış hizmetlere konmaz. Deneme bitince silin.")
        s.append(f"Tohum: {self.tohum} · Kademe: {self.kademe} · Okul adı: {self.okul}")
        s.append("")
        s.append(f"1. {DOSYA_OGRENCI}")
        s.append(f"   Öğrenci: {self.ogrenci} · Şube bloğu: {len(self.subeler)}")
        s.append(f"   Şubeler: {', '.join(self.subeler)}")
        s.append("   Önizlemede görülmeli:")
        s.append(f"   - {self.ogrenci} yeni öğrenci, 0 atlanan satır;")
        s.append(
            f"   - uyarılarda “e-Okul sınıf listesi biçimi algılandı: {len(self.subeler)} "
            "şube bloğu …” notu;"
        )
        s.append("   - I ve İ şubeleri AYRI iki şubedir (ikisi de listede).")
        s.append("")
        s.append(f"2. {DOSYA_OGRENCI_2} (1. dosya aktarıldıktan SONRA)")
        s.append(f"   Öğrenci: {self.donem2_toplam}")
        s.append("   Önizlemede görülmeli:")
        s.append(f"   - {self.gelen} yeni öğrenci (okula gelen);")
        s.append(f"   - {self.sube_degisen} güncellenen öğrenci (şube değiştiren);")
        s.append(f"   - {self.ayrilan} öğrenci ayrılış havuzuna eklenecek (okuldan ayrılan);")
        degismeyen = self.ogrenci - self.ayrilan - self.sube_degisen
        s.append(f"   - {degismeyen} değişmeyen öğrenci. Aktarım kimseyi silmez ve ayırmaz.")
        s.append("")
        s.append(f"3. {DOSYA_PERSONEL}")
        s.append(
            f"   Kişi: {self.personel} · öğretmen: {self.ogretmen} · "
            f"diğer personel: {self.diger_personel}"
        )
        s.append("   Önizlemede görülmeli:")
        s.append(f"   - {self.personel} yeni kişi;")
        s.append(
            f"   - satır {self.taninmayan_gorev_satiri}: görev tanınmadı, öğretmen sayıldı — "
            "“üye türünü denetleyin” uyarısı (görev metni saklanmaz);"
        )
        s.append("   - Branş ve kadro sütunları okunmaz.")
        s.append("")
        s.append(f"4. {DOSYA_KATALOG} (sayfa “{KATALOG_SAYFASI}”; “Örnek” sayfası okunmaz)")
        s.append(f"   Satır: {self.katalog_satir} · eser: {self.eser} · nüsha: {self.nusha}")
        s.append(
            f"   Bölümler ({len(self.bolumler)}; önizlemede “Bölüm listesinde bulunmayan "
            "değerler” olarak sorulur, her biri için “Yeni bölüm aç”):"
        )
        s.append(f"   {', '.join(self.bolumler)}")
        s.append("   Bölüm başına nüsha (etiket partisi seçimi için; azdan çoğa):")
        s.append(
            "   "
            + ", ".join(
                f"{ad} {sayi}"
                for ad, sayi in sorted(self.bolum_nusha.items(), key=lambda kv: (kv[1], kv[0]))
            )
        )
        s.append(
            f"   ({BOLUM_DANISMA} ve {BOLUM_DERS} bölümlerinin nüshaları ödünç verilmez: danışma "
            "kaynağıdır.)"
        )
        s.append("   Önizlemede görülmeli (bölüm kararlarından sonra):")
        s.append(f"   - yeni eser: {self.eser}")
        s.append(
            f"   - mevcut esere nüsha: {self.bolunmus_ek_satir} (aynı eserin ayrı satıra "
            "yazılmış nüshaları; eski kayıt no'lu)"
        )
        s.append(f"   - açılacak nüsha: {self.nusha}")
        s.append("   - şüpheli: 0 · aktarılmadı: 0 · yazılamadı: 0 · ISBN uyarısı: 0")
        s.append(f"   - ders kitabı varsayılanıyla danışma: {self.ders_kitabi_satiri} satır")
        s.append(
            f"   İçerik: ISBN-13'lü {self.isbn13_eser}, ISBN-10'lu {self.isbn10_eser}, "
            f"ISBN'siz {self.isbnsiz_eser} eser; danışma {self.danisma_eseri}, "
            f"ciltli süreli yayın {self.sureli_eser}, görsel-işitsel {self.gorsel_eser} eser; "
            f"eski kayıt no'lu {self.eski_kayitli_nusha} nüsha."
        )
        s.append("   Dijital kaynak (E-kitap, E-veri tabanı) Excel'le aktarılmaz; elle eklenir.")
        s.append("   Türkçe arama (Katalog → Eserler; aynı satırdaki iki yazım AYNI sayıyı verir):")
        for kucuk, buyuk in ARAMA_SINAMALARI:
            s.append(f"   - “{kucuk}” / “{buyuk}”: {self.arama.get(buyuk, 0)} eser")
        s.append("")
        s.append(f"5. {DOSYA_SORUNLU} (4. dosya UYGULANDIKTAN sonra önizleyin)")
        s.append("   Başlık satırı 4. satırdadır (üstünde kurum başlığı var); başlıklar")
        s.append("   şablonun eşanlamlılarıdır. Tanınmayan sütunlar: “Sıra”, “Fiyat”.")
        s.append("   Satır satır beklenen sonuç:")
        for satir in self.sorunlu:
            s.append(f"   - satır {satir.satir_no}: {satir.aciklama}")
        s.append("")
        return "\n".join(s)


# ---------------------------------------------------------------------------
# Öğrenci listesi
# ---------------------------------------------------------------------------
def _ad_soyad(rng: random.Random, *, cift_ad_orani: float = 0.15) -> tuple[str, str, str]:
    """(AD, SOYAD, cinsiyet) — BÜYÜK HARF, uydurma birleşim."""
    kiz = rng.random() < 0.5
    havuz = KIZ_ADLARI if kiz else ERKEK_ADLARI
    ad = rng.choice(havuz)
    if rng.random() < cift_ad_orani:
        ikinci = rng.choice(havuz)
        if ikinci != ad:
            ad = f"{ad} {ikinci}"
    return tr_upper(ad), tr_upper(rng.choice(SOYADLARI)), "Kız" if kiz else "Erkek"


def subeler(kademe: str) -> list[Sube]:
    seviyeler = KADEMELER[kademe][1]
    liste: list[Sube] = []
    for sira, seviye in enumerate(seviyeler):
        harfler = SUBELER_ILK if sira == 0 else SUBELER_DIGER
        liste.extend(Sube(seviye, harf) for harf in harfler)
    return liste


def ogrenci_bloklari(rng: random.Random, kademe: str, sayi: int) -> list[Blok]:
    """Şube şube öğrenciler; okul numaraları eşsiz ve 1-4 hanelidir."""
    sube_listesi = subeler(kademe)
    numaralar = rng.sample(range(1, max(9999, sayi * 4)), sayi)
    taban, artan = divmod(sayi, len(sube_listesi))
    bloklar: list[Blok] = []
    sira = 0
    for indis, sube in enumerate(sube_listesi):
        mevcut = taban + (1 if indis < artan else 0)
        ogrenciler: list[Ogrenci] = []
        for _ in range(mevcut):
            ad, soyad, cinsiyet = _ad_soyad(rng)
            ogrenciler.append(Ogrenci(numaralar[sira], ad, soyad, cinsiyet))
            sira += 1
        bloklar.append((sube, ogrenciler))
    return bloklar


def ikinci_donem(
    rng: random.Random, bloklar: Sequence[Blok], *, ayrilan: int, gelen: int, degisen: int
) -> tuple[list[Blok], dict[str, int]]:
    """Sonraki listeyi kurar: `ayrilan` kişi çıkar, `gelen` kişi girer, `degisen` şube değiştirir.

    Ayrılanlar AYRI şubelerden ve en az iki kişilik şubelerden seçilir: şube
    dosyada kalır, yani program onları "kanıtla" ayrılış havuzuna ekler (F1 ekleri 7).
    """
    yeni: list[tuple[Sube, list[Ogrenci]]] = [(sube, list(liste)) for sube, liste in bloklar]
    adaylar = [i for i, (_s, liste) in enumerate(yeni) if len(liste) >= 2]
    rng.shuffle(adaylar)
    ayrilan_sube = adaylar[:ayrilan]
    for indis in ayrilan_sube:
        liste = yeni[indis][1]
        liste.pop(rng.randrange(len(liste)))
    # Şube değiştirenler: aynı seviyede başka şubeye (ayrılanların şubeleri dışından;
    # kaynak şubede en az bir öğrenci kalır, şube dosyadan düşmez).
    tasinabilir = [i for i in adaylar[ayrilan:] if len(yeni[i][1]) >= 2]
    tasinan = 0
    for indis in tasinabilir:
        if tasinan >= degisen:
            break
        sube, liste = yeni[indis]
        hedefler = [j for j, (s, _l) in enumerate(yeni) if s.seviye == sube.seviye and j != indis]
        if not hedefler:
            continue
        ogrenci = liste.pop(rng.randrange(len(liste)))
        yeni[rng.choice(hedefler)][1].append(ogrenci)
        tasinan += 1
    # Gelenler: yeni ve eşsiz numaralarla rastgele şubelere.
    kullanilan = {o.numara for _s, liste in bloklar for o in liste}
    for _ in range(gelen):
        numara = rng.randrange(1, 9999)
        while numara in kullanilan:
            numara = rng.randrange(1, 9999)
        kullanilan.add(numara)
        ad, soyad, cinsiyet = _ad_soyad(rng)
        yeni[rng.randrange(len(yeni))][1].append(Ogrenci(numara, ad, soyad, cinsiyet))
    return list(yeni), {
        "ayrilan": len(ayrilan_sube),
        "gelen": gelen,
        "degisen": tasinan,
    }


def _hucre(sayfa: Worksheet, satir: int, sutun: int, deger: object) -> None:
    """e-Okul yerleşimindeki sütunlar 0 tabanlı yazılır (depodaki .xls örneğiyle aynı)."""
    sayfa.cell(row=satir, column=sutun + 1, value=deger)


def ogrenci_kitabi(
    bloklar: Sequence[Blok], kademe: str, ogretmen_adlari: Sequence[str]
) -> Workbook:
    """e-Okul OOG01001R020 yerleşimi: her şube için kurum başlığı → öğretmen satırları →
    sütun başlığı → öğrenciler → Kız/Erkek/Toplam dipnotu; en altta tarih ve rapor kodu.
    Sınıf ve şube YALNIZ blok başlığındadır (gerçek raporda da sütun yoktur)."""
    onek, _seviyeler, okul = KADEMELER[kademe]
    kitap = Workbook()
    sayfa = kitap.active
    assert sayfa is not None
    sayfa.title = "Sayfa1"
    satir = 1
    for sira, (sube, ogrenciler) in enumerate(bloklar):
        baslik = (
            "T.C.\nÖRNEK VALİLİĞİ\n"
            f"Örnek İlçe / {okul} Müdürlüğü\n"
            f"{onek} - {sube.seviye}. Sınıf / {sube.harf} Şubesi (ALANI YOK) Sınıf Listesi "
        )
        _hucre(sayfa, satir, 0, baslik)
        ogretmen = ogretmen_adlari[sira % len(ogretmen_adlari)] if ogretmen_adlari else ""
        _hucre(sayfa, satir + 1, 0, f"Sınıf Öğretmeni: {ogretmen}")
        _hucre(sayfa, satir + 1, 7, "Sınıf Başkanı:")
        _hucre(sayfa, satir + 2, 0, "Sınıf Müdür Yrd: ")
        _hucre(sayfa, satir + 2, 7, "Sınıf Başkan Yrd:")
        for sutun, metin in (
            (0, "S.No"),
            (1, "Öğrenci No"),
            (3, "Adı"),
            (7, "Soyadı"),
            (11, "Cinsiyeti"),
            (13, "Pansiyon Durum"),
        ):
            _hucre(sayfa, satir + 3, sutun, metin)
        for no, ogrenci in enumerate(ogrenciler, start=1):
            r = satir + 3 + no
            _hucre(sayfa, r, 0, no)
            _hucre(sayfa, r, 1, ogrenci.numara)
            _hucre(sayfa, r, 3, ogrenci.ad)
            _hucre(sayfa, r, 7, ogrenci.soyad)
            _hucre(sayfa, r, 11, ogrenci.cinsiyet)
        kiz = sum(1 for o in ogrenciler if o.cinsiyet == "Kız")
        dip = satir + 4 + len(ogrenciler)
        _hucre(sayfa, dip, 0, "Kız Öğrenci Sayısı        :")
        _hucre(sayfa, dip, 4, kiz)
        _hucre(sayfa, dip, 5, "Erkek Öğrenci Sayısı    :")
        _hucre(sayfa, dip, 9, len(ogrenciler) - kiz)
        _hucre(sayfa, dip, 11, "Toplam Öğrenci Sayısı    :")
        _hucre(sayfa, dip, 14, len(ogrenciler))
        satir = dip + 1
    _hucre(sayfa, satir, 11, _EOKUL_TARIH_SERISI)
    _hucre(sayfa, satir, 12, 0.755370370367018)
    _hucre(sayfa, satir, 14, 1)
    _hucre(sayfa, satir + 1, 1, "OOG01001R020")
    return kitap


# ---------------------------------------------------------------------------
# Personel listesi
# ---------------------------------------------------------------------------
def personel_listesi(rng: random.Random, sayi: int) -> list[Personel]:
    """Müdür, müdür yardımcıları, öğretmenler, diğer personel ve TEK tanınmayan görev."""
    adlar: set[str] = set()

    def yeni_ad() -> str:
        while True:
            ad, soyad, _c = _ad_soyad(rng, cift_ad_orani=0.1)
            tam = f"{ad} {soyad}"
            if tam not in adlar:
                adlar.add(tam)
                return tam

    kisiler: list[Personel] = [
        Personel(yeni_ad(), "Müdür", "KADROLU", rng.choice(BRANSLAR), OGRETMEN, True)
    ]
    kisiler.extend(
        Personel(yeni_ad(), "Müdür Yardımcısı", "KADROLU", rng.choice(BRANSLAR), OGRETMEN, True)
        for _ in range(max(1, sayi // 20))
    )
    diger_sayisi = max(2, sayi * 15 // 100)
    ogretmen_sayisi = sayi - len(kisiler) - diger_sayisi - 1
    agirliklar = [a for _g, _k, a in OGRETMEN_GOREVLERI]
    for _ in range(ogretmen_sayisi):
        gorev, kadro, _a = rng.choices(OGRETMEN_GOREVLERI, weights=agirliklar)[0]
        kisiler.append(Personel(yeni_ad(), gorev, kadro, rng.choice(BRANSLAR), OGRETMEN, True))
    kisiler.extend(
        Personel(yeni_ad(), rng.choice(DIGER_GOREVLER), "KADROLU", "", DIGER_PERSONEL, True)
        for _ in range(diger_sayisi)
    )
    kisiler.append(Personel(yeni_ad(), TANINMAYAN_GOREV, "KADROLU", "", OGRETMEN, False))
    rng.shuffle(kisiler)
    return kisiler


def personel_kitabi(kisiler: Sequence[Personel]) -> Workbook:
    """e-Okul OOK01001R1 yerleşimi: düz tablo + "Toplam Personel Sayısı" + tarih satırı."""
    kitap = Workbook()
    sayfa = kitap.active
    assert sayfa is not None
    sayfa.title = "Sayfa1"
    for sutun, metin in ((0, "ADI SOYADI"), (6, "GÖREVİ"), (8, "KADRO DURUMU"), (10, "BRANŞI")):
        _hucre(sayfa, 1, sutun, metin)
    for satir, kisi in enumerate(kisiler, start=2):
        _hucre(sayfa, satir, 0, kisi.ad_soyad)
        _hucre(sayfa, satir, 6, kisi.gorev)
        _hucre(sayfa, satir, 8, kisi.kadro)
        _hucre(sayfa, satir, 10, kisi.brans or None)
    son = len(kisiler) + 2
    _hucre(sayfa, son, 0, f"Toplam Personel Sayısı: {len(kisiler)}")
    _hucre(sayfa, son + 1, 12, _EOKUL_TARIH_SERISI)
    _hucre(sayfa, son + 1, 14, 0.898055555553583)
    return kitap


# ---------------------------------------------------------------------------
# Katalog
# ---------------------------------------------------------------------------
@dataclass
class _Eser:
    """Bir eser ve nüshalarının satıra dökülüşü (bölünmüş eser birden çok satırdır)."""

    satir: KatalogSatiri
    nusha: int = 1
    #: >1 ise eser bu kadar satıra, her satırda tek nüsha ve eski kayıt no ile yazılır.
    bolunmus: int = 0


class _IsbnKaynagi:
    """Eşsiz, sağlaması tutan UYDURMA ISBN'ler (978-605 ve eski 975 grubu)."""

    def __init__(self, rng: random.Random) -> None:
        self._on3 = iter(rng.sample(range(100000, 999999), 60000))
        self._on = iter(rng.sample(range(100000, 999999), 20000))

    def isbn13(self, *, tireli: bool = False) -> str:
        govde = f"978605{next(self._on3):06d}"
        numara = govde + isbn13_saglama(govde)
        if not tireli:
            return numara
        return f"{numara[:3]}-{numara[3:6]}-{numara[6:8]}-{numara[8:12]}-{numara[12]}"

    def isbn10(self, *, x_ile_biten: bool = False) -> str:
        while True:
            govde = f"975{next(self._on):06d}"
            saglama = isbn10_saglama(govde)
            if not x_ile_biten or saglama == "X":
                return govde + saglama


def _yazar_havuzu(rng: random.Random, sayi: int) -> list[str]:
    adlar: set[str] = set()
    while len(adlar) < sayi:
        ad = rng.choice(KIZ_ADLARI + ERKEK_ADLARI)
        adlar.add(f"{ad} {rng.choice(SOYADLARI)}")
    return sorted(adlar)


def _baslik_uret(adaylar: Iterable[str], alinan: set[str], adet: int) -> list[str]:
    """Kaba katlamada eşsiz ilk `adet` başlık (aynı adlı iki eser şüpheli olurdu)."""
    secilen: list[str] = []
    for aday in adaylar:
        if len(secilen) >= adet:
            break
        anahtar = kaba_katlama(aday)
        if anahtar in alinan:
            continue
        alinan.add(anahtar)
        secilen.append(aday)
    return secilen


def _kurgu_adaylari(rng: random.Random) -> list[tuple[str, str, str]]:
    govdeler = [f"{s} {a}" for s in SIFATLAR for a in ADLAR]
    govdeler += [f"{i} {y}" for i in ILGI_ADLARI for y in IYELIK_ADLARI]
    govdeler += [f"{yer} {b}" for yer in YERLER for b in BIR_ADLAR]
    rng.shuffle(govdeler)
    adaylar: list[tuple[str, str, str]] = []
    for govde in govdeler:
        ek, konu, dos = rng.choice(KURGU_TURLERI)
        adaylar.append((f"{govde}{ek}", konu, dos))
    # Havuz tükenirse ikinci tur: aynı gövde başka türle.
    for govde in govdeler:
        for ek, konu, dos in KURGU_TURLERI:
            adaylar.append((f"{govde}{ek}", konu, dos))
    return adaylar


def _bilgi_adaylari(rng: random.Random) -> list[tuple[str, str, str, str]]:
    adaylar: list[tuple[str, str, str, str]] = []
    for yonelme, konu, dos, bolum in BILIM_DALLARI:
        adaylar += [(f"{yonelme} {ek}", konu, dos, bolum) for ek in YONELME_EKLERI]
        adaylar += [(f"{konu} Üzerine {ek}", konu, dos, bolum) for ek in UZERINE_EKLERI]
        adaylar.append((f"Temel {konu} Bilgisi", konu, dos, bolum))
        adaylar.append((f"Gençler İçin {konu}", konu, dos, bolum))
        if konu != "Tarih":  # "Tarih Tarihi" olmasın
            adaylar.append((f"{konu} Tarihi", f"{konu}, tarih", dos, bolum))
        adaylar.append((f"Soru ve Cevaplarla {konu}", konu, dos, bolum))
    for donem in DONEMLER:
        adaylar += [(f"{donem} {ek}", "Tarih", "956.1", BOLUM_TARIH) for ek in DONEM_EKLERI]
    for ozne in DOGA_OZNELERI:
        adaylar += [(f"{ozne} {ek}", "Doğa bilimleri", "508", BOLUM_BILIM) for ek in DOGA_EKLERI]
    rng.shuffle(adaylar)
    return adaylar


def _cocuk_adaylari(rng: random.Random) -> list[str]:
    adaylar = [f"{o} {e}" for o in COCUK_OZNELERI for e in COCUK_EKLERI]
    rng.shuffle(adaylar)
    return adaylar


def _oran(eser: int, oran: float, tavan: int) -> int:
    return max(1, min(tavan, round(eser * oran)))


def katalog_eserleri(
    rng: random.Random, kademe: str, eser_sayisi: int, isbn: _IsbnKaynagi
) -> list[_Eser]:
    """`eser_sayisi` eşsiz eser: klasikler, ders kitapları, danışma, süreli yayın,
    görsel-işitsel, bilgi, çocuk ve kurgu. İlk iki eser her boyutta KLASIKLER[0:2]'dir.

    ISBN kaynağı dışarıdan verilir: sorunlu satırlar dosyası da AYNI kaynaktan çeker,
    böylece "yeni eser" beklenen bir sınama satırının numarası katalogdaki bir esere
    denk gelip satırı şüpheliye çeviremez.
    """
    yazarlar = _yazar_havuzu(rng, max(60, eser_sayisi // 5))
    cevirmenler = _yazar_havuzu(rng, 40)
    alinan: set[str] = set()
    eserler: list[_Eser] = []

    def kitap_satiri(
        ad: str, yazar: str, konu: str, dos: str, bolum: str, *, yil: int | None = None
    ) -> KatalogSatiri:
        yil = yil if yil is not None else rng.randint(1965, 2025)
        satir = KatalogSatiri(
            eser_adi=ad,
            yazar=yazar,
            yayinevi=rng.choice(YAYINEVLERI),
            baski=f"{rng.randint(1, 15)}. baskı" if rng.random() < 0.35 else "",
            yil=yil if rng.random() > 0.05 else None,
            konu=konu,
            sinif_kodu=dos if rng.random() > 0.10 else "",
            bolum=bolum,
        )
        zar = rng.random()
        if yil < 1985 or zar < 0.15:
            satir.isbn = ""
        elif zar < 0.25:
            satir.isbn = isbn.isbn10()
        else:
            satir.isbn = isbn.isbn13(tireli=rng.random() < 0.33)
        return satir

    # 1) Klasikler (ilk iki her boyutta; kalanı boyuta göre).
    klasik_sayisi = min(len(KLASIKLER), max(2, eser_sayisi // 40))
    for sira, klasik in enumerate(KLASIKLER[:klasik_sayisi]):
        alinan.add(kaba_katlama(klasik.ad))
        satir = KatalogSatiri(
            eser_adi=klasik.ad,
            yazar=klasik.yazar,
            cevirmen=rng.choice(cevirmenler) if klasik.ceviri else "",
            yayinevi=rng.choice(YAYINEVLERI),
            yil=klasik.yil or rng.randint(1995, 2024),
            isbn="" if klasik.isbn_yok else isbn.isbn13(),
            konu=klasik.konu,
            sinif_kodu=klasik.dos,
            dil=klasik.dil,
            bolum=klasik.bolum,
            danisma=EVET if klasik.danisma else "",
        )
        # İlk iki klasik tek nüsha: sorunlu satırlar dosyası onlara nüsha ekler/benzer.
        nusha = 1 if sira < 2 else rng.choice((1, 1, 2, 2, 3, 4))
        eserler.append(_Eser(satir, nusha=nusha))

    # 2) Ders kitapları (sınıf seti; danışma varsayılanı "Ders kitabı" türünden gelir).
    seviyeler = KADEMELER[kademe][1]
    dersler = [(d, s) for s in seviyeler for d in DERSLER]
    for ders, seviye in dersler[: _oran(eser_sayisi, 0.016, len(dersler))]:
        ad = f"{ders} {seviye}. Sınıf Ders Kitabı"
        alinan.add(kaba_katlama(ad))
        satir = KatalogSatiri(
            eser_adi=ad,
            yayinevi="Örnek Ders Kitapları Yayınevi",
            yil=rng.randint(2023, 2025),
            isbn=isbn.isbn13(),
            konu=f"{ders}, ders kitabı",
            tur=TUR_DERS,
            bolum=BOLUM_DERS,
        )
        eserler.append(_Eser(satir, nusha=rng.randint(10, 30)))

    # 3) Danışma kaynakları.
    for ad, konu, dos in DANISMA_ESERLERI[: _oran(eser_sayisi, 0.012, len(DANISMA_ESERLERI))]:
        alinan.add(kaba_katlama(ad))
        satir = KatalogSatiri(
            eser_adi=ad,
            yayinevi=rng.choice(YAYINEVLERI),
            yil=rng.randint(1998, 2024),
            isbn=isbn.isbn13(),
            konu=konu,
            sinif_kodu=dos,
            bolum=BOLUM_DANISMA,
            danisma=EVET,
        )
        eserler.append(_Eser(satir, nusha=rng.choice((1, 1, 2))))

    # 4) Ciltli süreli yayınlar (yıllık cilt; ISBN yok).
    ciltler = [(d, y) for y in DERGI_YILLARI for d in DERGILER]
    for (dergi, konu, yayinci), yil in ciltler[: _oran(eser_sayisi, 0.01, len(ciltler))]:
        ad = f"{dergi} — {yil} Cildi"
        alinan.add(kaba_katlama(ad))
        satir = KatalogSatiri(
            eser_adi=ad,
            yayinevi=yayinci,
            yil=yil,
            konu=konu,
            sinif_kodu="050",
            tur=TUR_SURELI,
            bolum=BOLUM_SURELI,
            ciltli=EVET,
        )
        eserler.append(_Eser(satir))

    # 5) Görsel-işitsel materyal.
    for ad, bicim, konu, dos in GORSEL_ESERLER[: _oran(eser_sayisi, 0.006, len(GORSEL_ESERLER))]:
        tam_ad = f"{ad} ({bicim})"
        alinan.add(kaba_katlama(tam_ad))
        satir = KatalogSatiri(
            eser_adi=tam_ad,
            yayinevi=rng.choice(YAYINEVLERI),
            yil=rng.randint(2005, 2024),
            konu=konu,
            sinif_kodu=dos,
            tur=TUR_GORSEL,
            bolum=BOLUM_GORSEL,
        )
        eserler.append(_Eser(satir))

    kalan = eser_sayisi - len(eserler)
    # 6) Bilgi ve çocuk kitapları (havuz tükenirse fark kurguya geçer).
    bilgi_hedef = round(kalan * 0.25)
    cocuk_hedef = round(kalan * 0.08)
    bilgi = _bilgi_adaylari(rng)
    bilgi_adlari = {ad: (konu, dos, bolum) for ad, konu, dos, bolum in bilgi}
    for ad in _baslik_uret((b[0] for b in bilgi), alinan, bilgi_hedef):
        konu, dos, bolum = bilgi_adlari[ad]
        eserler.append(_Eser(kitap_satiri(ad, rng.choice(yazarlar), konu, dos, bolum)))
    for ad in _baslik_uret(_cocuk_adaylari(rng), alinan, cocuk_hedef):
        eserler.append(
            _Eser(kitap_satiri(ad, rng.choice(yazarlar), "Çocuk edebiyatı", "894.353", BOLUM_COCUK))
        )
    # 7) Kurgu: kalan bütün eserler.
    kurgu = _kurgu_adaylari(rng)
    kurgu_bilgisi = {ad: (konu, dos) for ad, konu, dos in kurgu}
    for ad in _baslik_uret((k[0] for k in kurgu), alinan, eser_sayisi - len(eserler)):
        konu, dos = kurgu_bilgisi[ad]
        yazar = rng.choice(yazarlar)
        if rng.random() < 0.05:
            yazar = f"{yazar}, {rng.choice(yazarlar)}"
        eserler.append(_Eser(kitap_satiri(ad, yazar, konu, dos, BOLUM_EDEBIYAT)))
    if len(eserler) != eser_sayisi:  # pragma: no cover — havuzlar 10.000'i karşılar
        raise RuntimeError(f"Başlık havuzu yetmedi: {len(eserler)} / {eser_sayisi}")

    # Nüsha dağılımı ve eski kayıt no (klasikler, ders kitapları ve danışma dışında).
    for eser in eserler[klasik_sayisi:]:
        if eser.satir.tur != TUR_KITAP or eser.satir.bolum == BOLUM_DANISMA:
            continue
        zar = rng.random()
        if zar < 0.015:
            eser.bolunmus = rng.choice((2, 3))
        elif zar < 0.13:
            eser.nusha = 1
            eser.satir.eski_no = "eski"  # numara aşağıda verilir
        elif zar < 0.25:
            eser.nusha = rng.choice((2, 3))
        elif zar < 0.29:
            eser.nusha = rng.randint(4, 6)
        if rng.random() < 0.05:
            eser.satir.danisma = HAYIR  # açıkça "Hayır" yazılmış hücre de okunur
    return eserler


def katalog_satirlari(rng: random.Random, eserler: Sequence[_Eser]) -> list[KatalogSatiri]:
    """Eserleri satıra döker: bölünmüş eser k satır (tek nüsha + eski kayıt no).

    İlk iki satır KLASIKLER[0:2]'dir; kalan satırların sırası karıştırılır (okulun
    eski defterinden gelen liste de sırasızdır; aynı eserin satırları ayrık düşer).
    """
    eski_numaralar = iter(rng.sample(range(1, 30000), 20000))
    satirlar: list[KatalogSatiri] = []
    for eser in eserler:
        if eser.bolunmus:
            for _ in range(eser.bolunmus):
                kopya = KatalogSatiri(**{**eser.satir.__dict__})
                kopya.nusha = 1
                kopya.eski_no = str(next(eski_numaralar))
                satirlar.append(kopya)
            continue
        satir = eser.satir
        satir.nusha = eser.nusha
        if satir.eski_no == "eski":
            numara = next(eski_numaralar)
            satir.eski_no = f"E-{numara}" if numara % 5 == 0 else str(numara)
        satirlar.append(satir)
    bas, govde = satirlar[:2], satirlar[2:]
    rng.shuffle(govde)
    return bas + govde


def _baslik_satiri(sayfa: Worksheet, basliklar: Sequence[str], satir: int = 1) -> None:
    kalin = Font(bold=True)
    for sutun, metin in enumerate(basliklar, start=1):
        hucre = sayfa.cell(row=satir, column=sutun, value=metin)
        hucre.font = kalin


def katalog_kitabi(satirlar: Sequence[KatalogSatiri]) -> Workbook:
    """Şablonun doldurulmuş hâli: "Katalog" sayfası + okunmayan "Örnek" sayfası."""
    kitap = Workbook()
    sayfa = kitap.active
    assert sayfa is not None
    sayfa.title = KATALOG_SAYFASI
    _baslik_satiri(sayfa, KATALOG_BASLIKLARI)
    sayfa.freeze_panes = "A2"
    metin_sutunlari = {7, 9, 14}  # ISBN, Sınıflama Kodu, Eski Kayıt No (şablonda metin biçimli)
    for r, satir in enumerate(satirlar, start=2):
        for sutun, deger in enumerate(satir.hucreler(), start=1):
            if deger is None:
                continue
            hucre = sayfa.cell(row=r, column=sutun, value=deger)
            if sutun in metin_sutunlari:
                hucre.number_format = "@"
    genislikler = (44, 26, 22, 24, 10, 10, 20, 30, 14, 14, 22, 8, 18, 14, 10, 10)
    for harf, genislik in zip("ABCDEFGHIJKLMNOP", genislikler, strict=True):
        sayfa.column_dimensions[harf].width = genislik
    ornek = kitap.create_sheet("Örnek")
    _baslik_satiri(ornek, KATALOG_BASLIKLARI)
    ornek.append(["Örnek satır — bu sayfa içe aktarılmaz", "Örnek Yazar"])
    return kitap


# ---------------------------------------------------------------------------
# Sorunlu satırlar dosyası
# ---------------------------------------------------------------------------
SORUNLU_BASLIKLAR: tuple[str, ...] = (
    "Sıra",
    "Kitabın Adı",
    "Yazarı",
    "Çeviren",
    "Yayın Evi",
    "Baskı",
    "Basım Yılı",
    "ISBN No",
    "Konusu",
    "DOS Kodu",
    "Dili",
    "Türü",
    "Adet",
    "Raf",
    "Demirbaş No",
    "Ciltli",
    "Danışma",
    "Fiyat",
)
#: Başlık satırının Excel satır numarası (üstünde kurum başlığı ve açıklama var).
SORUNLU_BASLIK_SATIRI = 4


def sorunlu_satirlar(
    isbn: _IsbnKaynagi, ilk: KatalogSatiri, ikinci: KatalogSatiri
) -> list[SorunluSatir]:
    """Önizlemenin bütün kollarını gösteren satırlar (katalog dosyası UYGULANDIKTAN sonra).

    `ilk` ve `ikinci` katalog dosyasının ilk iki satırıdır (KLASIKLER[0:2]).
    """
    hatali = isbn.isbn13()
    hatali = hatali[:-1] + str((int(hatali[-1]) + 1) % 10)

    def satir(
        ad: str,
        *,
        yazar: str = "Örnek Yazar",
        yil: object = 2020,
        isbn_no: str = "",
        konu: str = "Türk edebiyatı, roman",
        tur: object = None,
        adet: object = 1,
        raf: str = BOLUM_EDEBIYAT,
        eski: object = None,
        ciltli: object = None,
        danisma: object = None,
    ) -> list[object]:
        return [
            ad or None,
            yazar or None,
            None,
            "Örnek Yayınevi",
            None,
            yil,
            isbn_no or None,
            konu,
            "894.353",
            "Türkçe",
            tur,
            adet,
            raf,
            eski,
            ciltli,
            danisma,
        ]

    tanimlar: list[tuple[list[object], str, bool, str]] = [
        (
            satir(ilk.eser_adi, yazar=ilk.yazar, isbn_no=ilk.isbn, konu=ilk.konu),
            "mevcut",
            False,
            f"“{ilk.eser_adi}” — katalogdaki esere nüsha eklenir (Mevcut esere nüsha)",
        ),
        (
            satir(f"{ilk.eser_adi} (Resimli Baskı)", yazar=ilk.yazar, isbn_no=ilk.isbn),
            "supheli",
            False,
            "ISBN katalogdaki bir eserle aynı, ad farklı — Şüpheli (karar bekler)",
        ),
        (
            satir(ikinci.eser_adi, yazar="Başka Bir Yazar", konu=ikinci.konu),
            "supheli",
            False,
            f"“{ikinci.eser_adi}” adı aynı, yazar farklı — Şüpheli (karar bekler)",
        ),
        (
            satir("Sınama — ISBN sağlaması hatalı", isbn_no=hatali),
            "yeni",
            True,
            "ISBN sağlaması tutmuyor — Yeni eser, “ISBN uyarısı” (kayıt engellenmez)",
        ),
        (
            satir("Sınama — on haneli ISBN, X ile biten", isbn_no=isbn.isbn10(x_ile_biten=True)),
            "yeni",
            False,
            "10 haneli ISBN (son hanesi X) — Yeni eser, uyarı yok",
        ),
        (
            satir("Sınama — nüsha sayısı aralık yazılmış", adet="2-3"),
            "aktarilmadi",
            False,
            "“2-3” sayı değil — Aktarılmadı",
        ),
        (
            satir("Sınama — nüsha sayısı 50'den fazla", adet=60),
            "aktarilmadi",
            False,
            "Nüsha sayısı 60 (en çok 50) — Aktarılmadı",
        ),
        (
            satir("Sınama — eski kayıt no ile üç nüsha", adet=3, eski="1452"),
            "aktarilmadi",
            False,
            "Eski kayıt no dolu ama nüsha sayısı 3 — Aktarılmadı",
        ),
        (
            satir(
                "Sınama — ciltsiz dergi", konu="Süreli yayın, bilim", tur="Dergi", raf=BOLUM_SURELI
            ),
            "yazilamadi",
            False,
            "“Dergi” = süreli yayın, ciltli değil — nüsha açılmaz, satır hatalı (yazılamadı)",
        ),
        (
            satir("Sınama — tanınmayan kaynak türü", tur="Roman"),
            "aktarilmadi",
            False,
            "Kaynak türü “Roman” tanınmadı (roman “Konu”ya yazılır) — Aktarılmadı",
        ),
        (
            satir("Sınama — danışma sütununda serbest metin", danisma="belki"),
            "aktarilmadi",
            False,
            "Danışma “belki” — Aktarılmadı (Evet/Hayır)",
        ),
        (
            satir("Sınama — yayın yılı okunamıyor", yil="bilinmiyor"),
            "yeni",
            True,
            "Yayın yılı okunamadı — Yeni eser, yıl boş ve uyarılı",
        ),
        (
            satir("", yazar="Adı Yazılmamış Satırın Yazarı"),
            "aktarilmadi",
            False,
            "Eser adı boş — Aktarılmadı",
        ),
        (
            satir("Sınama — ders kitabı konusu", konu="Ders kitabı"),
            "yeni",
            False,
            "Konu “Ders kitabı”, danışma boş — Yeni eser, danışma kaynağı sayılır",
        ),
        (
            satir("Sınama — listede olmayan bölüm", raf="Gezi Rafı"),
            "bolum",
            False,
            "Bölüm “Gezi Rafı” listede yok — önizlemede sorulur (Yeni bölüm aç / eşleştir)",
        ),
        (
            satir("Sınama — DVD yazılmış tür", konu="Belgesel", tur="DVD", raf=BOLUM_GORSEL),
            "yeni",
            False,
            "Tür “DVD” — Görsel-işitsel materyal olarak Yeni eser",
        ),
        (
            satir("Sınama — tireli ISBN", isbn_no=isbn.isbn13(tireli=True)),
            "yeni",
            False,
            "Tireli ISBN-13 — Yeni eser, uyarı yok",
        ),
        (
            satir("Sınama — Excel DOĞRU değeri", danisma=True),
            "yeni",
            False,
            "Danışma hücresi Excel DOĞRU değeri — Yeni eser, danışma kaynağı",
        ),
    ]
    # "Sıra" ve "Fiyat" programın tanımadığı sütunlardır: okunmaz, önizleme adlarını sayar.
    return [
        SorunluSatir(
            satir_no=SORUNLU_BASLIK_SATIRI + sira,
            hucreler=(sira, *hucreler, f"{10 + sira * 3},00 TL"),
            beklenen=beklenen,
            uyari=uyari,
            aciklama=aciklama,
        )
        for sira, (hucreler, beklenen, uyari, aciklama) in enumerate(tanimlar, start=1)
    ]


def sorunlu_kitabi(satirlar: Sequence[SorunluSatir]) -> Workbook:
    kitap = Workbook()
    sayfa = kitap.active
    assert sayfa is not None
    sayfa.title = "Liste"  # "Katalog" adı yok: program ilk sayfayı okur
    sayfa.cell(row=1, column=1, value="ÖRNEK ANADOLU LİSESİ KÜTÜPHANESİ")
    sayfa.cell(row=2, column=1, value="Eski defterden aktarılan liste (sınama satırları)")
    _baslik_satiri(sayfa, SORUNLU_BASLIKLAR, satir=SORUNLU_BASLIK_SATIRI)
    for satir in satirlar:
        for sutun, deger in enumerate(satir.hucreler, start=1):
            if deger is not None:
                sayfa.cell(row=satir.satir_no, column=sutun, value=deger)
    return kitap


# ---------------------------------------------------------------------------
# Hepsi
# ---------------------------------------------------------------------------
def _arama_sayilari(eserler: Sequence[_Eser]) -> dict[str, int]:
    """Programın arama anahtarının (ad + yazar + konu + ISBN) eşiyle sonuç sayıları."""
    anahtarlar = [
        " ".join(
            arama_katlamasi(parca) for parca in (e.satir.eser_adi, e.satir.yazar, e.satir.konu)
        )
        for e in eserler
    ]
    sayilar: dict[str, int] = {}
    for _kucuk, buyuk in ARAMA_SINAMALARI:
        terim = arama_katlamasi(buyuk)
        sayilar[buyuk] = sum(1 for anahtar in anahtarlar if terim in anahtar)
    return sayilar


def uret(
    cikti: Path,
    *,
    tohum: int = VARSAYILAN_TOHUM,
    kademe: str = "lise",
    ogrenci: int = VARSAYILAN_OGRENCI,
    personel: int = VARSAYILAN_PERSONEL,
    eser: int = VARSAYILAN_ESER,
) -> Ozet:
    """Bütün dosyaları `cikti` klasörüne yazar ve beklenen sayıları döndürür."""
    if kademe not in KADEMELER:
        raise ValueError(f"Bilinmeyen kademe: {kademe}")
    cikti.mkdir(parents=True, exist_ok=True)
    ozet = Ozet(tohum=tohum, kademe=kademe, okul=KADEMELER[kademe][2])

    # Her dosya kendi alt tohumuyla: bir dosyanın boyutu öbürünün içeriğini değiştirmez.
    # Kriptografik rastgelelik GEREKMEZ: amaç tekrarlanabilir uydurma veridir (S311).
    rng_personel = random.Random(f"{tohum}-personel")  # noqa: S311
    rng_ogrenci = random.Random(f"{tohum}-ogrenci")  # noqa: S311
    rng_katalog = random.Random(f"{tohum}-katalog")  # noqa: S311

    # Personel (öğrenci listesinin "Sınıf Öğretmeni" satırları buradan).
    kisiler = personel_listesi(rng_personel, personel)
    personel_kitabi(kisiler).save(cikti / DOSYA_PERSONEL)
    ozet.personel = len(kisiler)
    ozet.ogretmen = sum(1 for k in kisiler if k.uye_turu == OGRETMEN)
    ozet.diger_personel = sum(1 for k in kisiler if k.uye_turu == DIGER_PERSONEL)
    ozet.taninmayan_gorev_satiri = next(
        sira for sira, k in enumerate(kisiler, start=2) if not k.gorev_taninir
    )
    ogretmen_adlari = [k.ad_soyad for k in kisiler if k.uye_turu == OGRETMEN]

    # Öğrenciler: iki dönem.
    bloklar = ogrenci_bloklari(rng_ogrenci, kademe, ogrenci)
    ogrenci_kitabi(bloklar, kademe, ogretmen_adlari).save(cikti / DOSYA_OGRENCI)
    ozet.ogrenci = sum(len(liste) for _s, liste in bloklar)
    ozet.subeler = [s.etiket for s, _l in bloklar]
    bloklar2, fark = ikinci_donem(
        rng_ogrenci,
        bloklar,
        ayrilan=max(1, ogrenci // 60),
        gelen=max(1, ogrenci // 90),
        degisen=max(1, ogrenci // 150),
    )
    ogrenci_kitabi(bloklar2, kademe, ogretmen_adlari).save(cikti / DOSYA_OGRENCI_2)
    ozet.ayrilan, ozet.gelen, ozet.sube_degisen = fark["ayrilan"], fark["gelen"], fark["degisen"]
    ozet.donem2_toplam = sum(len(liste) for _s, liste in bloklar2)

    # Katalog.
    isbn_kaynagi = _IsbnKaynagi(rng_katalog)
    eserler = katalog_eserleri(rng_katalog, kademe, eser, isbn_kaynagi)
    satirlar = katalog_satirlari(rng_katalog, eserler)
    katalog_kitabi(satirlar).save(cikti / DOSYA_KATALOG)
    ozet.katalog_satir = len(satirlar)
    ozet.eser = len(eserler)
    ozet.nusha = sum(int(s.nusha) for s in satirlar)
    ozet.bolunmus_ek_satir = sum(e.bolunmus - 1 for e in eserler if e.bolunmus)
    ozet.ders_kitabi_satiri = sum(1 for s in satirlar if s.tur == TUR_DERS)
    ozet.danisma_eseri = sum(1 for e in eserler if e.satir.danisma == EVET)
    ozet.sureli_eser = sum(1 for e in eserler if e.satir.tur == TUR_SURELI)
    ozet.gorsel_eser = sum(1 for e in eserler if e.satir.tur == TUR_GORSEL)
    ozet.isbn13_eser = sum(1 for e in eserler if len(_rakamlar(e.satir.isbn)) == 13)
    ozet.isbn10_eser = sum(1 for e in eserler if len(_rakamlar(e.satir.isbn)) == 10)
    ozet.isbnsiz_eser = sum(1 for e in eserler if not e.satir.isbn)
    ozet.eski_kayitli_nusha = sum(1 for s in satirlar if s.eski_no)
    ozet.bolumler = list(dict.fromkeys(s.bolum for s in satirlar if s.bolum))
    for s in satirlar:
        if s.bolum:
            ozet.bolum_nusha[s.bolum] = ozet.bolum_nusha.get(s.bolum, 0) + int(s.nusha)
    ozet.arama = _arama_sayilari(eserler)

    # Sorunlu satırlar (ilk iki katalog satırına ve aynı ISBN kaynağına dayanır).
    ozet.sorunlu = sorunlu_satirlar(isbn_kaynagi, satirlar[0], satirlar[1])
    sorunlu_kitabi(ozet.sorunlu).save(cikti / DOSYA_SORUNLU)

    (cikti / DOSYA_OZET).write_text(ozet.metin(), encoding="utf-8")
    return ozet


def _rakamlar(metin: str) -> str:
    return "".join(k for k in metin if k.isdigit() or k in "Xx")


def _aralik(ad: str, sinir: tuple[int, int]) -> Callable[[str], int]:
    alt, ust = sinir

    def cevir(deger: str) -> int:
        try:
            sayi = int(deger)
        except ValueError as exc:
            raise argparse.ArgumentTypeError(f"{ad} bir tam sayı olmalı: {deger!r}") from exc
        if not alt <= sayi <= ust:
            raise argparse.ArgumentTypeError(f"{ad} {alt} ile {ust} arasında olmalı: {sayi}")
        return sayi

    return cevir


def argumanlar(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Saha kabulü için UYDURMA deneme verisi üretir (e-Okul öğrenci ve personel "
            "listeleri, katalog Excel'i). Dosyalar depoya konmaz."
        )
    )
    parser.add_argument(
        "--cikti",
        type=Path,
        default=Path(VARSAYILAN_CIKTI),
        help=f"çıktı klasörü (varsayılan: {VARSAYILAN_CIKTI}/ — .gitignore'dadır)",
    )
    parser.add_argument("--tohum", type=int, default=VARSAYILAN_TOHUM, help="rastgelelik tohumu")
    parser.add_argument(
        "--kademe", choices=sorted(KADEMELER), default="lise", help="okulun kademesi"
    )
    parser.add_argument(
        "--ogrenci",
        type=_aralik("öğrenci sayısı", OGRENCI_SINIRI),
        default=VARSAYILAN_OGRENCI,
        help=f"öğrenci sayısı ({OGRENCI_SINIRI[0]}-{OGRENCI_SINIRI[1]})",
    )
    parser.add_argument(
        "--personel",
        type=_aralik("personel sayısı", PERSONEL_SINIRI),
        default=VARSAYILAN_PERSONEL,
        help=f"personel sayısı ({PERSONEL_SINIRI[0]}-{PERSONEL_SINIRI[1]})",
    )
    parser.add_argument(
        "--eser",
        type=_aralik("eser sayısı", ESER_SINIRI),
        default=VARSAYILAN_ESER,
        help=f"katalogdaki eser sayısı ({ESER_SINIRI[0]}-{ESER_SINIRI[1]})",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    secenek = argumanlar(argv)
    ozet = uret(
        secenek.cikti,
        tohum=secenek.tohum,
        kademe=secenek.kademe,
        ogrenci=secenek.ogrenci,
        personel=secenek.personel,
        eser=secenek.eser,
    )
    sys.stdout.write(ozet.metin())
    sys.stdout.write(
        f"\nDosyalar yazıldı: {secenek.cikti}\n"
        "UYARI: Bu dosyalar UYDURMADIR ve yalnız deneme kurulumu içindir; depoya eklemeyin.\n"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
