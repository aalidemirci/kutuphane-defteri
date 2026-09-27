"""Taşınır Kütüphane Defteri dökümü ve yönetim hesabı cetveli hazırlığı (E11) — PDF + XLSX.

Tasarım §8.4, §10 E11; D1'in F10 kısmı (OYS "10/1-m" diyordu — öyle bir bent yoktur; doğru
atıf 34/2-c ve 34/3-a). Her atıf `docs/mevzuat/tasinir-mal-yonetmeligi.md`'den doğrulanır
(`tests/test_dokumler.py`, `tests/test_dokum_metinleri.py`).

**Taşınır Kütüphane Defteri dökümü** (TMY md. 9/1-ç: "Bu defter, kütüphanelerdeki yazma ve
basma nadir eserler ile kitap ve kitap dışı materyal için tutulur. Her bir taşınır için ayrı
kayıt yapılır."): her nüsha bir satırdır, kayıt no sırasıyla. Kayıttan düşülmüş ve
devredilmiş nüsha defterde kalır (çıkış tarihi ve durumuyla). **Ciltletilmemiş süreli yayın
GİRMEZ** (kod kapısı): dergi ve gazete gibi süreli yayınlar için Varlık İşlem Fişi düzenlenmez
(md. 10/1-a-4), cilt birliği sağlananlar ciltletildikten sonra kayda alınır (md. 15/4).
Nüsha kuralları (`services.catalog.ensure_copy_allowed`) böyle bir nüshanın açılmasına zaten
izin vermez; döküm kuralı yine de kendi süzgeciyle uygular (savunma derinliği — eski veri ya
da doğrudan yazım). Süzgeç TEKTİR (`selectors_sayim.ciltsiz_sureli_yayin_q`) ve yönetim hesabı
cetveli hazırlığı ile TMY 34/1 büyüklükleri de onu uygular: defter ile cetvel ayrışmaz (F10
düzeltme turu). Dijital kaynağın nüshası yoktur. Silinmiş (yanlış açılmış) kayıt girmez.

**Yönetim hesabı cetveli hazırlığı** (md. 34/2-c, 34/3-a): taşınır mal yönetim hesabının
büyüklükleri (md. 34/1) nüsha sayısıyla ve — birim fiyatı kayıtlı nüshalarda — tutarla.
YALNIZ "yıl sonu sayımı" işaretli ve ONAYLANMIŞ sayımdan basılır (F9 ekleri K6). Zincir: 34/3-a
"Sayım kurulu tarafından onaylanan Taşınır Sayım ve Döküm Cetveline dayanılarak …"; o cetvel
(32/9) kayıtların sayım sonuçlarıyla uygunluğu sağlandıktan sonra düzenlenir; uygunluk 32/7'nin
belgeleriyle sağlanır — programda sayımın onayıyla (noksan kayıttan düşülür, fazla kayda
alınır). İptal edilmiş sayım listelenmez; bir mali yılın tek yıl sonu sayımı olur
(`services.stocktake`, F10 düzeltme turu). Sayılar TEK
kaynaktandır: `selectors_sayim.tmy_34_1` (E10'un ekiyle aynı). Tutarlar aynı sınıflamayla
hesaplanır; test iki hesabın nüsha sayılarının eşitliğini sınar. Program TKYS'nin yerine
GEÇMEZ: resmî cetveller TKYS'dedir ve belge bunu söyler. 34/2-c'ye göre Kütüphane Yönetim
Hesabı Cetveli kütüphane olarak faaliyet gösteren harcama birimlerinindir; kütüphane
materyali bulunan idareler gerekli görürse ayrıca düzenleyebilir (34/3-a ikinci cümle).

Kişisel veri YOKTUR (defter nüsha kaydıdır; bağışçı — şifreli `source_note` — okunmaz).
"""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any, Final

from django.core.exceptions import ValidationError
from django.db.models import Q, QuerySet
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook

from apps.kutuphane import dokum_ortak as ortak
from apps.kutuphane import sayim_belgeleri, selectors_sayim
from apps.kutuphane.models import (
    TERMINAL_COPY_STATUSES,
    Acquisition,
    AcquisitionMethod,
    Copy,
    CopyStatus,
    StockTake,
    StockTakeItem,
    StockTakeOutcome,
    StockTakeResult,
    StockTakeStatus,
    StockTakeWriteOffPath,
)
from apps.kutuphane.services.stocktake import sayim_adi
from shared.pdf import html_to_pdf

SABLON: Final = "documents/dokum.html"
DEFTER_ADI: Final = "Taşınır Kütüphane Defteri dökümü"
HESAP_ADI: Final = "Yönetim hesabı cetveli hazırlığı"

# ---------------------------------------------------------------------------
# Mevzuat metinleri — docs/mevzuat ile BİREBİR (test sınar)
# ---------------------------------------------------------------------------
TMY_9_1_C: Final = (
    "Bu defter, kütüphanelerdeki yazma ve basma nadir eserler ile kitap ve kitap dışı materyal "
    "için tutulur. Her bir taşınır için ayrı kayıt yapılır."
)
TMY_10_1_A_4: Final = (
    "Dergi ve gazete gibi süreli yayınlar ile arşivlenme niteliği olmayan kütüphane materyalleri."
)
TMY_15_4: Final = (
    "Söz konusu yayınlardan cilt birliği sağlananlar, ciltletildikten sonra Varlık İşlem Fişi "
    "düzenlenerek kayıtlara alınır."
)
TMY_34_3_A: Final = (
    "Sayım kurulu tarafından onaylanan Taşınır Sayım ve Döküm Cetveline dayanılarak ilgisine "
    "göre Harcama Birimi Taşınır Mal Yönetim Hesabı Cetveli, Müze Yönetim Hesabı Cetveli veya "
    "Kütüphane Yönetim Hesabı Cetveli düzenlenir."
)
TMY_34_3_A_2: Final = (
    "Bünyesinde tarihi veya sanat değeri olan taşınırlar ile kütüphane materyalleri bulunan kamu "
    "idareleri, gerekli görülmesi hâlinde söz konusu cetvellerden ilgili olanını ayrıca "
    "düzenleyebilirler."
)

# ---------------------------------------------------------------------------
# Belge metinleri
# ---------------------------------------------------------------------------
DEFTER_ACIKLAMA: Final = (
    f"Kütüphane Defteri: “{TMY_9_1_C}” (Taşınır Mal Yönetmeliği md. 9/1-ç). Döküm her nüshayı "
    "kayıt no sırasıyla bir satırda gösterir; kayıttan düşülmüş ve devredilmiş nüsha çıkış "
    "tarihiyle defterde kalır."
)
SURELI_YAYIN_NOTU: Final = (
    "Ciltletilmemiş süreli yayınlar dökümde yoktur: dergi ve gazete gibi süreli yayınlar için "
    "Varlık İşlem Fişi düzenlenmez (Taşınır Mal Yönetmeliği md. 10/1-a-4); cilt birliği "
    "sağlananlar ciltletildikten sonra kayda alınır (md. 15/4)."
)
AKTARIM_NOTU: Final = (
    "“Mevcut koleksiyon (programa aktarım)” yolundaki nüshaların giriş tarihi programa "
    "aktarıldıkları gündür; taşınır kaydındaki asıl giriş tarihi değildir."
)
FIYAT_NOTU: Final = (
    "Birim fiyat edinimde yazılan değerdir; boş olan nüshanın değeri programda kayıtlı değildir."
)
DEFTER_DIPNOTU: Final = (
    "Bu döküm okulun kütüphane işlerini yürüttüğü yerel araçtaki kayıtlardan hazırlanmıştır; "
    "resmî taşınır kaydı Taşınır Kayıt ve Yönetim Sistemi'ndedir (TKYS)."
)
#: Sözleşmedeki not (E11): cetvel hazırlığının dayanağı ve resmî cetvelin yeri.
HESAP_NOTU: Final = (
    "Sayım kurulunca onaylanan Taşınır Sayım ve Döküm Cetveline dayanır; resmî cetveller "
    "TKYS'dedir."
)
HESAP_KAPSAM_NOTU: Final = (
    "Kütüphane Yönetim Hesabı Cetveli, kütüphane olarak faaliyet gösteren harcama birimlerinin "
    "taşınır mal yönetim hesabındadır (Taşınır Mal Yönetmeliği md. 34/2-c); kütüphane "
    f"materyali bulunan idareler için md. 34/3-a: “{TMY_34_3_A_2}”"
)
TUTAR_NOTU: Final = (
    "Tutar yalnız birim fiyatı programda kayıtlı nüshaların toplamıdır; fiyatı kayıtlı olmayan "
    "nüshalar ayrı sütunda sayılır. Sayım fazlasının değerini program yazmaz (md. 17/1)."
)
NOT_YEAR_END_MESSAGE: Final = (
    "Yönetim hesabı cetveli hazırlığı yalnız yıl sonu sayımından basılır. Sayımın ayrıntısında "
    "“Yıl sonu sayımı” işaretli değil."
)
#: F10 düzeltme turu: 32/9 Taşınır Sayım ve Döküm Cetvelini anlatır; yönetim hesabı cetvelinin
#: dayanağı 34/3-a'dır. Zincir: 34/3-a → 32/9 → 32/7 (programda: sayımın onayı).
NOT_APPROVED_MESSAGE: Final = (
    "Yönetim hesabı cetveli hazırlığı sayım onaylandıktan sonra basılır: yönetim hesabı cetveli "
    "sayım kurulunca onaylanan Taşınır Sayım ve Döküm Cetveline dayanır (Taşınır Mal "
    "Yönetmeliği md. 34/3-a); o cetvel de defter kayıtları sayım sonuçlarıyla uygun hâle "
    "getirildikten sonra düzenlenir (md. 32/7, 32/9). Programda kayıtlar sayımın onayıyla uygun "
    "hâle gelir: noksan kayıttan düşülür, fazla kayda alınır."
)
CANCELLED_MESSAGE: Final = "İptal edilmiş sayımdan yönetim hesabı cetveli hazırlığı basılmaz."


# ---------------------------------------------------------------------------
# Taşınır Kütüphane Defteri dökümü
# ---------------------------------------------------------------------------
def ciltsiz_sureli_yayin_q() -> Q:
    """Ciltletilmemiş süreli yayın nüshası (TMY 10/1-a-4, 15/4) — defterden DIŞLANIR.

    Tek kaynak `selectors_sayim.ciltsiz_sureli_yayin_q`'dur (cetvel hazırlığı da aynısını
    uygular)."""
    return selectors_sayim.ciltsiz_sureli_yayin_q()


def register_copies(*, year: int | None = None) -> QuerySet[Copy]:
    """Defterin nüshaları: kayıt no sırasıyla; silinmiş ve ciltsiz süreli yayın hariç.

    `year` verilirse yalnız o yıl giren (edinim tarihi) nüshalar.
    """
    qs = (
        Copy.objects.filter(work__deleted_at__isnull=True)
        .exclude(ciltsiz_sureli_yayin_q())
        .select_related("work")
        .order_by("accession_no")
    )
    if year is not None:
        qs = qs.filter(acquisition__date__year=year)
    return qs


@dataclass(frozen=True, slots=True)
class _Edinim:
    method: str
    method_display: str
    date: date
    unit_price: Decimal | None


def _edinimler() -> dict[int, _Edinim]:
    """Edinim alanları `values_list` ile (bağışçı — şifreli kaynak notu — çözülmez)."""
    yollar = dict(AcquisitionMethod.choices)
    return {
        int(pk): _Edinim(str(yol), str(yollar.get(yol, yol)), gun, fiyat)
        for pk, yol, gun, fiyat in Acquisition.all_objects.values_list(
            "pk", "method", "date", "unit_price"
        )
    }


@dataclass(frozen=True, slots=True)
class DefterSatiri:
    kayit_no: int
    barkod: str
    giris_tarihi: date
    giris_yolu: str
    kaynak_adi: str
    yazar: str
    yayinevi: str
    yayin_yili: int | None
    tkys_kodu: str
    eski_kayit_no: str
    birim_fiyat: Decimal | None
    durum: str
    cikis_tarihi: date | None
    nadir: bool


def register_rows(*, year: int | None = None) -> Iterator[DefterSatiri]:
    edinimler = _edinimler()
    cikislar = selectors_sayim.exit_dates()
    for n in register_copies(year=year):
        edinim = edinimler[int(n.acquisition_id)]
        cikis = None
        if n.status in TERMINAL_COPY_STATUSES:
            cikis = cikislar.get(int(n.pk)) or timezone.localdate(n.updated_at)
        yield DefterSatiri(
            kayit_no=int(n.accession_no),
            barkod=n.barcode,
            giris_tarihi=edinim.date,
            giris_yolu=edinim.method_display,
            kaynak_adi=n.work.title,
            yazar=n.work.authors,
            yayinevi=n.work.publisher,
            yayin_yili=n.work.publish_year,
            tkys_kodu=n.external_asset_ref,
            eski_kayit_no=n.old_register_no,
            birim_fiyat=edinim.unit_price,
            durum=str(n.get_status_display()),
            cikis_tarihi=cikis,
            nadir=bool(n.is_rare_or_manuscript),
        )


def library_register_context(*, year: int | None = None, on: date | None = None) -> dict[str, Any]:
    satirlar = list(register_rows(year=year))
    cikan = sum(1 for s in satirlar if s.cikis_tarihi is not None)
    fiyatli = [s.birim_fiyat for s in satirlar if s.birim_fiyat is not None]
    h = ortak.hucre
    tablo = ortak.tablo(
        "NÜSHALAR (KAYIT NO SIRASIYLA)",
        [
            ortak.sutun("Kayıt no", 8),
            ortak.sutun("Giriş tarihi", 7),
            ortak.sutun("Giriş yolu", 9),
            ortak.sutun("Kaynak adı", 17),
            ortak.sutun("Yazar", 11),
            ortak.sutun("Yayınevi, yıl", 9),
            ortak.sutun("TKYS kodu", 10),
            ortak.sutun("Eski kayıt no", 8),
            ortak.sutun("Birim fiyat", 7, "sayi"),
            ortak.sutun("Durum", 7),
            ortak.sutun("Çıkış tarihi", 7),
        ],
        [
            [
                h(s.kayit_no, "tek"),
                h(ortak.tarih(s.giris_tarihi)),
                h(s.giris_yolu),
                h(s.kaynak_adi + (" (nadir eser)" if s.nadir else ""), "kalin"),
                h(s.yazar),
                h(", ".join(p for p in (s.yayinevi, str(s.yayin_yili or "")) if p)),
                h(s.tkys_kodu, "tek"),
                h(s.eski_kayit_no, "tek"),
                h(ortak.para(s.birim_fiyat) if s.birim_fiyat is not None else "", "sayi"),
                h(s.durum),
                h(ortak.tarih(s.cikis_tarihi) if s.cikis_tarihi else ""),
            ]
            for s in satirlar
        ],
        "Defterde nüsha yok.",
    )
    info = [
        {"label": "Kapsam", "value": f"{year} yılında girenler" if year else "Bütün kayıt"},
        {"label": "Nüsha (dökümde)", "value": ortak.sayi(len(satirlar))},
        {"label": "Kayıttan düşülmüş ya da devredilmiş", "value": ortak.sayi(cikan)},
        {
            "label": "Birim fiyatı kayıtlı nüshaların toplamı",
            "value": f"{ortak.para(sum(fiyatli, Decimal(0)))} TL ({ortak.sayi(len(fiyatli))} nüsha)",
        },
    ]
    return {
        **ortak.antet(),
        "document_title": "TAŞINIR KÜTÜPHANE DEFTERİ DÖKÜMÜ",
        "subtitle": "Kütüphane materyali — Taşınır Mal Yönetmeliği md. 9/1-ç",
        "issued_on": ortak.tarih(ortak.bugun(on)),
        "info": info,
        "paragraphs": [DEFTER_ACIKLAMA],
        "tables": [tablo],
        "notes": [SURELI_YAYIN_NOTU, AKTARIM_NOTU, FIYAT_NOTU],
        "footnote": DEFTER_DIPNOTU,
        "landscape": True,
    }


def library_register_pdf(*, year: int | None = None, on: date | None = None) -> bytes:
    return html_to_pdf(render_to_string(SABLON, library_register_context(year=year, on=on)))


DEFTER_XLSX_SUTUNLARI: Final[tuple[tuple[str, int], ...]] = (
    ("Kayıt no", 13),
    ("Barkod", 13),
    ("Giriş tarihi", 12),
    ("Giriş yolu", 30),
    ("Kaynak adı", 44),
    ("Yazar(lar)", 28),
    ("Yayınevi", 24),
    ("Yayın yılı", 10),
    ("TKYS kodu", 18),
    ("Eski kayıt no", 14),
    ("Birim fiyat", 12),
    ("Durum", 28),
    ("Çıkış tarihi", 12),
    ("El yazması / nadir eser", 12),
)


def library_register_xlsx(*, year: int | None = None, on: date | None = None) -> bytes:
    """Defter dökümü (XLSX, yazma kipi): kayıt no ve fiyat SAYI, tarihler TARİH hücresidir."""
    gun = ortak.bugun(on)
    wb = Workbook(write_only=True)
    wb.properties.title = DEFTER_ADI
    ws = wb.create_sheet("Taşınır Kütüphane Defteri")
    ws.freeze_panes = "A3"
    kapsam = f"{year} yılında girenler" if year else "bütün kayıt"
    ws.append([ortak.wo_hucre(ws, f"{DEFTER_ADI} — {kapsam} · {ortak.tarih(gun)}", kalin_mi=True)])
    ortak.wo_baslik(ws, DEFTER_XLSX_SUTUNLARI)
    for s in register_rows(year=year):
        ws.append(
            [
                ortak.wo_hucre(ws, s.kayit_no, "0"),
                s.barkod,
                ortak.wo_hucre(ws, s.giris_tarihi, ortak.DATE_FORMAT),
                s.giris_yolu,
                s.kaynak_adi,
                s.yazar,
                s.yayinevi,
                s.yayin_yili,
                s.tkys_kodu,
                s.eski_kayit_no,
                ortak.wo_hucre(ws, s.birim_fiyat, ortak.MONEY_FORMAT),
                s.durum,
                ortak.wo_hucre(ws, s.cikis_tarihi, ortak.DATE_FORMAT),
                "Evet" if s.nadir else "",
            ]
        )
    notlar = wb.create_sheet("Notlar")
    notlar.column_dimensions["A"].width = 120
    for metin in (DEFTER_ACIKLAMA, SURELI_YAYIN_NOTU, AKTARIM_NOTU, FIYAT_NOTU, DEFTER_DIPNOTU):
        notlar.append([metin])
    cikti = BytesIO()
    wb.save(cikti)
    return cikti.getvalue()


def register_years() -> list[int]:
    """Defterde girişi olan yıllar (ekran seçicisi), yeniden eskiye."""
    yillar = {
        int(d.year)
        for d in register_copies().values_list("acquisition__date", flat=True).distinct()
        if d is not None
    }
    return sorted(yillar, reverse=True)


# ---------------------------------------------------------------------------
# Yönetim hesabı cetveli hazırlığı (TMY 34/2-c, 34/3-a)
# ---------------------------------------------------------------------------
def management_account_blocker(stocktake: StockTake) -> str:
    """Basılamıyorsa nedeni (kullanıcı metni); basılabiliyorsa boş metin."""
    if not stocktake.is_year_end:
        return NOT_YEAR_END_MESSAGE
    if stocktake.status == StockTakeStatus.CANCELLED:
        return CANCELLED_MESSAGE
    if stocktake.status != StockTakeStatus.APPROVED:
        return NOT_APPROVED_MESSAGE
    return ""


def year_end_stocktakes() -> list[dict[str, Any]]:
    """Ekran seçicisi: yıl sonu işaretli sayımlar ve basılabilirlikleri (yeniden eskiye).

    İptal edilmiş sayım listelenmez: hiç onaylanamaz, cetvel hazırlığına dayanak olamaz."""
    return [
        {
            "id": s.pk,
            "name": sayim_adi(s),
            "fiscal_year": s.fiscal_year,
            "status": s.status,
            "status_display": str(s.get_status_display()),
            "approved_on": s.approved_on,
            "available": management_account_blocker(s) == "",
            "reason": management_account_blocker(s),
        }
        for s in StockTake.objects.filter(is_year_end=True)
        .exclude(status=StockTakeStatus.CANCELLED)
        .order_by("-fiscal_year", "-pk")
    ]


@dataclass(slots=True)
class _Deger:
    """Bir büyüklüğün tutarı: birim fiyatı kayıtlı nüshaların toplamı + fiyatsız nüsha sayısı."""

    nusha: int = 0
    tutar: Decimal = Decimal("0")
    fiyatsiz: int = 0

    def ekle(self, fiyat: Decimal | None) -> None:
        self.nusha += 1
        if fiyat is None:
            self.fiyatsiz += 1
        else:
            self.tutar += fiyat


def management_account_values(stocktake: StockTake) -> dict[str, _Deger]:
    """34/1 büyüklüklerinin tutarları — `selectors_sayim.tmy_34_1` ile AYNI sınıflama.

    Anahtarlar: `devir`, `aktarim`, `giren:<yol>`, `cikan:<durum>`, `kayda_gore`,
    `gelecek` (sayımda bulunup onayda 27/1 ile düşülmeyen nüshalar + kayda alınan fazla).
    Test nüsha sayılarının `tmy_34_1` ile eşit olduğunu sınar (iki hesap ayrışamaz).
    """
    yil = stocktake.fiscal_year or timezone.localdate().year
    bas, son = date(yil, 1, 1), date(yil, 12, 31)
    cikislar = selectors_sayim.exit_dates()
    edinimler = _edinimler()
    degerler: dict[str, _Deger] = defaultdict(_Deger)
    nushalar = (
        Copy.objects.filter(work__deleted_at__isnull=True)
        .exclude(ciltsiz_sureli_yayin_q())
        .values_list("pk", "status", "acquisition_id", "updated_at")
    )
    for pk, durum, edinim_pk, guncel in nushalar:
        edinim = edinimler[int(edinim_pk)]
        giris, yol, fiyat = edinim.date, edinim.method, edinim.unit_price
        cikis: date | None = None
        if durum in TERMINAL_COPY_STATUSES:
            cikis = cikislar.get(int(pk)) or timezone.localdate(guncel)
        if giris > son:
            continue
        if giris < bas:
            if cikis is None or cikis >= bas:
                degerler["devir"].ekle(fiyat)
        elif yol == AcquisitionMethod.EXISTING_STOCK:
            degerler["aktarim"].ekle(fiyat)
        else:
            degerler[f"giren:{yol}"].ekle(fiyat)
        if cikis is not None and bas <= cikis <= son:
            degerler[f"cikan:{durum}"].ekle(fiyat)
        if cikis is None or cikis > son:
            degerler["kayda_gore"].ekle(fiyat)
    kalemler = StockTakeItem.objects.filter(stocktake=stocktake)
    bulunan = kalemler.filter(
        copy__isnull=False, result__in=(StockTakeResult.FOUND, StockTakeResult.BY_RECORD)
    ).exclude(
        outcome=StockTakeOutcome.WRITTEN_OFF, write_off_path=StockTakeWriteOffPath.DAMAGE_27_1
    )
    for (edinim_pk,) in bulunan.values_list("copy__acquisition_id"):
        degerler["gelecek"].ekle(edinimler[int(edinim_pk)].unit_price)
    for _ in kalemler.filter(copy__isnull=True, surplus_excluded=False):
        degerler["gelecek"].ekle(None)  # fazlanın değerini program yazmaz (17/1)
    return degerler


def _hesap_satirlari(stocktake: StockTake) -> list[dict[str, Any]]:
    """Cetvelin satırları: ad, nüsha (tmy_34_1), tutar, fiyatsız nüsha, girinti/kalınlık."""
    sayilar = selectors_sayim.tmy_34_1(stocktake)
    degerler = management_account_values(stocktake)
    yollar = dict(AcquisitionMethod.choices)
    durumlar = dict(CopyStatus.choices)

    def satir(
        ad: str, n: int | None, anahtar: str, *, girinti: bool = False, kalin: bool = False
    ) -> dict[str, Any]:
        deger = degerler.get(anahtar, _Deger())
        return {
            "ad": ad,
            "nusha": n,
            "tutar": deger.tutar if n else Decimal("0"),
            "fiyatsiz": deger.fiyatsiz if n else 0,
            "girinti": girinti,
            "kalin": kalin,
        }

    # Üst satırların (giren, çıkan) tutarı alt satırların toplamıdır.
    for grup, alt in (
        ("giren", sayilar["entered_by_method"]),
        ("cikan", sayilar["exited_by_status"]),
    ):
        birlesik = _Deger()
        for kod in alt:
            d = degerler.get(f"{grup}:{kod}")
            if d is not None:
                birlesik.nusha += d.nusha
                birlesik.tutar += d.tutar
                birlesik.fiyatsiz += d.fiyatsiz
        degerler[grup] = birlesik
    satirlar = [
        satir(
            "Önceki yıldan devir (yıl başında kayıtta olan)",
            sayilar["previous_year_carryover"],
            "devir",
            kalin=True,
        ),
        satir(
            "Yıl içinde giren (Yönetmelik Md. 10/5 yolları ve sayım fazlası)",
            sayilar["entered"],
            "giren",
            kalin=True,
        ),
    ]
    satirlar += [
        satir(str(yollar[y]), n, f"giren:{y}", girinti=True)
        for y, n in sayilar["entered_by_method"].items()
    ]
    satirlar.append(
        satir(
            "Programa aktarım (mevcut koleksiyonun kaydı; taşınır girişi değildir)",
            sayilar["program_transfer"],
            "aktarim",
        )
    )
    satirlar.append(
        satir("Yıl içinde çıkan (kayıttan düşme ve devir)", sayilar["exited"], "cikan", kalin=True)
    )
    satirlar += [
        satir(str(durumlar[d]), n, f"cikan:{d}", girinti=True)
        for d, n in sayilar["exited_by_status"].items()
    ]
    satirlar.append(
        satir(
            "Kayda göre yıl sonu (devir + aktarım + giren − çıkan)",
            sayilar["year_end_by_record"],
            "kayda_gore",
        )
    )
    satirlar.append(
        satir(
            sayim_belgeleri.GELECEK_YIL_SATIRI,
            sayilar["next_year_carryover"],
            "gelecek",
            kalin=True,
        )
    )
    satirlar.append(satir("Sayım fazlası", sayilar["surplus"], "_"))
    satirlar.append(satir("Sayım noksanı", sayilar["missing"], "_"))
    for s in satirlar[-2:]:
        s["tutar"], s["fiyatsiz"] = None, None
    return satirlar


def _hesap_dogrula(stocktake: StockTake) -> None:
    engel = management_account_blocker(stocktake)
    if engel:
        raise ValidationError(engel)


def management_account_context(stocktake: StockTake, *, on: date | None = None) -> dict[str, Any]:
    _hesap_dogrula(stocktake)
    satirlar = _hesap_satirlari(stocktake)
    h = ortak.hucre
    tablo = ortak.tablo(
        "TAŞINIR MAL YÖNETİM HESABININ BÜYÜKLÜKLERİ (MD. 34/1)",
        [
            ortak.sutun("Büyüklük", 52),
            ortak.sutun("Nüsha", 12, "sayi"),
            ortak.sutun("Tutar (TL)", 18, "sayi"),
            ortak.sutun("Birim fiyatı kayıtlı olmayan", 18, "sayi"),
        ],
        [
            [
                h(("— " if s["girinti"] else "") + s["ad"], "kalin" if s["kalin"] else ""),
                h(ortak.sayi(s["nusha"]), "sayi"),
                h(ortak.para(s["tutar"]) if s["tutar"] is not None else ortak.BOS, "sayi"),
                h(ortak.sayi(s["fiyatsiz"]) if s["fiyatsiz"] is not None else ortak.BOS, "sayi"),
            ]
            for s in satirlar
        ],
    )
    yil = stocktake.fiscal_year
    return {
        **ortak.antet(),
        "document_title": "KÜTÜPHANE YÖNETİM HESABI CETVELİ HAZIRLIĞI",
        "subtitle": f"{yil} mali yılı (1 Ocak - 31 Aralık) · Taşınır Mal Yönetmeliği md. 34/2-c, 34/3-a",
        "status_note": HESAP_NOTU,
        "issued_on": ortak.tarih(ortak.bugun(on)),
        "info": [
            {"label": "Mali yıl", "value": str(yil or ortak.BOS)},
            {"label": "Dayandığı yıl sonu sayımı", "value": sayim_adi(stocktake)},
            {"label": "Harcama yetkilisinin onayı", "value": ortak.tarih(stocktake.approved_on)},
        ],
        "paragraphs": [
            f"“{sayim_belgeleri.TMY34_1}” (Taşınır Mal Yönetmeliği md. 34/1).",
            f"“{TMY_34_3_A}” (md. 34/3-a).",
        ],
        "tables": [tablo],
        "notes": [HESAP_KAPSAM_NOTU, TUTAR_NOTU, sayim_belgeleri.GELECEK_YIL_NOTU],
        "footnote": selectors_sayim.TMY_34_1_NOTE,
        "landscape": False,
    }


def management_account_pdf(stocktake: StockTake, *, on: date | None = None) -> bytes:
    return html_to_pdf(render_to_string(SABLON, management_account_context(stocktake, on=on)))


def management_account_xlsx(stocktake: StockTake, *, on: date | None = None) -> bytes:
    """Cetvel hazırlığı (XLSX): nüsha ve tutar SAYI hücresidir."""
    _hesap_dogrula(stocktake)
    gun = ortak.bugun(on)
    wb = Workbook()
    wb.properties.title = HESAP_ADI
    ws = wb.active
    assert ws is not None  # yeni çalışma kitabında etkin sayfa daima vardır
    ws.title = "Yönetim hesabı hazırlığı"
    ws.append([f"Kütüphane yönetim hesabı cetveli hazırlığı — {stocktake.fiscal_year} mali yılı"])
    ortak.kalin(ws, 1)
    ws.append(
        [f"Dayandığı yıl sonu sayımı: {sayim_adi(stocktake)} · Düzenlenme: {ortak.tarih(gun)}"]
    )
    ws.append([HESAP_NOTU])
    ws.append([])
    ortak.baslik_satiri(
        ws,
        (("Büyüklük", 70), ("Nüsha", 12), ("Tutar (TL)", 16), ("Birim fiyatı kayıtlı olmayan", 16)),
    )
    for s in _hesap_satirlari(stocktake):
        ws.append([("— " if s["girinti"] else "") + s["ad"], s["nusha"], s["tutar"], s["fiyatsiz"]])
        ortak.bicimle(
            ws,
            int(ws.max_row),
            {2: ortak.NUMBER_FORMAT, 3: ortak.MONEY_FORMAT, 4: ortak.NUMBER_FORMAT},
        )
    ws.append([])
    for metin in (
        HESAP_KAPSAM_NOTU,
        TUTAR_NOTU,
        sayim_belgeleri.GELECEK_YIL_NOTU,
        selectors_sayim.TMY_34_1_NOTE,
    ):
        ws.append([metin])
    cikti = BytesIO()
    wb.save(cikti)
    return cikti.getvalue()


def library_register_filename(bicim: str = ortak.PDF, gun: date | None = None) -> str:
    return ortak.dosya_adi(DEFTER_ADI, gun, bicim=bicim)


def management_account_filename(bicim: str = ortak.PDF, gun: date | None = None) -> str:
    return ortak.dosya_adi(HESAP_ADI, gun, bicim=bicim)
