"""Sayım tutanağı (E10) — PDF ve XLSX (F9; tasarım §10 E10, §9-10, §9-11, F8 ekleri 13, 34).

Taşınır Mal Yönetmeliği md. 32: sayım, harcama yetkilisinin ya da görevlendirdiği kişinin
başkanlığında, taşınır kayıt yetkilisinin de katıldığı en az üç kişilik sayım kurulunca
yapılır (32/2); fark çıkan taşınırların sayımı bir kez daha tekrarlanır (32/6); noksan için
Kayıttan Düşme Teklif ve Onay Tutanağı ve Varlık İşlem Fişi, fazla için Varlık İşlem Fişi
düzenlettirilir (32/7). Her atıf `docs/mevzuat/` metninden doğrulanır; alıntılar ve sade
anlatımların dayandığı ifadeler testle depodaki metinde aranır
(`tests/test_sayim_belgeleri.py`).

**Tutanağın içeriği** (kod kapısı §14.1 F9):

- künye: mali yıl, başlangıç ve tamamlanma, ikinci sayım (32/6), harcama yetkilisinin onayı;
- SONUÇLAR: kayıtlara göre miktar (anlık görüntü), bulunan, kayda göre alınan, fazla,
  sayımda bulunan miktar (10/1-ğ), noksan;
- **iki seçenek AYRI satırlarda** (TMY 32/3 durdurması ve sayım için hizmet arası) ve
  iadenin her zaman açık olduğu üçüncü satır (`selectors_sayim.options_lines`);
- ödünçteki, teslimdeki ve onarımdaki nüsha için **kurulun seçimi ve dayanağı** (§9-11,
  AT-1: "32/5'e kıyasen; 23/4"; onarımda "sayım kurulunun kararıdır …, doğrudan hüküm
  yoktur" — F9 ekleri K2; `selectors_sayim.category_lines`);
- bölümlere göre sayım;
- **noksan düşüm teklifi (32/7)**, hasar önerisinin düşümü (27/1 + 10/1-e — F8 ekleri 34),
  kayıpta görünüp bulunanlar (uzlaştırma) ve sayım fazlası (TMY 17); her kalemde
  "Kayıp/hasar tutanağı" sütunu dosyanın tespit tarihini yazar — tutanağa bağlı liste;
- onaydan sonra harcama yetkilisinin onaylamadığı ve onayda durumu değişen kalemler;
- kapanış: notlar, SAYIM KURULU imzaları, harcama yetkilisinin OLUR'u, TKYS dipnotu;
- YENİ SAYFADA ek **"Taşınır Sayım ve Döküm Cetveline aktarılacak sayılar"** (TMY 34/1'in
  dört büyüklüğü, programa aktarım ayrı satırda, ödünç, teslim ve onarımın sayılışı ayrı
  satırda; gelecek yıla devirden onayda 27/1 ile düşülenler çıkarılır ve altında ayrı
  satırda gösterilir — F9 ekleri madde 25 (a); "Bu döküm Taşınır Sayım ve Döküm Cetveli
  değildir …" ibaresi — A8 kararı, F8 ekleri 13).

**ÖDÜNÇ ALANIN KİMLİĞİ BASILMAZ** (E10): sayım fazlası ve noksanına ilişkin sayfalar Varlık
İşlem Fişine eklenir ve muhasebe birimine gider (10/1-g, 32/8). Bu modül ödünç, teslim,
üyelik ve kişi tablosunu OKUMAZ; kalemler kişisizdir (`StockTakeItem`). Kişi adı yalnız
sayım kurulunun imzalarında ve harcama yetkililerinde durur — şifreli alanlardan yalnız
belgenin kendisine çözülür; uç yalnız yönetici kipindedir.

Program TKYS'nin yerine GEÇMEZ: taşınır kodu düzeyindeki Sayım Tutanağı, Kayıttan Düşme
Teklif ve Onay Tutanağı, Varlık İşlem Fişi ve Taşınır Sayım ve Döküm Cetveli Taşınır Kayıt
ve Yönetim Sistemi'nde düzenlenir; bu tutanak kütüphane materyalinin nüsha düzeyindeki
dökümüdür ve dipnot bunu söyler.

Basılabilirlik: taslakta basılmaz (anlık görüntü yok); iptal edilmiş sayımda basılmaz.
"Sürüyor"da ara döküm olarak ("TASLAK" ibaresi), "Tamamlandı"da imzaya (OLUR boş),
"Onaylandı"da onayla birlikte basılır. PDF'ler yalnız `shared.pdf.html_to_pdf` kapısından
üretilir (CLAUDE.md §3). Bu modül veritabanına YAZMAZ; hata iletilerinde kişi adı yoktur.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any, Final

from django.core.exceptions import ValidationError
from django.db.models import Count, Q
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.kutuphane import barcode as barcode_module
from apps.kutuphane import selectors_sayim
from apps.kutuphane.models import (
    AcquisitionMethod,
    CopyStatus,
    StockTake,
    StockTakeItem,
    StockTakeOutcome,
    StockTakeResult,
    StockTakeStatus,
)
from apps.okul.models import SchoolConfig
from shared.letterhead import letterhead_context
from shared.pdf import html_to_pdf
from shared.text import tr_upper

SABLON: Final = "documents/sayim_tutanagi.html"

# ---------------------------------------------------------------------------
# Belge adı ve adres adı (sözlük §2: E10 — "Sayım tutanağı")
# ---------------------------------------------------------------------------
TUTANAK_ADI: Final = "Sayım tutanağı"
#: Ekin adı (A8 kararı, F8 ekleri 13) — ekran ve kılavuz bu adı kullanır.
EK_ADI: Final = "Taşınır Sayım ve Döküm Cetveline aktarılacak sayılar"
SAYIM_TUTANAGI: Final = "sayim-tutanagi"

PDF: Final = "pdf"
XLSX: Final = "xlsx"


@dataclass(frozen=True)
class SayimBelgesi:
    """Bir sayım belgesinin tanımı: adres adı, belge adı, biçimleri."""

    slug: str
    ad: str
    bicimler: tuple[str, ...] = (PDF,)


#: Sayım belgeleri (ekran düğmeleri buradan kurulur).
SAYIM_BELGELERI: Final[tuple[SayimBelgesi, ...]] = (
    SayimBelgesi(SAYIM_TUTANAGI, TUTANAK_ADI, (PDF, XLSX)),
)
_BELGE_TANIMLARI: Final[dict[str, SayimBelgesi]] = {b.slug: b for b in SAYIM_BELGELERI}

# ---------------------------------------------------------------------------
# İletiler (kişisel veri YOK)
# ---------------------------------------------------------------------------
UNKNOWN_DOCUMENT_MESSAGE: Final = "Böyle bir sayım belgesi yok."
UNKNOWN_FORMAT_MESSAGE: Final = "Bu belge bu biçimde üretilmez."
DRAFT_MESSAGE: Final = "Sayım başladıktan sonra basılır."
CANCELLED_MESSAGE: Final = "İptal edilmiş sayımın tutanağı basılmaz."

#: Durum notları (başlığın altında, çerçeveli).
SURUYOR_NOTU: Final = (
    "TASLAK — Sayım sürüyor; sayılar basım anındaki kayıtlardandır ve kesin değildir."
)
ONAY_BEKLIYOR_NOTU: Final = "Sayım tamamlandı; harcama yetkilisinin onayı bekleniyor."
#: F9 düzeltme turu: kararı bekleyen sayım fazlası varken sayılar kesin değildir (kurul
#: "Kayda alınmayacak" derse sayımda bulunan miktar ve gelecek yıla devir değişir).
FAZLA_KARARI_NOTU: Final = (
    "{sayi} sayım fazlası kitabın kararı bekleniyor; “Sayımda bulunan miktar” ve ekteki "
    "“Gelecek yıla devir” bu kararlar verilince kesinleşir. Tutanağı kararlardan sonra imzaya "
    "basın."
)

# ---------------------------------------------------------------------------
# Mevzuat metinleri — docs/mevzuat ile BİREBİR (testler sınar)
# ---------------------------------------------------------------------------
#: TMY md. 17/1'in ilk cümlesi (sayım fazlası).
TMY17_1: Final = (
    "Yapılan sayım sonucunda fazla bulunan taşınırlar, Varlık İşlem Fişi düzenlenerek "
    "kayıtlara alınır."
)
#: TMY md. 34/1'in ikinci cümlesi (taşınır mal yönetim hesabının büyüklükleri).
TMY34_1: Final = (
    "Taşınır mal yönetim hesabında; önceki yıldan devredilen, yılı içinde giren, çıkan ve "
    "ertesi yıla devredilen taşınırlar ile yıl sonu sayımında bulunan fazla ve noksanlar "
    "gösterilir."
)

#: Açılış paragrafı: kurul (32/2) ve anlık görüntü.
ACILIS: Final = (
    "Okul kütüphanesindeki kütüphane materyalinin sayımı, harcama yetkilisinin ya da "
    "görevlendirdiği kişinin başkanlığında, taşınır kayıt yetkilisinin de katıldığı en az üç "
    "kişilik sayım kurulunca yapılmış ve sonuçları aşağıda gösterilmiştir (Taşınır Mal "
    "Yönetmeliği md. 32/2). Kayıtlara göre miktar, sayımın başladığı andaki kayıtlardır; "
    "sayım sırasında iade edilen ya da kütüphaneye dönen kitap bulunmuş sayılmıştır."
)
#: 32/6: ikinci sayım yapıldıysa. Program yalnız BULUNAMAYANLARI yeniden aratır (bilinçli
#: sınır a); not bunu söyler, "farklı çıkan bütün nüshalar" demez (F9 düzeltme turu).
IKINCI_SAYIM_NOTU: Final = (
    "İlk sayımda bulunamayan nüshaların sayımı bir kez daha tekrarlanmıştır; yine "
    "bulunamayanlar “Noksan” olarak yazılmıştır (Taşınır Mal Yönetmeliği md. 32/6). Sayım "
    "fazlası kitaplar ikinci kez sayılmamıştır: okutulan ya da elle yazılan kitap sayım "
    "sırasında kurulun elindedir."
)
#: 32/7 ve 10/1-e (noksanın kayıttan düşülmesi; onay harcama yetkilisinin takdiridir).
DUSME_NOTU: Final = (
    "Noksan için Kayıttan Düşme Teklif ve Onay Tutanağı ve Varlık İşlem Fişi, fazla için Varlık "
    "İşlem Fişi düzenlettirilerek kayıtların sayım sonuçlarıyla uygunluğu sağlanır (Taşınır Mal "
    "Yönetmeliği md. 32/7). Durumu belgeleyen tutanak bulunduğunda Kayıttan Düşme Teklif ve "
    "Onay Tutanağı komisyon kurulmadan harcama yetkilisince onaylanabilir (md. 10/1-e); bu "
    "tutanağın ve kayıp/hasar tutanağının o belge sayılıp sayılmayacağını harcama yetkilisi "
    "değerlendirir. Kasıt, kusur, ihmal veya tedbirsizlik olup olmadığını da harcama yetkilisi "
    "değerlendirir (md. 27/3)."
)
#: Hasar önerisinin yolu (F8 ekleri 34: 27/1 + 10/1-e, kayıp/hasar tutanağıyla, komisyonsuz).
HASAR_NOTU: Final = (
    "Hasar dosyasında kayıttan düşme önerilen kaynaklar, kullanılamaz hâle gelen taşınır olarak "
    "kayıtlardan çıkarılır (Taşınır Mal Yönetmeliği md. 27/1); durumu kayıp/hasar tutanağı "
    "belgeler."
)
#: Kayıp/hasar tutanağına bağlı liste.
TUTANAK_BAGI_NOTU: Final = (
    "“Kayıp/hasar tutanağı” sütununda tespit tarihi yazılı kalemlerde o tutanak, durumu "
    "belgeleyen belge olarak bu tutanağa eklenir."
)
#: TMY 17/1: fazlanın değeri (sade anlatım; ilk cümle alıntıdır).
FAZLA_NOTU: Final = (
    f"“{TMY17_1}” (Taşınır Mal Yönetmeliği md. 17/1). Kayda esas değer, aynı nitelikte son bir "
    "yıl içinde girişi yapılan taşınır varsa onun değeri, yoksa değer tespit komisyonunca "
    "belirlenecek değerdir; program değer yazmaz."
)
#: E10 kişisel veri notu (10/1-g, 32/8).
KIMLIK_NOTU: Final = (
    "Ödünçteki ve teslimdeki nüshalar kişi adı olmadan, sayıyla gösterilir: tutanağın sayım "
    "fazlası ve noksanına ilişkin sayfaları Varlık İşlem Fişine eklenir ve muhasebe birimine "
    "gönderilir (Taşınır Mal Yönetmeliği md. 10/1-g, 32/8)."
)
TKYS_DIPNOTU: Final = (
    "Bu tutanak kütüphane materyalinin nüsha düzeyindeki dökümüdür. Taşınır kodu düzeyindeki "
    "Sayım Tutanağı, Kayıttan Düşme Teklif ve Onay Tutanağı ve Varlık İşlem Fişi Taşınır Kayıt "
    "ve Yönetim Sistemi'nde (TKYS) düzenlenir."
)
#: Ekin 10/1-ğ notu. Madde 25 (a): sayımda bulunup onayda hasar nedeniyle (27/1) düşülen
#: nüsha sayımın kendi çıkışıdır; gelecek yıla devirden çıkarılır ve ayrı satırda durur.
GELECEK_YIL_NOTU: Final = (
    "Cetvelin “Gelecek Yıla Devir” sütunundaki miktar, yıl sonu Sayım Tutanağının “Sayımda "
    "Bulunan Miktar”ına eşit olmalıdır (Taşınır Mal Yönetmeliği md. 10/1-ğ). Sayımda bulunup "
    "onayda hasar nedeniyle kayıttan düşülen nüsha (md. 27/1) sayımın kendi çıkışıdır: gelecek "
    "yıla devirden çıkarılır ve altında ayrı satırda gösterilir; tutanağın “Sayımda bulunan "
    "miktar”ı değişmez. Kayda göre yıl sonuyla kalan fark onaylanmayan noksandan ya da sayım "
    "sürerken yapılan giriş ve çıkıştan doğar."
)

#: Ekin ara sayım notu (F9 düzeltme turu): cetvel YIL SONU hesabı içindir (10/1-ğ, 32/9);
#: program yıl sonu sayımını harcama yetkilisinin gerekli gördüğü sayımdan ayırmaz (32/1),
#: bu yüzden ek her sayımda bu cümleyi taşır. "Yıl sonu sayımı" işareti F10'dadır (F9
#: ekleri K6 — 25.09.2026 ana oturum kararı, F10 sözleşmesine devredildi).
ARA_SAYIM_NOTU: Final = (
    "Taşınır Sayım ve Döküm Cetveli taşınır kayıt yetkilisinin yıl sonu hesabı için düzenlenir "
    "(Taşınır Mal Yönetmeliği md. 10/1-ğ, 32/9). Bu sayım yıl sonu sayımı değilse (md. 32/1: "
    "harcama yetkilisinin gerekli gördüğü sayım), aşağıdaki sayılar yılın o güne kadarki "
    "durumunu gösterir ve cetvele aktarılmaz."
)

BOS_TARIH: Final = "…/…/……"
BELIRSIZ: Final = "Sayım tamamlanınca yazılır"
#: Tamamlanmış sayımda kararı bekleyen fazla varken (F9 düzeltme turu).
BELIRSIZ_FAZLA: Final = "Sayım fazlası kararları verilince yazılır"
#: Madde 25 (a): hasar önerisi onaylanmadan gelecek yıla devir kesin değildir.
BELIRSIZ_ONAY: Final = "Harcama yetkilisinin onayından sonra yazılır"
#: Ekin gelecek yıla devir satırı ve altındaki iki bileşen (madde 25 a).
GELECEK_YIL_SATIRI: Final = (
    "Gelecek yıla devir (sayımda bulunan − onayda hasar nedeniyle kayıttan düşülen)"
)
SAYIMDA_BULUNAN_SATIRI: Final = "Sayımda bulunan miktar (sayım tutanağı)"
HASAR_DUSULEN_SATIRI: Final = "Onayda hasar nedeniyle kayıttan düşülen (TMY md. 27/1)"
#: Sürerken noksan tablosunun boş metni: sonuç tamamlanınca yazılır (F9 düzeltme turu).
NOKSAN_SURERKEN: Final = "Noksan sayım tamamlanınca belirlenir."

# ---------------------------------------------------------------------------
# Belgedeki kısa adlar
# ---------------------------------------------------------------------------
SECENEK_DURUMU: Final[dict[bool, str]] = {True: "Seçildi", False: "Seçilmedi"}
IADE_DURUMU: Final = "Her zaman açık"
KURUL_ROLLERI: Final = ("Sayım kurulu başkanı", "Taşınır kayıt yetkilisi", "Üye")


# ---------------------------------------------------------------------------
# Ortak yardımcılar
# ---------------------------------------------------------------------------
def _gun(on: date | None = None) -> date:
    return on or timezone.localdate()


def _tarih(gun: date | None) -> str:
    return f"{gun:%d.%m.%Y}" if gun is not None else "—"


def _zaman(an: datetime | None) -> str:
    return f"{timezone.localtime(an):%d.%m.%Y %H:%M}" if an is not None else "—"


def _sayi(n: int | None) -> str:
    """Türkçe binlik ayraçlı sayı (1.234)."""
    if n is None:
        return "—"
    return f"{n:,}".replace(",", ".")


def belge_dosya_adi(ad: str, gun: date | None = None, *, bicim: str = PDF) -> str:
    """'Sayım-tutanağı_25.09.2026.pdf' — belge adı + yerel tarih (sözlük §3)."""
    return f"{ad.replace(' ', '-')}_{_gun(gun):%d.%m.%Y}.{bicim}"


def _tek_satir(metin: str) -> str:
    return " ".join((metin or "").split())


def _okul_adi(config: SchoolConfig) -> str:
    return " ".join((config.school_name or config.kisa_ad or "").split())


def _h(deger: Any, cls: str = "") -> dict[str, str]:
    """Tablo hücresi."""
    return {"v": "" if deger is None else str(deger), "c": cls}


def _sutun(baslik: str, genislik: int, cls: str = "") -> dict[str, Any]:
    return {"header": baslik, "width": genislik, "cls": cls}


def _tablo(
    baslik: str,
    sutunlar: list[dict[str, Any]],
    satirlar: list[list[dict[str, str]]],
    bos: str = "",
    *,
    not_: str = "",
) -> dict[str, Any]:
    assert sum(int(s["width"]) for s in sutunlar) == 100, "sütun genişlikleri %100 olmalı"
    return {"title": baslik, "columns": sutunlar, "rows": satirlar, "empty": bos, "note": not_}


def _imza_izgarasi(
    kisiler: Sequence[dict[str, str]], *, sutun: int = 3
) -> list[list[dict[str, str]]]:
    return [list(kisiler[i : i + sutun]) for i in range(0, len(kisiler), sutun)]


# ---------------------------------------------------------------------------
# Basılabilirlik
# ---------------------------------------------------------------------------
def stocktake_document_blocker(stocktake: StockTake, slug: str = SAYIM_TUTANAGI) -> str:
    """Belge basılamıyorsa nedeni (kullanıcı metni), basılabiliyorsa boş metin."""
    if slug not in _BELGE_TANIMLARI:
        return UNKNOWN_DOCUMENT_MESSAGE
    if stocktake.status == StockTakeStatus.DRAFT:
        return DRAFT_MESSAGE
    if stocktake.status == StockTakeStatus.CANCELLED:
        return CANCELLED_MESSAGE
    return ""


def stocktake_documents(stocktake: StockTake) -> list[dict[str, Any]]:
    """Sayımın belgeleri ve basılabilirlikleri (ekran düğmeleri buradan kurulur)."""
    sonuc = []
    for belge in SAYIM_BELGELERI:
        engel = stocktake_document_blocker(stocktake, belge.slug)
        sonuc.append(
            {
                "kind": belge.slug,
                "title": belge.ad,
                "formats": list(belge.bicimler),
                "available": engel == "",
                "reason": engel,
            }
        )
    return sonuc


def ensure_stocktake_document(stocktake: StockTake, slug: str, bicim: str = PDF) -> SayimBelgesi:
    """Belge tanımını döndürür; basılamıyorsa ya da biçim yoksa `ValidationError`."""
    belge = _BELGE_TANIMLARI.get(slug)
    if belge is None:
        raise ValidationError(UNKNOWN_DOCUMENT_MESSAGE)
    if bicim not in belge.bicimler:
        raise ValidationError({"kind": UNKNOWN_FORMAT_MESSAGE})
    engel = stocktake_document_blocker(stocktake, slug)
    if engel:
        raise ValidationError(engel)
    return belge


# ---------------------------------------------------------------------------
# Veri — kişisiz (kalemler, sonuç sayıları, bölümler)
# ---------------------------------------------------------------------------
def _kalemler(stocktake: StockTake) -> list[StockTakeItem]:
    """Sayımın bütün kalemleri — eser adının Türkçe sırasıyla, fazlalar sonda."""
    return list(selectors_sayim.items(stocktake))


def _barkod(kalem: StockTakeItem) -> str:
    if kalem.copy is not None:
        return barcode_module.format_barcode(kalem.copy.barcode)
    if kalem.surplus_barcode:
        return barcode_module.format_barcode(kalem.surplus_barcode)
    return "—"


def _eser_adi(kalem: StockTakeItem) -> str:
    if kalem.copy is not None:
        return kalem.copy.work.title
    if kalem.surplus_work is not None:
        return kalem.surplus_work.title
    if kalem.surplus_copy is not None:
        return kalem.surplus_copy.work.title
    return "—"


def _yazar(kalem: StockTakeItem) -> str:
    if kalem.copy is not None:
        return kalem.copy.work.authors
    if kalem.surplus_work is not None:
        return kalem.surplus_work.authors
    return ""


def _tkys(kalem: StockTakeItem) -> str:
    return (kalem.copy.external_asset_ref if kalem.copy is not None else "") or "—"


def _dosya(kalem: StockTakeItem) -> str:
    """Kayıp/hasar tutanağına bağ: dosyanın türü ve tespit tarihi (kişisiz)."""
    dosya = kalem.case
    if dosya is None or dosya.deleted_at is not None:
        return "—"
    return f"{dosya.get_case_type_display()} · tespit {_tarih(dosya.reported_on)}"


def _beklenen(kalem: StockTakeItem) -> str:
    etiket = str(kalem.get_expected_status_display()) if kalem.expected_status else "—"
    if kalem.class_section is not None:
        return f"{etiket} ({kalem.class_section.class_label})"
    return etiket


def _fazla_sonucu(kalem: StockTakeItem) -> str:
    if kalem.outcome == StockTakeOutcome.ENTERED and kalem.created_copy is not None:
        return f"Kayda alındı: {barcode_module.format_barcode(kalem.created_copy.barcode)}"
    if kalem.outcome == StockTakeOutcome.NOT_ENTERED and kalem.outcome_note:
        # Etiketi sayım sırasında bağlanan fazla (F9 düzeltme turu): neden yazılır.
        return f"Kayda alınmadı — {_tek_satir(kalem.outcome_note)}"
    if not kalem.outcome:
        bagli = selectors_sayim.bound_surplus_copy(kalem)
        if bagli is not None:
            return (
                f"Sayım sırasında kayda girdi: {barcode_module.format_barcode(bagli.barcode)} "
                "(onayda yeniden kayda alınmaz)"
            )
    if kalem.outcome == StockTakeOutcome.NOT_ENTERED or kalem.surplus_excluded:
        return "Kayda alınmayacak" if not kalem.outcome else "Kayda alınmadı"
    if kalem.surplus_work is not None:
        return f"Kayda alınacak: {kalem.surplus_work.title}"
    return "Karar bekliyor"


def _bolumler(stocktake: StockTake) -> list[dict[str, Any]]:
    """Bölümlere göre sayım (kişisiz, tek sorgu); bölümsüz nüshalar sonda."""
    qs = StockTakeItem.objects.filter(stocktake=stocktake, copy__isnull=False)
    satirlar = [
        {
            "name": satir["section__name"] or "Bölümsüz",
            "expected": satir["kayit"],
            "found": satir["bulunan"],
            "by_record": satir["kayda"],
            "missing": satir["noksan"],
            "_sira": (
                satir["section_id"] is None,
                satir["section__sort_order"] or 0,
                satir["section__name_sort_key"] or "",
            ),
        }
        for satir in qs.values(
            "section_id", "section__name", "section__sort_order", "section__name_sort_key"
        ).annotate(
            kayit=Count("pk"),
            bulunan=Count("pk", filter=Q(result=StockTakeResult.FOUND)),
            kayda=Count("pk", filter=Q(result=StockTakeResult.BY_RECORD)),
            noksan=Count("pk", filter=Q(result=StockTakeResult.MISSING)),
        )
    ]
    satirlar.sort(key=lambda s: s.pop("_sira"))
    return satirlar


@dataclass(frozen=True)
class _Gruplar:
    """Tutanağın kalem tabloları (kişisiz)."""

    noksan: list[StockTakeItem]
    hasar: list[StockTakeItem]
    uzlastirma: list[StockTakeItem]
    fazla: list[StockTakeItem]
    onaylanmayan: list[StockTakeItem]
    durumu_degisen: list[StockTakeItem]


def _gruplar(stocktake: StockTake, kalemler: Iterable[StockTakeItem]) -> _Gruplar:
    onayli = stocktake.status == StockTakeStatus.APPROVED
    liste = list(kalemler)

    def _teklif(k: StockTakeItem) -> bool:
        # Onaydan sonra teklif tabloları yalnız düşülenleri taşır; öbürleri ayrı tablodadır.
        return not onayli or k.outcome == StockTakeOutcome.WRITTEN_OFF

    return _Gruplar(
        noksan=[k for k in liste if k.result == StockTakeResult.MISSING and _teklif(k)],
        hasar=[k for k in liste if k.damage_write_off and _teklif(k)],
        uzlastirma=[
            k
            for k in liste
            if k.copy_id is not None
            and k.result == StockTakeResult.FOUND
            and (
                k.outcome == StockTakeOutcome.RECONCILED
                or (not onayli and k.status_at_completion == CopyStatus.LOST)
            )
        ],
        fazla=[k for k in liste if k.is_surplus],
        onaylanmayan=[k for k in liste if k.outcome == StockTakeOutcome.NOT_APPROVED],
        durumu_degisen=[k for k in liste if k.outcome == StockTakeOutcome.STATE_CHANGED],
    )


def sonuc_satirlari(stocktake: StockTake) -> list[tuple[str, int, bool]]:
    """SONUÇLAR tablosunun satırları: (kalem, nüsha, kalın mı). PDF ve XLSX aynı satırlar.

    Sayım Tutanağının sütunlarına karşılık gelir (TMY 32/4-32/6): kayıtlara göre miktar,
    sayımda bulunan miktar, fazla ve noksan; ödünçte ve teslimde kayda göre alınan ayrı
    satırdır (32/5'e kıyasen — "Kayıtlara Göre Kişilere Verilen Miktar"). Onarımda kayda
    göre alınan da AYRI satırdır: 32/5 ona atfedilmez, sayım kurulunun kararıdır (K2).
    """
    ozet = selectors_sayim.summary(stocktake)
    sonuclar = ozet["results"]
    bulunan = selectors_sayim.found_quantity(stocktake)
    onarimda = StockTakeItem.objects.filter(
        stocktake=stocktake,
        result=StockTakeResult.BY_RECORD,
        expected_status=CopyStatus.IN_REPAIR,
    ).count()
    satirlar: list[tuple[str, int, bool]] = [
        ("Kayıtlara göre miktar (sayımın başladığı andaki kayıt)", ozet["snapshot"], False),
        ("Kütüphanede ve yerinde sayılarak bulunan", sonuclar[StockTakeResult.FOUND], False),
        (
            "Kayda göre alınan (ödünçte ve teslimde — TMY md. 32/5'e kıyasen)",
            sonuclar[StockTakeResult.BY_RECORD] - onarimda,
            False,
        ),
    ]
    if onarimda:
        satirlar.append(("Kayda göre alınan (onarımda — sayım kurulunun kararı)", onarimda, False))
    satirlar += [
        ("Fazla (TMY md. 17)", bulunan["surplus"], False),
        ("Sayımda bulunan miktar (bulunan + kayda göre alınan + fazla)", bulunan["found"], True),
        ("Noksan (TMY md. 32/7)", sonuclar[StockTakeResult.MISSING], True),
    ]
    if sonuclar[StockTakeResult.EXITED]:
        satirlar.append(("Sayım sırasında kayıttan çıkan", sonuclar[StockTakeResult.EXITED], False))
    if sonuclar[StockTakeResult.PENDING]:
        satirlar.append(("Henüz sayılmayan", sonuclar[StockTakeResult.PENDING], False))
    if ozet["surplus_unresolved"] and stocktake.status != StockTakeStatus.APPROVED:
        satirlar.append(
            (
                "Kararı bekleyen sayım fazlası (fazlaya ve sayımda bulunan miktara katıldı; "
                "kesin değildir)",
                ozet["surplus_unresolved"],
                False,
            )
        )
    if ozet["surplus_excluded"]:
        satirlar.append(
            ("Okutulup kayda alınmayacak kitap (fazla sayılmaz)", ozet["surplus_excluded"], False)
        )
    if ozet["damage_write_off"]:
        satirlar.append(
            (
                "Hasar nedeniyle kayıttan düşme teklifi (TMY md. 27/1)",
                ozet["damage_write_off"],
                False,
            )
        )
    return satirlar


def _durumsuz(metin: str) -> str:
    """Seçenek metninin başındaki "Seçildi (…): " / "Seçilmedi: " öbeği atılır (Durum sütunu var).

    "Seçildi (okul kararı — 2026/15): sayım süresince …" → "Okul kararı — 2026/15: sayım
    süresince …"; "Seçilmedi: ödünç verme sürdü." → "Ödünç verme sürdü."
    """
    if metin.startswith("Seçildi (") and "): " in metin:
        ic, kalan = metin[len("Seçildi (") :].split("): ", 1)
        metin = f"{ic}: {kalan}"
    elif metin.startswith("Seçilmedi: "):
        metin = metin[len("Seçilmedi: ") :]
    return tr_upper(metin[:1]) + metin[1:]


def secenek_satirlari(stocktake: StockTake) -> list[dict[str, str]]:
    """İki seçenek ve iade — AYRI satırlar (§9-10). Durduran harcama yetkilisinin adı belgeye."""
    satirlar: list[dict[str, str]] = []
    for satir in selectors_sayim.options_lines(stocktake):
        metin = _durumsuz(str(satir["text"]))
        if satir["key"] == "tmy_32_3" and stocktake.tmy_stop:
            ad = _tek_satir(stocktake.tmy_stop_by_name)
            if ad:
                metin = f"{metin} Durduran harcama yetkilisi: {ad}."
        durum = IADE_DURUMU if satir["key"] == "return" else SECENEK_DURUMU[bool(satir["selected"])]
        satirlar.append(
            {
                "label": str(satir["label"]),
                "status": durum,
                "text": metin,
                "basis": str(satir["basis"]),
            }
        )
    return satirlar


def kurul_secimi_satirlari(stocktake: StockTake) -> list[dict[str, Any]]:
    """Ödünçteki, teslimdeki ve onarımdaki nüsha: kurulun seçimi, dayanağı ve sayıları (AYRI
    satırlar; onarım — F9 ekleri K2)."""
    satirlar = []
    for satir in selectors_sayim.category_lines(stocktake):
        dayanak = str(satir["dayanak"])
        if satir["fallback"]:
            kalip = selectors_sayim.BASIS_FALLBACK_TEXT[str(satir["category"])]
            dayanak += " " + kalip.format(sayi=_sayi(int(satir["fallback"])))
        sonuclar = satir["results"]
        satirlar.append(
            {
                "label": str(satir["label"]),
                "basis_display": str(satir["basis_display"]),
                "dayanak": dayanak,
                "count": int(satir["count"]),
                "found": int(sonuclar[StockTakeResult.FOUND]),
                "by_record": int(sonuclar[StockTakeResult.BY_RECORD]),
                "missing": int(sonuclar[StockTakeResult.MISSING]),
            }
        )
    return satirlar


def kurul_imzalari(stocktake: StockTake) -> list[dict[str, str]]:
    """Sayım kurulunun imza hücreleri (ŞİFRELİ alanlardan): başkan, taşınır kayıt yetkilisi, üyeler.

    Ad yazılmamışsa (eski ya da eksik kayıt) boş imza satırı basılır; kurul en az üç
    kişidir (32/2).
    """
    baskan, yetkili, uye = KURUL_ROLLERI
    uyeler = [_tek_satir(s) for s in (stocktake.committee_members or "").splitlines() if s.strip()]
    hucreler = [
        {"name": _tek_satir(stocktake.committee_chair), "role": baskan},
        {"name": _tek_satir(stocktake.committee_property_officer), "role": yetkili},
        *({"name": ad, "role": uye} for ad in uyeler),
    ]
    while len(hucreler) < 3:
        hucreler.append({"name": "", "role": uye})
    return hucreler


def _harcama_yetkilisi(stocktake: StockTake) -> dict[str, str]:
    onayli = stocktake.status == StockTakeStatus.APPROVED
    return {
        "heading": "OLUR",
        "date": _tarih(stocktake.approved_on) if onayli else BOS_TARIH,
        "name": _tek_satir(stocktake.approved_by_name) if onayli else "",
        "role": "Harcama yetkilisi",
    }


# ---------------------------------------------------------------------------
# PDF bağlamı
# ---------------------------------------------------------------------------
_NOKSAN_SUTUNLARI: Final = [
    _sutun("Sıra", 6, "sayi"),
    _sutun("Barkod", 15),
    _sutun("Kaynak adı", 24),
    _sutun("Yazar", 15),
    _sutun("Kayda göre durum", 13),
    _sutun("Kayıp/hasar tutanağı", 13),
    _sutun("TKYS kodu", 14),
]


def _noksan_satiri(sira: int, k: StockTakeItem) -> list[dict[str, str]]:
    return [
        _h(sira, "sayi"),
        _h(_barkod(k), "tek"),
        _h(_eser_adi(k)),
        _h(_yazar(k)),
        _h(_beklenen(k)),
        _h(_dosya(k)),
        _h(_tkys(k)),
    ]


def _kunye(stocktake: StockTake) -> list[tuple[str, str]]:
    ozet = selectors_sayim.summary(stocktake)
    info = [
        ("Mali yıl", str(stocktake.fiscal_year or "—")),
        ("Sayımın başlangıcı", _zaman(stocktake.started_at)),
    ]
    if stocktake.completed_at is not None:
        info.append(("Sayımın tamamlanması", _zaman(stocktake.completed_at)))
        info.append(
            (
                "İkinci sayım (TMY md. 32/6)",
                (
                    f"Yapıldı — {_sayi(ozet['found_in_round2'])} nüsha ikinci sayımda bulundu"
                    if stocktake.round == 2
                    else "Gerekmedi (noksan çıkmadı)"
                ),
            )
        )
    elif stocktake.round == 2:
        info.append(("İkinci sayım (TMY md. 32/6)", "Sürüyor"))
    info.append(
        (
            "Harcama yetkilisinin onayı",
            _tarih(stocktake.approved_on)
            if stocktake.status == StockTakeStatus.APPROVED
            else "Onay bekliyor",
        )
    )
    return info


def _tablolar(stocktake: StockTake, kalemler: Sequence[StockTakeItem]) -> list[dict[str, Any]]:
    gruplar = _gruplar(stocktake, kalemler)
    onayli = stocktake.status == StockTakeStatus.APPROVED
    tablolar = [
        _tablo(
            "SAYIM SONUÇLARI",
            [_sutun("Kalem", 80), _sutun("Nüsha", 20, "sayi")],
            [
                [_h(ad, "kalin" if kalin else ""), _h(_sayi(n), "sayi kalin" if kalin else "sayi")]
                for ad, n, kalin in sonuc_satirlari(stocktake)
            ],
        ),
        _tablo(
            "SAYIM SIRASINDAKİ SEÇENEKLER",
            [
                _sutun("Seçenek", 18),
                _sutun("Durum", 11),
                _sutun("Açıklama", 47),
                _sutun("Dayanak", 24),
            ],
            [
                [_h(s["label"], "kalin"), _h(s["status"]), _h(s["text"]), _h(s["basis"])]
                for s in secenek_satirlari(stocktake)
            ],
        ),
        _tablo(
            "ÖDÜNÇTEKİ, TESLİMDEKİ VE ONARIMDAKİ NÜSHALAR — SAYIM KURULUNUN SEÇİMİ",
            [
                _sutun("Nüsha", 16),
                _sutun("Kurulun seçimi", 13),
                _sutun("Dayanak", 35),
                _sutun("Sayı", 8, "sayi"),
                _sutun("Bulunan", 10, "sayi"),
                _sutun("Kayda göre", 9, "sayi"),
                _sutun("Noksan", 9, "sayi"),
            ],
            [
                [
                    _h(s["label"]),
                    _h(s["basis_display"]),
                    _h(s["dayanak"]),
                    _h(_sayi(s["count"]), "sayi"),
                    _h(_sayi(s["found"]), "sayi"),
                    _h(_sayi(s["by_record"]), "sayi"),
                    _h(_sayi(s["missing"]), "sayi"),
                ]
                for s in kurul_secimi_satirlari(stocktake)
            ],
            not_=KIMLIK_NOTU,
        ),
        _tablo(
            "BÖLÜMLERE GÖRE",
            [
                _sutun("Bölüm", 40),
                _sutun("Kayıtlara göre", 15, "sayi"),
                _sutun("Bulunan", 15, "sayi"),
                _sutun("Kayda göre alınan", 15, "sayi"),
                _sutun("Noksan", 15, "sayi"),
            ],
            [
                [
                    _h(b["name"]),
                    _h(_sayi(b["expected"]), "sayi"),
                    _h(_sayi(b["found"]), "sayi"),
                    _h(_sayi(b["by_record"]), "sayi"),
                    _h(_sayi(b["missing"]), "sayi"),
                ]
                for b in _bolumler(stocktake)
            ],
            "Sayımda nüsha yok.",
        ),
        _tablo(
            "SAYIM NOKSANI — KAYITTAN DÜŞME TEKLİFİ (TMY MD. 32/7)",
            _NOKSAN_SUTUNLARI,
            [_noksan_satiri(i, k) for i, k in enumerate(gruplar.noksan, start=1)],
            (
                "Kayıttan düşülen noksan yok."
                if onayli
                else NOKSAN_SURERKEN
                if stocktake.status == StockTakeStatus.IN_PROGRESS
                else "Noksan çıkmadı."
            ),
        ),
    ]
    if gruplar.hasar:
        tablolar.append(
            _tablo(
                "HASAR NEDENİYLE KAYITTAN DÜŞME TEKLİFİ (TMY MD. 27/1)",
                [
                    _sutun("Sıra", 6, "sayi"),
                    _sutun("Barkod", 15),
                    _sutun("Kaynak adı", 28),
                    _sutun("Yazar", 17),
                    _sutun("Kayıp/hasar tutanağı", 17),
                    _sutun("TKYS kodu", 17),
                ],
                [
                    [
                        _h(i, "sayi"),
                        _h(_barkod(k), "tek"),
                        _h(_eser_adi(k)),
                        _h(_yazar(k)),
                        _h(_dosya(k)),
                        _h(_tkys(k)),
                    ]
                    for i, k in enumerate(gruplar.hasar, start=1)
                ],
            )
        )
    if gruplar.uzlastirma:
        tablolar.append(
            _tablo(
                "KAYITTA KAYIP GÖRÜNÜP SAYIMDA BULUNAN KAYNAKLAR",
                [
                    _sutun("Sıra", 6, "sayi"),
                    _sutun("Barkod", 15),
                    _sutun("Kaynak adı", 29),
                    _sutun("Yazar", 17),
                    _sutun("Kayıp/hasar tutanağı", 17),
                    _sutun("Sonuç", 16),
                ],
                [
                    [
                        _h(i, "sayi"),
                        _h(_barkod(k), "tek"),
                        _h(_eser_adi(k)),
                        _h(_yazar(k)),
                        _h(_dosya(k)),
                        _h(
                            "Kayıp kaydı kapandı"
                            if k.outcome == StockTakeOutcome.RECONCILED
                            else "Onayda kayıp kaydı kapanır"
                        ),
                    ]
                    for i, k in enumerate(gruplar.uzlastirma, start=1)
                ],
            )
        )
    if gruplar.fazla:
        tablolar.append(
            _tablo(
                "SAYIM FAZLASI (TMY MD. 17)",
                [
                    _sutun("Sıra", 6, "sayi"),
                    _sutun("Okutulan kod", 16),
                    _sutun("Kaynak adı", 26),
                    _sutun("Açıklama", 26),
                    _sutun("Sonuç", 26),
                ],
                [
                    [
                        _h(i, "sayi"),
                        _h(_barkod(k), "tek"),
                        _h(_eser_adi(k)),
                        _h(_tek_satir(k.surplus_note) or "—"),
                        _h(_fazla_sonucu(k)),
                    ]
                    for i, k in enumerate(gruplar.fazla, start=1)
                ],
            )
        )
    if gruplar.onaylanmayan:
        tablolar.append(
            _tablo(
                "HARCAMA YETKİLİSİNİN ONAYLAMADIĞI KALEMLER (KAYITTA KALDI)",
                [
                    _sutun("Sıra", 6, "sayi"),
                    _sutun("Barkod", 15),
                    _sutun("Kaynak adı", 30),
                    _sutun("Yazar", 17),
                    _sutun("Onaylanmama gerekçesi", 32),
                ],
                [
                    [
                        _h(i, "sayi"),
                        _h(_barkod(k), "tek"),
                        _h(_eser_adi(k)),
                        _h(_yazar(k)),
                        _h(_tek_satir(k.outcome_note)),
                    ]
                    for i, k in enumerate(gruplar.onaylanmayan, start=1)
                ],
            )
        )
    if gruplar.durumu_degisen:
        tablolar.append(
            _tablo(
                "ONAYDA DURUMU DEĞİŞEN KALEMLER (KAYITTAN DÜŞÜLMEDİ)",
                [
                    _sutun("Sıra", 6, "sayi"),
                    _sutun("Barkod", 15),
                    _sutun("Kaynak adı", 38),
                    _sutun("Yazar", 18),
                    _sutun("Onaydaki durum", 23),
                ],
                [
                    [
                        _h(i, "sayi"),
                        _h(_barkod(k), "tek"),
                        _h(_eser_adi(k)),
                        _h(_yazar(k)),
                        _h(_tek_satir(k.outcome_note)),
                    ]
                    for i, k in enumerate(gruplar.durumu_degisen, start=1)
                ],
            )
        )
    return tablolar


def _notlar(stocktake: StockTake, kalemler: Sequence[StockTakeItem]) -> list[str]:
    gruplar = _gruplar(stocktake, kalemler)
    notlar: list[str] = []
    if stocktake.round == 2:
        notlar.append(IKINCI_SAYIM_NOTU)
    notlar.append(DUSME_NOTU)
    if gruplar.hasar:
        notlar.append(HASAR_NOTU)
    if any(k.case_id is not None for k in (*gruplar.noksan, *gruplar.hasar, *gruplar.uzlastirma)):
        notlar.append(TUTANAK_BAGI_NOTU)
    if gruplar.fazla:
        notlar.append(FAZLA_NOTU)
    return notlar


@dataclass(frozen=True)
class EkSatiri:
    """Ekin (TMY 34/1) bir satırı: büyüklük, basılan metin, ham sayı, girinti, kalınlık.

    `ham` kesin olmayan büyüklükte `None`'dır (metin "Sayım tamamlanınca yazılır" gibi).
    PDF metni basar; XLSX ham sayıyı SAYI olarak yazar (F9 düzeltme turu: metin olarak
    yazılan "1.234" Excel'de toplanmıyor, ondalık sanılabiliyordu).
    """

    ad: str
    metin: str
    ham: int | None
    girinti: bool = False
    kalin: bool = False


def _ek_satiri(ad: str, n: int | None, *, girinti: bool = False, kalin: bool = False) -> EkSatiri:
    return EkSatiri(ad, _sayi(n), n, girinti, kalin)


def ek_satirlari(stocktake: StockTake) -> dict[str, Any]:
    """Ekin (TMY 34/1) satırları (`EkSatiri`) ve kurul seçimi satırları.

    Sayılar `selectors_sayim.tmy_34_1`'dendir; PDF ve XLSX aynı satırları basar. Kesin
    olmayan büyüklük (sayım sürerken ya da tamamlanmış sayımda kararı bekleyen fazla
    varken) sayı yerine nedenini yazar.
    """
    sayilar = selectors_sayim.tmy_34_1(stocktake)
    yontem = dict(AcquisitionMethod.choices)
    durum = dict(CopyStatus.choices)
    satirlar: list[EkSatiri] = [
        _ek_satiri(
            "Önceki yıldan devir (yıl başında kayıtta olan)",
            sayilar["previous_year_carryover"],
            kalin=True,
        ),
        _ek_satiri(
            "Yıl içinde giren (Yönetmelik Md. 10/5 yolları ve sayım fazlası)",
            sayilar["entered"],
            kalin=True,
        ),
    ]
    satirlar += [
        _ek_satiri(str(yontem[y]), n, girinti=True) for y, n in sayilar["entered_by_method"].items()
    ]
    satirlar += [
        _ek_satiri(
            "Programa aktarım (mevcut koleksiyonun kaydı; taşınır girişi değildir)",
            sayilar["program_transfer"],
        ),
        _ek_satiri("Yıl içinde çıkan (kayıttan düşme ve devir)", sayilar["exited"], kalin=True),
    ]
    satirlar += [
        _ek_satiri(str(durum[d]), n, girinti=True) for d, n in sayilar["exited_by_status"].items()
    ]
    kesin = bool(sayilar["count_final"])
    devir_kesin = bool(sayilar["carryover_final"])
    belirsiz = (
        BELIRSIZ_FAZLA
        if stocktake.status in (StockTakeStatus.COMPLETED, StockTakeStatus.APPROVED)
        and sayilar["surplus_unresolved"]
        else BELIRSIZ
    )
    # Sayımda bulunan kesin ama hasar önerisi onay bekliyorsa devir onaydan sonra yazılır.
    devir_belirsiz = belirsiz if not kesin else BELIRSIZ_ONAY
    satirlar.append(
        _ek_satiri(
            "Kayda göre yıl sonu (devir + aktarım + giren − çıkan)", sayilar["year_end_by_record"]
        )
    )

    def _kesinse(
        ad: str, n: int | None, *, dogru: bool, bos: str, girinti: bool = False, kalin: bool = False
    ) -> EkSatiri:
        if dogru:
            return _ek_satiri(ad, n, girinti=girinti, kalin=kalin)
        return EkSatiri(ad, bos, None, girinti, kalin)

    # Madde 25 (a): gelecek yıla devirden onayda 27/1 ile düşülenler çıkarılır; bileşenleri
    # altında AYRI satırdadır (sayımda bulunan miktar değişmez — tutanağın SONUÇLAR'ıyla aynı).
    satirlar += [
        _kesinse(
            GELECEK_YIL_SATIRI,
            sayilar["next_year_carryover"],
            dogru=devir_kesin,
            bos=devir_belirsiz,
            kalin=True,
        ),
        _kesinse(
            SAYIMDA_BULUNAN_SATIRI,
            sayilar["found_quantity"],
            dogru=kesin,
            bos=belirsiz,
            girinti=True,
        ),
        _kesinse(
            HASAR_DUSULEN_SATIRI,
            sayilar["damage_written_off"],
            dogru=devir_kesin,
            bos=devir_belirsiz,
            girinti=True,
        ),
        _kesinse(
            "Fark (kayda göre yıl sonu − gelecek yıla devir)",
            sayilar["difference"],
            dogru=devir_kesin,
            bos=devir_belirsiz,
        ),
    ]
    satirlar += [
        _ek_satiri("Sayım fazlası", sayilar["surplus"]),
        _ek_satiri(
            "Sayım noksanı (kayıttan düşülen: " + _sayi(sayilar["missing_written_off"]) + ")",
            sayilar["missing"],
        ),
    ]
    return {
        "fiscal_year": int(sayilar["fiscal_year"]),
        "rows": satirlar,
        "categories": kurul_secimi_satirlari(stocktake),
        "note": str(sayilar["note"]),
        "final": kesin,
        "surplus_unresolved": int(sayilar["surplus_unresolved"]),
    }


def _ek(stocktake: StockTake) -> dict[str, Any]:
    ek = ek_satirlari(stocktake)
    return {
        "title": "EK: TAŞINIR SAYIM VE DÖKÜM CETVELİNE AKTARILACAK SAYILAR",
        "subtitle": f"{ek['fiscal_year']} mali yılı (1 Ocak - 31 Aralık) · kütüphane materyali, "
        "nüsha sayısıyla",
        "paragraphs": [
            f"“{TMY34_1}” (Taşınır Mal Yönetmeliği md. 34/1). Aşağıdaki sayılar taşınır kayıt "
            "yetkilisinin cetvele aktaracağı büyüklüklerdir.",
            ARA_SAYIM_NOTU,
        ],
        "tables": [
            _tablo(
                "",
                [_sutun("Büyüklük", 78), _sutun("Nüsha", 22, "sayi")],
                [
                    [
                        _h(("— " if s.girinti else "") + s.ad, "kalin" if s.kalin else ""),
                        _h(s.metin, "sayi kalin" if s.kalin else "sayi"),
                    ]
                    for s in ek["rows"]
                ],
            ),
            _tablo(
                "ÖDÜNÇTEKİ, TESLİMDEKİ VE ONARIMDAKİ NÜSHALARIN SAYILIŞI",
                [
                    _sutun("Nüsha", 28),
                    _sutun("Kurulun seçimi", 24),
                    _sutun("Sayı", 12, "sayi"),
                    _sutun("Bulunan", 12, "sayi"),
                    _sutun("Kayda göre", 12, "sayi"),
                    _sutun("Noksan", 12, "sayi"),
                ],
                [
                    [
                        _h(s["label"]),
                        _h(s["basis_display"]),
                        _h(_sayi(s["count"]), "sayi"),
                        _h(_sayi(s["found"]), "sayi"),
                        _h(_sayi(s["by_record"]), "sayi"),
                        _h(_sayi(s["missing"]), "sayi"),
                    ]
                    for s in ek["categories"]
                ],
            ),
        ],
        "notes": [GELECEK_YIL_NOTU, ek["note"]],
    }


def durum_notu(stocktake: StockTake) -> str:
    """Başlığın altındaki durum notu: sürerken TASLAK, tamamlanınca onay bekleniyor.

    Tamamlanmış sayımda kararı bekleyen sayım fazlası varsa not bunu da söyler: sayılar o
    kararlarla değişir (F9 düzeltme turu — imzalı tutanak onaydaki sayıdan farklı kalmasın).
    """
    if stocktake.status == StockTakeStatus.IN_PROGRESS:
        return SURUYOR_NOTU
    if stocktake.status != StockTakeStatus.COMPLETED:
        return ""
    kararsiz = int(selectors_sayim.found_quantity(stocktake)["surplus_unresolved"])
    if kararsiz:
        return f"{ONAY_BEKLIYOR_NOTU} {FAZLA_KARARI_NOTU.format(sayi=_sayi(kararsiz))}"
    return ONAY_BEKLIYOR_NOTU


def stocktake_report_context(stocktake: StockTake, *, on: date | None = None) -> dict[str, Any]:
    """Sayım tutanağının şablon bağlamı (basılabilirlik önce denetlenir)."""
    engel = stocktake_document_blocker(stocktake)
    if engel:
        raise ValidationError(engel)
    gun = _gun(on)
    kalemler = _kalemler(stocktake)
    config = SchoolConfig.load()
    return {
        **letterhead_context(
            school_name=_okul_adi(config),
            district=config.district,
            principal_name=config.principal_name,
        ),
        "document_title": "SAYIM TUTANAĞI",
        "subtitle": "Kütüphane materyali — Taşınır Mal Yönetmeliği md. 32",
        "status_note": durum_notu(stocktake),
        "issued_on": _tarih(gun),
        "info": [{"label": e, "value": v} for e, v in _kunye(stocktake)],
        "paragraphs": [ACILIS],
        "tables": _tablolar(stocktake, kalemler),
        "notes": _notlar(stocktake, kalemler),
        "signature_title": "SAYIM KURULU",
        "signature_rows": _imza_izgarasi(kurul_imzalari(stocktake)),
        "approval": _harcama_yetkilisi(stocktake),
        "footnote": TKYS_DIPNOTU,
        "annex": _ek(stocktake),
    }


def stocktake_report_pdf(stocktake: StockTake, *, on: date | None = None) -> bytes:
    return html_to_pdf(render_to_string(SABLON, stocktake_report_context(stocktake, on=on)))


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------
#: "Kalemler" sayfasının sütunları (başlık, genişlik) — kişisiz.
KALEM_XLSX_SUTUNLARI: Final[tuple[tuple[str, int], ...]] = (
    ("Sıra", 6),
    ("Barkod", 14),
    ("Kaynak adı", 44),
    ("Yazar(lar)", 28),
    ("Yer numarası", 16),
    ("Bölüm", 20),
    ("Sınıf kitaplığı (yerinde sayılan)", 14),
    ("Kayda göre durum", 26),
    ("Nasıl sayılır", 22),
    ("Sonuç", 24),
    ("Bulunduğu tur", 10),
    ("Nasıl bulundu", 26),
    ("Okutma zamanı", 17),
    ("Kayıp/hasar tutanağı", 22),
    ("Hasar nedeniyle düşme teklifi", 12),
    ("Onay sonucu", 30),
    ("Kayıttan düşme yolu", 34),
    ("Onay notu", 34),
    ("TKYS kodu", 18),
    ("Sayım fazlası açıklaması", 34),
    ("Sayım fazlası sonucu", 30),
)

_KALIN: Final = Font(bold=True)
_DOLGU: Final = PatternFill("solid", fgColor="E8EEF6")
_INCE: Final = Side(style="thin", color="B7C3D0")


def _baslik_satiri(ws: Worksheet, basliklar: Sequence[tuple[str, int]]) -> int:
    ws.append([ad for ad, _ in basliklar])
    satir = ws.max_row
    for sutun, (_, genislik) in enumerate(basliklar, start=1):
        hucre = ws.cell(row=satir, column=sutun)
        hucre.font = _KALIN
        hucre.fill = _DOLGU
        hucre.border = Border(bottom=_INCE)
        hucre.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(sutun)].width = genislik
    return int(satir)


def _ara_baslik(ws: Worksheet, metin: str) -> None:
    ws.append([])
    ws.append([metin])
    ws.cell(row=ws.max_row, column=1).font = _KALIN


def _kalem_satiri_xlsx(sira: int, k: StockTakeItem) -> list[Any]:
    eser = k.copy.work if k.copy is not None else None
    return [
        sira,
        _barkod(k),
        _eser_adi(k),
        _yazar(k),
        eser.call_number if eser is not None else "",
        k.section.name if k.section is not None else "",
        k.class_section.class_label if k.class_section is not None else "",
        str(k.get_expected_status_display()) if k.expected_status else "",
        selectors_sayim.item_basis_label(k),
        str(k.get_result_display()),
        k.found_in_round,
        str(k.get_found_via_display()) if k.found_via else "",
        _zaman(k.scanned_at) if k.scanned_at is not None else "",
        "" if _dosya(k) == "—" else _dosya(k),
        "Evet" if k.damage_write_off else "",
        str(k.get_outcome_display()) if k.outcome else "",
        str(k.get_write_off_path_display()) if k.write_off_path else "",
        _tek_satir(k.outcome_note),
        k.copy.external_asset_ref if k.copy is not None else "",
        _tek_satir(k.surplus_note),
        _fazla_sonucu(k) if k.is_surplus else "",
    ]


def stocktake_report_xlsx(stocktake: StockTake, *, on: date | None = None) -> bytes:
    """Sayım tutanağı (XLSX): özet, bütün kalemler ve ekin sayıları — PDF'le aynı kaynak."""
    engel = stocktake_document_blocker(stocktake)
    if engel:
        raise ValidationError(engel)
    gun = _gun(on)
    config = SchoolConfig.load()
    wb = Workbook()
    wb.properties.title = TUTANAK_ADI
    ozet = wb.active
    assert ozet is not None  # yeni çalışma kitabında etkin sayfa daima vardır
    ozet.title = TUTANAK_ADI
    ozet.append([_okul_adi(config) or "—"])
    ozet.append([TUTANAK_ADI])
    ozet.append([f"Düzenlenme tarihi: {_tarih(gun)} · Durum: {stocktake.get_status_display()}"])
    ozet["A1"].font = _KALIN
    ozet["A2"].font = Font(bold=True, size=13)
    for etiket, deger in _kunye(stocktake):
        # Mali yıl SAYI olarak yazılır (F9 düzeltme turu); öbür künye satırları metindir.
        yil = stocktake.fiscal_year if etiket == "Mali yıl" and stocktake.fiscal_year else None
        ozet.append([etiket, yil if yil is not None else deger])
    ozet.column_dimensions["A"].width = 58
    ozet.column_dimensions["B"].width = 60
    ozet.column_dimensions["C"].width = 70
    ozet.column_dimensions["D"].width = 40

    _ara_baslik(ozet, "Sayım sonuçları")
    for ad, n, _kalin in sonuc_satirlari(stocktake):
        ozet.append([ad, n])
    _ara_baslik(ozet, "Sayım sırasındaki seçenekler")
    ozet.append(["Seçenek", "Durum", "Açıklama", "Dayanak"])
    for s in secenek_satirlari(stocktake):
        ozet.append([s["label"], s["status"], s["text"], s["basis"]])
    _ara_baslik(ozet, "Ödünçteki, teslimdeki ve onarımdaki nüshalar — sayım kurulunun seçimi")
    ozet.append(["Nüsha", "Kurulun seçimi", "Dayanak", "Sayı", "Bulunan", "Kayda göre", "Noksan"])
    for secim in kurul_secimi_satirlari(stocktake):
        ozet.append(
            [
                secim["label"],
                secim["basis_display"],
                secim["dayanak"],
                secim["count"],
                secim["found"],
                secim["by_record"],
                secim["missing"],
            ]
        )
    ozet.append([KIMLIK_NOTU])
    ozet.append([TKYS_DIPNOTU])
    for satir in ozet.iter_rows():
        for hucre in satir:
            hucre.alignment = Alignment(vertical="top", wrap_text=True)

    kalem_sayfasi = wb.create_sheet("Kalemler")
    baslik = _baslik_satiri(kalem_sayfasi, KALEM_XLSX_SUTUNLARI)
    for i, k in enumerate(_kalemler(stocktake), start=1):
        kalem_sayfasi.append(_kalem_satiri_xlsx(i, k))
    kalem_sayfasi.freeze_panes = kalem_sayfasi.cell(row=baslik + 1, column=1)

    ek = ek_satirlari(stocktake)
    ek_sayfasi = wb.create_sheet("Cetvele aktarılacak sayılar")
    ek_sayfasi.append([EK_ADI])
    ek_sayfasi["A1"].font = Font(bold=True, size=13)
    ek_sayfasi.append([f"{ek['fiscal_year']} mali yılı · kütüphane materyali, nüsha sayısıyla"])
    ek_sayfasi.append([])
    _baslik_satiri(ek_sayfasi, (("Büyüklük", 70), ("Nüsha", 26)))
    for satir in ek["rows"]:
        # Sayılar SAYI hücresidir (binlik ayraçlı biçimle); kesin olmayan büyüklük nedenini yazar.
        ek_sayfasi.append(
            [
                ("— " if satir.girinti else "") + satir.ad,
                satir.ham if satir.ham is not None else satir.metin,
            ]
        )
        if satir.ham is not None:
            ek_sayfasi.cell(row=ek_sayfasi.max_row, column=2).number_format = "#,##0"
    _ara_baslik(ek_sayfasi, "Ödünçteki, teslimdeki ve onarımdaki nüshaların sayılışı")
    ek_sayfasi.append(["Nüsha", "Kurulun seçimi", "Sayı", "Bulunan", "Kayda göre", "Noksan"])
    for secim in ek["categories"]:
        ek_sayfasi.append(
            [
                secim["label"],
                secim["basis_display"],
                secim["count"],
                secim["found"],
                secim["by_record"],
                secim["missing"],
            ]
        )
    ek_sayfasi.append([])
    ek_sayfasi.append([GELECEK_YIL_NOTU])
    ek_sayfasi.append([ARA_SAYIM_NOTU])
    ek_sayfasi.append([ek["note"]])

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def stocktake_document(stocktake: StockTake, slug: str, bicim: str = PDF) -> bytes:
    """Belgeyi üretir (basılabilirlik ve biçim önce denetlenir)."""
    ensure_stocktake_document(stocktake, slug, bicim)
    if bicim == XLSX:
        return stocktake_report_xlsx(stocktake)
    return stocktake_report_pdf(stocktake)


def stocktake_document_filename(slug: str, bicim: str = PDF, gun: date | None = None) -> str:
    return belge_dosya_adi(_BELGE_TANIMLARI[slug].ad, gun, bicim=bicim)
