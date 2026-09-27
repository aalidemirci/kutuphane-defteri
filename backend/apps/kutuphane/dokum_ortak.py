"""F10 dökümlerinin (D kolu) ortak yardımcıları — PDF bağlamı ve XLSX biçimi.

Dışa aktarım dosyası, alfabetik katalog dökümü (E17), Taşınır Kütüphane Defteri dökümü ve
yönetim hesabı cetveli hazırlığı (E11) ve kişi dökümü (KVKK md. 11) aynı küçük yardımcıları
kullanır: yerel tarih, Türkçe binlik ayraçlı sayı, dosya adı (belge adı + yerel tarih —
sözlük §3), antet bağlamı, tablo sözlüğü (`documents/_sayim_tablosu.html` biçimi) ve Excel
başlık satırı. PDF'ler YALNIZ `shared.pdf.html_to_pdf` kapısından üretilir (CLAUDE.md §3).

**XLSX'te sayılar sayıdır:** sayı, tarih ve para hücreleri metne çevrilmeden yazılır ve
biçimle gösterilir (`NUMBER_FORMAT`, `DATE_FORMAT`, `MONEY_FORMAT`).
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Final

from django.utils import timezone
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.okul.models import SchoolConfig
from shared.letterhead import letterhead_context

PDF: Final = "pdf"
XLSX: Final = "xlsx"

NUMBER_FORMAT: Final = "#,##0"
MONEY_FORMAT: Final = "#,##0.00"
DATE_FORMAT: Final = "DD.MM.YYYY"

BOS: Final = "—"

_KALIN: Final = Font(bold=True)
_DOLGU: Final = PatternFill("solid", fgColor="E8EEF6")
_INCE: Final = Side(style="thin", color="B7C3D0")


def bugun(on: date | None = None) -> date:
    """Belgenin tarihi — yerel gün (UTC'den tarih türetilmez, CLAUDE.md §2-9)."""
    return on or timezone.localdate()


def tarih(gun: date | None) -> str:
    return f"{gun:%d.%m.%Y}" if gun is not None else BOS


def zaman(an: datetime | None) -> str:
    return f"{timezone.localtime(an):%d.%m.%Y %H:%M}" if an is not None else BOS


def sayi(n: int | None) -> str:
    """Türkçe binlik ayraçlı sayı (1.234)."""
    if n is None:
        return BOS
    return f"{n:,}".replace(",", ".")


def para(tutar: Decimal | None) -> str:
    """Türkçe para biçimi (1.234,50); boşsa tire."""
    if tutar is None:
        return BOS
    metin = f"{tutar:,.2f}"
    return metin.replace(",", "\x00").replace(".", ",").replace("\x00", ".")


def tek_satir(metin: str | None) -> str:
    return " ".join((metin or "").split())


def dosya_adi(ad: str, gun: date | None = None, *, bicim: str = PDF) -> str:
    """'Alfabetik-katalog-dökümü_26.09.2026.pdf' — belge adı + yerel tarih (sözlük §3)."""
    return f"{ad.replace(' ', '-')}_{bugun(gun):%d.%m.%Y}.{bicim}"


def okul_adi(config: SchoolConfig | None = None) -> str:
    yapilandirma = config or SchoolConfig.load()
    return " ".join((yapilandirma.school_name or yapilandirma.kisa_ad or "").split())


def antet(config: SchoolConfig | None = None) -> dict[str, str]:
    yapilandirma = config or SchoolConfig.load()
    return letterhead_context(
        school_name=okul_adi(yapilandirma),
        district=yapilandirma.district,
        principal_name=yapilandirma.principal_name,
    )


def hucre(deger: Any, cls: str = "") -> dict[str, str]:
    """PDF tablo hücresi (`documents/_sayim_tablosu.html`)."""
    return {"v": "" if deger is None else str(deger), "c": cls}


def sutun(baslik: str, genislik: int, cls: str = "") -> dict[str, Any]:
    return {"header": baslik, "width": genislik, "cls": cls}


def tablo(
    baslik: str,
    sutunlar: list[dict[str, Any]],
    satirlar: list[list[dict[str, str]]],
    bos: str = "",
    *,
    not_: str = "",
) -> dict[str, Any]:
    """PDF tablosu; sütun genişliklerinin toplamı %100 olmalıdır (sayfa bütçesi)."""
    toplam = sum(int(s["width"]) for s in sutunlar)
    if toplam != 100:
        raise ValueError(f"Sütun genişlikleri %100 olmalı (şu an %{toplam}).")
    return {"title": baslik, "columns": sutunlar, "rows": satirlar, "empty": bos, "note": not_}


def baslik_satiri(ws: Worksheet, basliklar: Sequence[tuple[str, int]]) -> int:
    """Kalın, dolgulu başlık satırı yazar ve sütun genişliklerini kurar → satır no."""
    ws.append([ad for ad, _ in basliklar])
    satir = int(ws.max_row)
    for sira, (_, genislik) in enumerate(basliklar, start=1):
        h = ws.cell(row=satir, column=sira)
        h.font = _KALIN
        h.fill = _DOLGU
        h.border = Border(bottom=_INCE)
        h.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(sira)].width = genislik
    return satir


def kalin(ws: Worksheet, satir: int, sutun_no: int = 1) -> None:
    ws.cell(row=satir, column=sutun_no).font = _KALIN


def bicimle(ws: Worksheet, satir: int, bicimler: dict[int, str]) -> None:
    """Satırın verilen sütunlarına sayı/tarih biçimi uygular (değer sayı kalır)."""
    for sutun_no, bicim in bicimler.items():
        h = ws.cell(row=satir, column=sutun_no)
        if h.value is not None and h.value != "":
            h.number_format = bicim


# ---------------------------------------------------------------------------
# `write_only` çalışma kitabı (büyük listeler: hücre nesnesi bellekte birikmez)
# ---------------------------------------------------------------------------
def wo_baslik(ws: Any, basliklar: Sequence[tuple[str, int]], *, ilk_sutun: int = 1) -> None:
    """Yazma kipindeki sayfaya kalın başlık satırı ekler ve sütun genişliklerini kurar.

    Genişlikler satırlardan ÖNCE kurulmalıdır (yazma kipi sütunları başta yazar).
    """
    satir = []
    for sira, (ad, genislik) in enumerate(basliklar, start=ilk_sutun):
        ws.column_dimensions[get_column_letter(sira)].width = genislik
        h = WriteOnlyCell(ws, value=ad)
        h.font = _KALIN
        h.fill = _DOLGU
        h.alignment = Alignment(vertical="center", wrap_text=True)
        satir.append(h)
    ws.append(satir)


def wo_hucre(ws: Any, deger: Any, bicim: str = "", *, kalin_mi: bool = False) -> Any:
    """Yazma kipinde biçimli hücre; değer sayı/tarih KALIR. Boş değer düz yazılır."""
    if (not bicim and not kalin_mi) or deger is None or deger == "":
        return deger
    h = WriteOnlyCell(ws, value=deger)
    if bicim:
        h.number_format = bicim
    if kalin_mi:
        h.font = _KALIN
    return h
