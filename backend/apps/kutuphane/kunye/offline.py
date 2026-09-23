"""Çevrimdışı künye yolu (§8.5, U13'ün ikinci yarısı) — dosyayla gidip gelen künye.

Kütüphane masasında internet yoksa özellik ölmez: künyesi eksik eserlerin
**ISBN listesi dışa aktarılır** → internetli **başka bir cihazda** doldurulur →
dosya geri aktarılır ve aynı ön izleme/onay ekranından geçer.

**Ayrı cihaz demektir** (Yönerge 11/18): kurum bilgisayarına telefon, mobil
modem ya da kişisel erişim noktası bağlanarak internet alınamaz. Taşımada
Yönerge 10/4-10/5'teki taşınabilir bellek kuralları geçerlidir. Bu iki cümle
kılavuza ve dosyanın "Bilgi" sayfasına yazılır.

**Dosya kişisel veri taşımaz** (§8.5): satırlar kitap künyesidir; ödünç, üye,
bağışçı ve demirbaş bilgisi girmez.

**Sayfa adı "Katalog" DEĞİLDİR, "Künye"dir.** Bilinçlidir: katalog Excel içe
aktarımı yalnız "Katalog" sayfasını okur (`import_schema.CATALOG_SHEET`) ve
"Nüsha Sayısı" boş satırı bir nüsha sayar. Bu dosya var olan eserlerin
künyesini tamamlamak içindir; yanlışlıkla içe aktarıma verilirse her satır için
yeni nüsha açılırdı. Ayrı sayfa adı o kazayı baştan keser.

**Çevirmen sütunu yoktur ve gelen dosyada varsa yok sayılır** (§8.5-5):
çevirmen dışarıdan doldurulmaz, kullanıcıya sorulur.

Geri okunan dosya da **güvenilmeyen girdidir**: boyut ve satır tavanı vardır,
her hücre `temizlik`ten (NFC, denetim karakteri, HTML varlığı, uzunluk) geçer.

**Yazılan hücre de metindir, formül değildir.** openpyxl "=" ile başlayan bir
dizgeyi formül diye işaretler; bu dosyanın tek varlık nedeni internete bağlı
BAŞKA bir cihaza taşınmaktır ve formül orada açılınca çalışırdı. Katalog metni
kullanıcı onayından geçmeden de dolabiliyor (toplu Excel içe aktarımı eser
adını doğrudan yazar), bu yüzden hücre türü metne SABİTLENİR (`_yaz`).
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date
from typing import Any

from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.kutuphane import isbn as isbn_module
from apps.kutuphane import selectors
from apps.kutuphane.import_schema import COLUMNS_BY_KEY, fold, match_header
from apps.kutuphane.kunye import temizlik
from apps.kutuphane.kunye.oneri import (
    ALAN_TAVANLARI,
    UYARI_CEVIRMEN,
    KunyeOnerisi,
    gecerli_yil,
)
from apps.kutuphane.models import Work

#: Belgenin adı (indirme adı ve ekrandaki kart başlığı).
DOCUMENT_NAME = "ISBN Künye Listesi"
#: Veri sayfası — "Katalog" DEĞİL (modül başlığındaki gerekçe).
SHEET_NAME = "Künye"
INFO_SHEET_NAME = "Bilgi"

#: Dosyadaki eser kimliği sütunu (sözlükte yoktur; eşleşmeyi kesinleştirir).
WORK_NO_HEADER = "Eser No"

#: Dışa aktarılan künye sütunları (sözlükteki anahtarlar; sıra ekran sırasıdır).
EXPORT_KEYS: tuple[str, ...] = (
    "isbn",
    "title",
    "authors",
    "publisher",
    "edition",
    "publish_year",
    "subjects",
    "classification_code",
    "language",
)

#: Güvenilmeyen girdi tavanları.
MAX_DOSYA_BAYT = 5 * 1024 * 1024
MAX_SATIR = 5000
MAX_SUTUN = 40

DURUM_ESLESTI = "eslesti"
DURUM_ESER_YOK = "eser_yok"
DURUM_COKLU_ESER = "coklu_eser"
DURUM_ISBN_YOK = "isbn_yok"

KAYNAK_ADI = "Çevrimdışı künye dosyası"

#: Dayanağın TAM adı. Depodaki metin `docs/mevzuat/meb-bilgi-ve-sistem-
#: guvenligi-yonergesi.md`'dir ve kısa adı `docs/mevzuat/BENIOKU.md`'de
#: "Yönerge = Bilgi ve Sistem Güvenliği Yönergesi" diye sabitlenmiştir
#: (CLAUDE.md §2-13: atıf depodaki metinden doğrulanır). Okul bu dosyayı
#: BTR'ye ya da müdürlüğe gösterdiğinde ad tutmalıdır.
YONERGE_ADI = "Millî Eğitim Bakanlığı Bilgi ve Sistem Güvenliği Yönergesi"

BILGI_SATIRLARI: tuple[str, ...] = (
    "Bu dosya, künyesi eksik eserlerin ISBN listesidir.",
    "Boş hücreleri internete bağlı BAŞKA bir cihazda doldurun, dosyayı programa geri yükleyin.",
    "Kurum bilgisayarına telefon, mobil modem ya da kişisel erişim noktası bağlanarak "
    f"internet alınamaz ({YONERGE_ADI} 11/18).",
    "Taşınabilir bellekle taşırken Yönerge 10/4 ve 10/5 kuralları geçerlidir.",
    "“Eser No” ve “ISBN” sütunlarını DEĞİŞTİRMEYİN: eşleşme bu iki sütundan yapılır.",
    "Çevirmen sütunu yoktur; çeviri eserde çevirmeni programda elle yazın.",
    "Dosyaya öğrenci, veli ya da personel bilgisi yazmayın; bu dosya yalnız kitap künyesi taşır.",
    "Geri yüklediğinizde program size önce önizleme gösterir; onaylamadan hiçbir alan değişmez.",
)

#: Excel'in formül ya da komut sayabileceği başlangıçlar (modül başlığındaki
#: gerekçe). "=" openpyxl'i doğrudan formül kipine sokar; öbürleri Excel'in
#: kendi yorumlayıcısına karşı savunmadır.
_FORMUL_BASLANGICLARI: tuple[str, ...] = ("=", "+", "-", "@", "\t", "\r")

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_KEY_FILL = PatternFill("solid", fgColor="F8CBAD")
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)


def export_filename(today: date | None = None) -> str:
    """İndirme adı: 'ISBN-Künye-Listesi_23.09.2026.xlsx' (yerel tarih, UTC değil)."""
    gun = today or timezone.localdate()
    return f"{DOCUMENT_NAME.replace(' ', '-')}_{gun:%d.%m.%Y}.xlsx"


def _yaz(ws: Worksheet, *, satir: int, sutun: int, deger: Any) -> Any:
    """Hücreyi yazar ve metin değerini FORMÜL olmaktan çıkarır (modül başlığı)."""
    hucre = ws.cell(row=satir, column=sutun, value=deger)
    if isinstance(deger, str) and deger.startswith(_FORMUL_BASLANGICLARI):
        hucre.data_type = "s"
    return hucre


def _basliklar() -> list[str]:
    return [WORK_NO_HEADER, *(COLUMNS_BY_KEY[anahtar].header for anahtar in EXPORT_KEYS)]


def _eser_satiri(eser: Work) -> list[Any]:
    degerler: dict[str, Any] = {
        "isbn": eser.isbn13 or eser.isbn,
        "title": eser.title,
        "authors": eser.authors,
        "publisher": eser.publisher,
        "edition": eser.edition,
        "publish_year": eser.publish_year,
        "subjects": eser.subjects,
        "classification_code": eser.classification_code,
        "language": eser.language,
    }
    return [eser.pk, *(degerler[anahtar] for anahtar in EXPORT_KEYS)]


def _kunye_sayfasi(ws: Worksheet, eserler: Iterable[Work]) -> None:
    ws.title = SHEET_NAME
    basliklar = _basliklar()
    for sutun, baslik in enumerate(basliklar, start=1):
        hucre = _yaz(ws, satir=1, sutun=sutun, deger=baslik)
        hucre.font = _HEADER_FONT
        # Eşleşme anahtarları (Eser No, ISBN) ayrı renkte: değiştirilmemeleri gerekir.
        hucre.fill = _KEY_FILL if sutun <= 2 else _HEADER_FILL
        hucre.alignment = _HEADER_ALIGNMENT
        ws.column_dimensions[get_column_letter(sutun)].width = max(12, len(baslik) + 6)
    ws.column_dimensions["B"].number_format = "@"  # ISBN metin kalsın
    ws.column_dimensions["C"].width = 40
    ws.freeze_panes = "A2"
    for satir, eser in enumerate(eserler, start=2):
        for sutun, deger in enumerate(_eser_satiri(eser), start=1):
            _yaz(ws, satir=satir, sutun=sutun, deger=deger)


def _bilgi_sayfasi(ws: Worksheet) -> None:
    ws.title = INFO_SHEET_NAME
    ws.column_dimensions["A"].width = 110
    baslik = _yaz(ws, satir=1, sutun=1, deger=DOCUMENT_NAME)
    baslik.font = _HEADER_FONT
    for satir, metin in enumerate(BILGI_SATIRLARI, start=3):
        hucre = _yaz(ws, satir=satir, sutun=1, deger=metin)
        hucre.alignment = Alignment(wrap_text=True, vertical="top")


def build_export(eserler: Iterable[Work] | None = None) -> bytes:
    """Künyesi eksik eserlerin ISBN listesini xlsx olarak üretir."""
    kayitlar = list(eserler) if eserler is not None else list(selectors.works_missing_metadata())
    wb = Workbook()
    wb.properties.title = DOCUMENT_NAME
    kunye = wb.active
    assert kunye is not None  # yeni çalışma kitabında etkin sayfa daima vardır
    _kunye_sayfasi(kunye, kayitlar)
    _bilgi_sayfasi(wb.create_sheet(INFO_SHEET_NAME))
    wb.active = 0
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Doldurulmuş dosyanın geri alınması
# ---------------------------------------------------------------------------
@dataclass(frozen=True, slots=True)
class CevrimdisiSatir:
    """Dosyanın bir satırı: hangi esere, hangi öneri, hangi durumda."""

    satir_no: int
    isbn13: str
    durum: str
    #: Eşleşen eser (yoksa `None`). Nesne TUTULUR, kimlik değil: ön izleme
    #: satır başına yeniden sorgu atmasın (5000 satırlık dosyada N+1 olurdu).
    work: Work | None
    oneri: KunyeOnerisi

    @property
    def work_id(self) -> int | None:
        return self.work.pk if self.work is not None else None

    @property
    def work_title(self) -> str:
        return self.work.title if self.work is not None else ""


@dataclass(frozen=True, slots=True)
class CevrimdisiOnizleme:
    """Önizleme sonucu — YAZMA YOKTUR, kullanıcı onayı ayrı adımdır (§8.5-5)."""

    satirlar: tuple[CevrimdisiSatir, ...]
    tarih: date
    atlanan_sutunlar: tuple[str, ...]

    @property
    def kaynak_etiketi(self) -> str:
        return f"{KAYNAK_ADI}, {self.tarih:%d.%m.%Y}"

    @property
    def sayilar(self) -> dict[str, int]:
        sayac = {DURUM_ESLESTI: 0, DURUM_ESER_YOK: 0, DURUM_COKLU_ESER: 0, DURUM_ISBN_YOK: 0}
        for satir in self.satirlar:
            sayac[satir.durum] = sayac.get(satir.durum, 0) + 1
        return sayac


class CevrimdisiDosyaHatasi(ValueError):
    """Dosya okunamadı ya da beklenen sütunları taşımıyor (kullanıcıya gösterilir)."""


def _xlsx_satirlari(icerik: bytes) -> list[list[Any]]:
    try:
        wb = load_workbook(io.BytesIO(icerik), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl bozuk dosyada çeşitli hata verir
        raise CevrimdisiDosyaHatasi(
            "Dosya Excel dosyası olarak açılamadı. Programın verdiği dosyayı doldurup yükleyin."
        ) from exc
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
        satirlar: list[list[Any]] = []
        for satir in ws.iter_rows(max_row=MAX_SATIR + 1, max_col=MAX_SUTUN, values_only=True):
            satirlar.append(list(satir))
        return satirlar
    finally:
        wb.close()


def _csv_satirlari(icerik: bytes) -> list[list[Any]]:
    metin = icerik.decode("utf-8-sig", errors="replace")
    ornek = metin[:4096]
    ayrac = ";" if ornek.count(";") > ornek.count(",") else ","
    okuyucu = csv.reader(io.StringIO(metin), delimiter=ayrac)
    return [satir for _, satir in zip(range(MAX_SATIR + 1), okuyucu, strict=False)]


def _satirlari_coz(icerik: bytes) -> list[list[Any]]:
    if not icerik:
        raise CevrimdisiDosyaHatasi("Dosya boş.")
    if len(icerik) > MAX_DOSYA_BAYT:
        raise CevrimdisiDosyaHatasi("Dosya çok büyük (en çok 5 MB).")
    # xlsx bir zip'tir: 'PK' imzası dosya türünü ayırmanın en ucuz yoludur.
    if icerik[:2] == b"PK":
        return _xlsx_satirlari(icerik)
    return _csv_satirlari(icerik)


def _baslik_haritasi(satir: Sequence[Any]) -> tuple[dict[int, str], list[str]]:
    """Sütun sırası → sözlük anahtarı; tanınmayan ve YOK SAYILAN başlıklar ayrıca döner."""
    harita: dict[int, str] = {}
    atlanan: list[str] = []
    for sira, hucre in enumerate(satir[:MAX_SUTUN]):
        metin = temizlik.metin_temizle(hucre, tavan=120)
        if not metin:
            continue
        if fold(metin) == fold(WORK_NO_HEADER):
            harita[sira] = "work_id"
            continue
        anahtar = match_header(metin)
        if anahtar == "translator":
            # §8.5-5: çevirmen dışarıdan doldurulmaz. Sütun varsa YOK SAYILIR.
            atlanan.append(metin)
            continue
        if anahtar in EXPORT_KEYS:
            harita[sira] = anahtar
        else:
            atlanan.append(metin)
    return harita, atlanan


def _satirdan_oneri(degerler: dict[str, Any], *, isbn13: str) -> KunyeOnerisi:
    tavan = ALAN_TAVANLARI
    return KunyeOnerisi(
        isbn=isbn13,
        title=temizlik.metin_temizle(degerler.get("title"), tavan=tavan["title"]),
        authors=temizlik.metin_temizle(degerler.get("authors"), tavan=tavan["authors"]),
        publisher=temizlik.metin_temizle(degerler.get("publisher"), tavan=tavan["publisher"]),
        edition=temizlik.metin_temizle(degerler.get("edition"), tavan=tavan["edition"]),
        publish_year=gecerli_yil(temizlik.yil_ayikla(degerler.get("publish_year"))),
        subjects=temizlik.metin_temizle(degerler.get("subjects"), tavan=tavan["subjects"]),
        classification_code=temizlik.metin_temizle(
            degerler.get("classification_code"), tavan=tavan["classification_code"]
        ),
        language=temizlik.metin_temizle(degerler.get("language"), tavan=tavan["language"]),
        uyarilar=(UYARI_CEVIRMEN,),
    )


def _eser_eslestir(
    isbn13: str, eser_no: int | None, eserler: dict[str, list[Work]]
) -> tuple[str, Work | None]:
    adaylar = eserler.get(isbn13, [])
    if not adaylar:
        return DURUM_ESER_YOK, None
    if len(adaylar) == 1:
        return DURUM_ESLESTI, adaylar[0]
    # Aynı ISBN'li birden çok eser: "Eser No" varsa kesinleştirir, yoksa
    # kullanıcıya sorulur — program sessizce ilkini seçmez (§6.2).
    if eser_no is not None:
        for aday in adaylar:
            if aday.pk == eser_no:
                return DURUM_ESLESTI, aday
    return DURUM_COKLU_ESER, None


def _veri_satirlari(satirlar: list[list[Any]]) -> Iterator[tuple[int, list[Any]]]:
    for sira, satir in enumerate(satirlar[1:], start=2):
        if sira > MAX_SATIR + 1:
            return
        if any(str(hucre or "").strip() for hucre in satir):
            yield sira, satir


def onizleme(icerik: bytes, *, bugun: date | None = None) -> CevrimdisiOnizleme:
    """Doldurulmuş dosyayı okur ve künye ÖNERİLERİ üretir — hiçbir şey yazmaz.

    Eşleşme ISBN üzerindendir (§8.5); aynı ISBN'li birden çok eser varsa satır
    "birden çok eser" durumuna düşer ve kullanıcıya sorulur.
    """
    satirlar = _satirlari_coz(icerik)
    if not satirlar:
        raise CevrimdisiDosyaHatasi("Dosyada satır yok.")
    harita, atlanan = _baslik_haritasi(satirlar[0])
    if "isbn" not in harita.values():
        raise CevrimdisiDosyaHatasi(
            "Dosyada “ISBN” sütunu bulunamadı. Programın verdiği dosyayı doldurup yükleyin."
        )

    ham_satirlar = list(_veri_satirlari(satirlar))
    numaralar = {
        isbn13
        for _sira, satir in ham_satirlar
        for sutun, anahtar in harita.items()
        if anahtar == "isbn"
        and sutun < len(satir)
        and (isbn13 := isbn_module.to_isbn13(satir[sutun]))
    }
    eserler: dict[str, list[Work]] = {}
    for eser in selectors.works_by_isbn13(numaralar):
        eserler.setdefault(eser.isbn13, []).append(eser)

    sonuc: list[CevrimdisiSatir] = []
    for sira, satir in ham_satirlar:
        degerler = {
            anahtar: satir[sutun] for sutun, anahtar in harita.items() if sutun < len(satir)
        }
        isbn13 = isbn_module.to_isbn13(degerler.get("isbn"))
        eser_no = temizlik.sayi_ayikla(degerler.get("work_id"))
        if not isbn13:
            sonuc.append(
                CevrimdisiSatir(
                    satir_no=sira,
                    isbn13="",
                    durum=DURUM_ISBN_YOK,
                    work=None,
                    oneri=_satirdan_oneri(degerler, isbn13=""),
                )
            )
            continue
        durum, eslesen = _eser_eslestir(isbn13, eser_no, eserler)
        sonuc.append(
            CevrimdisiSatir(
                satir_no=sira,
                isbn13=isbn13,
                durum=durum,
                work=eslesen,
                oneri=_satirdan_oneri(degerler, isbn13=isbn13),
            )
        )
    return CevrimdisiOnizleme(
        satirlar=tuple(sonuc),
        tarih=bugun or timezone.localdate(),
        atlanan_sutunlar=tuple(dict.fromkeys(atlanan)),
    )
