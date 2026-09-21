"""Katalog Excel şablonu (tasarım §8.1) — şablonu openpyxl ile açıp sözlükle karşılaştırır.

Asıl sabitlenen gidiş-dönüştür: "Katalog" sayfasının başlıkları sözlükle birebir
aynıdır ve her biri `match_header` ile kendi anahtarına döner — F3 içe aktarımı
programın kendi verdiği dosyayı hiç dokunulmadan okuyabilmelidir. "Örnek"
sayfası sözlüğün kurallarına uyar (örnek, kuralı çiğneyen bir satırı öğretmesin).
"""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

import pytest
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from apps.kutuphane import excel_template
from apps.kutuphane.import_schema import (
    CATALOG_SHEET,
    COLUMNS,
    COLUMNS_SHEET,
    DEFAULT_COPIES,
    EXAMPLE_SHEET,
    MAX_COPIES_PER_ROW,
    ColumnKind,
    defaults_to_reference,
    match_header,
    resource_type_value,
    yes_no_value,
)

_SON_SATIR = 1_048_576


@pytest.fixture(scope="module")
def kitap() -> Workbook:
    return load_workbook(BytesIO(excel_template.build_catalog_template()))


def _harf(sira: int) -> str:
    return str(get_column_letter(sira))


def _baslik_satiri(ws: Worksheet) -> list[Any]:
    return [hucre.value for hucre in ws[1]]


def _ornek_kayitlari(kitap: Workbook) -> list[dict[str | None, Any]]:
    ws = kitap[EXAMPLE_SHEET]
    anahtarlar = [match_header(baslik) for baslik in _baslik_satiri(ws)]
    assert None not in anahtarlar
    return [
        dict(zip(anahtarlar, satir, strict=True))
        for satir in ws.iter_rows(min_row=2, values_only=True)
    ]


def test_sayfa_adlari_ve_sirasi(kitap: Workbook) -> None:
    assert kitap.sheetnames == [CATALOG_SHEET, COLUMNS_SHEET, EXAMPLE_SHEET]
    assert (CATALOG_SHEET, COLUMNS_SHEET, EXAMPLE_SHEET) == ("Katalog", "Sütunlar", "Örnek")
    # Dosya açılınca doldurulacak sayfa görünür.
    assert kitap.active is not None
    assert kitap.active.title == CATALOG_SHEET


def test_katalog_basliklari_sozlukle_birebir_ve_geri_eslenir(kitap: Workbook) -> None:
    basliklar = _baslik_satiri(kitap[CATALOG_SHEET])
    assert basliklar == [kolon.header for kolon in COLUMNS]
    assert [match_header(b) for b in basliklar] == [kolon.key for kolon in COLUMNS]


def test_katalog_sayfasinda_basliktan_baska_satir_yok(kitap: Workbook) -> None:
    """İçe aktarım yalnız Katalog sayfasını okur: not ya da örnek satır kitap sanılırdı."""
    assert kitap[CATALOG_SHEET].max_row == 1


def test_baslik_satiri_dondurulmus(kitap: Workbook) -> None:
    assert kitap[CATALOG_SHEET].freeze_panes == "A2"
    assert kitap[COLUMNS_SHEET].freeze_panes == "A2"


def test_zorunlu_sutun_gorsel_olarak_ayrisir_ve_notu_zorunlu_der(kitap: Workbook) -> None:
    ws = kitap[CATALOG_SHEET]
    zorunlu_renkler = set()
    istege_bagli_renkler = set()
    for sira, kolon in enumerate(COLUMNS, start=1):
        hucre = ws.cell(row=1, column=sira)
        assert hucre.font.bold
        assert hucre.comment is not None
        assert kolon.description in hucre.comment.text
        renk = hucre.fill.fgColor.rgb
        if kolon.required:
            zorunlu_renkler.add(renk)
            assert hucre.comment.text.startswith("Zorunlu sütun.")
        else:
            istege_bagli_renkler.add(renk)
            assert not hucre.comment.text.startswith("Zorunlu")
    assert len(zorunlu_renkler) == 1
    assert zorunlu_renkler.isdisjoint(istege_bagli_renkler)


def test_sutun_genislikleri_basligi_sigdirir(kitap: Workbook) -> None:
    ws = kitap[CATALOG_SHEET]
    for sira, kolon in enumerate(COLUMNS, start=1):
        genislik = ws.column_dimensions[_harf(sira)].width
        assert genislik >= len(kolon.header) + 2, kolon.key


def test_evet_hayir_kaynak_turu_ve_nusha_sutunlarinda_veri_dogrulama(kitap: Workbook) -> None:
    ws = kitap[CATALOG_SHEET]
    dogrulamalar = {str(dv.sqref): dv for dv in ws.data_validations.dataValidation}
    beklenen_sayi = 0
    for sira, kolon in enumerate(COLUMNS, start=1):
        aralik = f"{_harf(sira)}2:{_harf(sira)}{_SON_SATIR}"
        if kolon.kind is ColumnKind.YES_NO:
            dv = dogrulamalar[aralik]
            assert (dv.type, dv.formula1) == ("list", '"Evet,Hayır"')
        elif kolon.kind is ColumnKind.CHOICE:
            dv = dogrulamalar[aralik]
            assert dv.type == "list"
            assert dv.formula1 == '"' + ",".join(kolon.choices) + '"'
        elif kolon.kind is ColumnKind.INTEGER:
            dv = dogrulamalar[aralik]
            assert (dv.type, dv.operator) == ("whole", "between")
            assert (dv.formula1, dv.formula2) == ("1", str(MAX_COPIES_PER_ROW))
        else:
            assert aralik not in dogrulamalar
            continue
        beklenen_sayi += 1
        assert dv.allow_blank
        assert dv.showErrorMessage
        assert kolon.header in (dv.error or "")
    assert len(dogrulamalar) == beklenen_sayi == 4


def test_metin_sutunlari_metin_bicimindedir(kitap: Workbook) -> None:
    """ISBN bilimsel gösterime, "001.4" ya da "894.353" sayıya dönmesin."""
    ws = kitap[CATALOG_SHEET]
    for sira, kolon in enumerate(COLUMNS, start=1):
        bicim = ws.column_dimensions[_harf(sira)].number_format
        if kolon.key in ("isbn", "classification_code", "old_register_no"):
            assert bicim == "@", kolon.key
        if kolon.kind in (ColumnKind.INTEGER, ColumnKind.YEAR, ColumnKind.YES_NO):
            assert bicim != "@", kolon.key


def test_sutunlar_sayfasi_sozlugu_aynen_tasir(kitap: Workbook) -> None:
    ws = kitap[COLUMNS_SHEET]
    assert _baslik_satiri(ws) == [
        "Sütun",
        "Zorunlu mu?",
        "Ne yazılır?",
        "Kabul edilen değerler",
        "Boş bırakılırsa",
        "Örnek",
    ]
    satirlar = list(ws.iter_rows(min_row=2, max_row=len(COLUMNS) + 1, values_only=True))
    assert satirlar == [
        (
            kolon.header,
            "Evet" if kolon.required else "Hayır",
            kolon.description,
            kolon.accepted_values,
            kolon.blank_rule,
            kolon.example or "—",
        )
        for kolon in COLUMNS
    ]


def test_sutunlar_sayfasi_notlari_kisisel_veri_ve_elli_nusha_kuralini_soyler(
    kitap: Workbook,
) -> None:
    ws = kitap[COLUMNS_SHEET]
    metin = " ".join(
        str(deger)
        for satir in ws.iter_rows(min_row=len(COLUMNS) + 2, values_only=True)
        for deger in satir
        if deger
    )
    assert "Notlar" in metin
    for not_ in excel_template.DICTIONARY_NOTES:
        assert not_ in metin
    assert "kişisel veri yazmayın" in metin
    assert f"en çok {MAX_COPIES_PER_ROW} nüsha" in metin
    assert "“Katalog” sayfasını okur" in metin


def test_ornek_sayfasi_katalogla_ayni_basliklari_tasir(kitap: Workbook) -> None:
    assert _baslik_satiri(kitap[EXAMPLE_SHEET]) == _baslik_satiri(kitap[CATALOG_SHEET])


def test_ornek_satirlar_sozlugun_kurallarina_uyar(kitap: Workbook) -> None:
    kayitlar = _ornek_kayitlari(kitap)
    assert 2 <= len(kayitlar) <= 5
    for kayit in kayitlar:
        assert kayit["title"]
        # ISBN uydurulmaz.
        assert kayit["isbn"] is None
        # Yayınevi açıkça örnektir (gerçek bir baskıyı taklit etmez).
        assert kayit["publisher"] == "Örnek Yayınevi"
        yes_no_value(kayit["is_bound_periodical"])
        yes_no_value(kayit["is_reference"])
        resource_type_value(kayit["resource_type"])
        nusha = kayit["copies"] or DEFAULT_COPIES
        assert 1 <= nusha <= MAX_COPIES_PER_ROW
        if kayit["old_register_no"]:
            assert nusha == 1, "eski kayıt no tek nüshaya aittir"


def test_ornek_nusha_yazmanin_iki_yolunu_ve_danismayi_gosterir(kitap: Workbook) -> None:
    kayitlar = _ornek_kayitlari(kitap)
    # 1) Toplu satır: birden çok nüsha, eski kayıt no yok.
    assert any((k["copies"] or 1) > 1 and not k["old_register_no"] for k in kayitlar)
    # 2) Aynı eserin eski kayıt numaralı nüshaları ayrı satırlarda.
    eski_nolu = [k for k in kayitlar if k["old_register_no"]]
    assert len(eski_nolu) >= 2
    assert len({k["title"] for k in eski_nolu}) == 1
    assert len({k["old_register_no"] for k in eski_nolu}) == len(eski_nolu)
    # 3) Danışma kaynağı.
    assert any(yes_no_value(k["is_reference"]) for k in kayitlar)
    # Ders kitabı örneği yok: danışma burada açıkça işaretlenmiştir.
    assert not any(defaults_to_reference(k["resource_type"], k["subjects"]) for k in kayitlar)


def test_sozluk_ornekleri_ornek_sayfasindan_gelir(kitap: Workbook) -> None:
    """Sütunlar sayfasındaki örnek değer, Örnek sayfasında o sütunda gerçekten geçer."""
    kayitlar = _ornek_kayitlari(kitap)
    for kolon in COLUMNS:
        if not kolon.example or kolon.kind is ColumnKind.YES_NO:
            continue
        degerler = {str(k[kolon.key]) for k in kayitlar if k[kolon.key] is not None}
        assert kolon.example in degerler, kolon.key


def test_ornek_sayfasinda_kod_ve_numaralar_metin_olarak_durur(kitap: Workbook) -> None:
    kayitlar = _ornek_kayitlari(kitap)
    assert {k["classification_code"] for k in kayitlar} >= {"894.353", "843.4"}
    assert all(isinstance(k["old_register_no"], str) for k in kayitlar if k["old_register_no"])


def test_dosya_adi_belge_adi_ve_tarihtir() -> None:
    assert (
        excel_template.template_filename(date(2026, 9, 21))
        == "Katalog-Excel-Şablonu_21.09.2026.xlsx"
    )


def test_dosya_adindaki_tarih_yerel_tarihtir(monkeypatch: pytest.MonkeyPatch) -> None:
    """UTC'de 20 Eylül 21:30 = İstanbul'da 21 Eylül 00:30 → dosya 21.09.2026 taşır."""
    monkeypatch.setattr(timezone, "now", lambda: datetime(2026, 9, 20, 21, 30, tzinfo=UTC))
    assert excel_template.template_filename() == "Katalog-Excel-Şablonu_21.09.2026.xlsx"


def test_calisma_kitabi_ozellikleri(kitap: Workbook) -> None:
    assert kitap.properties.title == excel_template.TEMPLATE_DOCUMENT_NAME
    assert kitap.properties.creator == "Kütüphane Defteri"
