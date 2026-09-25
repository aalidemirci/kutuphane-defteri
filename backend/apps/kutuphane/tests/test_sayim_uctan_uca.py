"""Sayım — uçtan uca kod kapısı (F9; tasarım §14.1 F9, §9-10, §9-11, §10 E10; D4, D17, D18).

Tek bir sentetik sayım kapının maddelerinden SIRAYLA geçer (öteki testler maddeleri tek
tek sınar; bu test aynı sayımda birbirlerini bozmadıklarını gösterir):

1. Taslakta iki AYRI seçenek (TMY 32/3 durdurması + sayım için hizmet arası) ve kurulun
   seçimi (ödünç kayda göre; sınıf kitaplığı yerinde; öğretmene teslim kayda göre) →
   başlat: anlık görüntü.
2. Sürerken: edinim ve yeni nüsha reddedilir (32/3), programa aktarım TMY'siz iletiyle
   reddedilir (K4), kayıp bildirimi reddedilir (D4), yeni ödünç ve yeni teslim reddedilir
   (hizmet arası — madde 27); iade GEÇER ve o turda "bulundu" sayılır (Md. 23/1-c); anlık
   görüntü iadeden etkilenmez.
3. Okutma (rafta, hasar önerili, yerinde sayılan sınıf kitaplığı), üye kartı kaydedilmez,
   etiketsiz kitap sayım fazlası olarak yazılır.
4. Tamamla: ilk turda noksan → ikinci sayım (32/6) → "Tamamlandı"; kilitler sürer (D17).
5. Onaydan önce bir noksan kalem kütüphaneye döner (yerinde sayılıp bulunamayan sınıf
   kitaplığı nüshası teslimden geri alınır — geri alma da iade gibi hiç kilitlenmez).
   Hizmet arası onaya dek yeni ödüncü ve teslimi durdurduğu için noksan bir nüshanın ödüncü
   ya da teslimi o arada açılıp kapanamaz; onay anında kütüphaneye dönüş yolları teslimden
   geri alma ve kayıp dosyasında bulunmadır (bulunma 32/3 durdurmasının kapsamında
   değildir — F9 ekleri K1).
6. Onay: dönen kalem DÜŞÜLMEZ (D17); noksan 32/7 ile, kayıp önerili nüsha "Kayıp (kayıttan
   düşüldü)" diye 32/7 ile, hasar önerisi 27/1 + 10/1-e ile komisyonsuz düşülür; sayım
   fazlası TMY 17 ile kayda girer (kilitler önce kalktığı için edinim kapıdan geçer).
7. Kilitler kalkar: edinim ve ödünç açılır.
8. E10 (PDF + XLSX) ve eki "Taşınır Sayım ve Döküm Cetveline aktarılacak sayılar" basılır
   (gelecek yıla devirden onayda 27/1 ile düşülen çıkarılır — madde 25 a);
   ödünç alanın, kayıp/hasar dosyası kişisinin ve teslim alan öğretmenin kimliği HİÇBİR
   çıktıda — belgede ve sayım yanıtlarında — yoktur (TMY 10/1-g, 32/8).

Bütün kişi ve kurum adları UYDURMADIR ("Deneme …", "Örnek …"; CLAUDE.md §2-12).
"""

from __future__ import annotations

import io
import re
from typing import Any

import pytest
from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from apps.kutuphane import sayim_belgeleri as belgeler
from apps.kutuphane import selectors_sayim
from apps.kutuphane.models import (
    AcquisitionMethod,
    CaseResolution,
    CopyStatus,
    CountBasis,
    StockTakeFoundVia,
    StockTakeOutcome,
    StockTakeResult,
    StockTakeStatus,
    StockTakeWriteOffPath,
)
from apps.kutuphane.services import (
    catalog,
    circulation,
    deliveries,
    loss_damage,
    stocktake,
    tmy_kapisi,
)
from apps.kutuphane.services.circulation import DolasimReddi
from apps.kutuphane.tests.dolasim_ortak import odunc_nushasi, odunc_ver, ogrenci, uye
from apps.kutuphane.tests.ortak import edinim, eser
from apps.kutuphane.tests.sayim_ortak import (
    HARCAMA_YETKILISI,
    KURUL,
    durdurma_alanlari,
    kalem,
    okut,
    taslak,
    tazele_sayim,
)
from apps.kutuphane.tests.teslim_ortak import ogretmen, sube, tazele_dosya, tazele_nusha, teslim_et
from apps.okul.models import SchoolConfig

pytestmark = pytest.mark.django_db

#: Sentetik kişiler — tutanağın ve sayım yanıtlarının hiçbir yerinde görünmemeli.
YASAK_ADLAR = (
    "Denemeuctan",
    "Kayitgoreoglu",
    "Denemeiadeli",
    "Donusoglu",
    "Denemekayipli",
    "Bildirimoglu",
    "Denemeogretmenli",
    "Teslimalanoglu",
)
YASAK_OKUL_NOLARI = ("918301", "918302", "918303")
HIZMET_ARASI_KARARI = "Okul müdürlüğünün 2026/31 sayılı kararı"

_T_ARALIGI = re.compile(r"T (?=[a-zçğıöşü])")


def _pdf_metni(icerik: bytes) -> str:
    ham = "\n".join(s.extract_text() or "" for s in PdfReader(io.BytesIO(icerik)).pages)
    return " ".join(_T_ARALIGI.sub("T", ham).split())


def _xlsx_metni(icerik: bytes) -> tuple[list[str], str]:
    kitap = load_workbook(io.BytesIO(icerik))
    parcalar: list[str] = []
    for sayfa in kitap.worksheets:
        for satir in sayfa.iter_rows(values_only=True):
            parcalar.extend(str(h) for h in satir if h is not None)
    return kitap.sheetnames, " ".join(parcalar)


def _durdurma_reddi(exc: pytest.ExceptionInfo[ValidationError]) -> bool:
    return exc.value.code == tmy_kapisi.RED_KODU and "TMY 32/3 durdurması" in str(exc.value)


def _hucreler(baglam: dict[str, Any], baslik_parcasi: str) -> list[list[str]]:
    tablolar = baglam["tables"]
    for tablo in tablolar:
        if baslik_parcasi in tablo["title"]:
            return [[h["v"] for h in satir] for satir in tablo["rows"]]
    raise AssertionError(f"{baslik_parcasi} tablosu yok: {[t['title'] for t in tablolar]}")


@pytest.fixture(autouse=True)
def okul() -> SchoolConfig:
    config = SchoolConfig.load()
    config.school_name = "Örnek Anadolu Lisesi"
    config.district = "Örnek İlçe"
    config.principal_name = "Örnek Müdür"
    config.save()
    return config


def test_sayim_uctan_uca_durdurma_hizmet_arasi_iade_onay_ve_tutanak() -> None:
    # ------------------------------------------------------------ kurgu (sayımdan önce)
    kayda_gore_uye = uye(
        ogrenci(first_name="Denemeuctan", last_name="Kayitgoreoglu", student_number="918301")
    )
    iadeli_uye = uye(
        ogrenci(first_name="Denemeiadeli", last_name="Donusoglu", student_number="918302")
    )
    kayipli_uye = uye(
        ogrenci(first_name="Denemekayipli", last_name="Bildirimoglu", student_number="918303")
    )
    kart_nolari = [u.card_no for u in (kayda_gore_uye, iadeli_uye, kayipli_uye)]

    rafta = odunc_nushasi(title="Uçtan Uca Rafta")
    ikinci_turda = odunc_nushasi(title="Uçtan Uca İkinci Sayımda Bulunan")
    noksan = odunc_nushasi(title="Uçtan Uca Noksan")
    oduncte = odunc_ver(kayda_gore_uye, odunc_nushasi(title="Uçtan Uca Ödünçte"))
    iade_edilecek = odunc_ver(iadeli_uye, odunc_nushasi(title="Uçtan Uca İade Edilen"))

    # Kayıp önerisi (F8 ekleri 34: 32/7 noksan düşüm teklifine bağlanır) — ödünçteyken kayıp.
    kayip_odunc = odunc_ver(kayipli_uye, odunc_nushasi(title="Uçtan Uca Kayıp Önerili"))
    kayip_dosyasi = loss_damage.report_lost(
        copy=kayip_odunc.copy, responsible_note="Denemekayipli Bildirimoglu bildirdi"
    )
    loss_damage.resolve_case(kayip_dosyasi, resolution=CaseResolution.WRITE_OFF_PROPOSED)
    # Hasar önerisi (F8 ekleri 34: 27/1 + 10/1-e, komisyonsuz) — iadede hasarlı çıkmış.
    hasarli = odunc_nushasi(title="Uçtan Uca Hasar Önerili")
    hasar_dosyasi = loss_damage.open_damage_case(
        copy=hasarli, membership=iadeli_uye, responsible_note="Denemeiadeli Donusoglu"
    )
    loss_damage.resolve_case(hasar_dosyasi, resolution=CaseResolution.WRITE_OFF_PROPOSED)

    sinif_yok, sinif_var = teslim_et(
        [odunc_nushasi(title="Uçtan Uca Sınıfta Yok"), odunc_nushasi(title="Uçtan Uca Sınıfta")],
        section=sube(10, "B"),
    ).deliveries
    ogretmende = teslim_et(
        [odunc_nushasi(title="Uçtan Uca Öğretmende")],
        personnel=ogretmen(first_name="Denemeogretmenli", last_name="Teslimalanoglu"),
    ).deliveries[0]
    fazla_eseri = eser(title="Uçtan Uca Fazla Eseri")
    mevcut_eser = eser(title="Uçtan Uca Durdurmada Eser")
    mevcut_edinim = edinim(method=AcquisitionMethod.MINISTRY)  # gerçek taşınır girişi
    aktarim_edinimi = edinim()  # "Mevcut koleksiyon (programa aktarım)"
    teslim_edilmeyecek = odunc_nushasi(title="Uçtan Uca Teslim Edilmeyen")

    # ------------------------------------------------------------ 1. taslak ve başlatma
    sayim = taslak(
        **durdurma_alanlari(),
        service_pause=True,
        service_pause_decision=HIZMET_ARASI_KARARI,
        loan_basis=CountBasis.BY_RECORD,
        section_delivery_basis=CountBasis.IN_PLACE,
        teacher_delivery_basis=CountBasis.BY_RECORD,
    )
    assert selectors_sayim.tmy_stop_active() is False  # taslakta kilit yok
    sayim = stocktake.start_stocktake(sayim)
    assert sayim.status == StockTakeStatus.IN_PROGRESS
    assert selectors_sayim.tmy_stop_active() and selectors_sayim.service_pause_active()
    assert kalem(sayim, kayip_odunc.copy).expected_status == CopyStatus.LOST
    assert kalem(sayim, sinif_yok.copy).basis == CountBasis.IN_PLACE
    assert kalem(sayim, ogretmende.copy).basis == CountBasis.BY_RECORD

    # ------------------------------------------------------------ 2. sürerken kapılar
    with pytest.raises(ValidationError) as exc:
        edinim(method=AcquisitionMethod.MINISTRY)
    assert _durdurma_reddi(exc) and "edinim ve yeni nüsha kaydı yapılamaz" in str(exc.value)
    with pytest.raises(ValidationError) as exc:
        catalog.create_copy(work=mevcut_eser, acquisition=mevcut_edinim)
    assert _durdurma_reddi(exc)
    # K4 (b): programa aktarım kapalı ama iletisi TMY'ye dayanmaz.
    with pytest.raises(ValidationError) as exc:
        catalog.create_copy(work=mevcut_eser, acquisition=aktarim_edinimi)
    assert exc.value.code == tmy_kapisi.RED_KODU_AKTARIM
    assert "TMY" not in str(exc.value) and "programa aktarım" in str(exc.value)
    with pytest.raises(ValidationError) as exc:
        loss_damage.report_lost(copy=noksan)  # D4: kayıp bildirimi de kapsamda
    assert _durdurma_reddi(exc)
    with pytest.raises(DolasimReddi) as odunc_reddi:
        odunc_ver(iadeli_uye, rafta)
    assert odunc_reddi.value.code == circulation.RED_HIZMET_ARASI
    assert odunc_reddi.value.message == circulation.SERVICE_PAUSE_MESSAGE
    assert tazele_nusha(rafta).status == CopyStatus.AVAILABLE
    # Madde 27: hizmet arası yeni teslimi de durdurur (aynı kod ve ileti).
    with pytest.raises(DolasimReddi) as teslim_reddi:
        teslim_et([teslim_edilmeyecek], section=sube(10, "B"))
    assert teslim_reddi.value.code == circulation.RED_HIZMET_ARASI
    assert tazele_nusha(teslim_edilmeyecek).status == CopyStatus.AVAILABLE

    # İade GEÇER (Md. 23/1-c) — ne durdurma ne hizmet arası iadeyi durdurur.
    circulation.return_copy(copy=iade_edilecek.copy)
    assert tazele_nusha(iade_edilecek.copy).status == CopyStatus.AVAILABLE
    # Anlık görüntü sayım sırasındaki değişiklikten etkilenmez.
    assert kalem(sayim, iade_edilecek.copy).expected_status == CopyStatus.ON_LOAN

    # ------------------------------------------------------------ 3. okutma
    sonuclar = okut(sayim, rafta, hasarli, sinif_var.copy, teslim_edilmeyecek)
    assert [s.code for s in sonuclar] == [stocktake.OKUTMA_BULUNDU] * 4
    (kart,) = stocktake.scan_many(sayim, [kayda_gore_uye.card_no])
    assert (kart.code, kart.barcode, kart.item) == (stocktake.OKUTMA_GECERSIZ, "", None)
    fazla = stocktake.add_surplus(sayim, note="Etiketsiz kitap, danışma rafı", work=fazla_eseri)

    ilerleme = selectors_sayim.progress(sayim)
    assert ilerleme["class_libraries"] == [
        {"class_section": sinif_var.section_id, "label": "10/B", "expected": 2, "found": 1}
    ]

    # ------------------------------------------------------------ 4. tamamla (32/6)
    ilk = stocktake.complete_stocktake(sayim)
    # Noksan: kayıtta rafta olup bulunamayan iki nüsha, kayıptaki nüsha, yerinde sayılıp
    # bulunamayan sınıf kitaplığı nüshası. İade edilen "bulundu"dur, ödünçteki ve
    # öğretmendeki kayda göre alınır.
    assert (ilk.second_round, ilk.missing) == (True, 4)
    assert okut(sayim, ikinci_turda)[0].message == stocktake.SCAN_FOUND_ROUND2
    son = stocktake.complete_stocktake(sayim)
    assert (son.second_round, son.missing) == (False, 3)
    sayim = tazele_sayim(sayim)
    assert sayim.status == StockTakeStatus.COMPLETED

    iade_kalemi = kalem(sayim, iade_edilecek.copy)
    assert (iade_kalemi.result, iade_kalemi.found_via) == (
        StockTakeResult.FOUND,
        StockTakeFoundVia.RETURN,
    )
    assert kalem(sayim, ikinci_turda).found_in_round == 2
    for kayda_gore in (oduncte.copy, ogretmende.copy):
        assert kalem(sayim, kayda_gore).result == StockTakeResult.BY_RECORD
    for bulunamayan in (noksan, kayip_odunc.copy, sinif_yok.copy):
        assert kalem(sayim, bulunamayan).result == StockTakeResult.MISSING
    assert kalem(sayim, kayip_odunc.copy).case_id == kayip_dosyasi.pk
    assert kalem(sayim, hasarli).damage_write_off is True

    # D17: "Tamamlandı" ile "Onaylandı" arasında kilitler sürer; iade yine açık.
    with pytest.raises(ValidationError) as exc:
        edinim(method=AcquisitionMethod.MINISTRY)
    assert _durdurma_reddi(exc)
    with pytest.raises(DolasimReddi):
        odunc_ver(iadeli_uye, rafta)

    # ------------------------------------------------------------ 5. onaydan önce dönen noksan
    assert deliveries.take_back(sinif_yok.copy) is not None
    assert tazele_nusha(sinif_yok.copy).status == CopyStatus.AVAILABLE

    # ------------------------------------------------------------ 6. onay (D17, 32/7, 27/1, 17)
    sonuc = stocktake.approve_stocktake(
        sayim, approved_by_name=HARCAMA_YETKILISI, approved_on=timezone.localdate()
    )
    assert (
        sonuc.written_off,
        sonuc.damage_written_off,
        sonuc.state_changed,
        sonuc.not_approved,
        sonuc.reconciled,
        sonuc.surplus_entered,
    ) == (2, 1, 1, 0, 0, 1)

    donen = kalem(sayim, sinif_yok.copy)
    assert (donen.outcome, donen.result, donen.found_via) == (
        StockTakeOutcome.STATE_CHANGED,
        StockTakeResult.FOUND,
        StockTakeFoundVia.RETURN,
    )
    assert donen.outcome_note == "Onayda durumu: Rafta."
    assert tazele_nusha(sinif_yok.copy).status == CopyStatus.AVAILABLE  # düşülmedi

    for dusulen, hedef in (
        (noksan, CopyStatus.WITHDRAWN_MISSING),
        (kayip_odunc.copy, CopyStatus.WITHDRAWN_LOST),
    ):
        k = kalem(sayim, dusulen)
        assert (k.outcome, k.write_off_path) == (
            StockTakeOutcome.WRITTEN_OFF,
            StockTakeWriteOffPath.MISSING_32_7,
        )
        assert tazele_nusha(dusulen).status == hedef
    hasar_kalemi = kalem(sayim, hasarli)
    assert (hasar_kalemi.outcome, hasar_kalemi.write_off_path) == (
        StockTakeOutcome.WRITTEN_OFF,
        StockTakeWriteOffPath.DAMAGE_27_1,
    )
    assert tazele_nusha(hasarli).status == CopyStatus.WITHDRAWN_DAMAGED
    assert loss_damage.oneri_geri_alinabilir(tazele_dosya(kayip_dosyasi)) is False
    # Bulunan, ödünçteki ve öğretmendeki nüsha kayıtta kalır.
    assert tazele_nusha(rafta).status == CopyStatus.AVAILABLE
    assert tazele_nusha(oduncte.copy).status == CopyStatus.ON_LOAN
    assert tazele_nusha(ogretmende.copy).status == CopyStatus.DELIVERED

    fazla.refresh_from_db()
    assert fazla.outcome == StockTakeOutcome.ENTERED and fazla.created_copy is not None
    assert fazla.created_copy.acquisition.method == AcquisitionMethod.INVENTORY_FOUND

    # ------------------------------------------------------------ 7. kilitler kalkar
    sayim = tazele_sayim(sayim)
    assert sayim.status == StockTakeStatus.APPROVED
    assert not selectors_sayim.tmy_stop_active() and not selectors_sayim.service_pause_active()
    assert selectors_sayim.locking_stocktake() is None
    edinim()
    catalog.create_copy(work=mevcut_eser, acquisition=mevcut_edinim)
    yeni_odunc = odunc_ver(uye(), rafta)
    assert tazele_nusha(yeni_odunc.copy).status == CopyStatus.ON_LOAN

    # ------------------------------------------------------------ 8. E10 ve eki
    baglam = belgeler.stocktake_report_context(sayim)
    secenekler = _hucreler(baglam, "SAYIM SIRASINDAKİ SEÇENEKLER")
    assert [s[0] for s in secenekler] == ["TMY 32/3 durdurması", "Sayım için hizmet arası", "İade"]
    assert [s[1] for s in secenekler] == ["Seçildi", "Seçildi", "Her zaman açık"]
    assert HIZMET_ARASI_KARARI in secenekler[1][2]
    noksan_satirlari = _hucreler(baglam, "TMY MD. 32/7")
    assert {s[2] for s in noksan_satirlari} == {"Uçtan Uca Noksan", "Uçtan Uca Kayıp Önerili"}
    (hasar_satiri,) = _hucreler(baglam, "HASAR NEDENİYLE KAYITTAN DÜŞME TEKLİFİ (TMY MD. 27/1)")
    assert hasar_satiri[2] == "Uçtan Uca Hasar Önerili"
    (degisen_satir,) = _hucreler(baglam, "ONAYDA DURUMU DEĞİŞEN")
    assert degisen_satir[2] == "Uçtan Uca Sınıfta Yok"
    assert degisen_satir[4] == "Onayda durumu: Rafta."
    kurul_secimi = {s[0]: s for s in _hucreler(baglam, "SAYIM KURULUNUN SEÇİMİ")}
    assert "32/5'e kıyasen; 23/4" in kurul_secimi["Ödünçteki nüsha"][2]
    assert kurul_secimi["Sınıf kitaplığına teslim edilen nüsha"][1] == "Yerinde sayılır"
    assert baglam["approval"]["name"] == HARCAMA_YETKILISI
    assert baglam["annex"]["title"] == "EK: TAŞINIR SAYIM VE DÖKÜM CETVELİNE AKTARILACAK SAYILAR"
    # Madde 25 (a): gelecek yıla devir = sayımda bulunan − onayda 27/1 ile düşülen (hasarlı).
    sayilar = selectors_sayim.tmy_34_1(sayim)
    assert sayilar["damage_written_off"] == 1
    assert sayilar["next_year_carryover"] == sayilar["found_quantity"] - 1

    pdf = _pdf_metni(belgeler.stocktake_report_pdf(sayim))
    sayfa_adlari, xlsx = _xlsx_metni(belgeler.stocktake_report_xlsx(sayim))
    assert sayfa_adlari == ["Sayım tutanağı", "Kalemler", "Cetvele aktarılacak sayılar"]
    assert "EK: TAŞINIR SAYIM VE DÖKÜM CETVELİNE AKTARILACAK SAYILAR" in pdf
    assert "Bu döküm Taşınır Sayım ve Döküm Cetveli değildir" in pdf
    assert belgeler.EK_ADI in xlsx
    for metin in (pdf, xlsx):
        for kitap in ("Uçtan Uca Noksan", "Uçtan Uca Kayıp Önerili", "Uçtan Uca Hasar Önerili"):
            assert kitap in metin, kitap
    # Kişi adı yalnız kurulda ve harcama yetkilisinde (imza bloğu).
    for ad in (*KURUL.values(), HARCAMA_YETKILISI):
        assert ad in pdf, ad
    # "sayım kilidi" ve "dondurma" denmez (sözlük).
    assert "kilidi" not in pdf.casefold() and "dondur" not in pdf.casefold()

    # Sayım yanıtları (ekranın gördüğü her şey) ve belge uçları.
    istemci = APIClient()
    kok = f"/api/v1/library/stocktakes/{sayim.pk}/"
    yanitlar = [
        istemci.get(kok),
        istemci.get(f"{kok}items/", {"limit": 200}),
        istemci.get(f"{kok}progress/"),
        istemci.get(f"{kok}tmy-34-1/"),
        istemci.get(f"{kok}documents/"),
        istemci.get("/api/v1/library/stocktakes/"),
        istemci.get("/api/v1/library/stocktakes/state/"),
    ]
    assert [y.status_code for y in yanitlar] == [200] * len(yanitlar)
    json_metni = str([y.json() for y in yanitlar])
    pdf_yaniti = istemci.get(f"{kok}documents/sayim-tutanagi/")
    xlsx_yaniti = istemci.get(f"{kok}documents/sayim-tutanagi/?kind=xlsx")
    assert (pdf_yaniti.status_code, xlsx_yaniti.status_code) == (200, 200)
    uc_pdf = _pdf_metni(b"".join(pdf_yaniti.streaming_content))  # type: ignore[attr-defined]
    _, uc_xlsx = _xlsx_metni(b"".join(xlsx_yaniti.streaming_content))  # type: ignore[attr-defined]
    for cikti, metin in (
        ("pdf", pdf),
        ("xlsx", xlsx),
        ("uç pdf", uc_pdf),
        ("uç xlsx", uc_xlsx),
        ("json", json_metni),
    ):
        for yasak in (*YASAK_ADLAR, *YASAK_OKUL_NOLARI, *kart_nolari):
            assert yasak not in metin, f"{cikti}: {yasak}"
    for disposition in (pdf_yaniti["Content-Disposition"], xlsx_yaniti["Content-Disposition"]):
        assert KURUL["committee_chair"] not in disposition
