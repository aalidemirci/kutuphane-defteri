"""Çevrimdışı künye yolu: ISBN listesi dışa aktarılır, doldurulur, geri alınır (§8.5).

Kanıtlananlar: liste yalnız künyesi eksik ve ISBN'i olan eserleri taşır · dosya
kişisel veri taşımaz · sayfa adı "Katalog" DEĞİLDİR (yanlışlıkla içe aktarıma
verilip nüsha açmasın) · çevirmen sütunu yoktur ve gelen dosyada varsa yok
sayılır · dosya güvenilmeyen girdidir (boyut, satır, NFC, denetim karakteri) ·
geri alma hiçbir şey yazmaz.
"""

from __future__ import annotations

import io
import unicodedata

import pytest
from openpyxl import Workbook, load_workbook

from apps.kutuphane.kunye import offline
from apps.kutuphane.models import Work
from apps.kutuphane.tests import ortak

pytestmark = pytest.mark.django_db

ISBN = "9789753638029"
ISBN_IKI = "9789750718533"


def kunye_dosyasi(basliklar: list[str], satirlar: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = offline.SHEET_NAME
    ws.append(basliklar)
    for satir in satirlar:
        ws.append(satir)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def sayfa(icerik: bytes) -> list[list[object]]:
    wb = load_workbook(io.BytesIO(icerik))
    try:
        return [list(satir) for satir in wb[offline.SHEET_NAME].iter_rows(values_only=True)]
    finally:
        wb.close()


# ===================================================== Dışa aktarım
class TestDisaAktarim:
    def test_yalniz_kunyesi_eksik_ve_isbnli_eserler_yazilir(self) -> None:
        eksik = ortak.eser(title="Eksik künye", isbn=ISBN)
        ortak.eser(
            title="Tam künye",
            isbn=ISBN_IKI,
            publisher="Örnek Yayınevi",
            publish_year=2020,
            subjects="Roman",
            classification_code="894.353",
        )
        ortak.eser(title="ISBN'siz eser")

        satirlar = sayfa(offline.build_export())

        assert [satir[0] for satir in satirlar[1:]] == [eksik.pk]

    def test_basliklar_sozlukten_gelir_ve_cevirmen_yoktur(self) -> None:
        ortak.eser(title="Eksik künye", isbn=ISBN)

        basliklar = sayfa(offline.build_export())[0]

        assert basliklar[0] == offline.WORK_NO_HEADER
        assert basliklar[1] == "ISBN"
        assert "Çevirmen" not in basliklar
        assert "Eser Adı" in basliklar

    def test_sayfa_adi_katalog_degildir(self) -> None:
        """Katalog içe aktarımı yalnız "Katalog" sayfasını okur; bu dosya oraya girmez."""
        from apps.kutuphane.import_schema import CATALOG_SHEET

        wb = load_workbook(io.BytesIO(offline.build_export()))
        try:
            assert offline.SHEET_NAME in wb.sheetnames
            assert CATALOG_SHEET not in wb.sheetnames
        finally:
            wb.close()

    def test_bilgi_sayfasi_ayri_cihaz_kuralini_yazar(self) -> None:
        wb = load_workbook(io.BytesIO(offline.build_export()))
        try:
            metin = " ".join(
                str(hucre)
                for satir in wb[offline.INFO_SHEET_NAME].iter_rows(values_only=True)
                for hucre in satir
                if hucre
            )
        finally:
            wb.close()

        assert "BAŞKA bir cihazda" in metin
        assert "11/18" in metin  # mobil modem yasağı
        assert "öğrenci, veli ya da personel bilgisi yazmayın" in metin

    def test_esitle_baslayan_eser_adi_formul_olmaz(self) -> None:
        """Dosya internete bağlı BAŞKA bir cihazda açılır: formül orada çalışırdı.

        Katalog metni kullanıcı onayından geçmeden de dolabiliyor (toplu Excel
        içe aktarımı eser adını doğrudan yazar), yani "=" ile başlayan bir
        başlık katalogda durup bu dosyaya geçebilir.
        """
        ortak.eser(title="=cmd|' /C calc'!A0", isbn=ISBN)

        wb = load_workbook(io.BytesIO(offline.build_export()))
        try:
            hucre = wb[offline.SHEET_NAME].cell(row=2, column=3)
            assert hucre.data_type == "s"
            assert hucre.value == "=cmd|' /C calc'!A0"
        finally:
            wb.close()

    def test_bilgi_sayfasi_yonergeyi_depodaki_adiyla_anar(self) -> None:
        """CLAUDE.md §2-13: atıf `docs/mevzuat/`teki metinden doğrulanır."""
        metin = " ".join(offline.BILGI_SATIRLARI)

        assert "Millî Eğitim Bakanlığı Bilgi ve Sistem Güvenliği Yönergesi" in metin
        assert "Bilişim Kaynakları Kullanım Yönergesi" not in metin

    def test_dosya_adi_yerel_tarihlidir(self) -> None:
        from datetime import date

        assert offline.export_filename(date(2026, 9, 23)) == "ISBN-Künye-Listesi_23.09.2026.xlsx"


# ===================================================== Geri alma
class TestGeriAlma:
    def test_isbnle_eslesir_ve_oneri_uretir(self) -> None:
        eser = ortak.eser(title="Eksik künye", isbn=ISBN)
        icerik = kunye_dosyasi(
            [offline.WORK_NO_HEADER, "ISBN", "Eser Adı", "Yayınevi", "Yayın Yılı"],
            [[eser.pk, ISBN, "Kürk Mantolu Madonna", "Yapı Kredi Yayınları", "2015"]],
        )

        onizleme = offline.onizleme(icerik)

        (satir,) = onizleme.satirlar
        assert satir.durum == offline.DURUM_ESLESTI
        assert satir.work_id == eser.pk
        assert satir.oneri.publisher == "Yapı Kredi Yayınları"
        assert satir.oneri.publish_year == 2015

    def test_hicbir_sey_yazilmaz(self) -> None:
        eser = ortak.eser(title="Eksik künye", isbn=ISBN)
        oncesi = Work.objects.filter(pk=eser.pk).values().first()
        icerik = kunye_dosyasi(
            ["ISBN", "Eser Adı", "Yayınevi"], [[ISBN, "Başka ad", "Başka yayınevi"]]
        )

        offline.onizleme(icerik)

        assert Work.objects.filter(pk=eser.pk).values().first() == oncesi

    def test_cevirmen_sutunu_yok_sayilir(self) -> None:
        """§8.5-5: çevirmen dışarıdan doldurulmaz, kullanıcıya sorulur."""
        ortak.eser(title="Eksik künye", isbn=ISBN)
        icerik = kunye_dosyasi(
            ["ISBN", "Eser Adı", "Çevirmen"], [[ISBN, "Kürk Mantolu Madonna", "Uydurma Çevirmen"]]
        )

        onizleme = offline.onizleme(icerik)

        (satir,) = onizleme.satirlar
        assert "Çevirmen" in onizleme.atlanan_sutunlar
        assert "translator" not in satir.oneri.alan_sozlugu()
        assert "Uydurma Çevirmen" not in str(satir.oneri.alan_sozlugu())

    def test_bilinmeyen_isbn_eser_yok_durumuna_duser(self) -> None:
        icerik = kunye_dosyasi(["ISBN", "Eser Adı"], [[ISBN_IKI, "Katalogda olmayan kitap"]])

        onizleme = offline.onizleme(icerik)

        assert onizleme.satirlar[0].durum == offline.DURUM_ESER_YOK
        assert onizleme.sayilar[offline.DURUM_ESER_YOK] == 1

    def test_ayni_isbnli_iki_eserde_kullaniciya_sorulur(self) -> None:
        ortak.eser(title="Birinci baskı", isbn=ISBN)
        ortak.eser(title="İkinci baskı", isbn=ISBN)
        icerik = kunye_dosyasi(["ISBN", "Eser Adı"], [[ISBN, "Kürk Mantolu Madonna"]])

        onizleme = offline.onizleme(icerik)

        assert onizleme.satirlar[0].durum == offline.DURUM_COKLU_ESER
        assert onizleme.satirlar[0].work_id is None

    def test_eser_no_cokluyu_cozer(self) -> None:
        ortak.eser(title="Birinci baskı", isbn=ISBN)
        ikinci = ortak.eser(title="İkinci baskı", isbn=ISBN)
        icerik = kunye_dosyasi(
            [offline.WORK_NO_HEADER, "ISBN", "Eser Adı"], [[ikinci.pk, ISBN, "Kürk Mantolu"]]
        )

        onizleme = offline.onizleme(icerik)

        assert onizleme.satirlar[0].durum == offline.DURUM_ESLESTI
        assert onizleme.satirlar[0].work_id == ikinci.pk

    def test_isbnsiz_satir_isaretlenir(self) -> None:
        icerik = kunye_dosyasi(["ISBN", "Eser Adı"], [["", "Numarasız kitap"]])

        onizleme = offline.onizleme(icerik)

        assert onizleme.satirlar[0].durum == offline.DURUM_ISBN_YOK

    def test_bos_satirlar_atlanir(self) -> None:
        icerik = kunye_dosyasi(["ISBN", "Eser Adı"], [[None, None], [ISBN, "Kitap"]])

        onizleme = offline.onizleme(icerik)

        assert len(onizleme.satirlar) == 1


# ===================================================== Güvenilmeyen girdi
class TestGuvenilmeyenDosya:
    def test_gelen_metin_nfce_cevrilir_ve_temizlenir(self) -> None:
        ortak.eser(title="Eksik künye", isbn=ISBN)
        ayrisik = unicodedata.normalize("NFD", "İletişim")
        yon_degistirme = chr(0x202E)  # kaynak koda düz yazılmaz: satırı ters çevirir
        icerik = kunye_dosyasi(
            ["ISBN", "Eser Adı", "Yayınevi"],
            [[ISBN, f"Kürk{yon_degistirme}Mantolu", ayrisik]],
        )

        (satir,) = offline.onizleme(icerik).satirlar

        assert satir.oneri.publisher == "İletişim"
        assert unicodedata.is_normalized("NFC", satir.oneri.publisher)
        assert satir.oneri.title == "Kürk Mantolu"

    def test_isbn_sutunu_yoksa_turkce_hata(self) -> None:
        icerik = kunye_dosyasi(["Eser Adı", "Yayınevi"], [["Kitap", "Yayınevi"]])

        with pytest.raises(offline.CevrimdisiDosyaHatasi, match="ISBN"):
            offline.onizleme(icerik)

    def test_bos_dosya_reddedilir(self) -> None:
        with pytest.raises(offline.CevrimdisiDosyaHatasi):
            offline.onizleme(b"")

    def test_buyuk_dosya_reddedilir(self) -> None:
        with pytest.raises(offline.CevrimdisiDosyaHatasi, match="büyük"):
            offline.onizleme(b"P" * (offline.MAX_DOSYA_BAYT + 1))

    def test_excel_olmayan_govde_turkce_hata_verir(self) -> None:
        with pytest.raises(offline.CevrimdisiDosyaHatasi):
            offline.onizleme(b"PK\x03\x04 bozuk zip")

    def test_csv_dosyasi_da_okunur(self) -> None:
        """İnternetli cihazda Google Sheets kullanan okul CSV geri getirir."""
        ortak.eser(title="Eksik künye", isbn=ISBN)
        icerik = f"ISBN;Eser Adı;Yayınevi\n{ISBN};Kürk Mantolu Madonna;İletişim\n".encode()

        (satir,) = offline.onizleme(icerik).satirlar

        assert satir.durum == offline.DURUM_ESLESTI
        assert satir.oneri.publisher == "İletişim"

    def test_satir_tavani_asilmaz(self) -> None:
        basliklar = "ISBN,Eser Adı"
        satirlar = "\n".join(f"{ISBN},Kitap {i}" for i in range(offline.MAX_SATIR + 50))
        icerik = f"{basliklar}\n{satirlar}\n".encode()

        onizleme = offline.onizleme(icerik)

        assert len(onizleme.satirlar) <= offline.MAX_SATIR
