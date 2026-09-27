"""Alfabetik katalog dökümü (E17) ve Taşınır Kütüphane Defteri dökümü + yönetim hesabı
cetveli hazırlığı (E11) — F10 kod kapısı.

Kapı maddeleri: **ciltletilmemiş süreli yayın E11'e girmez** (TMY 10/1-a-4, 15/4) · E11'in
yönetim hesabı hazırlığı yalnız yıl sonu işaretli ve onaylanmış sayımdan basılır (F9 ekleri
K6) · E17 üç eksende Türk alfabesiyle sıralıdır (Md. 11/1) · her atıf depodaki metinden
doğrulanır · XLSX'te sayılar sayıdır · sayfa bütçesi gerçek uzunlukta veriyle.

Bütün veriler uydurmadır (CLAUDE.md §2-12).
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from django.core.exceptions import ValidationError
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import load_workbook

from apps.kutuphane import katalog_dokumu as kd
from apps.kutuphane import selectors_sayim
from apps.kutuphane import tmy_dokumleri as tmy
from apps.kutuphane.models import (
    AcquisitionMethod,
    Copy,
    CopyStatus,
    ResourceType,
)
from apps.kutuphane.services import catalog
from apps.kutuphane.tests import ortak
from apps.kutuphane.tests.disa_aktarim_ortak import (
    EN_UZUN_ESER,
    UZUN_TKYS,
    UZUN_YAYINEVI,
    UZUN_YAZAR,
    hucre_satirlari,
    madde,
    pdf_metni,
    pdf_sayfalari,
    yatay_tasmalar,
)
from apps.kutuphane.tests.dolasim_ortak import odunc_nushasi
from apps.kutuphane.tests.sayim_ortak import baslat, okut, onayla, tamamla

pytestmark = pytest.mark.django_db

_TMY = "tasinir-mal-yonetmeligi.md"
_YONETMELIK = "meb-okul-kutuphaneleri-yonetmeligi.md"


def _hucreler(tablo: dict[str, Any]) -> list[list[str]]:
    return [[h["v"] for h in satir] for satir in tablo["rows"]]


# ---------------------------------------------------------------------------
# E17 — Alfabetik katalog dökümü
# ---------------------------------------------------------------------------
class TestAlfabetikKatalog:
    def test_uc_eksen_turk_alfabesiyle_siralanir(self) -> None:
        for ad, yazar, konu in (
            ("Zeytin", "Deneme Zeybek", "şiir"),
            ("Çiçek", "Deneme Çelik", "roman, Türk edebiyatı"),
            ("Cam", "Deneme Can", "deneme"),
            ("Ilık Rüzgâr", "Deneme Ilgaz", "öykü"),
            ("İnce Yol", "Deneme İnce", "Işık"),
            ("Şafak", "", ""),
        ):
            ortak.nusha(ortak.eser(title=ad, authors=yazar, subjects=konu))
        basliklar = [s.baslik for s in kd.catalog_rows(kd.AXIS_TITLE)]
        assert basliklar == ["Cam", "Çiçek", "Ilık Rüzgâr", "İnce Yol", "Şafak", "Zeytin"]
        yazarlar = [s.baslik for s in kd.catalog_rows(kd.AXIS_AUTHOR)]
        assert yazarlar == [
            "Can, Deneme",
            "Çelik, Deneme",
            "Ilgaz, Deneme",
            "İnce, Deneme",
            "Zeybek, Deneme",
            kd.YAZARSIZ,
        ]
        konular = [(s.baslik, s.kaynak_adi) for s in kd.catalog_rows(kd.AXIS_SUBJECT)]
        assert konular == [
            ("deneme", "Cam"),
            ("Işık", "İnce Yol"),
            ("öykü", "Ilık Rüzgâr"),
            ("roman", "Çiçek"),
            ("şiir", "Zeytin"),
            ("Türk edebiyatı", "Çiçek"),
            (kd.KONUSUZ, "Şafak"),
        ]

    def test_cok_yazarli_eserde_vd(self) -> None:
        ortak.nusha(ortak.eser(title="Ortak Kitap", authors="Deneme Bir, Deneme İki"))
        (satir,) = kd.catalog_rows(kd.AXIS_AUTHOR)
        assert satir.baslik == "Bir, Deneme vd."

    def test_kapsam_kayitta_nushasi_olan_ve_dijital(self) -> None:
        ortak.nusha(ortak.eser(title="Rafta Olan"))
        dusulen = ortak.nusha(ortak.eser(title="Yalnız Düşülmüş"))
        Copy.objects.filter(pk=dusulen.pk).update(status=CopyStatus.WITHDRAWN_WEEDED)
        silinen = ortak.nusha(ortak.eser(title="Yalnız Silinmiş"))
        catalog.delete_copy(silinen)
        kayip = ortak.nusha(ortak.eser(title="Kayıpta Olan"))
        Copy.objects.filter(pk=kayip.pk).update(status=CopyStatus.LOST)
        ortak.eser(title="Dijital Kaynak", resource_type=ResourceType.EBOOK)
        ortak.eser(title="Nüshasız Basılı")
        satirlar = {s.kaynak_adi: s for s in kd.catalog_rows(kd.AXIS_TITLE)}
        assert set(satirlar) == {"Rafta Olan", "Kayıpta Olan", "Dijital Kaynak"}
        assert satirlar["Dijital Kaynak"].nusha is None
        assert satirlar["Rafta Olan"].nusha == 1

    def test_md_11_1_alintisi_depodaki_metinle_birebir(self) -> None:
        assert kd.MD_11_1 in madde(_YONETMELIK, 11)
        assert (
            "Katalogları; yazar adı, kaynak adı ve konularına göre alfabetik olarak düzenler."
            in (madde(_YONETMELIK, 8))
        )

    def test_pdf_her_eksen_ve_xlsx_uc_sayfa_sayilar_sayi(self) -> None:
        ortak.nusha(ortak.eser(title="Çalıkuşu", authors="Reşat Nuri Güntekin", publish_year=2019))
        for eksen, ad in kd.AXES.items():
            metin = pdf_metni(kd.catalog_listing_pdf(eksen))
            assert "ALFABETİK KATALOG DÖKÜMÜ" in metin and ad in metin
            assert "Çalıkuşu" in metin
        kitap = load_workbook(BytesIO(kd.catalog_listing_xlsx()))
        assert kitap.sheetnames == list(kd.AXES.values())
        ws = kitap["Kaynak adına göre"]
        assert [h.value for h in ws[2]][:3] == ["Sıra", "Başlık", "Kaynak adı"]
        assert ws["A3"].value == 1 and ws["F3"].value == 2019 and ws["J3"].value == 1

    def test_bolum_suzgeci(self) -> None:
        edebiyat = ortak.bolum(name="Edebiyat")
        ortak.nusha(ortak.eser(title="Bölümlü", section=edebiyat))
        ortak.nusha(ortak.eser(title="Bölümsüz"))
        assert [s.kaynak_adi for s in kd.catalog_rows(kd.AXIS_TITLE, section_id=edebiyat.pk)] == [
            "Bölümlü"
        ]
        # F10 düzeltme turu: bölüm bölüm basan kullanıcı bölümü yazılmamış eserleri de basar.
        bolumsuz = kd.catalog_rows(kd.AXIS_TITLE, section_id=kd.SECTION_NONE)
        assert [s.kaynak_adi for s in bolumsuz] == ["Bölümsüz"]
        baglam = kd.catalog_listing_context(kd.AXIS_TITLE, section_id=kd.SECTION_NONE)
        assert {"label": "Bölüm", "value": kd.SECTION_NONE_LABEL} in baglam["info"]

    def test_kurum_yazari_ters_cevrilmez_adiyla_siralanir(self) -> None:
        """F10 düzeltme turu: "Millî Eğitim Bakanlığı" "Bakanlığı, Millî Eğitim" diye basılıp
        B harfinde sıralanıyordu."""
        for yazar, ad in (
            ("Reşat Nuri Güntekin", "Çalıkuşu"),
            ("Türk Dil Kurumu", "Örnek Sözlük"),
            ("Millî Eğitim Bakanlığı", "Örnek Ders Kitabı"),
            ("Yaşar Kemal", "İnce Memed"),
            ("Deneme Bakanoğlu", "Örnek Kişi"),
        ):
            ortak.nusha(ortak.eser(title=ad, authors=yazar))
        basliklar = [s.baslik for s in kd.catalog_rows(kd.AXIS_AUTHOR)]
        assert basliklar == [
            "Bakanoğlu, Deneme",
            "Güntekin, Reşat Nuri",
            "Kemal, Yaşar",
            "Millî Eğitim Bakanlığı",
            "Türk Dil Kurumu",
        ]
        from apps.kutuphane import keys

        assert keys.build_call_number("371.3", "Millî Eğitim Bakanlığı") == "371.3 MİL"
        assert keys.build_call_number("813", "Reşat Nuri Güntekin") == "813 GÜN"
        assert not keys.is_corporate_author("Kurumu")

    def test_sayfa_butcesi_ve_yatay_tasma_gercek_uzunlukta_veriyle(self) -> None:
        for sira in range(60):
            ortak.nusha(
                ortak.eser(
                    title=f"{EN_UZUN_ESER} {sira:02d}",
                    authors=UZUN_YAZAR,
                    publisher=UZUN_YAYINEVI,
                    publish_year=2024,
                    subjects="Türk edebiyatı, karşılaştırmalı edebiyat incelemesi, şiir",
                    call_number="894.3510 YAZ 2024 c.1 n.2",
                )
            )
        for eksen in kd.AXES:
            html = render_to_string(kd.SABLON, kd.catalog_listing_context(eksen))
            assert yatay_tasmalar(html) == [], eksen
        sayfalar = pdf_sayfalari(kd.catalog_listing_pdf(kd.AXIS_TITLE))
        # 60 uzun satır: satır başına en çok üç-dört satırlık hücre → sayfada en az 12 eser.
        assert len(sayfalar) <= 6, len(sayfalar)

    def test_bolum_ve_nusha_sutunlari_sozcugu_bolmez(self) -> None:
        # Bütünleştirme örneğinde görüldü: dar Bölüm sütunu "Edebiy/at", "Başvur/u" diye
        # sözcük ortasından bölüyordu; sayfa bütçesi testi bölümsüz veriyle koştuğu için
        # yakalamamıştı. Bölüm adları gerçek okul bölümlerinin uzunluğunda.
        for sira, ad in enumerate(("Başvuru Kaynakları", "Tarih ve Coğrafya", "Edebiyat")):
            ortak.nusha(ortak.eser(title=f"Deneme eseri {sira}", section=ortak.bolum(name=ad)))
        ortak.eser(title="Deneme dergi dizini", resource_type=ResourceType.EBOOK)
        for eksen in (kd.AXIS_TITLE, kd.AXIS_AUTHOR):
            html = render_to_string(kd.SABLON, kd.catalog_listing_context(eksen))
            sozcukler = {s for satir in hucre_satirlari(html) for s in satir.split()}
            # "Bölüm" yalnız sütun başlığında geçer (bölüm süzgeci yok); ötekiler bölüm adı.
            for sozcuk in ("Bölüm", "Başvuru", "Kaynakları", "Coğrafya", "Edebiyat"):
                assert sozcuk in sozcukler, (eksen, sozcuk)
            assert kd.DIJITAL_NUSHA in sozcukler, eksen


# ---------------------------------------------------------------------------
# E11 — Taşınır Kütüphane Defteri dökümü
# ---------------------------------------------------------------------------
class TestKutuphaneDefteri:
    def test_ciltsiz_sureli_yayin_girmez_ciltli_girer(self) -> None:
        """Kod kapısı: ciltletilmemiş süreli yayın defter dökümüne GİRMEZ (10/1-a-4, 15/4).

        Nüsha kuralları böyle bir nüshayı açtırmaz; döküm kendi süzgeciyle yine dışlar
        (eski veri ya da doğrudan yazım — savunma derinliği)."""
        dergi = ortak.eser(title="Örnek Dergisi", authors="", resource_type=ResourceType.PERIODICAL)
        ciltli = ortak.nusha(dergi, is_bound_periodical=True)
        ciltsiz = Copy.objects.create(
            work=dergi,
            acquisition=ciltli.acquisition,
            accession_no=2099000001,
            barcode="2099000001",
            is_bound_periodical=False,
        )
        kitap = ortak.nusha(ortak.eser(title="Kitap"))
        barkodlar = [s.barkod for s in tmy.register_rows()]
        assert ciltli.barcode in barkodlar and kitap.barcode in barkodlar
        assert ciltsiz.barcode not in barkodlar
        for icerik in (tmy.library_register_pdf(),):
            assert "2099000001" not in pdf_metni(icerik)
        kitap_x = load_workbook(BytesIO(tmy.library_register_xlsx()))
        degerler = [
            str(h)
            for s in kitap_x.worksheets
            for r in s.iter_rows(values_only=True)
            for h in r
            if h
        ]
        assert "2099000001" not in degerler and 2099000001 not in [
            h for s in kitap_x.worksheets for r in s.iter_rows(values_only=True) for h in r
        ]

    def test_kayit_no_sirasi_cikis_tarihi_ve_silinmis_nusha(self) -> None:
        eser = ortak.eser(title="Defterdeki")
        ilk, ikinci, ucuncu = (ortak.nusha(eser) for _ in range(3))
        Copy.objects.filter(pk=ikinci.pk).update(status=CopyStatus.WITHDRAWN_WEEDED)
        catalog.delete_copy(ucuncu)
        satirlar = list(tmy.register_rows())
        assert [s.barkod for s in satirlar] == [ilk.barcode, ikinci.barcode]
        assert satirlar[0].cikis_tarihi is None and satirlar[1].cikis_tarihi is not None

    def test_yil_suzgeci_edinim_yilidir(self) -> None:
        eski = ortak.edinim(date=date(2025, 5, 1))
        yeni = ortak.edinim(date=date(2026, 5, 1))
        a = ortak.nusha(ortak.eser(title="Eski"), eski)
        b = ortak.nusha(ortak.eser(title="Yeni"), yeni)
        assert [s.barkod for s in tmy.register_rows(year=2026)] == [b.barcode]
        assert tmy.register_years() == [2026, 2025]
        del a

    def test_xlsx_sayilar_sayi_tarihler_tarih(self) -> None:
        satin = ortak.edinim(
            method=AcquisitionMethod.PURCHASE, date=date(2026, 2, 10), unit_price=Decimal("125.50")
        )
        ortak.nusha(ortak.eser(title="Fiyatlı"), satin)
        ws = load_workbook(BytesIO(tmy.library_register_xlsx()))["Taşınır Kütüphane Defteri"]
        satir = [h.value for h in ws[3]]
        assert isinstance(satir[0], int)
        assert isinstance(satir[2], datetime)
        assert Decimal(str(satir[10])) == Decimal("125.5")
        assert ws["C3"].number_format == "DD.MM.YYYY"

    def test_atiflar_depodaki_metinle_birebir(self) -> None:
        dokuz = madde(_TMY, 9)
        assert tmy.TMY_9_1_C in dokuz and "ç) Kütüphane Defteri:" in dokuz
        on = madde(_TMY, 10)
        assert tmy.TMY_10_1_A_4 in on and "süreli yayınlardan ciltletilmiş olanlar hariç" in on
        assert tmy.TMY_15_4 in madde(_TMY, 15)
        otuz_dort = madde(_TMY, 34)
        assert tmy.TMY_34_3_A in otuz_dort and tmy.TMY_34_3_A_2 in otuz_dort
        assert "müze ve kütüphane olarak faaliyet gösteren harcama birimlerinde" in otuz_dort
        # K6: yönetim hesabı yıl sonu sayımının tutanağını içerir; cetvelin tanımı 10/1-k'dedir
        # (OYS'nin "10/1-m" atfı yoktur — D1).
        assert "a) Yıl sonu sayımına ilişkin Sayım Tutanağı." in otuz_dort
        assert "k) Müze/Kütüphane Yönetim Hesabı Cetveli:" in on
        assert " m) " not in on
        otuz_iki = madde(_TMY, 32)
        assert "Kayıtların sayım sonuçlarıyla uygunluğu sağlandıktan sonra" in otuz_iki

    def test_sayfa_butcesi_ve_yatay_tasma_gercek_uzunlukta_veriyle(self) -> None:
        satin = ortak.edinim(
            method=AcquisitionMethod.PURCHASE,
            date=date(2026, 2, 10),
            unit_price=Decimal("12345.50"),
        )
        eser = ortak.eser(title=EN_UZUN_ESER, authors=UZUN_YAZAR, publisher=UZUN_YAYINEVI)
        for _ in range(60):
            n = ortak.nusha(eser, satin, external_asset_ref=UZUN_TKYS, old_register_no="9" * 40)
            Copy.objects.filter(pk=n.pk).update(status=CopyStatus.WITHDRAWN_MISSING)
        html = render_to_string(tmy.SABLON, tmy.library_register_context())
        assert yatay_tasmalar(html) == []
        sayfalar = pdf_sayfalari(tmy.library_register_pdf())
        assert len(sayfalar) <= 12, len(sayfalar)
        assert "TAŞINIR KÜTÜPHANE DEFTERİ DÖKÜMÜ" in sayfalar[0]

    def test_olagan_kodlar_ve_tarihler_bolunmeden_basilir(self) -> None:
        # Bütünleştirme örneğinde görüldü: dar sütunlar olağan uzunluktaki TKYS kodunu
        # ("255.01.02.03/.04") ve eski kayıt no'yu ("1998/0/0427") bölüyordu; sayfa bütçesi
        # testi yalnız aşırı uzun değerlerle koştuğu için olağan değerin bölünmesini sınamıyordu.
        satin = ortak.edinim(
            method=AcquisitionMethod.PURCHASE,
            date=date(2026, 2, 10),
            unit_price=Decimal("12345.50"),
        )
        n = ortak.nusha(
            ortak.eser(title="Deneme eseri"),
            satin,
            external_asset_ref="255.01.02.03.04",
            old_register_no="1998/00427",
        )
        Copy.objects.filter(pk=n.pk).update(status=CopyStatus.WITHDRAWN_MISSING)
        html = render_to_string(tmy.SABLON, tmy.library_register_context())
        sozcukler = {s for satir in hucre_satirlari(html) for s in satir.split()}
        giris = n.acquisition.date
        cikis = timezone.localdate(Copy.all_objects.get(pk=n.pk).updated_at)
        for deger in (
            "255.01.02.03.04",
            "1998/00427",
            str(n.accession_no),
            f"{giris:%d.%m.%Y}",
            f"{cikis:%d.%m.%Y}",
            "12.345,50",
        ):
            assert deger in sozcukler, (deger, sorted(sozcukler))


# ---------------------------------------------------------------------------
# E11 — Yönetim hesabı cetveli hazırlığı (yıl sonu sayımına dayanır — K6)
# ---------------------------------------------------------------------------
class TestYonetimHesabi:
    def test_isaretsiz_ya_da_onaysiz_sayimdan_basilmaz(self) -> None:
        odunc_nushasi(title="Sayılan")
        ara = tamamla(baslat())
        onayla(ara)
        with pytest.raises(ValidationError) as hata:
            tmy.management_account_pdf(ara)
        assert tmy.NOT_YEAR_END_MESSAGE in str(hata.value)
        assert tmy.year_end_stocktakes() == []

    def test_ciltsiz_sureli_yayin_cetvel_hazirligina_da_girmez(self) -> None:
        """F10 düzeltme turu: defterin süzgeci cetvel hazırlığında ve 34/1'de de uygulanır —
        defter ile cetvel ayrışmaz (savunma derinliği: doğrudan yazılmış eski veri)."""
        dergi = ortak.eser(title="Örnek Dergisi", authors="", resource_type=ResourceType.PERIODICAL)
        ciltli = ortak.nusha(dergi, is_bound_periodical=True)
        Copy.objects.create(
            work=dergi,
            acquisition=ciltli.acquisition,
            accession_no=2099000001,
            barcode="2099000001",
            is_bound_periodical=False,
        )
        kitap = ortak.nusha(ortak.eser(title="Kitap"))
        sayim = baslat(is_year_end=True)
        okut(sayim, ciltli, kitap)
        tamamla(sayim)
        onayla(sayim)
        sayim.refresh_from_db()
        sayilar = selectors_sayim.tmy_34_1(sayim)
        defter = len(list(tmy.register_rows()))
        assert defter == 2
        aktarim_ve_giren = sayilar["program_transfer"] + sayilar["entered"]
        assert aktarim_ve_giren + sayilar["previous_year_carryover"] == defter
        assert sayilar["year_end_by_record"] == defter
        degerler = tmy.management_account_values(sayim)
        assert degerler["kayda_gore"].nusha == defter
        assert sum(d.nusha for k, d in degerler.items() if k.startswith("cikan:")) == 0

    def test_iptal_edilmis_yil_sonu_sayimi_listelenmez(self) -> None:
        from apps.kutuphane.services import stocktake

        odunc_nushasi(title="Sayılan")
        sayim = baslat(is_year_end=True)
        stocktake.cancel_stocktake(sayim)
        sayim.refresh_from_db()
        assert tmy.year_end_stocktakes() == []
        assert tmy.management_account_blocker(sayim) == tmy.CANCELLED_MESSAGE

    def test_bir_mali_yilin_tek_yil_sonu_sayimi_olur(self) -> None:
        """F10 düzeltme turu: aynı mali yıl için ikinci "yıl sonu sayımı" işareti reddedilir
        (TMY 32/1, 34/2-a); iptal edilmiş sayım sayılmaz, başka yıl serbesttir."""
        from apps.kutuphane.services import stocktake

        odunc_nushasi(title="Sayılan")
        ilk = tamamla(baslat(is_year_end=True, fiscal_year=2026))
        onayla(ilk)
        with pytest.raises(ValidationError) as hata:
            stocktake.create_stocktake(is_year_end=True, fiscal_year=2026)
        assert "2026 mali yılının yıl sonu sayımı zaten var" in str(hata.value)
        ikinci = stocktake.create_stocktake(fiscal_year=2026)
        with pytest.raises(ValidationError):
            stocktake.update_stocktake(ikinci, is_year_end=True)
        stocktake.update_stocktake(ikinci, fiscal_year=2027, is_year_end=True)
        ikinci.refresh_from_db()
        assert ikinci.is_year_end is True
        with pytest.raises(ValidationError):
            stocktake.update_stocktake(ikinci, fiscal_year=2026)

    def test_onaysiz_iletisi_zinciri_metinden_dogru_atar(self) -> None:
        """F10 düzeltme turu: 32/9 Taşınır Sayım ve Döküm Cetvelidir; yönetim hesabı
        cetvelinin dayanağı 34/3-a'dır ve uygunluk 32/7'nin belgeleriyle sağlanır."""
        assert "md. 34/3-a" in tmy.NOT_APPROVED_MESSAGE
        assert "md. 32/7, 32/9" in tmy.NOT_APPROVED_MESSAGE
        assert "Sayım kurulu tarafından onaylanan Taşınır Sayım ve Döküm Cetveline" in madde(
            _TMY, 34
        )
        otuz_iki = madde(_TMY, 32)
        assert "sayım kurulu tarafından Taşınır Sayım ve Döküm Cetveli düzenlenir" in otuz_iki
        assert "defter kayıtlarının sayım sonuçlarıyla uygunluğu sağlanır" in otuz_iki

    def test_yil_sonu_ama_onaysiz_sayimdan_basilmaz(self) -> None:
        odunc_nushasi(title="Sayılan")
        sayim = tamamla(baslat(is_year_end=True))
        with pytest.raises(ValidationError) as hata:
            tmy.management_account_xlsx(sayim)
        assert tmy.NOT_APPROVED_MESSAGE in str(hata.value)
        (secenek,) = tmy.year_end_stocktakes()
        assert secenek["available"] is False and secenek["reason"] == tmy.NOT_APPROVED_MESSAGE

    def test_onayli_yil_sonu_sayimindan_basilir_sayilar_tmy_34_1_ile_ayni(self) -> None:
        satin = ortak.edinim(
            method=AcquisitionMethod.PURCHASE, date=date(2026, 2, 10), unit_price=Decimal("100")
        )
        bulunan = ortak.nusha(ortak.eser(title="Bulunan"), satin)
        noksan = ortak.nusha(ortak.eser(title="Noksan"), satin)
        fiyatsiz = odunc_nushasi(title="Fiyatsız")
        sayim = baslat(is_year_end=True, fiscal_year=2026)
        okut(sayim, bulunan, fiyatsiz)
        tamamla(sayim)
        onayla(sayim)
        sayim.refresh_from_db()

        sayilar = selectors_sayim.tmy_34_1(sayim)
        degerler = tmy.management_account_values(sayim)
        assert degerler["gelecek"].nusha == sayilar["next_year_carryover"] == 2
        assert degerler["gelecek"].tutar == Decimal("100") and degerler["gelecek"].fiyatsiz == 1
        assert degerler["kayda_gore"].nusha == sayilar["year_end_by_record"]
        giren = sum(d.nusha for k, d in degerler.items() if k.startswith("giren:"))
        assert giren == sayilar["entered"]
        cikan = sum(d.nusha for k, d in degerler.items() if k.startswith("cikan:"))
        assert cikan == sayilar["exited"] == 1

        baglam = tmy.management_account_context(sayim)
        assert baglam["status_note"] == tmy.HESAP_NOTU
        satirlar = {s[0]: s for s in _hucreler(baglam["tables"][0])}
        assert satirlar["Sayım noksanı"][1] == "1"
        metin = pdf_metni(tmy.management_account_pdf(sayim))
        assert "KÜTÜPHANE YÖNETİM HESABI CETVELİ HAZIRLIĞI" in metin
        assert "resmî cetveller TKYS'dedir" in metin
        ws = load_workbook(BytesIO(tmy.management_account_xlsx(sayim))).active
        assert ws is not None
        nushalar = [
            r[1] for r in ws.iter_rows(min_row=6, max_col=2, values_only=True) if r[1] is not None
        ]
        assert all(isinstance(n, int) for n in nushalar)
        del noksan
