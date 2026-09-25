"""Komisyon belgeleri: ayıklama (E7), nadir eserler (E8), bağış ön kaydı ve sonucu (E16) — F8.

Kapılar:
- E7 belgeleri TMY yoluna göre basılır: teklif listesi her zaman, ayıklama tutanağı ve
  kayıttan düşme teklif listesi komisyon kararından sonra, imha tutanağı YALNIZ imha
  kararı işaretli onaylı teklifte (TMY 28/5), devir listesi (PDF + XLSX) devir kalemi
  varsa. Kural sunucudadır (`weeding_document_blocker`) ve uç 400 + gerekçe döner.
- "İmha" sözcüğü yalnız imha tutanağında geçer (sözlük §1).
- Bağış değerlendirme sonucu (25.09.2026 kullanıcı kararı, tasarım F8 ekleri 13) yalnız
  karar uygulanmış ön kayıtta basılır; "tutanak" demez, TKYS dipnotunu taşır.
- Mevzuat alıntıları `docs/mevzuat/` metniyle BİREBİRDİR.
- Komisyon, TMY komisyonu, harcama yetkilisi ve bağışçı adları şifreli alandan yalnız
  belgenin kendisine çözülür; indirme adında ve hata iletisinde ad yoktur.
- Sayfa bütçesi GERÇEK UZUNLUKTA veriyle koşar (CLAUDE.md §3): en uzun okul adı, sekiz
  kişilik komisyon, çeviri künyeli yazar, 64 karakterlik TKYS kodu.

Bütün kişi ve kurum adları uydurmadır ("Deneme …", "Örnek …").
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from apps.kutuphane import barcode as barcode_module
from apps.kutuphane import komisyon_belgeleri as belgeler
from apps.kutuphane import selectors_ayiklama
from apps.kutuphane.models import Copy, WeedingBatch
from apps.kutuphane.services import donations, rare_works, weeding
from apps.kutuphane.tests.ayiklama_ortak import (
    DEVRALAN_OKUL,
    HARCAMA_YETKILISI,
    TMY_KOMISYONU,
    ayiklama_karari,
    kalem_ekle,
    onaya_kadar,
    teklif,
)
from apps.kutuphane.tests.ortak import eser, nusha
from apps.kutuphane.tests.teslim_ortak import etkin_yil
from apps.okul.kip import KIP
from apps.okul.models import SchoolConfig

pytestmark = pytest.mark.django_db

EN_UZUN_OKUL = ("Örnek Mesleki ve Teknik Anadolu Lisesi Uzun Adlı Kampüs Yerleşkesi " * 3)[:255]
EN_UZUN_ESER = (
    "Uzun Adlı Örnek Kaynak: Çağdaş Türk Şiirinde Doğa, Şehir ve İnsan Üzerine "
    "Karşılaştırmalı Bir İnceleme — Birinci Cilt, Genişletilmiş Üçüncü Baskı"
)
UZUN_YAZAR = (
    "Deneme Yazaroğlu; Hazırlayan: Örnek Derleyicioğlu, İkinci Hazırlayıcı; "
    "Çeviren: Üçüncü Çevirmenoğlu ve Dördüncü Çevirmen; Resimleyen: Beşinci Çizeroğlu"
)
UZUN_TKYS = "255.01.02.03.04.05.06.07.08.09-" + "9" * 33
BASKAN = "Deneme Komisyonbaşkanı"
UYELER = [
    "Deneme Birincikomisyonüyesi Uzunsoyadlıoğlu",
    "Deneme İkincikomisyonüyesi Uzunsoyadlıoğlu",
    "Deneme Üçüncükomisyonüyesi Uzunsoyadlıoğlu",
    "Deneme Dördüncükomisyonüyesi Uzunsoyadlıoğlu",
    "Deneme Beşincikomisyonüyesi Uzunsoyadlıoğlu",
    "Deneme Altıncıkomisyonüyesi Uzunsoyadlıoğlu",
    "Deneme Yedincikomisyonüyesi Uzunsoyadlıoğlu",
]
BAGISCI = "Deneme Bağışçıadı"

_T_ARALIGI = re.compile(r"T (?=[a-zçğıöşü])")


def _metin(icerik: bytes) -> str:
    ham = "\n".join(s.extract_text() or "" for s in PdfReader(io.BytesIO(icerik)).pages)
    return " ".join(_T_ARALIGI.sub("T", ham).split())


def _sayfalar(icerik: bytes) -> list[str]:
    return [
        " ".join(_T_ARALIGI.sub("T", s.extract_text() or "").split())
        for s in PdfReader(io.BytesIO(icerik)).pages
    ]


def _mevzuat(dosya: str) -> str:
    yerel = Path(__file__).resolve().parents[4]
    for kok in (yerel, Path("/repo")):
        yol = kok / "docs" / "mevzuat" / dosya
        if yol.is_file():
            return " ".join(yol.read_text(encoding="utf-8").split())
    pytest.fail(f"{dosya} bulunamadı.")


def _sablon(ad: str) -> str:
    yerel = Path(__file__).resolve().parents[3] / "templates" / "documents" / ad
    for yol in (yerel, Path("/app/templates/documents") / ad):
        if yol.is_file():
            return yol.read_text(encoding="utf-8")
    pytest.fail(f"{ad} bulunamadı.")


@pytest.fixture(autouse=True)
def okul() -> SchoolConfig:
    config = SchoolConfig.load()
    config.school_name = "Örnek Anadolu Lisesi"
    config.district = "Örnek İlçe"
    config.principal_name = "Örnek Müdür"
    config.save()
    etkin_yil()  # teklif ve nadir eser listesi etkin ders yılına açılır
    return config


def _kitap(title: str = "Deneme Eseri", *, uzun: bool = False, **alanlar: Any) -> Copy:
    if uzun:
        return nusha(
            eser(
                title=EN_UZUN_ESER,
                authors=UZUN_YAZAR,
                publisher="Örnek Uzun Adlı Yayınevi Basım Dağıtım",
                publish_year=1987,
            ),
            external_asset_ref=UZUN_TKYS,
            **alanlar,
        )
    return nusha(eser(title=title, authors="Deneme Yazar", publish_year=1990), **alanlar)


def _komisyon_karari(**alanlar: Any) -> Any:
    alanlar.setdefault("chair_name", BASKAN)
    alanlar.setdefault("chair_title", "İlçe Millî Eğitim Şube Müdürü")
    alanlar.setdefault("participants_text", "\n".join(UYELER[:3]))
    return ayiklama_karari(**alanlar)


def _karma_teklif() -> tuple[WeedingBatch, dict[str, Copy]]:
    """Dört gerekçenin her biri ve iki devir yolu (taslak)."""
    batch = teklif()
    kitaplar = {
        "yipranmis": _kitap("Yıpranmış Eser"),
        "eskimis": _kitap("Bilimsel Değeri Kalmamış Eser"),
        "duzey": _kitap("Düzeye Uygun Olmayan Eser"),
        "olcut": _kitap("Ölçüte Uygun Olmayan Eser"),
    }
    kalem_ekle(batch, kitaplar["yipranmis"], "WORN")
    kalem_ekle(batch, kitaplar["eskimis"], "OBSOLETE")
    kalem_ekle(batch, kitaplar["duzey"], "LEVEL_MISMATCH", transfer_target=DEVRALAN_OKUL)
    kalem_ekle(batch, kitaplar["olcut"], "CRITERIA_MISMATCH", criterion="10_4")
    return batch, kitaplar


def _t(batch: WeedingBatch) -> WeedingBatch:
    """Teklifin güncel hâli (servisler kopyayı değil veritabanını değiştirir)."""
    guncel = selectors_ayiklama.get_weeding_batch(batch.pk)
    assert guncel is not None
    return guncel


def _durumlar(batch: WeedingBatch) -> dict[str, bool]:
    return {b["kind"]: b["available"] for b in belgeler.weeding_documents(_t(batch))}


# ============================================================ basılabilirlik


class TestBasilabilirlik:
    def test_taslakta_yalniz_teklif_listesi(self) -> None:
        batch, _ = _karma_teklif()
        assert _durumlar(batch) == {
            "teklif-listesi": True,
            "ayiklama-tutanagi": False,
            "kayittan-dusme": False,
            "imha-tutanagi": False,
            "devir-listesi": False,
        }
        engel = belgeler.weeding_document_blocker(_t(batch), belgeler.AYIKLAMA_TUTANAGI)
        assert engel == belgeler.NEEDS_DECISION_MESSAGE

    def test_kalemsiz_teklifte_hicbiri(self) -> None:
        batch = teklif()
        assert not any(_durumlar(batch).values())
        assert (
            belgeler.weeding_document_blocker(_t(batch), "teklif-listesi") == "Teklifte kalem yok."
        )

    def test_karardan_sonra_tutanak_dusme_ve_devir(self) -> None:
        batch, _ = _karma_teklif()
        weeding.submit_batch(batch)
        weeding.bind_decision(batch, commission_decision=_komisyon_karari())
        assert _durumlar(batch) == {
            "teklif-listesi": True,
            "ayiklama-tutanagi": True,
            "kayittan-dusme": True,
            "imha-tutanagi": False,
            "devir-listesi": True,
        }

    def test_imha_tutanagi_yalniz_imha_kararinda_ve_28_yolunda(self) -> None:
        batch, _ = _karma_teklif()
        onaya_kadar(batch, decision=_komisyon_karari())
        assert _durumlar(batch)["imha-tutanagi"] is False
        assert (
            belgeler.weeding_document_blocker(_t(batch), belgeler.IMHA_TUTANAGI)
            == belgeler.NEEDS_DESTRUCTION_MESSAGE
        )

        imhali = teklif()
        kalem_ekle(imhali, _kitap("Hurdaya Ayrılacak"), "OBSOLETE")
        onaya_kadar(imhali, decision=_komisyon_karari(), destruction_decided=True)
        assert _durumlar(imhali)["imha-tutanagi"] is True

    def test_yalniz_devir_kaleminde_kayittan_dusme_yok_yalniz_dusmede_devir_yok(self) -> None:
        devir = teklif()
        kalem_ekle(devir, _kitap(), "LEVEL_MISMATCH", transfer_target=DEVRALAN_OKUL)
        weeding.submit_batch(devir)
        weeding.bind_decision(devir, commission_decision=_komisyon_karari())
        assert _durumlar(devir)["kayittan-dusme"] is False
        assert _durumlar(devir)["devir-listesi"] is True

        dusme = teklif()
        kalem_ekle(dusme, _kitap("Başka"), "WORN")
        weeding.submit_batch(dusme)
        weeding.bind_decision(dusme, commission_decision=_komisyon_karari())
        assert _durumlar(dusme)["devir-listesi"] is False
        assert _durumlar(dusme)["kayittan-dusme"] is True

    def test_komisyonun_ayiklamadigi_devir_kalemi_devir_listesine_girmez(self) -> None:
        batch = teklif()
        kalem = kalem_ekle(batch, _kitap(), "LEVEL_MISMATCH", transfer_target=DEVRALAN_OKUL)
        kalem_ekle(batch, _kitap("Başka"), "WORN")
        weeding.submit_batch(batch)
        weeding.bind_decision(
            batch,
            commission_decision=_komisyon_karari(),
            kept={kalem.pk: "Okul kitaplığında kalsın."},
        )
        assert _durumlar(batch)["devir-listesi"] is False

    def test_bilinmeyen_belge_ve_bicim(self) -> None:
        batch, _ = _karma_teklif()
        with pytest.raises(Exception, match="Böyle bir ayıklama belgesi yok"):
            belgeler.ensure_weeding_document(batch, "yok")
        with pytest.raises(Exception, match="Bu belge bu biçimde üretilmez"):
            belgeler.ensure_weeding_document(batch, belgeler.TEKLIF_LISTESI, belgeler.XLSX)


# ============================================================ E7 içerik


class TestAyiklamaBelgeleri:
    def test_teklif_listesi_gerekce_bendi_ve_tmy_yolu(self) -> None:
        batch, kitaplar = _karma_teklif()
        metin = _metin(belgeler.proposal_list_pdf(_t(batch)))

        assert "AYIKLAMA TEKLİF LİSTESİ" in metin
        assert "Seçim ve Ayıklama Komisyonunun değerlendirmesine sunulur" in metin
        for kitap in kitaplar.values():
            assert barcode_module.format_barcode(kitap.barcode) in metin
        assert "Aşırı kullanımdan yıpranmış (Md. 12/1-a)" in metin
        assert "Bilimsel değeri kalmamış (Md. 12/1-b)" in metin
        assert "Kurumun düzeyine uygun değil (Md. 12/1-c; 10/1-b)" in metin
        # Md. 12/1 devri 10/1-b'ye uygun olmayan kaynak için öngörür; not bunu söyler.
        assert "yaş ve gelişim düzeyine uygun olmayan kaynağın (Md. 10/1-b)" in metin
        assert "(Md. 10/4)" in metin
        assert "Kayıttan düşme (TMY md. 27/1)" in metin
        assert "Hurdaya ayırma (TMY md. 28)" in metin
        assert f"Devir: MEB okulu (TMY md. 24/2) — {DEVRALAN_OKUL}" in metin
        assert "Teklif eden — Kütüphane yöneticisi" in metin

    def test_ayiklama_tutanagi_komisyon_adlari_ve_ayiklanmayanlar(self) -> None:
        batch, kitaplar = _karma_teklif()
        weeding.submit_batch(batch)
        kalemler = {k.copy_id: k for k in batch.items.all()}
        weeding.bind_decision(
            batch,
            commission_decision=_komisyon_karari(decision_no="2027/4"),
            kept={kalemler[kitaplar["eskimis"].pk].pk: "Kaynak hâlâ kullanılıyor."},
        )
        metin = _metin(belgeler.weeding_report_pdf(_t(batch)))

        assert "AYIKLAMA TUTANAĞI" in metin
        assert "tarihli ve 2027/4 sayılı kararıyla" in metin
        assert "AYIKLANMASINA KARAR VERİLEN KAYNAKLAR" in metin
        assert "KOMİSYONUN AYIKLANMASINA KARAR VERMEDİĞİ KAYNAKLAR" in metin
        assert "Kaynak hâlâ kullanılıyor." in metin
        assert "SEÇİM VE AYIKLAMA KOMİSYONU" in metin
        assert BASKAN in metin and "İlçe Millî Eğitim Şube Müdürü" in metin
        for uye in UYELER[:3]:
            assert uye in metin
        assert belgeler.MD12_1_DEVIR in metin

    def test_kayittan_dusme_27_yolunda_komisyonsuz_onay_notu(self) -> None:
        batch = teklif()
        kalem_ekle(batch, _kitap(), "WORN")
        onaya_kadar(batch, decision=_komisyon_karari(), tmy_commission_members="")
        metin = _metin(belgeler.write_off_pdf(_t(batch)))

        assert "KAYITTAN DÜŞME TEKLİF LİSTESİ" in metin
        assert "Kayıttan Düşme Teklif ve Onay Tutanağına esas hazırlık çıktısı" in metin
        assert "KULLANILMAZ HÂLE GELEN KAYNAKLAR (TMY MD. 27/1)" in metin
        # 10/1-e yorumu sonuç cümlesi değil, harcama yetkilisinin takdiridir.
        assert "komisyon kurulması gerekmeksizin harcama yetkilisince onaylanabilir" in metin
        assert "harcama yetkilisi değerlendirir" in metin
        assert "harcama yetkilisince onaylanır" not in metin
        assert "md. 27/3" in metin and "md. 5/8" in metin
        assert "OLUR" in metin and HARCAMA_YETKILISI in metin
        assert f"{timezone.localdate():%d.%m.%Y}" in metin
        assert "Taşınır Kayıt ve Yönetim Sistemi'nde (TKYS) düzenlenir" in metin
        assert "KOMİSYON" not in metin.replace("KOMİSYONUN", "")

    def test_kayittan_dusme_28_yolunda_tmy_komisyonu_ve_onaylanmayan(self) -> None:
        batch = teklif()
        kalem = kalem_ekle(batch, _kitap("Onaylanmayan"), "OBSOLETE")
        kalem_ekle(batch, _kitap("Hurdaya Ayrılan"), "CRITERIA_MISMATCH", criterion="10_1_D")
        onaya_kadar(
            batch,
            decision=_komisyon_karari(),
            not_approved={kalem.pk: "Ekonomik değeri var; kullanılabilir."},
        )
        metin = _metin(belgeler.write_off_pdf(_t(batch)))

        assert "HURDAYA AYRILAN KAYNAKLAR (TMY MD. 28)" in metin
        assert "KAYITTAN DÜŞÜLMEYEN KAYNAKLAR" in metin
        assert "Ekonomik değeri var; kullanılabilir." in metin
        assert "Türkçenin doğru ve güzel kullanımını desteklemiyor (Md. 10/1-d)" in metin
        for uye in TMY_KOMISYONU.splitlines():
            assert uye in metin
        assert "komisyon kurulması gerekmeksizin" not in metin

    def test_onaydan_once_harcama_yetkilisi_bos_ve_28_yolunda_uc_bos_imza(self) -> None:
        batch = teklif()
        kalem_ekle(batch, _kitap(), "OBSOLETE")
        weeding.submit_batch(batch)
        weeding.bind_decision(batch, commission_decision=_komisyon_karari())
        baglam = belgeler.write_off_context(_t(batch))

        assert baglam["approval"]["name"] == "" and baglam["approval"]["date"] == "…/…/……"
        assert [k["name"] for satir in baglam["signature_rows"] for k in satir] == ["", "", ""]
        assert "Onay bekliyor" in _metin(belgeler.write_off_pdf(_t(batch)))

    def test_imha_tutanagi_sablonu(self) -> None:
        batch = teklif()
        kitap = _kitap("Hurdaya Ayrılan")
        kalem_ekle(batch, kitap, "OBSOLETE")
        kalem_ekle(batch, _kitap("Yıpranan"), "WORN")
        onaya_kadar(batch, decision=_komisyon_karari(), destruction_decided=True)
        metin = _metin(belgeler.destruction_report_pdf(_t(batch)))

        assert "İMHA TUTANAĞI" in metin
        assert "İmha tarihi" in metin and "İmha yeri" in metin and "İmha yöntemi" in metin
        assert belgeler.TMY28_5_IMHA in metin
        assert barcode_module.format_barcode(kitap.barcode) in metin
        assert "Yıpranan" not in metin  # 27/1 yolundaki kaynak imha edilmez
        for uye in TMY_KOMISYONU.splitlines():
            assert uye in metin

    def test_devir_listesi_pdf_ve_xlsx(self) -> None:
        batch = teklif()
        kitap = _kitap("Devredilecek Eser")
        kalem_ekle(batch, kitap, "LEVEL_MISMATCH", transfer_target=DEVRALAN_OKUL)
        onaya_kadar(batch, decision=_komisyon_karari())
        metin = _metin(belgeler.transfer_list_pdf(_t(batch)))

        assert "DEVİR LİSTESİ" in metin
        assert DEVRALAN_OKUL in metin and f"Devralan — {DEVRALAN_OKUL}" in metin
        assert belgeler.MD12_1_DEVIR in metin
        assert "(Taşınır Mal Yönetmeliği md. 24/2)" in metin
        assert HARCAMA_YETKILISI in metin
        assert "Taşınır Kayıt ve Yönetim Sistemi'nde (TKYS) yapılır" in metin

        wb = load_workbook(io.BytesIO(belgeler.transfer_list_xlsx(_t(batch))))
        ws = wb.active
        assert ws is not None
        satirlar = [list(r) for r in ws.iter_rows(values_only=True)]
        baslik = [ad for ad, _ in belgeler.DEVIR_XLSX_SUTUNLARI]
        sira = satirlar.index(baslik)
        veri = satirlar[sira + 1]
        assert veri[:3] == [1, barcode_module.format_barcode(kitap.barcode), "Devredilecek Eser"]
        assert veri[8] == DEVRALAN_OKUL and veri[9] == "Devir: MEB okulu (TMY md. 24/2)"
        # Biçim ve bölme BAŞLIK satırındadır (boş ayraç satırında değil): kalın + dolgulu
        # başlık, donmuş bölge başlığın bir altından başlar.
        baslik_satiri = sira + 1  # iter_rows 0'dan, sayfa 1'den sayar
        for sutun in range(1, len(baslik) + 1):
            hucre = ws.cell(row=baslik_satiri, column=sutun)
            assert hucre.font.bold, hucre.coordinate
            assert hucre.fill.fgColor.rgb.endswith("E8EEF6"), hucre.coordinate
        bos = ws.cell(row=baslik_satiri - 1, column=1)
        assert bos.value is None and not bos.font.bold
        assert ws.freeze_panes == f"A{baslik_satiri + 1}"

    def test_birden_cok_devralanda_sutun_ve_31_notu(self) -> None:
        batch = teklif()
        kalem_ekle(batch, _kitap("Birinci"), "LEVEL_MISMATCH", transfer_target=DEVRALAN_OKUL)
        kalem_ekle(
            batch,
            _kitap("İkinci"),
            "LEVEL_MISMATCH",
            tmy_path="TMY_31",
            transfer_target="Deneme Halk Kütüphanesi",
        )
        onaya_kadar(batch, decision=_komisyon_karari())
        metin = _metin(belgeler.transfer_list_pdf(_t(batch)))

        assert "Devralacak" in metin and "Devralan okul ya da kurum" in metin
        assert "Deneme Halk Kütüphanesi" in metin
        assert "(Taşınır Mal Yönetmeliği md. 31/1)" in metin and "(md. 24/1)" in metin

    def test_kurumu_yazilmamis_devir_kaleminde_devralan_bos_kalir(self) -> None:
        """Karardan sonra kurum yazılmamış olabilir: "Devralan — —" basılmaz (F8 düzeltme)."""
        batch = teklif()
        kalem_ekle(batch, _kitap("Kurumsuz Devir"), "LEVEL_MISMATCH")
        weeding.submit_batch(batch)
        weeding.bind_decision(batch, commission_decision=_komisyon_karari())
        baglam = belgeler.transfer_list_context(_t(batch))
        roller = [k["role"] for satir in baglam["signature_rows"] for k in satir]
        assert roller == ["Devreden — Kütüphane yöneticisi", "Devralan okul ya da kurum"]
        bilgi = {b["label"]: b["value"] for b in baglam["info"]}
        assert bilgi["Devralacak okul ya da kurum"] == belgeler.KURUM_YAZILMADI
        metin = _metin(belgeler.transfer_list_pdf(_t(batch)))
        assert "Devralan — —" not in metin and "— —" not in metin

    def test_hurdaya_ayirmada_ilk_uye_isin_uzmani(self) -> None:
        """TMY 28/1: komisyonun biri işin uzmanıdır — ilk imza satırı bu rolle basılır."""
        batch = teklif()
        kalem_ekle(batch, _kitap("Hurdaya Ayrılan"), "OBSOLETE")
        weeding.submit_batch(batch)
        weeding.bind_decision(batch, commission_decision=_komisyon_karari())
        bos = belgeler.write_off_context(_t(batch))
        assert [k["role"] for satir in bos["signature_rows"] for k in satir] == [
            belgeler.UZMAN_ROLU,
            belgeler.UYE_ROLU,
            belgeler.UYE_ROLU,
        ]
        onayli = teklif()
        kalem_ekle(onayli, _kitap("İkinci Hurda"), "OBSOLETE")
        onaya_kadar(onayli, decision=_komisyon_karari())
        dolu = belgeler.write_off_context(_t(onayli))
        hucreler = [k for satir in dolu["signature_rows"] for k in satir]
        assert hucreler[0] == {
            "name": TMY_KOMISYONU.splitlines()[0],
            "role": belgeler.UZMAN_ROLU,
        }
        # 27/1 yolunda (isteğe bağlı komisyon) "işin uzmanı" rolü yoktur.
        yipranma = teklif()
        kalem_ekle(yipranma, _kitap("Yıpranan"), "WORN")
        onaya_kadar(yipranma, decision=_komisyon_karari())
        roller = {
            k["role"]
            for satir in belgeler.write_off_context(_t(yipranma))["signature_rows"]
            for k in satir
        }
        assert roller == {belgeler.UYE_ROLU}

    def test_yipranma_kalemi_27_1_lafziyla_ve_5_8_ile_basilir(self) -> None:
        """Kalemler yalnız Md. 12/1 gerekçelidir (F8 ekleri 34 — hasar önerisi ayıklamaya
        konmaz): 27/1 paragrafı TMY'nin lafzıyla, 27/3 ve 5/8 ile; dayanak yalnız Ayıklama
        tutanağıdır, Kayıp/hasar tutanağı anılmaz."""
        batch = teklif()
        kalem_ekle(batch, _kitap("Sade Yıpranan"), "WORN")
        onaya_kadar(batch, decision=_komisyon_karari(), tmy_commission_members="")
        baglam = belgeler.write_off_context(_t(batch))
        paragraflar = " ".join(baglam["paragraphs"])
        notlar = " ".join(baglam["notes"])
        gerekceler = [s[-1]["v"] for tb in baglam["tables"] for s in tb["rows"]]
        dayanak = {b["label"]: b["value"] for b in baglam["info"]}["Dayanak"]

        assert "Yıpranma, kırılma veya bozulma gibi nedenlerle" in paragraflar
        assert "md. 27/3" in paragraflar and "md. 5/8" in paragraflar
        assert gerekceler == ["Aşırı kullanımdan yıpranmış (Md. 12/1-a)"]
        assert dayanak.endswith("kararı ve Ayıklama tutanağı")
        assert "harcama yetkilisince onaylanabilir" in notlar
        metin = _metin(belgeler.write_off_pdf(_t(batch)))
        assert "hasar" not in metin.casefold()

    def test_imha_tutanagi_yalniz_imha_kararinin_kapsadigi_kalemleri_basar(self) -> None:
        """TMY 28/5 kalem düzeyinde: kapsam dışı 28 kalemi imha tutanağına girmez."""
        batch = teklif()
        imha = kalem_ekle(batch, _kitap("İmha Edilecek"), "OBSOLETE")
        kalem_ekle(batch, _kitap("Ekonomik Değeri Olan"), "OBSOLETE")
        onaya_kadar(
            batch,
            decision=_komisyon_karari(),
            destruction_decided=True,
            destruction_items=[imha.pk],
        )
        metin = _metin(belgeler.destruction_report_pdf(_t(batch)))
        assert "İmha Edilecek" in metin and "Ekonomik Değeri Olan" not in metin
        baglam = belgeler.destruction_report_context(_t(batch))
        assert len(baglam["tables"][0]["rows"]) == 1
        assert {b["label"]: b["value"] for b in baglam["info"]}["Kaynak sayısı"] == "1"

    def test_imha_sozcugu_yalniz_imha_tutanaginda(self) -> None:
        batch, _ = _karma_teklif()
        onaya_kadar(batch, decision=_komisyon_karari(), destruction_decided=True)
        for uretici in (
            belgeler.proposal_list_pdf,
            belgeler.weeding_report_pdf,
            belgeler.write_off_pdf,
            belgeler.transfer_list_pdf,
        ):
            assert "imha" not in _metin(uretici(_t(batch))).casefold(), uretici.__name__
        assert "imha" in _metin(belgeler.destruction_report_pdf(_t(batch))).casefold()


# ============================================================ mevzuat alıntıları


class TestMevzuatAlintilari:
    def test_yonetmelik_alintilari_depodaki_metinle_birebir(self) -> None:
        yonetmelik = _mevzuat("meb-okul-kutuphaneleri-yonetmeligi.md")
        for alinti in (belgeler.MD12_1_DEVIR, belgeler.MD12_2, belgeler.MD10_3):
            assert alinti in yonetmelik, alinti

    def test_tmy_hukumleri_depodaki_metinde(self) -> None:
        tmy = _mevzuat("tasinir-mal-yonetmeligi.md")
        assert belgeler.TMY28_5_IMHA in tmy
        assert "Bu işleme ilişkin ayrıca bir imha tutanağı düzenlenir." in tmy
        assert "komisyon kurulması gerekmeksizin harcama yetkilisince onaylanır" in tmy
        assert (
            "Varlık İşlem Fişinin bir nüshası devredilen harcama biriminin taşınır kayıt "
            "yetkilisine verilir" in tmy
        )
        assert "Varlık İşlem Fişinin bir nüshası taşınırın devredildiği idareye verilir" in tmy
        assert "diğer kamu idarelerine bedelsiz olarak devredebilir" in tmy
        assert "kasıt, kusur, ihmal veya tedbirsizlik olup olmadığı harcama yetkilisince" in tmy
        assert "olağan kullanımından kaynaklanan yıpranma" in tmy
        assert "en az üç kişiden oluşan komisyon tarafından değerlendirilir" in tmy
        assert "Hurdaya ayrılan veya imha edilen taşınırlar Varlık İşlem Fişi" in tmy

    def test_sablonlarda_text_transform_yok(self) -> None:
        for ad in ("komisyon_listesi.html", "yil_sonu_raporu.html"):
            assert "text-transform:" not in _sablon(ad).replace(" ", ""), ad


# ============================================================ E8, E16


class TestNadirEserListesi:
    def test_karara_bagli_liste_ve_gonderim(self) -> None:
        liste = rare_works.create_submission(commission_decision=_komisyon_karari())
        kitap = _kitap("El Yazması Divan", is_rare_or_manuscript=True)
        rare_works.add_copies(liste, copy_ids=[kitap.pk])
        rare_works.send_submission(liste, sent_on=timezone.localdate(), sent_document_no="E-77")
        guncel = selectors_ayiklama.get_rare_works_submission(liste.pk)
        assert guncel is not None
        metin = _metin(belgeler.rare_works_pdf(guncel))

        assert "EL YAZMASI VE NADİR ESERLER LİSTESİ" in metin
        assert belgeler.MD12_2 in metin
        assert "El Yazması Divan" in metin
        assert "E-77 sayılı yazıyla" in metin
        assert BASKAN in metin and UYELER[0] in metin
        assert "imha" not in metin.casefold()

    def test_karar_baglanmamis_liste_bos_imza_satirlari(self) -> None:
        liste = rare_works.create_submission()
        rare_works.add_copies(liste, copy_ids=[_kitap(is_rare_or_manuscript=True).pk])
        baglam = belgeler.rare_works_context(liste)

        assert {"label": "Seçim ve Ayıklama Komisyonu kararı", "value": "Bağlanmadı"} in baglam[
            "info"
        ]
        assert [k["role"] for s in baglam["signature_rows"] for k in s] == [
            "Komisyon başkanı",
            "Üye",
            "Üye",
        ]

    def test_bos_liste_basilmaz(self) -> None:
        liste = rare_works.create_submission()
        with pytest.raises(Exception, match="Listede eser yok"):
            belgeler.rare_works_pdf(liste)


class TestBagisOnKayitListesi:
    def test_liste_bagisci_ve_bos_karar_sutunu(self) -> None:
        kayit = donations.create_intake(
            donor_name=BAGISCI,
            received_date=timezone.localdate(),
            items=[
                {"title": "Bağış Eseri Bir", "authors": "Deneme Yazar", "copies": 2},
                {"title": "Bağış Eseri İki", "isbn": "9789750718533"},
            ],
        )
        metin = _metin(belgeler.donation_intake_pdf(kayit))

        assert "BAĞIŞ ÖN KAYIT LİSTESİ" in metin
        assert BAGISCI in metin
        assert belgeler.MD10_3 in metin
        assert "Kabul / Ret" in metin
        assert "2 kaynak, 3 nüsha" in metin
        assert "Hazırlayan — Kütüphane yöneticisi" in metin

    def test_karar_sonucu_listeye_yazilmaz(self) -> None:
        kayit = donations.create_intake(
            received_date=timezone.localdate(),
            items=[{"title": "Kabul Edilen"}, {"title": "Reddedilen"}],
        )
        kalemler = list(kayit.items.order_by("pk"))
        donations.apply_decision(
            kayit,
            commission_decision=ayiklama_karari(decision_type="DONATION_REVIEW"),
            accepted_ids=[kalemler[0].pk],
            rejected={kalemler[1].pk: "Deneme ret gerekçesi"},
        )
        metin = _metin(belgeler.donation_intake_pdf(kayit))

        assert "Deneme ret gerekçesi" not in metin
        assert "Kabul edildi" not in metin and "Reddedildi" not in metin
        assert "Karar işlendi" in metin


def _kararli_bagis(**alanlar: Any) -> Any:
    """Bir kabul, bir ret: kararı uygulanmış bağış ön kaydı (tarihli ve sayılı karar)."""
    alanlar.setdefault("donor_name", BAGISCI)
    kayit = donations.create_intake(
        received_date=timezone.localdate(),
        items=alanlar.pop(
            "items",
            [
                {"title": "Kabul Edilen Eser", "authors": "Deneme Yazar", "copies": 2},
                {"title": "Reddedilen Eser", "authors": "Başka Yazar"},
            ],
        ),
        **alanlar,
    )
    kalemler = list(kayit.items.order_by("pk"))
    donations.apply_decision(
        kayit,
        commission_decision=ayiklama_karari(decision_type="DONATION_REVIEW", decision_no="2026/7"),
        accepted_ids=[kalemler[0].pk],
        rejected={k.pk: "Deneme ret gerekçesi: yaş düzeyine uygun değil" for k in kalemler[1:]},
    )
    kayit.refresh_from_db()
    return kayit


class TestBagisDegerlendirmeSonucu:
    """25.09.2026 kullanıcı kararı (tasarım F8 ekleri 13): kararın ardından kalem kalem döküm;
    VİF'e (TMY 16/1) dayanak ve bağışçıya bilgi. "Bağış kabul tutanağı" TMY'de yoktur."""

    def test_karar_uygulanmadan_basilmaz(self) -> None:
        kayit = donations.create_intake(
            received_date=timezone.localdate(), items=[{"title": "Bekleyen"}]
        )
        assert belgeler.donation_result_blocker(kayit) == belgeler.DONATION_NOT_DECIDED_MESSAGE
        with pytest.raises(Exception, match="Komisyon kararı uygulandıktan sonra"):
            belgeler.donation_result_pdf(kayit)
        donations.cancel_intake(kayit, reason="Geri verildi")
        kayit.refresh_from_db()
        assert belgeler.donation_result_blocker(kayit) == belgeler.DONATION_NOT_DECIDED_MESSAGE

    def test_kabul_ret_gerekce_karar_ve_bagisci(self) -> None:
        kayit = _kararli_bagis()
        assert belgeler.donation_result_blocker(kayit) == ""
        metin = _metin(belgeler.donation_result_pdf(kayit))
        karar = kayit.commission_decision

        assert "BAĞIŞ DEĞERLENDİRME SONUCU" in metin
        assert BAGISCI in metin
        assert f"{karar.decision_date:%d.%m.%Y} tarihli ve 2026/7 sayılı" in metin
        assert "KABUL EDİLEN KAYNAKLAR" in metin and "REDDEDİLEN KAYNAKLAR" in metin
        assert "Kabul Edilen Eser" in metin and "Reddedilen Eser" in metin
        assert "Deneme ret gerekçesi: yaş düzeyine uygun değil" in metin
        assert "1 kaynak, 2 nüsha" in metin and "1 kaynak, 1 nüsha" in metin
        assert f"{kayit.acquisition.date:%d.%m.%Y}" in metin
        assert belgeler.MD10_3 in metin
        assert "Hazırlayan — Kütüphane yöneticisi" in metin
        # Resmî belge sanısı doğmaz: "tutanak" demez, TKYS'yi anar.
        assert "tutana" not in metin.casefold()
        assert "TKYS" in metin and "md. 16/1" in metin and "md. 13/2-c" in metin

    def test_hepsi_reddedilince_edinim_yok_ve_bos_kabul_tablosu(self) -> None:
        kayit = donations.create_intake(
            received_date=timezone.localdate(), items=[{"title": "Tek Kalem"}]
        )
        (kalem,) = kayit.items.all()
        donations.apply_decision(
            kayit,
            commission_decision=ayiklama_karari(decision_type="DONATION_REVIEW"),
            accepted_ids=[],
            rejected={kalem.pk: "Koleksiyonda yeterli nüsha var"},
        )
        kayit.refresh_from_db()
        baglam = belgeler.donation_result_context(kayit)
        etiketler = [b["label"] for b in baglam["info"]]

        assert "Edinim tarihi" not in etiketler and "Bağışçı" not in etiketler
        assert baglam["tables"][0]["rows"] == []
        assert baglam["tables"][0]["empty"] == "Komisyon hiçbir kaynağı kabul etmedi."
        assert len(baglam["tables"]) == 2

    def test_tmy_notlari_fikra_metnine_dayanir(self) -> None:
        """Notların sade anlatımı TMY 16/1 ve 13/2-c'nin depodaki metninde (CLAUDE.md §2-13)."""
        tmy = _mevzuat("tasinir-mal-yonetmeligi.md")
        for ifade in (
            "taşınır kayıt yetkilisi tarafından Varlık İşlem Fişi düzenlenerek kayıtlara alınır",
            "Varlık İşlem Fişinin bir nüshası bağış ve yardım edene verilir veya gönderilir",
            "ispat edici bir belge ile değeri belirtilmiş ise bu değer, belli bir değeri yoksa "
            "değer tespit komisyonunca belirlenen değer",
        ):
            assert ifade in tmy, ifade
        assert "bağış kabul tutanağı" not in tmy.casefold()


# ============================================================ sayfa bütçesi (gerçek uzunluk)


class TestSayfaButcesi:
    @pytest.fixture(autouse=True)
    def uzun_okul(self, okul: SchoolConfig) -> None:
        okul.school_name = EN_UZUN_OKUL
        okul.district = "Örnek Uzun Adlı İlçe"
        okul.save()

    def _uzun_karar(self) -> Any:
        return _komisyon_karari(
            chair_name="Deneme Komisyonbaşkanı Uzunsoyadlıoğlu",
            participants_text="\n".join(UYELER),
            decision_no="2027/123456789",
        )

    def test_tek_kalemli_tutanak_ve_dusme_listesi_tek_sayfa(self) -> None:
        batch = teklif()
        kalem_ekle(batch, _kitap(uzun=True), "OBSOLETE")
        onaya_kadar(
            batch,
            decision=self._uzun_karar(),
            tmy_commission_members="\n".join(UYELER[:5]),
            destruction_decided=True,
        )
        for uretici in (
            belgeler.weeding_report_pdf,
            belgeler.write_off_pdf,
            belgeler.destruction_report_pdf,
        ):
            sayfalar = _sayfalar(uretici(_t(batch)))
            assert len(sayfalar) == 1, (uretici.__name__, len(sayfalar))

    def test_uzun_teklif_listesinde_satir_ve_imza_bolunmez(self) -> None:
        batch = teklif()
        for _ in range(30):
            kalem_ekle(batch, _kitap(uzun=True), "WORN")
        sayfalar = _sayfalar(belgeler.proposal_list_pdf(_t(batch)))

        assert 2 <= len(sayfalar) <= 12
        # Başlık satırı her sayfada yinelenir; imza son sayfada, notuyla birlikte.
        assert all("Kaynak adı" in s for s in sayfalar[:-1])
        assert "Teklif eden — Kütüphane yöneticisi" in sayfalar[-1]
        assert "Önerilen yol" in sayfalar[-1]

    def test_uzun_tutanakta_komisyon_imzalari_birlikte(self) -> None:
        batch = teklif()
        for _ in range(12):
            kalem_ekle(batch, _kitap(uzun=True), "LEVEL_MISMATCH", transfer_target=DEVRALAN_OKUL)
        onaya_kadar(batch, decision=self._uzun_karar())
        sayfalar = _sayfalar(belgeler.weeding_report_pdf(_t(batch)))

        son = sayfalar[-1]
        assert "SEÇİM VE AYIKLAMA KOMİSYONU" in son
        assert all(uye in son for uye in UYELER)
        devir = _sayfalar(belgeler.transfer_list_pdf(_t(batch)))
        assert "Devreden — Kütüphane yöneticisi" in devir[-1] and "OLUR" in devir[-1]

    def test_nadir_eser_ve_bagis_listesi_tek_kalemde_tek_sayfa(self) -> None:
        liste = rare_works.create_submission(commission_decision=self._uzun_karar())
        rare_works.add_copies(liste, copy_ids=[_kitap(uzun=True, is_rare_or_manuscript=True).pk])
        assert len(_sayfalar(belgeler.rare_works_pdf(liste))) == 1

        kayit = donations.create_intake(
            donor_name="Deneme Uzunadlı Bağışçı Kültür ve Eğitim Derneği Şubesi",
            received_date=timezone.localdate(),
            items=[
                {
                    "title": EN_UZUN_ESER,
                    "authors": UZUN_YAZAR,
                    "publisher": "Örnek Uzun Adlı Yayınevi Basım Dağıtım",
                    "publish_year": 1987,
                    "isbn": "9789750718533",
                    "copies": 50,
                }
            ],
        )
        assert len(_sayfalar(belgeler.donation_intake_pdf(kayit))) == 1

    def test_bagis_degerlendirme_sonucu_bir_kabul_bir_retle_tek_sayfa(self) -> None:
        """Gerçek uzunlukta: uzun bağışçı, iki uzun künye, 255 karakterlik ret gerekçesi."""
        uzun = {
            "title": EN_UZUN_ESER,
            "authors": UZUN_YAZAR,
            "publisher": "Örnek Uzun Adlı Yayınevi Basım Dağıtım",
            "publish_year": 1987,
            "isbn": "9789750718533",
            "copies": 50,
        }
        kayit = donations.create_intake(
            donor_name="Deneme Uzunadlı Bağışçı Kültür ve Eğitim Derneği Şubesi",
            received_date=timezone.localdate(),
            items=[uzun, {**uzun, "isbn": ""}],
        )
        kabul, ret = list(kayit.items.order_by("pk"))
        karar = self._uzun_karar()
        karar.decision_type = "DONATION_REVIEW"
        karar.save(update_fields=["decision_type"])
        donations.apply_decision(
            kayit,
            commission_decision=karar,
            accepted_ids=[kabul.pk],
            rejected={ret.pk: ("Deneme uzun ret gerekçesi " * 12)[:255]},
            work_links={kabul.pk: None},
        )
        kayit.refresh_from_db()
        sayfalar = _sayfalar(belgeler.donation_result_pdf(kayit))
        assert len(sayfalar) == 1, sayfalar[-1]
        assert "Hazırlayan — Kütüphane yöneticisi" in sayfalar[-1]


# ============================================================ uçlar


@pytest.fixture
def istemci() -> APIClient:
    return APIClient()


class TestUclar:
    def test_belge_listesi_ve_pdf_xlsx_yanitlari(self, istemci: APIClient) -> None:
        batch = teklif()
        kalem_ekle(batch, _kitap(), "LEVEL_MISMATCH", transfer_target=DEVRALAN_OKUL)
        onaya_kadar(batch, decision=_komisyon_karari())
        kok = f"/api/v1/library/weeding-batches/{batch.pk}/documents/"

        liste = istemci.get(kok)
        assert liste.status_code == 200
        assert [b["kind"] for b in liste.json()] == [
            "teklif-listesi",
            "ayiklama-tutanagi",
            "kayittan-dusme",
            "imha-tutanagi",
            "devir-listesi",
        ]
        devir = next(b for b in liste.json() if b["kind"] == "devir-listesi")
        assert devir == {
            "kind": "devir-listesi",
            "title": "Devir listesi",
            "formats": ["pdf", "xlsx"],
            "available": True,
            "reason": "",
        }

        pdf = istemci.get(f"{kok}ayiklama-tutanagi/")
        assert pdf.status_code == 200 and pdf["Content-Type"] == "application/pdf"
        assert pdf["Cache-Control"] == "no-store"
        assert (
            "Ayıklama-tutanağı_" in pdf["Content-Disposition"]
            or "Ay%C4%B1klama" in pdf["Content-Disposition"]
        )
        assert BASKAN not in pdf["Content-Disposition"]

        xlsx = istemci.get(f"{kok}devir-listesi/?kind=xlsx")
        assert xlsx.status_code == 200
        assert xlsx["Content-Type"].startswith("application/vnd.openxmlformats")
        assert b"".join(xlsx.streaming_content)[:2] == b"PK"  # type: ignore[attr-defined]

    def test_basilamayan_belge_400_gerekceyle_bilinmeyen_404(self, istemci: APIClient) -> None:
        batch = teklif()
        kalem_ekle(batch, _kitap(), "WORN")
        kok = f"/api/v1/library/weeding-batches/{batch.pk}/documents/"

        yanit = istemci.get(f"{kok}ayiklama-tutanagi/")
        assert yanit.status_code == 400
        assert yanit.json()["message"] == belgeler.NEEDS_DECISION_MESSAGE
        assert istemci.get(f"{kok}teklif-listesi/?kind=xlsx").status_code == 400
        assert istemci.get(f"{kok}teklif-listesi/?kind=docx").status_code == 400
        assert istemci.get(f"{kok}bilinmeyen/").status_code == 404
        assert istemci.get("/api/v1/library/weeding-batches/99999/documents/").status_code == 404

    def test_nadir_eser_ve_bagis_uclari(self, istemci: APIClient) -> None:
        liste = rare_works.create_submission()
        assert (
            istemci.get(f"/api/v1/library/rare-works-submissions/{liste.pk}/pdf/").status_code
            == 400
        )
        rare_works.add_copies(liste, copy_ids=[_kitap(is_rare_or_manuscript=True).pk])
        yanit = istemci.get(f"/api/v1/library/rare-works-submissions/{liste.pk}/pdf/")
        assert yanit.status_code == 200 and yanit["Content-Type"] == "application/pdf"

        kayit = donations.create_intake(
            donor_name=BAGISCI, received_date=timezone.localdate(), items=[{"title": "Bir"}]
        )
        yanit = istemci.get(f"/api/v1/library/donation-intakes/{kayit.pk}/pdf/")
        assert yanit.status_code == 200 and yanit["Cache-Control"] == "no-store"
        assert (
            "Ba%C4%9F%C4%B1%C5%9F" in yanit["Content-Disposition"]
            or "Bağış" in yanit["Content-Disposition"]
        )
        assert "Bağışçıadı" not in yanit["Content-Disposition"]

    def test_bagis_degerlendirme_sonucu_ucu(self, istemci: APIClient) -> None:
        bekleyen = donations.create_intake(
            received_date=timezone.localdate(), items=[{"title": "Bekleyen"}]
        )
        yanit = istemci.get(f"/api/v1/library/donation-intakes/{bekleyen.pk}/result-pdf/")
        assert yanit.status_code == 400
        assert yanit.json()["message"] == belgeler.DONATION_NOT_DECIDED_MESSAGE

        kayit = _kararli_bagis()
        yanit = istemci.get(f"/api/v1/library/donation-intakes/{kayit.pk}/result-pdf/")
        assert yanit.status_code == 200 and yanit["Content-Type"] == "application/pdf"
        assert yanit["Cache-Control"] == "no-store"
        assert "Bağışçıadı" not in yanit["Content-Disposition"]
        assert istemci.get("/api/v1/library/donation-intakes/99999/result-pdf/").status_code == 404

    def test_gorevli_kipinde_belge_uclari_403(self, istemci: APIClient) -> None:
        batch = teklif()
        kalem_ekle(batch, _kitap(), "WORN")
        kayit = donations.create_intake(received_date=timezone.localdate(), items=[{"title": "B"}])
        KIP.gorevliye_gec()
        for yol in (
            f"/api/v1/library/weeding-batches/{batch.pk}/documents/",
            f"/api/v1/library/weeding-batches/{batch.pk}/documents/teklif-listesi/",
            f"/api/v1/library/donation-intakes/{kayit.pk}/pdf/",
            f"/api/v1/library/donation-intakes/{kayit.pk}/result-pdf/",
        ):
            yanit = istemci.get(yol)
            assert yanit.status_code == 403 and yanit.json()["code"] == "kip_yetkisiz", yol
