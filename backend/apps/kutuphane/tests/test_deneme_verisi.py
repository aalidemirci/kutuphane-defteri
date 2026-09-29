"""Uydurma deneme verisi üreticisi (`scripts/deneme_verisi.py`) — içe aktarıcılarla uyum.

Saha kabul protokolü (`docs/saha-kabulu.md` §1.2) programı gerçek listeler yerine
bu üreticinin dosyalarıyla sınar; protokoldeki "ne görülmeli" sayıları üreticinin
`OZET.txt`'sinden gelir. Bu test iki sözü birlikte tutar (F12 sözleşmesi, S kolu):

1. Dosyalar programın GERÇEK içe aktarıcılarından geçer: e-Okul öğrenci ve
   personel listeleri F1 aktarımından (`apps.okul.services.imports`), katalog F3
   aktarımından (`apps.kutuphane.services.import_service`).
2. OZET'teki sayılar önizlemenin göstereceği sayılardır (önizleme = uygulama).

Varsayılan kapıda veritabanı yazan testler küçük boyutla koşar (aynı üretici, aynı
biçim); tam boyutun (750 öğrenci, 60 personel, 2.000 eser) veritabanısız ayrıştırma
ve plan denetimleri de varsayılan kapıdadır. Tam boyutun uygulanması `yavas`
işaretlidir (`KD_YAVAS=1`).

KVKK (CLAUDE.md §2-12): üretici T.C. kimlik numarası, telefon ve e-posta üretmez;
veri dosyaları depoya girmez (`.gitignore` testi); üretici program modüllerini ve
yeni bir bağımlılığı içe aktarmaz (kaynak taraması).
"""

from __future__ import annotations

import ast
import functools
import importlib.util
import os
import re
import sys
from datetime import date
from pathlib import Path
from types import ModuleType
from typing import Any

import pytest
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from apps.kutuphane import import_schema, keys, selectors
from apps.kutuphane.models import CatalogImportSource, Copy, ResourceType, Section, Work
from apps.kutuphane.services import import_service as ia
from apps.okul import eokul
from apps.okul.excel_ogrenci import read_sheet
from apps.okul.models import MemberKind, Personnel, SchoolYear, Student, StudentStatus
from apps.okul.services import imports

_URETICI = Path("scripts") / "deneme_verisi.py"

#: Varsayılan kapıdaki küçük boyut (DB yazan testler).
KUCUK = {"ogrenci": 60, "personel": 12, "eser": 120}


def _depo_koku() -> Path:
    """Depo kökü: yerelde `depo/backend/...`, konteynerde `/app` + `/repo`."""
    yerel = Path(__file__).resolve().parents[4]
    for kok in (yerel, Path("/repo")):
        if (kok / _URETICI).is_file():
            return kok
    pytest.fail(f"{_URETICI} bulunamadı (depo kökü: {yerel} ya da /repo).")


@functools.cache
def _uretici() -> ModuleType:
    """Üreticiyi dosya yolundan yükler (`scripts/` bir paket değildir)."""
    yol = _depo_koku() / _URETICI
    spec = importlib.util.spec_from_file_location("kd_deneme_verisi", yol)
    assert spec is not None and spec.loader is not None
    modul = importlib.util.module_from_spec(spec)
    # dataclass'lar modülü `sys.modules` üzerinden çözer.
    sys.modules[spec.name] = modul
    spec.loader.exec_module(modul)
    return modul


@pytest.fixture(scope="module")
def kucuk(tmp_path_factory: pytest.TempPathFactory) -> tuple[Any, Path]:
    klasor = tmp_path_factory.mktemp("deneme-kucuk")
    return _uretici().uret(klasor, **KUCUK), klasor


@pytest.fixture(scope="module")
def tam(tmp_path_factory: pytest.TempPathFactory) -> tuple[Any, Path]:
    klasor = tmp_path_factory.mktemp("deneme-tam")
    return _uretici().uret(klasor), klasor


@pytest.fixture
def aktif_yil() -> SchoolYear:
    yil: SchoolYear = SchoolYear.objects.create(
        name="2026-2027",
        start_date=date(2026, 9, 1),
        end_date=date(2027, 6, 30),
        is_active=True,
    )
    return yil


def _hucreler(yol: Path) -> list[list[Any]]:
    """Kitabın bütün sayfalarının bütün hücre değerleri (sayfa sırasıyla)."""
    kitap = load_workbook(yol, read_only=True, data_only=True)
    try:
        return [
            list(satir) for sayfa in kitap.worksheets for satir in sayfa.iter_rows(values_only=True)
        ]
    finally:
        kitap.close()


def _katalog_uygula(ozet: Any, klasor: Path) -> ia.CatalogImportReport:
    U = _uretici()
    parsed, ozet_hash = ia.rows_from_file(
        (klasor / U.DOSYA_KATALOG).read_bytes(), source=CatalogImportSource.EXCEL
    )
    return ia.apply_import(
        parsed, payload_sha256=ozet_hash, file_name=U.DOSYA_KATALOG, new_sections=ozet.bolumler
    )


# ---------------------------------------------------------------------------
# Üreticinin kendisi (veritabanısız)
# ---------------------------------------------------------------------------
def test_uretici_yalniz_standart_kitaplik_ve_openpyxl_kullanir() -> None:
    """Yeni Python bağımlılığı yok (CLAUDE.md §2-10); program modülleri de içe aktarılmaz."""
    agac = ast.parse((_depo_koku() / _URETICI).read_text(encoding="utf-8"))
    kokler: set[str] = set()
    for dugum in ast.walk(agac):
        if isinstance(dugum, ast.Import):
            kokler.update(ad.name.split(".")[0] for ad in dugum.names)
        elif isinstance(dugum, ast.ImportFrom) and dugum.module and dugum.level == 0:
            kokler.add(dugum.module.split(".")[0])
    izinli = set(sys.stdlib_module_names) | {"openpyxl", "__future__"}
    assert kokler <= izinli, kokler - izinli
    assert not kokler & {"apps", "config", "shared", "desktop", "django", "katalog"}


def test_veri_dosyalari_depoya_girmez() -> None:
    """Varsayılan çıktı klasörü ve `*.xlsx` `.gitignore`'dadır; OZET.txt de klasörle kalır."""
    U = _uretici()
    satirlar = {
        satir.strip()
        for satir in (_depo_koku() / ".gitignore").read_text(encoding="utf-8").splitlines()
    }
    assert f"{U.VARSAYILAN_CIKTI}/" in satirlar
    assert "*.xlsx" in satirlar


def test_ayni_tohum_ayni_veriyi_uretir(tmp_path: Path) -> None:
    U = _uretici()
    birinci = U.uret(tmp_path / "bir", tohum=7, **KUCUK)
    ikinci = U.uret(tmp_path / "iki", tohum=7, **KUCUK)
    assert birinci.metin() == ikinci.metin()
    for ad in (
        U.DOSYA_OGRENCI,
        U.DOSYA_OGRENCI_2,
        U.DOSYA_PERSONEL,
        U.DOSYA_KATALOG,
        U.DOSYA_SORUNLU,
    ):
        assert _hucreler(tmp_path / "bir" / ad) == _hucreler(tmp_path / "iki" / ad), ad
    assert (tmp_path / "bir" / U.DOSYA_OZET).read_text(encoding="utf-8") == birinci.metin()


def test_kimlik_numarasi_ve_iletisim_bilgisi_uretilmez(tam: tuple[Any, Path]) -> None:
    """11 haneli sayı (T.C. kimlik numarası biçimi), e-posta ve telefon YOK (KVKK)."""
    _ozet, klasor = tam
    telefon = re.compile(r"(?<!\d)0?5\d{2}[ -]?\d{3}[ -]?\d{2}[ -]?\d{2}(?!\d)")
    for yol in sorted(klasor.glob("*.xlsx")):
        for satir in _hucreler(yol):
            for deger in satir:
                if isinstance(deger, bool) or deger is None:
                    continue
                if isinstance(deger, int | float):
                    assert abs(deger) < 10**10, (yol.name, "11 haneli sayı")
                    continue
                metin = str(deger)
                # ISBN'ler 10-13 hanedir ama TİRESİZ 11 haneli dizi hiç üretilmez.
                assert not re.search(r"(?<!\d)\d{11}(?!\d)", metin), (yol.name, "11 hane")
                assert "@" not in metin, (yol.name, "e-posta")
                if not re.fullmatch(r"97[89][\d-]{10,14}", metin):
                    assert not telefon.search(metin), (yol.name, "telefon")


def test_komut_satiri_calisir_ve_sinirlari_denetler(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    U = _uretici()
    kod = U.main(["--cikti", str(tmp_path), "--ogrenci", "50", "--personel", "10", "--eser", "40"])
    assert kod == 0
    for ad in (U.DOSYA_OGRENCI, U.DOSYA_PERSONEL, U.DOSYA_KATALOG, U.DOSYA_OZET):
        assert (tmp_path / ad).is_file(), ad
    cikti = capsys.readouterr().out
    assert "UYDURMADIR" in cikti and "depoya eklemeyin" in cikti
    with pytest.raises(SystemExit) as hata:
        U.main(["--cikti", str(tmp_path), "--eser", "5"])
    assert hata.value.code == 2


def test_arama_katlamasi_programin_katlamasiyla_ayni(tam: tuple[Any, Path]) -> None:
    """OZET'in arama sayıları programın katlamasıyla aynı kuralla hesaplanır."""
    U = _uretici()
    _ozet, klasor = tam
    grid = read_sheet((klasor / U.DOSYA_KATALOG).read_bytes(), sheet_name=U.KATALOG_SAYFASI)
    for satir in grid[1:]:
        for hucre in satir[:3] + satir[7:8]:  # eser adı, yazar, çevirmen, konu
            if isinstance(hucre, str):
                assert U.arama_katlamasi(hucre) == keys.fold_search(hucre), hucre
    for kucuk_yazim, buyuk_yazim in U.ARAMA_SINAMALARI:
        assert U.arama_katlamasi(kucuk_yazim) == U.arama_katlamasi(buyuk_yazim)


# ---------------------------------------------------------------------------
# e-Okul listeleri (F1 aktarımı)
# ---------------------------------------------------------------------------
def test_ogrenci_listesi_e_okul_yerlesimindedir(kucuk: tuple[Any, Path]) -> None:
    U = _uretici()
    ozet, klasor = kucuk
    grid = read_sheet((klasor / U.DOSYA_OGRENCI).read_bytes())
    assert eokul.sinif_listesi_mi(grid)
    bloklar = [eokul.blok_sinifi(satir) for satir in grid if eokul.blok_sinifi(satir)]
    assert bloklar == ozet.subeler
    assert {"9/I", "9/İ"} <= set(bloklar), "I ve İ şubeleri ayrı iki bloktur"
    personel = read_sheet((klasor / U.DOSYA_PERSONEL).read_bytes())
    assert any(eokul.sayac_veya_kod_mu(satir) for satir in personel), "sayaç dipnotu"


@pytest.mark.django_db
def test_ogrenci_listesi_aktarilir(kucuk: tuple[Any, Path], aktif_yil: SchoolYear) -> None:
    U = _uretici()
    ozet, klasor = kucuk
    veri = (klasor / U.DOSYA_OGRENCI).read_bytes()

    onizleme = imports.preview_students_file(file_bytes=veri, file_name=U.DOSYA_OGRENCI)
    assert onizleme.dry_run and onizleme.created_students == ozet.ogrenci
    assert Student.objects.count() == 0

    rapor = imports.commit_students_file(file_bytes=veri, file_name=U.DOSYA_OGRENCI)
    assert rapor.created_students == ozet.ogrenci
    assert rapor.skipped == []
    assert [u for u in rapor.warnings if u.field != "header"] == [], "satır uyarısı beklenmez"
    assert any(f"{len(ozet.subeler)} şube bloğu" in u.issue for u in rapor.warnings)
    assert {o.class_label for o in Student.objects.all()} == set(ozet.subeler)


@pytest.mark.django_db
def test_ikinci_donem_listesi_havuz_ve_degisiklikleri_verir(
    kucuk: tuple[Any, Path], aktif_yil: SchoolYear
) -> None:
    U = _uretici()
    ozet, klasor = kucuk
    imports.commit_students_file(
        file_bytes=(klasor / U.DOSYA_OGRENCI).read_bytes(), file_name=U.DOSYA_OGRENCI
    )
    veri = (klasor / U.DOSYA_OGRENCI_2).read_bytes()
    rapor = imports.preview_students_file(file_bytes=veri, file_name=U.DOSYA_OGRENCI_2)
    assert rapor.created_students == ozet.gelen
    assert rapor.updated_students == ozet.sube_degisen
    assert rapor.pool_added_students == ozet.ayrilan
    assert rapor.unchanged_students == ozet.ogrenci - ozet.ayrilan - ozet.sube_degisen
    assert rapor.skipped == []

    imports.commit_students_file(file_bytes=veri, file_name=U.DOSYA_OGRENCI_2)
    havuz = Student.objects.filter(leave_candidate_since__isnull=False)
    assert havuz.count() == ozet.ayrilan
    # Aktarım kimseyi ayırmaz ve silmez (F1 ekleri 7).
    assert Student.objects.filter(status=StudentStatus.ACTIVE).count() == ozet.ogrenci + ozet.gelen


@pytest.mark.django_db
def test_personel_listesi_aktarilir(kucuk: tuple[Any, Path]) -> None:
    U = _uretici()
    ozet, klasor = kucuk
    rapor = imports.commit_personnel_file(
        file_bytes=(klasor / U.DOSYA_PERSONEL).read_bytes(), file_name=U.DOSYA_PERSONEL
    )
    assert rapor.created_personnel == ozet.personel
    assert rapor.skipped == []
    assert Personnel.objects.filter(member_kind=MemberKind.TEACHER).count() == ozet.ogretmen
    assert Personnel.objects.filter(member_kind=MemberKind.STAFF).count() == ozet.diger_personel
    uyarilar = [u.row_number for u in rapor.warnings if u.field == "member_kind"]
    assert uyarilar == [ozet.taninmayan_gorev_satiri]
    # Sayaç dipnotu kişi olarak yazılmadı.
    assert not any("Sayısı" in p.full_name for p in Personnel.objects.all())


# ---------------------------------------------------------------------------
# Katalog (F3 aktarımı)
# ---------------------------------------------------------------------------
def test_katalog_basliklari_sablonla_birebir(kucuk: tuple[Any, Path]) -> None:
    """Üretici şablonun başlıklarını elle yazar; şablon değişirse burası kırmızıya döner."""
    U = _uretici()
    _ozet, klasor = kucuk
    assert U.KATALOG_SAYFASI == import_schema.CATALOG_SHEET
    assert list(U.KATALOG_BASLIKLARI) == [sutun.header for sutun in import_schema.COLUMNS]
    grid = read_sheet((klasor / U.DOSYA_KATALOG).read_bytes(), sheet_name=U.KATALOG_SAYFASI)
    assert grid[0] == list(U.KATALOG_BASLIKLARI)


def test_tam_boy_katalog_temiz_ayrisir(tam: tuple[Any, Path]) -> None:
    """2.000 eserlik dosyanın her satırı içe aktarılabilir; uyarı ve tanınmayan sütun yok."""
    U = _uretici()
    ozet, klasor = tam
    parsed, _ozet_hash = ia.rows_from_file(
        (klasor / U.DOSYA_KATALOG).read_bytes(), source=CatalogImportSource.EXCEL
    )
    assert len(parsed.rows) == ozet.katalog_satir, "“Örnek” sayfası okunmamalı"
    assert parsed.unknown_headers == [] and parsed.missing_columns == []
    sorunlu = [(r.row_number, r.errors, r.warnings) for r in parsed.rows if r.errors or r.warnings]
    assert sorunlu == []
    assert sum(r.copies for r in parsed.rows) == ozet.nusha
    assert sum(1 for r in parsed.rows if r.reference_by_textbook) == ozet.ders_kitabi_satiri
    assert sum(1 for r in parsed.rows if r.resource_type == "PERIODICAL") == ozet.sureli_eser
    assert all(r.is_bound_periodical for r in parsed.rows if r.resource_type == "PERIODICAL")
    assert sum(1 for r in parsed.rows if r.old_register_no) == ozet.eski_kayitli_nusha
    assert ozet.eser == U.VARSAYILAN_ESER


@pytest.mark.django_db
def test_tam_boy_katalog_plani_ozetle_ayni(tam: tuple[Any, Path]) -> None:
    """Eşleşme planı (yazmadan): N yeni eser, bölünmüş eserin öbür satırları mevcut esere.

    Tam boyutun uygulanması `yavas` testindedir; plan, önizlemenin ve uygulamanın
    kovaları için kullandığı kodun kendisidir (`import_service._plan`).
    """
    U = _uretici()
    ozet, klasor = tam
    parsed, _ozet_hash = ia.rows_from_file(
        (klasor / U.DOSYA_KATALOG).read_bytes(), source=CatalogImportSource.EXCEL
    )
    plan = ia._plan(parsed.rows, decisions={}, section_map={}, new_sections=ozet.bolumler)
    assert sum(1 for p in plan.rows if p.creates_work) == ozet.eser
    kovalar = [p.report.bucket for p in plan.rows]
    assert kovalar.count(ia.BUCKET_EXISTING) == ozet.bolunmus_ek_satir
    assert ia.BUCKET_SUSPECT not in kovalar and ia.BUCKET_SKIPPED not in kovalar
    assert plan.pending_decisions == [] and plan.unknown_sections == []
    assert ozet.bolunmus_ek_satir > 0, "tam boyutta aynı eserin ayrı satırları bulunmalı"


@pytest.mark.django_db
def test_katalog_onizleme_ve_uygulama_ozetle_ayni(kucuk: tuple[Any, Path]) -> None:
    U = _uretici()
    ozet, klasor = kucuk
    parsed, ozet_hash = ia.rows_from_file(
        (klasor / U.DOSYA_KATALOG).read_bytes(), source=CatalogImportSource.EXCEL
    )
    # Bölüm kararı verilmeden önizleme bütün bölümleri "listede yok" diye sorar.
    ilk = ia.preview_import(parsed, payload_sha256=ozet_hash)
    assert sorted(k["value"] for k in ilk.unknown_sections) == sorted(ozet.bolumler)

    onizleme = ia.preview_import(parsed, payload_sha256=ozet_hash, new_sections=ozet.bolumler)
    beklenen = {
        "new_works": ozet.eser,
        "existing_matches": ozet.bolunmus_ek_satir,
        "copies_created": ozet.nusha,
        "suspect": 0,
        "skipped_rows": 0,
        "error_rows": 0,
        "reference_defaults": ozet.ders_kitabi_satiri,
        "sections_created": len(ozet.bolumler),
    }
    assert {k: onizleme.stats[k] for k in beklenen} == beklenen
    assert Work.objects.count() == 0, "önizleme yazdı"

    rapor = _katalog_uygula(ozet, klasor)
    assert rapor.stats == onizleme.stats
    assert Work.objects.count() == ozet.eser
    assert Copy.objects.count() == ozet.nusha
    assert Section.objects.count() == len(ozet.bolumler)
    assert Work.objects.filter(resource_type=ResourceType.PERIODICAL).count() == ozet.sureli_eser
    assert Work.objects.filter(resource_type=ResourceType.AV_MATERIAL).count() == ozet.gorsel_eser
    assert Copy.objects.exclude(old_register_no="").count() == ozet.eski_kayitli_nusha
    # Bölüm başına nüsha (protokol §8.3 etiket partisini bununla seçer — F12 düzeltme turu).
    gercek = {
        bolum.name: Copy.objects.filter(section=bolum).count() for bolum in Section.objects.all()
    }
    assert gercek == ozet.bolum_nusha
    assert sum(ozet.bolum_nusha.values()) == ozet.nusha
    ozet_metni = (klasor / U.DOSYA_OZET).read_text(encoding="utf-8")
    assert "Bölüm başına nüsha" in ozet_metni
    for ad, sayi in ozet.bolum_nusha.items():
        assert f"{ad} {sayi}" in ozet_metni, ad

    # Aynı dosya ikinci kez uygulanamaz (engel).
    with pytest.raises(ValidationError, match="zaten"):
        _katalog_uygula(ozet, klasor)

    # Protokolün Türkçe arama sayıları (OZET) programın aramasıyla aynı.
    for kucuk_yazim, buyuk_yazim in U.ARAMA_SINAMALARI:
        beklenen_sayi = ozet.arama[buyuk_yazim]
        assert selectors.works(q=kucuk_yazim).count() == beklenen_sayi, kucuk_yazim
        assert selectors.works(q=buyuk_yazim).count() == beklenen_sayi, buyuk_yazim


@pytest.mark.django_db
def test_sorunlu_satirlar_beklenen_sonuclari_verir(kucuk: tuple[Any, Path]) -> None:
    """Sınama dosyası katalog uygulandıktan sonra önizlenir; her satır OZET'te yazdığı gibi."""
    U = _uretici()
    ozet, klasor = kucuk
    _katalog_uygula(ozet, klasor)

    parsed, ozet_hash = ia.rows_from_file(
        (klasor / U.DOSYA_SORUNLU).read_bytes(), source=CatalogImportSource.EXCEL
    )
    assert parsed.header_row == U.SORUNLU_BASLIK_SATIRI
    assert sorted(parsed.unknown_headers) == ["Fiyat", "Sıra"]
    assert parsed.missing_columns == []
    rapor = ia.preview_import(parsed, payload_sha256=ozet_hash)
    satirlar = {r.row: r for r in rapor.rows}
    assert sorted(satirlar) == [s.satir_no for s in ozet.sorunlu]

    for beklenen in ozet.sorunlu:
        r = satirlar[beklenen.satir_no]
        durum = (beklenen.satir_no, beklenen.beklenen, r.bucket, r.issues)
        if beklenen.beklenen == "mevcut":
            assert r.bucket == ia.BUCKET_EXISTING and r.copies_created == 1, durum
        elif beklenen.beklenen == "supheli":
            assert r.bucket == ia.BUCKET_SUSPECT and r.needs_decision, durum
        elif beklenen.beklenen == "aktarilmadi":
            assert r.bucket == ia.BUCKET_SKIPPED and r.issues, durum
        elif beklenen.beklenen == "yazilamadi":
            assert r.bucket == ia.BUCKET_NEW and r.copies_created == 0 and r.issues, durum
        else:  # "yeni" ve "bolum"
            assert r.bucket == ia.BUCKET_NEW and r.copies_created == 1, durum
            assert bool(r.issues) == beklenen.uyari, durum

    bolum_satirlari = [s.satir_no for s in ozet.sorunlu if s.beklenen == "bolum"]
    assert [(k["value"], k["rows"]) for k in rapor.unknown_sections] == [
        ("Gezi Rafı", bolum_satirlari)
    ]
    say = {kod: sum(1 for s in ozet.sorunlu if s.beklenen == kod) for kod in ("mevcut", "supheli")}
    assert rapor.stats["existing_matches"] == say["mevcut"]
    assert rapor.stats["suspect"] == say["supheli"]
    assert rapor.stats["skipped_rows"] == sum(
        1 for s in ozet.sorunlu if s.beklenen == "aktarilmadi"
    )
    assert rapor.stats["error_rows"] == sum(1 for s in ozet.sorunlu if s.beklenen == "yazilamadi")
    assert rapor.stats["reference_defaults"] == 1


# ---------------------------------------------------------------------------
# Tam boy uygulama (gecelik kapı)
# ---------------------------------------------------------------------------
@pytest.mark.yavas
@pytest.mark.skipif(os.environ.get("KD_YAVAS") != "1", reason="KD_YAVAS=1 ile koşar.")
@pytest.mark.django_db
def test_tam_boy_deneme_verisi_uctan_uca_aktarilir(
    tam: tuple[Any, Path], aktif_yil: SchoolYear
) -> None:
    """750 öğrenci + 60 personel + 2.000 eser: protokoldeki sayılar aynen gerçekleşir."""
    U = _uretici()
    ozet, klasor = tam
    ogrenci = imports.commit_students_file(
        file_bytes=(klasor / U.DOSYA_OGRENCI).read_bytes(), file_name=U.DOSYA_OGRENCI
    )
    assert ogrenci.created_students == ozet.ogrenci and ogrenci.skipped == []
    donem2 = imports.preview_students_file(
        file_bytes=(klasor / U.DOSYA_OGRENCI_2).read_bytes(), file_name=U.DOSYA_OGRENCI_2
    )
    assert (donem2.created_students, donem2.updated_students, donem2.pool_added_students) == (
        ozet.gelen,
        ozet.sube_degisen,
        ozet.ayrilan,
    )
    personel = imports.commit_personnel_file(
        file_bytes=(klasor / U.DOSYA_PERSONEL).read_bytes(), file_name=U.DOSYA_PERSONEL
    )
    assert personel.created_personnel == ozet.personel
    rapor = _katalog_uygula(ozet, klasor)
    assert rapor.stats["new_works"] == ozet.eser
    assert rapor.stats["copies_created"] == ozet.nusha
    assert Copy.objects.count() == ozet.nusha
