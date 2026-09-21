"""Katalog şablonu indirme ucu — URL/ad OYS'den korunur, dosya adı belge adı + tarih."""

from __future__ import annotations

from collections.abc import Iterable
from datetime import date
from io import BytesIO
from typing import cast
from urllib.parse import unquote

import pytest
from django.http import StreamingHttpResponse
from django.urls import resolve, reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.kutuphane import views
from apps.kutuphane.import_schema import CATALOG_SHEET, COLUMNS, COLUMNS_SHEET, EXAMPLE_SHEET

XLSX_TURU = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
YOL = "/api/v1/library/import/template/"


def _govde(yanit: object) -> bytes:
    assert isinstance(yanit, StreamingHttpResponse)
    return b"".join(cast(Iterable[bytes], yanit.streaming_content))


def test_url_adi_ve_oneki_oys_ile_ayni() -> None:
    assert reverse("library-import-template") == YOL
    assert getattr(resolve(YOL).func, "view_class", None) is views.CatalogImportTemplateView


@pytest.mark.django_db
def test_sablon_ucu_excel_eki_dondurur(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(timezone, "localdate", lambda: date(2026, 9, 21))

    yanit = APIClient().get(YOL)

    assert yanit.status_code == 200
    assert yanit["Content-Type"] == XLSX_TURU
    disposition = unquote(yanit["Content-Disposition"])
    assert disposition.startswith("attachment")
    assert "Katalog-Excel-Şablonu_21.09.2026.xlsx" in disposition
    kitap = load_workbook(BytesIO(_govde(yanit)))
    assert kitap.sheetnames == [CATALOG_SHEET, COLUMNS_SHEET, EXAMPLE_SHEET]
    assert [h.value for h in kitap[CATALOG_SHEET][1]] == [k.header for k in COLUMNS]


@pytest.mark.django_db
def test_sablon_ucu_yalniz_okur() -> None:
    """Uç yalnız GET kabul eder; şablon indirmek hiçbir şey yazmaz."""
    yanit = APIClient().post(YOL, {}, format="json")
    assert yanit.status_code == 405
