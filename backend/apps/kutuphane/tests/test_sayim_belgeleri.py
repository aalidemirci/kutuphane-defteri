"""Sayım tutanağı (E10) — PDF ve XLSX (F9; tasarım §10 E10, §14.1 F9 kod kapısı).

Kapılar:
- basılabilirlik sunucudadır: taslakta ve iptal edilmiş sayımda basılmaz; sürerken ara
  döküm ("TASLAK"), tamamlanınca imzaya (OLUR boş), onaydan sonra onayla basılır;
- iki seçenek (TMY 32/3 durdurması ve sayım için hizmet arası) ve iade AYRI satırlarda;
- sonuçlar (bulunan, noksan, fazla), noksan düşüm teklifi 32/7, hasar düşümü 27/1 ve
  kayıp/hasar tutanağına bağlı liste, kayıpta görünüp bulunanlar, sayım fazlası (TMY 17);
- ödünçteki ve teslimdeki nüsha için kurulun seçimi ve dayanağı ("32/5'e kıyasen; 23/4");
- **ödünç alanın kimliği YOK** — sentetik adla PDF ve XLSX metin taraması;
- ek "Taşınır Sayım ve Döküm Cetveline aktarılacak sayılar" (TMY 34/1) yeni sayfada, "Bu
  döküm Taşınır Sayım ve Döküm Cetveli değildir" ibaresiyle;
- mevzuat alıntıları ve sade anlatımların dayandığı ifadeler `docs/mevzuat/` metninde;
- sayfa bütçesi GERÇEK UZUNLUKTA veriyle (CLAUDE.md §3): en uzun okul adı, sekiz kişilik
  kurul, çeviri künyeli yazar, 64 karakterlik TKYS kodu;
- uçlar yalnız yönetici kipinde; indirme adında kişi adı yok.

Bütün kişi ve kurum adları uydurmadır ("Deneme …", "Örnek …").
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

import pytest
from django.core.exceptions import ValidationError
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader
from rest_framework.test import APIClient

from apps.kutuphane import barcode as barcode_module
from apps.kutuphane import sayim_belgeleri as belgeler
from apps.kutuphane.models import CaseResolution, Copy, CountBasis, StockTake
from apps.kutuphane.services import loss_damage, stocktake
from apps.kutuphane.tests.dolasim_ortak import odunc_nushasi, odunc_ver, ogrenci, uye
from apps.kutuphane.tests.ortak import eser, nusha
from apps.kutuphane.tests.sayim_ortak import (
    HARCAMA_YETKILISI,
    KURUL,
    baslat,
    durdurma_alanlari,
    kalem,
    okut,
    onayla,
    tamamla,
    taslak,
    tazele_sayim,
)
from apps.kutuphane.tests.teslim_ortak import ogretmen, sube, teslim_et
from apps.okul.kip import KIP
from apps.okul.models import SchoolConfig

pytestmark = pytest.mark.django_db

EN_UZUN_OKUL = ("Örnek Mesleki ve Teknik Anadolu Lisesi Uzun Adlı Kampüs Yerleşkesi " * 3)[:255]
EN_UZUN_ESER = (
    "Uzun Adlı Örnek Kaynak: Çağdaş Türk Şiirinde Doğa, Şehir ve İnsan Üzerine "
    "Karşılaştırmalı Bir İnceleme — Birinci Cilt, Genişletilmiş Üçüncü Baskı"
)
UZUN_YAZAR = (
    "Deneme Yazaroğlu; Hazırlayan: Örnek Derleyicioğlu, İkinci Hazırlayıcı; "
    "Çeviren: Üçüncü Çevirmenoğlu ve Dördüncü Çevirmen; Resimleyen: Beşinci Çizeroğlu"
)
UZUN_TKYS = "255.01.02.03.04.05.06.07.08.09-" + "9" * 33
#: Sekiz kişilik kurul (TMY 32/2 en az üç kişi): başkan + taşınır kayıt yetkilisi + altı üye.
UZUN_KURUL: dict[str, str] = {
    "committee_chair": "Deneme Kurulbaşkanı Uzunsoyadlıoğlu",
    "committee_property_officer": "Deneme Taşınırkayıtyetkilisi Uzunsoyadlıoğlu",
    "committee_members": "\n".join(
        f"Deneme {sira}kurulüyesi Uzunsoyadlıoğlu"
        for sira in ("Birinci", "İkinci", "Üçüncü", "Dördüncü", "Beşinci", "Altıncı")
    ),
}
#: Sentetik kişi — tutanağın hiçbir yerinde görünmemeli (E10 kod kapısı).
ODUNC_ALAN_AD = "Denemesayimbelge"
ODUNC_ALAN_SOYAD = "Kimlikyokoglu"
TESLIM_ALAN_AD = "Denemeteslimsayim"

_T_ARALIGI = re.compile(r"T (?=[a-zçğıöşü])")


def _metin(icerik: bytes) -> str:
    ham = "\n".join(s.extract_text() or "" for s in PdfReader(io.BytesIO(icerik)).pages)
    return " ".join(_T_ARALIGI.sub("T", ham).split())


def _sayfalar(icerik: bytes) -> list[str]:
    return [
        " ".join(_T_ARALIGI.sub("T", s.extract_text() or "").split())
        for s in PdfReader(io.BytesIO(icerik)).pages
    ]


def _xlsx_metni(icerik: bytes) -> str:
    kitap = load_workbook(io.BytesIO(icerik))
    parcalar: list[str] = []
    for sayfa in kitap.worksheets:
        parcalar.append(sayfa.title)
        for satir in sayfa.iter_rows(values_only=True):
            parcalar.extend(str(h) for h in satir if h is not None)
    return " ".join(parcalar)


def _oku(yol: Path) -> str:
    yerel = Path(__file__).resolve().parents[4]
    for kok in (yerel, Path("/repo")):
        dosya = kok / yol
        if dosya.is_file():
            return dosya.read_text(encoding="utf-8")
    pytest.fail(f"{yol} bulunamadı.")


_TMY = Path("docs") / "mevzuat" / "tasinir-mal-yonetmeligi.md"


def _fikra(madde: int, fikra: int) -> str:
    metin = _oku(_TMY)
    cikis = re.search(
        rf'<a id="madde-{madde}"></a>(.*?)(?=<a id="madde-|\Z)', metin, flags=re.DOTALL
    )
    assert cikis is not None, f"TMY md. {madde} yok"
    govde = " ".join(cikis.group(1).split())
    bolum = re.search(rf"\({fikra}\) (.*?)(?= \({fikra + 1}\) |\Z)", govde)
    assert bolum is not None, f"TMY md. {madde}/{fikra} yok"
    return bolum.group(1)


def _sablon(ad: str) -> str:
    yerel = Path(__file__).resolve().parents[3] / "templates" / "documents" / ad
    for yol in (yerel, Path("/app/templates/documents") / ad):
        if yol.is_file():
            return yol.read_text(encoding="utf-8")
    pytest.fail(f"{ad} bulunamadı.")


@pytest.fixture(autouse=True)
def okul() -> SchoolConfig:
    config = SchoolConfig.load()
    config.school_name = "Örnek Anadolu Lisesi"
    config.district = "Örnek İlçe"
    config.principal_name = "Örnek Müdür"
    config.save()
    return config


def _t(sayim: StockTake) -> StockTake:
    return tazele_sayim(sayim)


def _tablo(baglam: dict[str, Any], baslik_parcasi: str) -> dict[str, Any]:
    for t in baglam["tables"]:
        if baslik_parcasi in t["title"]:
            bulunan: dict[str, Any] = t
            return bulunan
    raise AssertionError(f"{baslik_parcasi} tablosu yok: {[t['title'] for t in baglam['tables']]}")


def _hucreler(tablo: dict[str, Any]) -> list[list[str]]:
    return [[h["v"] for h in satir] for satir in tablo["rows"]]


def _uzun_kitap(**alanlar: Any) -> Copy:
    return nusha(
        eser(
            title=EN_UZUN_ESER,
            authors=UZUN_YAZAR,
            publisher="Örnek Uzun Adlı Yayınevi Basım Dağıtım",
            publish_year=1987,
        ),
        external_asset_ref=UZUN_TKYS,
        **alanlar,
    )


# ============================================================ basılabilirlik


class TestBasilabilirlik:
    def test_taslakta_ve_iptal_edilmis_sayimda_basilmaz(self) -> None:
        sayim = taslak()
        assert belgeler.stocktake_document_blocker(sayim) == belgeler.DRAFT_MESSAGE
        assert belgeler.stocktake_documents(sayim) == [
            {
                "kind": "sayim-tutanagi",
                "title": "Sayım tutanağı",
                "formats": ["pdf", "xlsx"],
                "available": False,
                "reason": belgeler.DRAFT_MESSAGE,
            }
        ]
        with pytest.raises(ValidationError, match="Sayım başladıktan sonra"):
            belgeler.stocktake_report_pdf(sayim)
        stocktake.cancel_stocktake(sayim)
        assert belgeler.stocktake_document_blocker(_t(sayim)) == belgeler.CANCELLED_MESSAGE
        with pytest.raises(ValidationError, match="İptal edilmiş"):
            belgeler.stocktake_report_xlsx(_t(sayim))

    def test_surerken_tamamlaninca_ve_onaydan_sonra_basilir(self) -> None:
        odunc_nushasi(title="Sayılan")
        sayim = baslat()
        for asama in ("suruyor", "tamamlandi", "onaylandi"):
            if asama == "tamamlandi":
                tamamla(sayim)
            elif asama == "onaylandi":
                onayla(_t(sayim))
            guncel = _t(sayim)
            assert belgeler.stocktake_document_blocker(guncel) == "", asama
            assert belgeler.stocktake_documents(guncel)[0]["available"] is True

    def test_bilinmeyen_belge_ve_bicim(self) -> None:
        sayim = baslat()
        with pytest.raises(ValidationError, match="Böyle bir sayım belgesi yok"):
            belgeler.stocktake_document(sayim, "bilinmeyen")
        with pytest.raises(ValidationError, match="bu biçimde üretilmez"):
            belgeler.stocktake_document(sayim, belgeler.SAYIM_TUTANAGI, "docx")
        assert belgeler.stocktake_document_filename(
            belgeler.SAYIM_TUTANAGI, "xlsx", timezone.localdate()
        ) == (f"Sayım-tutanağı_{timezone.localdate():%d.%m.%Y}.xlsx")


# ============================================================ içerik


class TestTutanakIcerigi:
    def test_kurul_tarih_ve_iki_secenek_ayri_satirlarda(self) -> None:
        odunc_nushasi(title="Rafta Bulunan")
        sayim = baslat(
            **durdurma_alanlari(), service_pause=True, service_pause_decision="2026/15 sayılı karar"
        )
        baglam = belgeler.stocktake_report_context(_t(sayim))
        secenekler = _tablo(baglam, "SAYIM SIRASINDAKİ SEÇENEKLER")
        satirlar = _hucreler(secenekler)

        # Kod kapısı: iki seçenek ve iade ÜÇ AYRI satırda.
        assert [s[0] for s in satirlar] == [
            "TMY 32/3 durdurması",
            "Sayım için hizmet arası",
            "İade",
        ]
        assert [s[1] for s in satirlar] == ["Seçildi", "Seçildi", "Her zaman açık"]
        assert satirlar[0][3] == "TMY 32/3"
        assert f"Durduran harcama yetkilisi: {HARCAMA_YETKILISI}." in satirlar[0][2]
        assert "2026/15 sayılı karar" in satirlar[1][2] and "yeni ödünç" in satirlar[1][2]
        assert "23/1-c" in satirlar[2][3]

        metin = _metin(belgeler.stocktake_report_pdf(_t(sayim)))
        assert "SAYIM TUTANAĞI" in metin and "SAYIM KURULU" in metin
        for ad in KURUL.values():
            assert ad in metin, ad
        for rol in ("Sayım kurulu başkanı", "Taşınır kayıt yetkilisi", "Üye"):
            assert rol in metin
        assert f"{timezone.localtime(sayim.started_at):%d.%m.%Y}" in metin
        assert belgeler.SURUYOR_NOTU in metin
        # "sayım kilidi" ve "dondurma" denmez (sözlük §1).
        assert "kilidi" not in metin.casefold() and "dondur" not in metin.casefold()

    def test_secilmeyen_secenekler_de_satirda_durur(self) -> None:
        sayim = baslat()
        satirlar = _hucreler(
            _tablo(belgeler.stocktake_report_context(sayim), "SAYIM SIRASINDAKİ SEÇENEKLER")
        )
        assert [s[1] for s in satirlar] == ["Seçilmedi", "Seçilmedi", "Her zaman açık"]
        assert "Durduran harcama yetkilisi" not in satirlar[0][2]

    def test_sonuclar_ve_32_7_noksan_teklifi(self) -> None:
        bulunan = odunc_nushasi(title="Bulunan Kitap")
        noksan = nusha(eser(title="Bulunamayan Kitap"), external_asset_ref="255.01.02")
        sayim = baslat()
        okut(sayim, bulunan)
        stocktake.add_surplus(sayim, note="Etiketsiz kitap, 3. raf", work=eser(title="Fazla Eser"))
        tamamla(sayim)
        guncel = _t(sayim)
        baglam = belgeler.stocktake_report_context(guncel)

        sonuclar = {s[0]: s[1] for s in _hucreler(_tablo(baglam, "SAYIM SONUÇLARI"))}
        assert sonuclar["Kayıtlara göre miktar (sayımın başladığı andaki kayıt)"] == "2"
        assert sonuclar["Kütüphanede ve yerinde sayılarak bulunan"] == "1"
        assert sonuclar["Fazla (TMY md. 17)"] == "1"
        assert sonuclar["Sayımda bulunan miktar (bulunan + kayda göre alınan + fazla)"] == "2"
        assert sonuclar["Noksan (TMY md. 32/7)"] == "1"

        noksan_tablosu = _tablo(baglam, "KAYITTAN DÜŞME TEKLİFİ (TMY MD. 32/7)")
        (satir,) = _hucreler(noksan_tablosu)
        assert satir[1] == barcode_module.format_barcode(noksan.barcode)
        assert satir[2] == "Bulunamayan Kitap" and satir[6] == "255.01.02"
        assert satir[4] == "Rafta"
        fazla = _hucreler(_tablo(baglam, "SAYIM FAZLASI (TMY MD. 17)"))
        assert fazla[0][2] == "Fazla Eser" and fazla[0][4] == "Kayda alınacak: Fazla Eser"

        metin = _metin(belgeler.stocktake_report_pdf(guncel))
        assert belgeler.ONAY_BEKLIYOR_NOTU in metin
        assert "Bulunan Kitap" not in metin  # bulunan kalemler sayıyla, listesiz
        assert "Bulunamayan Kitap" in metin
        assert "md. 32/7" in metin and belgeler.TMY17_1 in metin
        assert "Onay bekliyor" in metin

    def test_ikinci_sayim_32_6_kunyede_ve_notta(self) -> None:
        kitap = odunc_nushasi(title="İkinci Turda Bulunan")
        sayim = baslat()
        assert stocktake.complete_stocktake(sayim).second_round is True
        okut(sayim, kitap)
        tamamla(sayim)
        baglam = belgeler.stocktake_report_context(_t(sayim))
        kunye = {b["label"]: b["value"] for b in baglam["info"]}
        assert kunye["İkinci sayım (TMY md. 32/6)"].startswith("Yapıldı — 1 nüsha")
        assert belgeler.IKINCI_SAYIM_NOTU in baglam["notes"]

    def test_oduncteki_ve_teslimdeki_nusha_kurul_secimi_ve_dayanagi(self) -> None:
        odunc_ver(uye())
        teslim_et([odunc_nushasi(title="Şubede")], section=sube(10, "B"))
        teslim_et([odunc_nushasi(title="Öğretmende")], personnel=ogretmen())
        sayim = baslat(loan_basis=CountBasis.COLLECT)
        tamamla(sayim)
        baglam = belgeler.stocktake_report_context(_t(sayim))
        satirlar = {s[0]: s for s in _hucreler(_tablo(baglam, "SAYIM KURULUNUN SEÇİMİ"))}
        odunc = satirlar["Ödünçteki nüsha"]
        assert odunc[1] == "Sayımdan önce toplanır"
        assert "32/5'e kıyasen; 23/4" in odunc[2]
        # Toplanamayan ödünç kayda göre alınır ve işaretlidir.
        assert "Toplanamayan 1 nüsha kayda göre alındı." in odunc[2]
        assert odunc[3:] == ["1", "0", "1", "0"]
        sinif = satirlar["Sınıf kitaplığına teslim edilen nüsha"]
        assert sinif[1] == "Yerinde sayılır" and "32/5 birinci cümleye kıyasen" in sinif[2]
        assert sinif[6] == "1"  # yerinde sayılıp bulunmayan: noksan
        ogrt = satirlar["Öğretmene teslim edilen nüsha"]
        assert ogrt[1] == "Kayda göre alınır" and "32/5 ikinci cümle" in ogrt[2]
        assert _tablo(baglam, "SAYIM KURULUNUN SEÇİMİ")["note"] == belgeler.KIMLIK_NOTU
        # Yerinde sayılan sınıf kitaplığının noksanında şube etiketi (kişisel veri değil).
        noksan = _hucreler(_tablo(baglam, "TMY MD. 32/7"))
        assert any("Sınıf kitaplığında (10/B)" in s[4] for s in noksan)

    def test_onarimdaki_nusha_ayri_satirda_ve_dayanaksiz_ama_acik(self) -> None:
        """K2 (25.09.2026): onarımda kayda göre alınan SONUÇLAR'da ve kurul seçiminde AYRI
        satırdadır; 32/5 atfedilmez, dayanak metni doğrudan hüküm olmadığını söyler."""
        onarimda = odunc_nushasi(title="Ciltçideki Kitap")
        loss_damage.send_to_repair(onarimda)
        sayim = tamamla(baslat(repair_basis=CountBasis.COLLECT))
        baglam = belgeler.stocktake_report_context(_t(sayim))
        tablo = _tablo(baglam, "SAYIM KURULUNUN SEÇİMİ")
        assert "ONARIMDAKİ" in tablo["title"]
        onarim = {s[0]: s for s in _hucreler(tablo)}["Onarımdaki nüsha"]
        assert onarim[1] == "Sayımdan önce geri alınır"
        assert "doğrudan hüküm yoktur" in onarim[2] and "32/5" not in onarim[2]
        assert "Onarımdan geri alınamayan 1 nüsha kayda göre alındı." in onarim[2]
        assert onarim[3:] == ["1", "0", "1", "0"]
        sonuclar = {s[0]: s[1] for s in _hucreler(_tablo(baglam, "SAYIM SONUÇLARI"))}
        assert sonuclar["Kayda göre alınan (onarımda — sayım kurulunun kararı)"] == "1"
        assert sonuclar["Kayda göre alınan (ödünçte ve teslimde — TMY md. 32/5'e kıyasen)"] == "0"
        assert sonuclar["Noksan (TMY md. 32/7)"] == "0"

    def test_ekte_gelecek_yila_devirden_hasar_dusumu_ayri_satirda(self) -> None:
        """Madde 25 (a): onaydan önce devir "onaydan sonra yazılır"; onaydan sonra sayımda
        bulunan değişmez, hasar nedeniyle düşülen altında ayrı satırda ve devirden çıkarılır."""
        hasarli = odunc_nushasi(title="Hasar Önerili Ek")
        odunc_nushasi(title="Rafta Ek")
        dosya = loss_damage.open_damage_case(copy=hasarli)
        loss_damage.resolve_case(dosya, resolution=CaseResolution.WRITE_OFF_PROPOSED)
        sayim = baslat()
        okut(sayim, *Copy.objects.all())
        tamamla(sayim)
        satirlar = {
            s[0]: s[1]
            for s in _hucreler(belgeler.stocktake_report_context(_t(sayim))["annex"]["tables"][0])
        }
        assert satirlar[belgeler.GELECEK_YIL_SATIRI] == belgeler.BELIRSIZ_ONAY
        assert satirlar[f"— {belgeler.SAYIMDA_BULUNAN_SATIRI}"] == "2"
        onayla(_t(sayim))
        satirlar = {
            s[0]: s[1]
            for s in _hucreler(belgeler.stocktake_report_context(_t(sayim))["annex"]["tables"][0])
        }
        assert satirlar[f"— {belgeler.SAYIMDA_BULUNAN_SATIRI}"] == "2"
        assert satirlar[f"— {belgeler.HASAR_DUSULEN_SATIRI}"] == "1"
        assert satirlar[belgeler.GELECEK_YIL_SATIRI] == "1"
        assert satirlar["Fark (kayda göre yıl sonu − gelecek yıla devir)"] == "0"
        assert "gelecek yıla devirden çıkarılır" in belgeler.GELECEK_YIL_NOTU

    def test_hasar_dususu_ve_kayip_hasar_tutanagina_bagli_liste(self) -> None:
        hasarli = odunc_nushasi(title="Hasarlı Kitap")
        dosya = loss_damage.open_damage_case(copy=hasarli)
        loss_damage.resolve_case(dosya, resolution=CaseResolution.WRITE_OFF_PROPOSED)
        kayip = odunc_nushasi(title="Kayıp Kitap")
        kayip_dosyasi = loss_damage.report_lost(copy=kayip)
        sayim = baslat()
        okut(sayim, hasarli)
        tamamla(sayim)
        baglam = belgeler.stocktake_report_context(_t(sayim))

        (hasar,) = _hucreler(
            _tablo(baglam, "HASAR NEDENİYLE KAYITTAN DÜŞME TEKLİFİ (TMY MD. 27/1)")
        )
        assert hasar[2] == "Hasarlı Kitap"
        assert hasar[4] == f"Hasar · tespit {dosya.reported_on:%d.%m.%Y}"
        (noksan,) = _hucreler(_tablo(baglam, "TMY MD. 32/7"))
        assert noksan[2] == "Kayıp Kitap" and noksan[4] == "Kayıp"
        assert noksan[5] == f"Kayıp · tespit {kayip_dosyasi.reported_on:%d.%m.%Y}"
        assert belgeler.HASAR_NOTU in baglam["notes"]
        assert belgeler.TUTANAK_BAGI_NOTU in baglam["notes"]
        assert belgeler.DUSME_NOTU in baglam["notes"]

        onayla(_t(sayim))
        metin = _metin(belgeler.stocktake_report_pdf(_t(sayim)))
        assert "Hasarlı Kitap" in metin and "Kayıp Kitap" in metin
        assert HARCAMA_YETKILISI in metin and "OLUR" in metin
        assert "md. 27/1" in metin and "md. 10/1-e" in metin

    def test_kayipta_gorunup_bulunan_ve_fazlanin_kayda_alinisi(self) -> None:
        kayip = odunc_nushasi(title="Bulunan Kayıp")
        loss_damage.report_lost(copy=kayip)
        sayim = baslat()
        okut(sayim, kayip)
        fazla = stocktake.add_surplus(sayim, note="Etiketsiz", work=eser(title="Fazla Kayda"))
        tamamla(sayim)
        (uzlasma,) = _hucreler(
            _tablo(belgeler.stocktake_report_context(_t(sayim)), "KAYITTA KAYIP GÖRÜNÜP")
        )
        assert uzlasma[2] == "Bulunan Kayıp" and uzlasma[5] == "Onayda kayıp kaydı kapanır"

        onayla(_t(sayim))
        baglam = belgeler.stocktake_report_context(_t(sayim))
        (uzlasma,) = _hucreler(_tablo(baglam, "KAYITTA KAYIP GÖRÜNÜP"))
        assert uzlasma[5] == "Kayıp kaydı kapandı"
        fazla.refresh_from_db()
        assert fazla.created_copy is not None
        (satir,) = _hucreler(_tablo(baglam, "SAYIM FAZLASI"))
        assert satir[4] == (
            f"Kayda alındı: {barcode_module.format_barcode(fazla.created_copy.barcode)}"
        )

    def test_onaydan_sonra_olur_onaylanmayan_ve_durumu_degisen_kalemler(self) -> None:
        dusulen = odunc_nushasi(title="Düşülen")
        reddedilen = odunc_nushasi(title="Onaylanmayan")
        degisen = odunc_nushasi(title="Sonradan Ödünç")
        sayim = tamamla(baslat())
        onceki = belgeler.stocktake_report_context(_t(sayim))
        assert onceki["approval"] == {
            "heading": "OLUR",
            "date": belgeler.BOS_TARIH,
            "name": "",
            "role": "Harcama yetkilisi",
        }
        odunc_ver(uye(), degisen)  # hizmet arası yok: ödünç açık (D17)
        onayla(_t(sayim), not_approved={kalem(sayim, reddedilen).pk: "Ciltçide; dönüşü bekleniyor"})
        baglam = belgeler.stocktake_report_context(_t(sayim))

        assert baglam["approval"]["name"] == HARCAMA_YETKILISI
        assert baglam["approval"]["date"] == f"{timezone.localdate():%d.%m.%Y}"
        assert baglam["status_note"] == ""
        (noksan,) = _hucreler(_tablo(baglam, "TMY MD. 32/7"))
        assert noksan[2] == "Düşülen"
        (red,) = _hucreler(_tablo(baglam, "HARCAMA YETKİLİSİNİN ONAYLAMADIĞI"))
        assert red[2] == "Onaylanmayan" and red[4] == "Ciltçide; dönüşü bekleniyor"
        (durum,) = _hucreler(_tablo(baglam, "ONAYDA DURUMU DEĞİŞEN"))
        assert durum[2] == "Sonradan Ödünç" and durum[4] == "Onayda durumu: Ödünçte."
        kunye = {b["label"]: b["value"] for b in baglam["info"]}
        assert kunye["Harcama yetkilisinin onayı"] == f"{timezone.localdate():%d.%m.%Y}"
        assert dusulen.pk != reddedilen.pk

    def test_odunc_alanin_kimligi_pdf_ve_xlsx_te_yok(self) -> None:
        """E10 kod kapısı: fazla/noksan sayfaları VİF'e bağlanıp muhasebe birimine gider."""
        kisi = ogrenci(
            first_name=ODUNC_ALAN_AD, last_name=ODUNC_ALAN_SOYAD, student_number="918273"
        )
        uyelik = uye(kisi)
        odunc = odunc_ver(uyelik)
        kayipta = odunc_ver(uyelik, odunc_nushasi(title="Kayba Dönüşen"))
        loss_damage.report_lost(copy=kayipta.copy, responsible_note=f"{ODUNC_ALAN_AD} bildirdi")
        hasarli = odunc_nushasi(title="Hasar Önerili")
        dosya = loss_damage.open_damage_case(copy=hasarli, membership=uyelik)
        loss_damage.resolve_case(dosya, resolution=CaseResolution.WRITE_OFF_PROPOSED)
        teslim_et(
            [odunc_nushasi(title="Öğretmende Duran")],
            personnel=ogretmen(first_name=TESLIM_ALAN_AD, last_name="Kimlikyok"),
        )
        sayim = baslat(**durdurma_alanlari(), service_pause=True)
        okut(sayim, hasarli)
        onayla(tamamla(sayim))
        guncel = _t(sayim)
        assert kalem(guncel, odunc.copy).result == "BY_RECORD"

        pdf = _metin(belgeler.stocktake_report_pdf(guncel))
        xlsx = _xlsx_metni(belgeler.stocktake_report_xlsx(guncel))
        for metin in (pdf, xlsx):
            for yasak in (
                ODUNC_ALAN_AD,
                ODUNC_ALAN_SOYAD,
                TESLIM_ALAN_AD,
                "918273",
                uyelik.card_no,
            ):
                assert yasak not in metin, yasak
            assert "Kayba Dönüşen" in metin and "Hasar Önerili" in metin
        # Kişi adı yalnız kurulda ve harcama yetkililerinde.
        assert KURUL["committee_chair"] in pdf and HARCAMA_YETKILISI in pdf


# ============================================================ ek (TMY 34/1 — A8 kararı)


class TestEk:
    def test_ek_dort_buyukluk_ve_cetvel_degildir_ibaresi(self) -> None:
        odunc_nushasi(title="Sayılan")
        sayim = baslat()
        ek = belgeler.stocktake_report_context(_t(sayim))["annex"]
        assert ek["title"] == "EK: TAŞINIR SAYIM VE DÖKÜM CETVELİNE AKTARILACAK SAYILAR"
        satirlar = {s[0]: s[1] for s in _hucreler(ek["tables"][0])}
        for buyukluk in (
            "Önceki yıldan devir (yıl başında kayıtta olan)",
            "Yıl içinde giren (Yönetmelik Md. 10/5 yolları ve sayım fazlası)",
            "Programa aktarım (mevcut koleksiyonun kaydı; taşınır girişi değildir)",
            "Yıl içinde çıkan (kayıttan düşme ve devir)",
            belgeler.GELECEK_YIL_SATIRI,
            f"— {belgeler.SAYIMDA_BULUNAN_SATIRI}",
            f"— {belgeler.HASAR_DUSULEN_SATIRI}",
        ):
            assert buyukluk in satirlar, buyukluk
        # Sayım sürerken gelecek yıla devir kesin değildir.
        assert satirlar[belgeler.GELECEK_YIL_SATIRI] == belgeler.BELIRSIZ
        assert "— Sayım fazlası (kayda giriş)" in satirlar
        assert "— Sayım noksanı (kayıttan düşüldü)" in satirlar
        assert (
            "Bu döküm Taşınır Sayım ve Döküm Cetveli değildir; resmî cetvel TKYS'de düzenlenir."
            in ek["notes"]
        )
        assert f"“{belgeler.TMY34_1}”" in ek["paragraphs"][0]
        assert [s[0] for s in _hucreler(ek["tables"][1])] == [
            "Ödünçteki nüsha",
            "Sınıf kitaplığına teslim edilen nüsha",
            "Öğretmene teslim edilen nüsha",
            "Onarımdaki nüsha",
        ]

        onayla(tamamla(_t(sayim)))
        satirlar = {
            s[0]: s[1]
            for s in _hucreler(belgeler.stocktake_report_context(_t(sayim))["annex"]["tables"][0])
        }
        assert satirlar[belgeler.GELECEK_YIL_SATIRI] == "0"
        assert satirlar[f"— {belgeler.HASAR_DUSULEN_SATIRI}"] == "0"
        assert satirlar["Sayım noksanı (kayıttan düşülen: 1)"] == "1"

    def test_ek_yeni_sayfada_ve_xlsx_te_ayri_sayfada(self) -> None:
        odunc_nushasi(title="Sayılan")
        sayim = tamamla(baslat())
        sayfalar = _sayfalar(belgeler.stocktake_report_pdf(_t(sayim)))
        son = sayfalar[-1]
        assert son.startswith("EK: TAŞINIR SAYIM VE DÖKÜM CETVELİNE AKTARILACAK SAYILAR"), son
        assert "SAYIM KURULU" not in son
        assert "Bu döküm Taşınır Sayım ve Döküm Cetveli değildir" in son
        assert "SAYIM KURULU" in sayfalar[-2] and "OLUR" in sayfalar[-2]

        kitap = load_workbook(io.BytesIO(belgeler.stocktake_report_xlsx(_t(sayim))))
        assert kitap.sheetnames == [
            "Sayım tutanağı",
            "Kalemler",
            "Cetvele aktarılacak sayılar",
        ]
        ek = kitap["Cetvele aktarılacak sayılar"]
        assert ek["A1"].value == belgeler.EK_ADI
        kalemler = kitap["Kalemler"]
        assert [h.value for h in kalemler[1]][:3] == ["Sıra", "Barkod", "Kaynak adı"]
        assert kalemler.freeze_panes == "A2"
        assert kalemler["C2"].value == "Sayılan" and kalemler["J2"].value == "Noksan"


# ============================================================ mevzuat


class TestMevzuat:
    def test_alintilar_depodaki_fikrayla_birebir(self) -> None:
        assert belgeler.TMY17_1 in _fikra(17, 1)
        assert belgeler.TMY34_1 in _fikra(34, 1)

    @pytest.mark.parametrize(
        ("atif", "metin", "madde", "fikra", "ifadeler"),
        [
            (
                "md. 32/2",
                belgeler.ACILIS,
                32,
                2,
                (
                    "kendisinin veya görevlendireceği bir kişinin başkanlığında",
                    "taşınır kayıt yetkilisinin de katılımıyla",
                    "en az üç kişiden oluşturulan sayım kurulu",
                ),
            ),
            (
                "md. 32/6",
                belgeler.IKINCI_SAYIM_NOTU,
                32,
                6,
                ("bir kez daha tekrarlanır", "“Fazla” veya “Noksan” sütununa kaydedilir"),
            ),
            (
                "md. 32/7",
                belgeler.DUSME_NOTU,
                32,
                7,
                (
                    "Kayıttan Düşme Teklif ve Onay Tutanağı ve Varlık İşlem Fişi",
                    "fazla olduğunun tespit edilmesi hâlinde ise Varlık İşlem Fişi",
                    "defter kayıtlarının sayım sonuçlarıyla uygunluğu sağlanır",
                ),
            ),
            (
                "md. 10/1-e",
                belgeler.DUSME_NOTU,
                10,
                1,
                (
                    "sayımda noksan çıkması",
                    "durumu belgeleyen tutanak, rapor ve benzeri belgelerin bulunması",
                    "komisyon kurulması gerekmeksizin harcama yetkilisince onaylanır",
                ),
            ),
            (
                "md. 27/3",
                belgeler.DUSME_NOTU,
                27,
                3,
                ("kasıt, kusur, ihmal veya tedbirsizlik olup olmadığı harcama yetkilisince",),
            ),
            (
                "md. 27/1",
                belgeler.HASAR_NOTU,
                27,
                1,
                ("kullanılamaz hâle gelen taşınırlar", "kayıtlardan çıkarılır"),
            ),
            (
                "md. 10/1-g",
                belgeler.KIMLIK_NOTU,
                10,
                1,
                (
                    "Sayım Tutanağının sayım fazlası veya noksanına ilişkin sayfalarının",
                    "Varlık İşlem Fişi ekine",
                    "muhasebe birimine gönderilecek nüshasına bağlanır",
                ),
            ),
            (
                "32/8",
                belgeler.KIMLIK_NOTU,
                32,
                8,
                ("muhasebe birimine gönderilir",),
            ),
            (
                "md. 10/1-ğ",
                belgeler.GELECEK_YIL_NOTU,
                10,
                1,
                (
                    "“Gelecek Yıla Devir” sütununda",
                    "“Sayımda Bulunan Miktar”",
                    "eşit olması gerekir",
                ),
            ),
            (
                "md. 17/1",
                belgeler.FAZLA_NOTU,
                17,
                1,
                (
                    "aynı nitelikte son bir yıl içinde girişi yapılan taşınır varsa bu değer",
                    "değer tespit komisyonu tarafından belirlenecek değer",
                ),
            ),
        ],
    )
    def test_sade_anlatim_atif_yapilan_fikrada(
        self, atif: str, metin: str, madde: int, fikra: int, ifadeler: tuple[str, ...]
    ) -> None:
        assert atif in metin, atif
        govde = _fikra(madde, fikra)
        for ifade in ifadeler:
            assert ifade in govde, f"TMY md. {madde}/{fikra}: {ifade}"

    def test_dayanak_32_5_ve_23_4_metinde(self) -> None:
        """Kurul seçiminin dayanakları (`selectors_sayim.BASIS_DAYANAK`) atıf yaptıkları fıkrada."""
        assert "ortak kullanım alanlarında bulunan taşınırlar" in _fikra(32, 5)
        assert "Kayıtlara Göre Kişilere Verilen Miktar" in _fikra(32, 5)
        assert "ödünç takip sistemleri ile takip edilir" in _fikra(23, 4)
        assert "ortak kullanım alanlarına Dayanıklı Taşınırlar Listesi" in _fikra(23, 6)

    def test_sablonlarda_text_transform_yok(self) -> None:
        for ad in ("sayim_tutanagi.html", "_sayim_tablosu.html"):
            assert "text-transform:" not in _sablon(ad).replace(" ", ""), ad


# ============================================================ sayfa bütçesi (gerçek uzunluk)


class TestSayfaButcesi:
    @pytest.fixture(autouse=True)
    def uzun_okul(self, okul: SchoolConfig) -> None:
        okul.school_name = EN_UZUN_OKUL
        okul.district = "Örnek Uzun Adlı İlçe"
        okul.save()

    def _uzun_sayim(self, noksan_sayisi: int) -> StockTake:
        for _ in range(noksan_sayisi):
            _uzun_kitap()
        odunc_ver(uye())
        teslim_et([odunc_nushasi(title="Şubede")], section=sube(12, "Ç"))
        teslim_et([odunc_nushasi(title="Öğretmende")], personnel=ogretmen())
        sayim = baslat(
            **UZUN_KURUL,
            **{
                **durdurma_alanlari(),
                "tmy_stop_by_name": "Deneme Harcamayetkilisi Uzunsoyadlıoğlu",
            },
            service_pause=True,
            service_pause_decision=("Okul müdürlüğünün 2026/123 sayılı sayım kararı " * 4)[:120],
        )
        stocktake.add_surplus(sayim, note=("Etiketsiz uzun açıklamalı kitap " * 10)[:255])
        fazla = stocktake.add_surplus(sayim, note="İkinci etiketsiz", work=eser(title=EN_UZUN_ESER))
        assert fazla is not None
        tamamla(sayim)
        for k in _t(sayim).items.filter(copy__isnull=True, surplus_work__isnull=True):
            stocktake.update_surplus(k, excluded=True)
        return _t(sayim)

    def test_kisa_listeli_tutanak_en_cok_uc_sayfa_ek_tek_sayfa(self) -> None:
        """Bütün tabloları dolu (noksan, fazla, kayda alınmayan, üç kurul seçimi) ama kısa
        listeli tutanak: ana metin en çok üç sayfa, ek tek sayfa; kapanış bölünmez."""
        sayim = self._uzun_sayim(1)
        onayla(sayim, approved_by_name="Deneme Harcamayetkilisi Uzunsoyadlıoğlu")
        sayfalar = _sayfalar(belgeler.stocktake_report_pdf(_t(sayim)))

        assert len(sayfalar) <= 4, len(sayfalar)
        assert sayfalar[-1].startswith("EK:"), sayfalar[-1][:80]
        assert "Bu döküm Taşınır Sayım ve Döküm Cetveli değildir" in sayfalar[-1]
        imza = sayfalar[-2]
        # Kapanış bölünmez: notlar, kurulun sekiz imzası ve OLUR aynı sayfada.
        assert "SAYIM KURULU" in imza and "OLUR" in imza and "Harcama yetkilisi" in imza
        for ad in (UZUN_KURUL["committee_chair"], UZUN_KURUL["committee_property_officer"]):
            assert ad in imza, ad
        for ad in UZUN_KURUL["committee_members"].splitlines():
            assert ad in imza, ad
        assert "Noksan için Kayıttan Düşme Teklif" in imza
        assert "TKYS" in imza

    def test_uzun_noksan_listesinde_baslik_yinelenir_imzalar_birlikte(self) -> None:
        sayim = self._uzun_sayim(30)
        sayfalar = _sayfalar(belgeler.stocktake_report_pdf(sayim))

        assert 3 <= len(sayfalar) <= 14, len(sayfalar)
        ana = sayfalar[:-1]
        # Noksan tablosunun satırı taşıyan sayfalar (TKYS kodu yalnız noksan tablosunda basılır).
        tablolu = [s for s in ana if UZUN_TKYS[:12] in s]
        assert len(tablolu) >= 2
        # Başlık satırı her sayfada yinelenir.
        assert all("Kayda göre durum" in s for s in tablolu)
        assert "SAYIM KURULU" in ana[-1] and "OLUR" in ana[-1]
        assert all(ad in ana[-1] for ad in UZUN_KURUL["committee_members"].splitlines())
        assert sayfalar[-1].startswith("EK:")
        assert "Bu döküm Taşınır Sayım ve Döküm Cetveli değildir" in sayfalar[-1]


# ============================================================ uçlar


@pytest.fixture
def istemci() -> APIClient:
    return APIClient()


class TestUclar:
    def test_belge_listesi_pdf_ve_xlsx_yanitlari(self, istemci: APIClient) -> None:
        odunc_nushasi(title="Uçtaki Kitap")
        sayim = tamamla(baslat())
        kok = f"/api/v1/library/stocktakes/{sayim.pk}/documents/"

        liste = istemci.get(kok)
        assert liste.status_code == 200
        assert liste.json() == [
            {
                "kind": "sayim-tutanagi",
                "title": "Sayım tutanağı",
                "formats": ["pdf", "xlsx"],
                "available": True,
                "reason": "",
            }
        ]
        pdf = istemci.get(f"{kok}sayim-tutanagi/")
        assert pdf.status_code == 200 and pdf["Content-Type"] == "application/pdf"
        assert pdf["Cache-Control"] == "no-store"
        assert (
            "Say%C4%B1m-tutana%C4%9F%C4%B1_" in pdf["Content-Disposition"]
            or "Sayım-tutanağı_" in pdf["Content-Disposition"]
        )
        assert KURUL["committee_chair"] not in pdf["Content-Disposition"]
        xlsx = istemci.get(f"{kok}sayim-tutanagi/?kind=xlsx")
        assert xlsx.status_code == 200
        assert xlsx["Content-Type"].startswith("application/vnd.openxmlformats")
        assert b"".join(xlsx.streaming_content)[:2] == b"PK"  # type: ignore[attr-defined]

    def test_taslakta_400_gerekceyle_bilinmeyen_404(self, istemci: APIClient) -> None:
        sayim = taslak()
        kok = f"/api/v1/library/stocktakes/{sayim.pk}/documents/"
        yanit = istemci.get(f"{kok}sayim-tutanagi/")
        assert yanit.status_code == 400
        assert yanit.json()["message"] == belgeler.DRAFT_MESSAGE
        assert istemci.get(kok).json()[0]["available"] is False
        assert istemci.get(f"{kok}sayim-tutanagi/?kind=docx").status_code == 400
        assert istemci.get(f"{kok}bilinmeyen/").status_code == 404
        assert istemci.get("/api/v1/library/stocktakes/99999/documents/").status_code == 404

    def test_gorevli_kipinde_belge_uclari_403(self, istemci: APIClient) -> None:
        sayim = baslat()
        KIP.gorevliye_gec()
        for yol in (
            f"/api/v1/library/stocktakes/{sayim.pk}/documents/",
            f"/api/v1/library/stocktakes/{sayim.pk}/documents/sayim-tutanagi/",
        ):
            yanit = istemci.get(yol)
            assert yanit.status_code == 403 and yanit.json()["code"] == "kip_yetkisiz", yol


# ============================================================ F9 düzeltme turu (25.09.2026)


def _yatay_tasmalar(html: str) -> list[tuple[str, float]]:
    """WeasyPrint düzeninde tablo hücresinin içerik kutusunu aşan metinler (metin, aşım px).

    Sayfa bütçesi testleri yalnız sayfa sayısına bakar; `nowrap` barkod hücresinin yandaki
    sütunun üstüne taşmasını ancak düzen kutuları gösterir.
    """
    from weasyprint import HTML
    from weasyprint.formatting_structure import boxes
    from weasyprint.text.fonts import FontConfiguration

    belge = HTML(string=html).render(font_config=FontConfiguration())
    tasmalar: list[tuple[str, float]] = []

    def gez(kutu: Any, hucre: Any) -> None:
        if isinstance(kutu, boxes.TableCellBox):
            hucre = kutu
        if isinstance(kutu, boxes.TextBox) and hucre is not None:
            sag = hucre.content_box_x() + hucre.width
            asim = kutu.position_x + kutu.width - sag
            if asim > 0.5:
                tasmalar.append((kutu.text, round(asim, 1)))
        for cocuk in getattr(kutu, "children", []) or []:
            gez(cocuk, hucre)

    for sayfa in belge.pages:
        gez(sayfa._page_box, None)
    return tasmalar


class TestDuzeltmeTuru:
    def test_barkod_ve_uzun_fazla_kodu_hucreden_tasmaz(self) -> None:
        """10 haneli barkod sütuna sığar; 13, 20 ve 32 haneli fazla kodu satır satır kırılır."""
        kitaplar = [_uzun_kitap() for _ in range(3)]
        hasarli = odunc_nushasi(title=EN_UZUN_ESER)
        dosya = loss_damage.open_damage_case(copy=hasarli)
        loss_damage.resolve_case(dosya, resolution=CaseResolution.WRITE_OFF_PROPOSED)
        sayim = baslat(**durdurma_alanlari())
        okut(sayim, hasarli)
        sonuclar = stocktake.scan_many(sayim, ["8690504123456", "1234567890" * 2, "9" * 32])
        assert [s.code for s in sonuclar] == [stocktake.OKUTMA_FAZLA] * 3
        tamamla(sayim)
        for k in _t(sayim).items.filter(copy__isnull=True):
            stocktake.update_surplus(k, excluded=True, note="Kütüphaneye ait değil")
        onayla(
            _t(sayim),
            not_approved={kalem(sayim, kitaplar[0]).pk: ("Gerekçe metni uzun. " * 12)[:255]},
        )
        html = render_to_string(belgeler.SABLON, belgeler.stocktake_report_context(_t(sayim)))
        assert _yatay_tasmalar(html) == []

    def test_ikinci_sayim_notu_yalniz_bulunamayanlari_soyler(self) -> None:
        assert belgeler.IKINCI_SAYIM_NOTU.startswith("İlk sayımda bulunamayan nüshaların")
        assert "farklı çıkan nüshaların" not in belgeler.IKINCI_SAYIM_NOTU
        assert "ikinci kez sayılmamıştır" in belgeler.IKINCI_SAYIM_NOTU

    def test_surerken_noksan_tablosu_sonuc_yazmaz(self) -> None:
        nusha(eser(title="Henüz Okutulmayan"))
        baglam = belgeler.stocktake_report_context(_t(baslat()))
        tablo = _tablo(baglam, "KAYITTAN DÜŞME TEKLİFİ (TMY MD. 32/7)")
        assert tablo["rows"] == [] and tablo["empty"] == belgeler.NOKSAN_SURERKEN

    def test_karar_bekleyen_fazla_varken_imza_tutanagi_kesin_demez(self) -> None:
        kitap = odunc_nushasi(title="Rafta Bulunan")
        sayim = baslat()
        okut(sayim, kitap)
        fazla = stocktake.add_surplus(sayim, note="Etiketsiz")
        tamamla(sayim)
        baglam = belgeler.stocktake_report_context(_t(sayim))
        assert baglam["status_note"].startswith(belgeler.ONAY_BEKLIYOR_NOTU)
        assert belgeler.FAZLA_KARARI_NOTU.format(sayi="1") in baglam["status_note"]
        satirlar = {s[0]: s[1] for s in _hucreler(baglam["annex"]["tables"][0])}
        assert satirlar[belgeler.GELECEK_YIL_SATIRI] == belgeler.BELIRSIZ_FAZLA
        assert satirlar[f"— {belgeler.SAYIMDA_BULUNAN_SATIRI}"] == belgeler.BELIRSIZ_FAZLA
        sonuclar = [s[0] for s in _hucreler(_tablo(baglam, "SAYIM SONUÇLARI"))]
        assert any(s.startswith("Kararı bekleyen sayım fazlası") for s in sonuclar)

        stocktake.update_surplus(fazla, excluded=True, note="Kütüphaneye ait değil")
        baglam = belgeler.stocktake_report_context(_t(sayim))
        assert baglam["status_note"] == belgeler.ONAY_BEKLIYOR_NOTU
        satirlar = {s[0]: s[1] for s in _hucreler(baglam["annex"]["tables"][0])}
        assert satirlar[belgeler.GELECEK_YIL_SATIRI] == "1"

    def test_ek_ara_sayim_notu_ve_atiflari(self) -> None:
        odunc_nushasi(title="Sayılan")
        ek = belgeler.stocktake_report_context(_t(baslat()))["annex"]
        assert ek["paragraphs"][1] == belgeler.ARA_SAYIM_NOTU
        assert "yıl sonu hesaplarına ilişkin işlemlerinde" in _fikra(10, 1)
        assert "yıl sonu hesabını oluşturur" in _fikra(32, 9)
        assert "harcama yetkilisinin gerekli gördüğü durum ve zamanlarda" in _fikra(32, 1)
        for atif in ("md. 10/1-ğ, 32/9", "md. 32/1"):
            assert atif in belgeler.ARA_SAYIM_NOTU

    def test_xlsx_ek_sayilari_ve_mali_yil_sayi_hucresidir(self) -> None:
        for _ in range(2):
            odunc_nushasi(title="Sayılan")
        sayim = tamamla(baslat())
        onayla(sayim)
        kitap = load_workbook(io.BytesIO(belgeler.stocktake_report_xlsx(_t(sayim))))
        ek = kitap["Cetvele aktarılacak sayılar"]
        degerler = {
            str(a.value): b for a, b in ek.iter_rows(min_row=4, max_col=2) if a.value is not None
        }
        for buyukluk in (
            "Önceki yıldan devir (yıl başında kayıtta olan)",
            "Yıl içinde çıkan (kayıttan düşme ve devir)",
            belgeler.GELECEK_YIL_SATIRI,
            f"— {belgeler.SAYIMDA_BULUNAN_SATIRI}",
            f"— {belgeler.HASAR_DUSULEN_SATIRI}",
        ):
            hucre = degerler[buyukluk]
            assert isinstance(hucre.value, int), buyukluk
            assert hucre.number_format == "#,##0", buyukluk
        ozet = kitap["Sayım tutanağı"]
        yil = {str(a.value): b.value for a, b in ozet.iter_rows(max_col=2) if a.value}
        assert yil["Mali yıl"] == _t(sayim).fiscal_year

    def test_sayimda_baglanan_bos_etiket_fazlasi_tutanakta_yeniden_kayda_alinmaz(self) -> None:
        from apps.kutuphane.services import barcode_reservations
        from apps.kutuphane.tests.ortak import edinim

        etiket = barcode_reservations.reserve(1).numbers.get().barcode
        hedef = eser(title="Sayımda Bağlanan")
        sayim = baslat()
        stocktake.scan_many(sayim, [etiket])
        bagli = barcode_reservations.bind_label(label=etiket, work=hedef, acquisition=edinim())
        barkod = barcode_module.format_barcode(bagli.barcode)
        baglam = belgeler.stocktake_report_context(_t(sayim))
        fazla = _hucreler(_tablo(baglam, "SAYIM FAZLASI (TMY MD. 17)"))
        assert fazla[0][4].startswith(f"Sayım sırasında kayda girdi: {barkod}")
        onayla(tamamla(sayim))
        baglam = belgeler.stocktake_report_context(_t(sayim))
        fazla = _hucreler(_tablo(baglam, "SAYIM FAZLASI (TMY MD. 17)"))
        assert fazla[0][4].startswith("Kayda alınmadı — Etiket sayım sırasında")
        assert barkod in fazla[0][4]
