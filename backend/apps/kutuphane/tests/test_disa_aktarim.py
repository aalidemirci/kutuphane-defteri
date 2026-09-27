"""Dışa aktarım ve "dışa aktarım dosyası" içe aktarma kipi (tasarım §8.4, U1; F10 kod kapısı).

Kapının maddesi: **gidiş-dönüş aynı kataloğu verir** — dışa aktar → boş kurulumda içe aktar
→ aynı katalog (eser, nüsha, barkod, durum, bölüm, bayraklar) ve ikinci dışa aktarım ilk
dosyayla aynı satırları taşır. Ayrıca: numara asla yeniden kullanılmaz (silinmiş nüsha,
boş barkod aralığı, sayacın gerisi, kart biçimi, dosya içi tekrar), sayaçlar ilerler, şema
içe aktarım sözlüğünün üst kümesidir, dosyada kişisel veri yoktur, XLSX'te sayılar sayıdır.

Bütün veriler uydurmadır (CLAUDE.md §2-12).
"""

from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

import pytest
from django.core.exceptions import ValidationError
from django.utils import timezone
from openpyxl import load_workbook

from apps.kutuphane import disa_aktarim, import_schema
from apps.kutuphane import export_schema as sema
from apps.kutuphane.models import (
    CatalogImportSource,
    CommissionDecision,
    Copy,
    CopyCounter,
    CopyStatus,
    Work,
)
from apps.kutuphane.services import barcode_reservations, numbering
from apps.kutuphane.services import export_import as ei
from apps.kutuphane.services import import_service as ia
from apps.kutuphane.tests import ortak
from apps.kutuphane.tests.disa_aktarim_ortak import (
    BAGISCI,
    CIKIS_TARIHI,
    KOMISYON_BASKANI,
    bos_kurulum,
    butun_metin,
    katalog_goruntusu,
    katalog_kur,
    sayfa_satirlari,
)
from apps.kutuphane.tests.dolasim_ortak import odunc_ver, ogrenci, personel, uye
from apps.okul.excel_ogrenci import ParserError

pytestmark = pytest.mark.django_db


def _ayristir(icerik: bytes) -> tuple[ia.ParsedFile, str]:
    return ia.rows_from_file(icerik, source=CatalogImportSource.EXPORT)


def _uygula(icerik: bytes, **secenekler: Any) -> ia.CatalogImportReport:
    parsed, ozet = _ayristir(icerik)
    return ia.apply_import(parsed, payload_sha256=ozet, file_name="dosya.xlsx", **secenekler)


def _onizle(icerik: bytes, **secenekler: Any) -> ia.CatalogImportReport:
    parsed, ozet = _ayristir(icerik)
    return ia.preview_import(parsed, payload_sha256=ozet, file_name="dosya.xlsx", **secenekler)


def _duzenle(icerik: bytes, sayfa: str, degisiklik: dict[tuple[int, str], Any]) -> bytes:
    """Dosyanın bir sayfasında (satır, başlık) → değer değişikliği yapar."""
    kitap = load_workbook(BytesIO(icerik))
    ws = kitap[sayfa]
    basliklar = [h.value for h in ws[1]]
    for (satir, baslik), deger in degisiklik.items():
        ws.cell(row=satir, column=basliklar.index(baslik) + 1, value=deger)
    cikti = BytesIO()
    kitap.save(cikti)
    return cikti.getvalue()


def _katalog_satirlari(icerik: bytes) -> list[dict[str, Any]]:
    satirlar = sayfa_satirlari(icerik, sema.CATALOG_SHEET)
    basliklar = [str(h) for h in satirlar[0]]
    return [dict(zip(basliklar, s, strict=False)) for s in satirlar[1:]]


# ---------------------------------------------------------------------------
# Gidiş-dönüş (kod kapısı)
# ---------------------------------------------------------------------------
class TestGidisDonus:
    def test_disa_aktar_bos_kurulumda_ice_aktar_ayni_katalog(self) -> None:
        kurgu = katalog_kur()
        once = katalog_goruntusu()
        ilk_dosya = disa_aktarim.build_export_xlsx()

        bos_kurulum()
        assert Work.all_objects.count() == 0 and Copy.all_objects.count() == 0

        onizleme = _onizle(ilk_dosya)
        assert onizleme.dry_run
        assert Work.objects.count() == 0, "önizleme hiçbir kayıt yazmaz"
        rapor = _uygula(ilk_dosya)
        assert rapor.stats["error_rows"] == 0, [r.issues for r in rapor.rows if r.issues]
        assert rapor.stats["skipped_rows"] == 0
        assert rapor.stats["copies_created"] == onizleme.stats["copies_created"]
        assert rapor.stats["new_works"] == onizleme.stats["new_works"]

        sonra = katalog_goruntusu()
        # Ödünçte ve sınıf kitaplığında olan nüsha "Rafta" açılır (kayıt dosyada yok).
        beklenen = _durumlari_rafa_cek(once)
        assert sonra == beklenen
        assert rapor.stats["status_reset"] == 2

        # İkinci dışa aktarım aynı satırları taşır (Durum sütunundaki bilinen fark dışında).
        ikinci_dosya = disa_aktarim.build_export_xlsx()
        ilk = _katalog_satirlari(ilk_dosya)
        ikinci = _katalog_satirlari(ikinci_dosya)
        assert len(ilk) == len(ikinci)
        for a, b in zip(ilk, ikinci, strict=True):
            if a["Durum"] in ("Ödünçte", "Sınıf kitaplığında"):
                a = {**a, "Durum": "Rafta"}
            assert a == b
        assert sayfa_satirlari(ilk_dosya, sema.SECTIONS_SHEET) == sayfa_satirlari(
            ikinci_dosya, sema.SECTIONS_SHEET
        )
        del kurgu

    def test_numara_korunur_sayac_silinen_ve_ayrilan_numaralarin_otesine_gecer(self) -> None:
        kurgu = katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        en_buyuk = max(
            [int(k) for k in Copy.all_objects.values_list("barcode", flat=True)]
            + [int(k) for k in kurgu["ayrilmis"]]
        )
        bos_kurulum()
        _uygula(dosya)
        yeni_no, yeni_barkod = numbering.next_copy_identity()
        assert int(yeni_barkod) > en_buyuk
        # Silinen nüshanın ve bağlanmamış boş etiketin numarası yeni kurulumda verilmez.
        assert not Copy.all_objects.filter(barcode=kurgu["silinen_barkod"]).exists()
        assert yeni_barkod != kurgu["silinen_barkod"] and yeni_barkod not in kurgu["ayrilmis"]
        del yeni_no

    def test_kayittan_cikis_tarihi_korunur(self) -> None:
        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        bos_kurulum()
        _uygula(dosya)
        for n in Copy.objects.filter(
            status__in=(CopyStatus.WITHDRAWN_WEEDED, CopyStatus.TRANSFERRED)
        ):
            assert timezone.localdate(n.updated_at) == CIKIS_TARIHI

    def test_bagisin_karari_tarih_ve_sayisiyla_yeniden_kurulur_adsiz(self) -> None:
        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        bos_kurulum()
        rapor = _uygula(dosya)
        assert rapor.stats["decisions_created"] == 1
        karar = CommissionDecision.objects.get()
        assert (karar.decision_date, karar.decision_no) == (date(2026, 3, 1), "2026/3")
        assert karar.chair_name == ""
        assert karar.notes == ei.EXPORT_DECISION_NOTE

    def test_ayni_dosya_ikinci_kez_uygulanamaz(self) -> None:
        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        bos_kurulum()
        _uygula(dosya)
        with pytest.raises(ValidationError) as hata:
            _uygula(dosya)
        assert "zaten içe aktarıldı" in str(hata.value)
        # Katalog artık dolu: önizleme de reddeder ve dosyanın uygulandığını söyler.
        with pytest.raises(ValidationError) as hata:
            _onizle(dosya)
        assert ei.NOT_EMPTY_MESSAGE in str(hata.value)
        assert "zaten içe aktarıldı" in str(hata.value)

    def test_dolu_kataloga_aktarilmaz_eserler_cogalmaz(self) -> None:
        """F10 düzeltme turu: eşleştirme yapılmadığı için dolu kataloğa aktarım nüshasız
        eserleri (e-kitap) çoğaltıyor, kısmen aktarılmış eserin nüshasını ikiz esere
        bağlıyordu. Katalogda canlı eser varken önizleme de uygulama da reddedilir."""
        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        eser_sayisi, nusha_sayisi = Work.all_objects.count(), Copy.all_objects.count()
        for islem in (_onizle, _uygula):
            with pytest.raises(ValidationError) as hata:
                islem(dosya)
            assert ei.NOT_EMPTY_MESSAGE in str(hata.value)
        assert Work.all_objects.count() == eser_sayisi
        assert Copy.all_objects.count() == nusha_sayisi
        assert Work.objects.filter(title="Örnek E-Kitap").count() == 1

    def test_kayitli_barkod_satiri_aktarilmaz(self) -> None:
        """Katalogda canlı eser yok ama numara bu kurulumda bir nüshada kayıtlı."""
        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        Work.objects.all().update(deleted_at=timezone.now())
        rapor = _onizle(dosya)
        iletiler = " ".join(i for r in rapor.rows for i in r.issues)
        assert "bu kurulumda kayıtlı bir nüshada" in iletiler
        assert rapor.stats["copies_created"] == 0


def _durumlari_rafa_cek(goruntu: dict[str, Any]) -> dict[str, Any]:
    eserler = []
    for eser in goruntu["eserler"]:
        nushalar = tuple(
            sorted(
                (n[0], n[1], CopyStatus.AVAILABLE if n[2] in ei.RESET_STATUSES else n[2], *n[3:])
                for n in eser[-1]
            )
        )
        eserler.append((*eser[:-1], nushalar))
    return {**goruntu, "eserler": sorted(eserler, key=repr)}


# ---------------------------------------------------------------------------
# Asla yeniden kullanım yok — çakışma denetimleri
# ---------------------------------------------------------------------------
def _tek_eserli_dosya(**nusha_alanlari: Any) -> bytes:
    ortak.nusha(ortak.eser(title="Örnek Tek Eser", authors="Deneme Yazar"), **nusha_alanlari)
    dosya = disa_aktarim.build_export_xlsx()
    bos_kurulum()
    return dosya


class TestNumaraCakismasi:
    def test_bos_barkod_araliginda_ayrilmis_numara_reddedilir(self) -> None:
        dosya = _tek_eserli_dosya()
        barkod = _katalog_satirlari(dosya)[0]["Barkod"]
        # Hedef kurulumda aynı numarayı ayır: sayaç sıfırdan başlar, ilk numara aynıdır.
        aralik = barcode_reservations.reserve(1)
        assert aralik.numbers.get().barcode == barkod
        CopyCounter.objects.all().delete()  # sayaç denetimini değil ayrılmış numarayı sına
        rapor = _onizle(dosya)
        assert rapor.stats["copies_created"] == 0
        assert ei.RESERVED_MESSAGE.format(kod=barkod) in rapor.rows[0].issues

    def test_sayacin_gerisindeki_numara_reddedilir(self) -> None:
        dosya = _tek_eserli_dosya()
        barkod = _katalog_satirlari(dosya)[0]["Barkod"]
        CopyCounter.objects.create(year=int(barkod[:4]), last_no=int(barkod[4:]))
        rapor = _onizle(dosya)
        assert ei.COUNTER_MESSAGE.format(kod=barkod) in rapor.rows[0].issues

    def test_silinmis_nushanin_numarasi_reddedilir(self) -> None:
        nusha = ortak.nusha(ortak.eser(title="Örnek Silinecek", authors="Deneme Yazar"))
        dosya = disa_aktarim.build_export_xlsx()
        from apps.kutuphane.services import catalog

        catalog.delete_copy(nusha)
        catalog.delete_work(nusha.work)  # katalog boş olmalı (dışa aktarım yalnız boş kataloga)
        CopyCounter.objects.all().delete()
        rapor = _onizle(dosya)
        assert ei.COPY_DELETED_MESSAGE.format(kod=nusha.barcode) in rapor.rows[0].issues

    def test_kart_bicimindeki_kod_ve_bozuk_barkod_reddedilir(self) -> None:
        dosya = _tek_eserli_dosya()
        kart = _duzenle(dosya, sema.CATALOG_SHEET, {(2, "Barkod"): "94718263"})
        rapor = _onizle(kart)
        assert ei.CARD_SHAPED_MESSAGE.format(kod="94718263") in rapor.rows[0].issues
        bozuk = _duzenle(dosya, sema.CATALOG_SHEET, {(2, "Barkod"): "1234"})
        rapor = _onizle(bozuk)
        assert ei.BARCODE_FORMAT_MESSAGE.format(kod="1234") in rapor.rows[0].issues
        assert Copy.all_objects.count() == 0

    def test_kayit_no_barkodla_uyusmazsa_reddedilir(self) -> None:
        dosya = _tek_eserli_dosya()
        bozuk = _duzenle(dosya, sema.CATALOG_SHEET, {(2, "Kayıt No"): 2026999999})
        rapor = _onizle(bozuk)
        assert rapor.stats["skipped_rows"] == 1

    def test_dosyada_tekrarlanan_barkodun_yalniz_ilki_aktarilir(self) -> None:
        eser = ortak.eser(title="Örnek İkili", authors="Deneme Yazar")
        ortak.nusha(eser)
        ikinci = ortak.nusha(eser)
        dosya = disa_aktarim.build_export_xlsx()
        bos_kurulum()
        ilk_barkod = _katalog_satirlari(dosya)[0]["Barkod"]
        tekrar = _duzenle(dosya, sema.CATALOG_SHEET, {(3, "Barkod"): ilk_barkod})
        tekrar = _duzenle(tekrar, sema.CATALOG_SHEET, {(3, "Kayıt No"): int(ilk_barkod)})
        rapor = _onizle(tekrar)
        assert rapor.stats["copies_created"] == 1
        assert ei.DUPLICATE_BARCODE_MESSAGE.format(kod=ilk_barkod) in rapor.rows[1].issues
        # Hepsi ya da hiçbiri: aktarılamayan satır varken uygulama reddedilir.
        with pytest.raises(ValidationError) as hata:
            _uygula(tekrar)
        assert ei.ROWS_NOT_IMPORTED_MESSAGE.format(sayi=1) in str(hata.value)
        assert Copy.all_objects.count() == 0
        del ikinci


class TestHepsiYaDaHicbiri:
    def test_reddedilen_satirin_numarasi_yakilmaz_duzeltilmis_dosya_uygulanir(self) -> None:
        """F10 düzeltme turu: ilk uygulamada reddedilen satır, sayaç dosyadaki "Numara
        Sayaçları"na ilerlediği için bir daha aktarılamıyordu (kitap yeni etiket istiyordu).
        Artık aktarılamayan satır varken uygulama reddedilir; sayaç ilerlemez, hiçbir şey
        yazılmaz; dosya düzeltilince bütün satırlar kendi numarasıyla aktarılır."""
        eser = ortak.eser(title="Örnek İki Nüshalı", authors="Deneme Yazar")
        ortak.nusha(eser)
        ortak.nusha(eser)
        dosya = disa_aktarim.build_export_xlsx()
        barkodlar = sorted(Copy.all_objects.values_list("barcode", flat=True))
        bos_kurulum()
        bozuk = _duzenle(dosya, sema.CATALOG_SHEET, {(3, "Edinim Tarihi"): "31.02.2026"})
        rapor = _onizle(bozuk)
        assert rapor.stats["copies_created"] == 1 and rapor.stats["skipped_rows"] == 1
        assert ei.not_imported_rows(rapor) == 1
        with pytest.raises(ValidationError) as hata:
            _uygula(bozuk)
        assert ei.ROWS_NOT_IMPORTED_MESSAGE.format(sayi=1) in str(hata.value)
        assert Copy.all_objects.count() == 0 and Work.all_objects.count() == 0
        assert not CopyCounter.objects.filter(last_no__gt=0).exists()
        rapor = _uygula(dosya)
        assert rapor.stats["copies_created"] == 2
        assert sorted(Copy.objects.values_list("barcode", flat=True)) == barkodlar


class TestKaynakSecimiKarari:
    def test_bagis_disi_edinimin_karari_geri_yuklenir_ve_edinimleri_ayirir(self) -> None:
        """F10 düzeltme turu: satın almanın "Kaynak seçimi" kararı dosyaya yazılıyor ama geri
        yüklemede sessizce düşüyordu; aynı gün ve fiyatlı, kararı farklı iki satın alma tek
        edinime birleşiyordu."""
        from apps.kutuphane.models import Acquisition, AcquisitionMethod, CommissionDecisionType

        kararlar = [
            ortak.karar(
                decision_type=CommissionDecisionType.SELECTION,
                decision_date=date(2026, 2, 1),
                decision_no=no,
                chair_name=KOMISYON_BASKANI,
            )
            for no in ("2026/2", "2026/4")
        ]
        for i, karar in enumerate(kararlar):
            satin = ortak.edinim(
                method=AcquisitionMethod.PURCHASE,
                date=date(2026, 2, 10),
                unit_price=Decimal("80"),
                commission_decision=karar,
            )
            ortak.nusha(ortak.eser(title=f"Örnek Satın Alınan {i}", authors="Deneme Yazar"), satin)
        ilk_dosya = disa_aktarim.build_export_xlsx()
        bos_kurulum()
        rapor = _uygula(ilk_dosya)
        assert rapor.stats["decisions_created"] == 2
        edinimler = Acquisition.objects.filter(method=AcquisitionMethod.PURCHASE)
        assert edinimler.count() == 2
        yeni = sorted(
            (a.commission_decision.decision_type, a.commission_decision.decision_no)
            for a in edinimler.select_related("commission_decision")
            if a.commission_decision is not None
        )
        assert yeni == [
            (CommissionDecisionType.SELECTION, "2026/2"),
            (CommissionDecisionType.SELECTION, "2026/4"),
        ]
        assert all(k.chair_name == "" for k in CommissionDecision.objects.all())
        ikinci = _katalog_satirlari(disa_aktarim.build_export_xlsx())
        assert sorted(s["Komisyon Kararı Sayısı"] for s in ikinci) == ["2026/2", "2026/4"]


# ---------------------------------------------------------------------------
# Dosya biçimi
# ---------------------------------------------------------------------------
class TestDosyaBicimi:
    def test_sayfalar_ve_bilgi(self) -> None:
        katalog_kur()
        kitap = load_workbook(BytesIO(disa_aktarim.build_export_xlsx()))
        assert kitap.sheetnames == list(sema.SHEETS)
        bilgi = {
            str(s[0]): s[1] for s in kitap[sema.INFO_SHEET].iter_rows(values_only=True) if s[0]
        }
        assert bilgi[sema.INFO_KIND] == sema.EXPORT_FILE_KIND
        assert bilgi[sema.INFO_VERSION] == sema.EXPORT_SCHEMA_VERSION
        assert isinstance(bilgi[sema.INFO_DATE], datetime)
        assert bilgi[sema.INFO_WORKS] == Work.objects.count()

    def test_ust_kume_ilk_sutunlar_ice_aktarim_sozlugudur(self) -> None:
        katalog_kur()
        basliklar = sayfa_satirlari(disa_aktarim.build_export_xlsx(), sema.CATALOG_SHEET)[0]
        tabani = [s.header for s in import_schema.COLUMNS]
        assert list(basliklar[: len(tabani)]) == tabani
        assert list(basliklar) == list(sema.HEADERS)
        # Sözleşmedeki ek sütunlar (§8.4) — çevirmen, baskı ve bölüm sözlükte zaten var.
        for baslik in (
            "Barkod",
            "Kayıt No",
            "Eski Kayıt No",
            "TKYS Kodu",
            "Durum",
            "Edinim Yolu",
            "Edinim Tarihi",
            "Çevirmen",
            "Baskı",
            "Bölüm",
            "Danışma Kaynağı",
            "Ciltli Süreli Yayın",
            "Piyasada Mevcudu Yok",
            "El Yazması / Nadir Eser",
        ):
            assert baslik in basliklar, baslik

    def test_ek_basliklar_olagan_ice_aktarimda_taninmaz(self) -> None:
        for sutun in sema.EXTRA_COLUMNS:
            assert import_schema.match_header(sutun.header) is None, sutun.header

    def test_sayilar_sayi_tarihler_tarih_hucresidir(self) -> None:
        katalog_kur()
        satirlar = _katalog_satirlari(disa_aktarim.build_export_xlsx())
        fiyatli = next(s for s in satirlar if s["Edinim Yolu"] == "Satın alma")
        assert isinstance(fiyatli["Kayıt No"], int)
        assert isinstance(fiyatli["Eser No"], int)
        assert isinstance(fiyatli["Edinim Tarihi"], datetime)
        assert Decimal(str(fiyatli["Birim Fiyat"])) == Decimal("125.5")
        assert isinstance(fiyatli["Barkod"], str) and len(fiyatli["Barkod"]) == 10
        yili = next(s for s in satirlar if s["Eser Adı"] == "Çalıkuşu")
        assert yili["Yayın Yılı"] == 2019

    def test_nushasiz_eser_barkodu_bos_tek_satirdir(self) -> None:
        katalog_kur()
        satirlar = _katalog_satirlari(disa_aktarim.build_export_xlsx())
        (ekitap,) = [s for s in satirlar if s["Eser Adı"] == "Örnek E-Kitap"]
        assert ekitap["Barkod"] in (None, "")
        assert ekitap["Nüsha Sayısı"] == 0
        assert ekitap["Kaynak Türü"] == "E-kitap"

    def test_kisisel_veri_yoktur(self) -> None:
        katalog_kur()
        kisi = ogrenci(first_name="Deneme", last_name="Gizliadlı", student_number="918273")
        uyelik = uye(kisi)
        odunc_ver(uyelik)
        uye(personel(first_name="Deneme", last_name="Öğretmenadlı"))
        metin = butun_metin(disa_aktarim.build_export_xlsx())
        for yasak in (
            BAGISCI,
            KOMISYON_BASKANI,
            "Gizliadlı",
            "Öğretmenadlı",
            "918273",
            uyelik.card_no,
        ):
            assert yasak not in metin, yasak

    def test_uye_ozeti_kisisiz_sayilardir(self) -> None:
        uye(ogrenci(class_level=9, class_section="A"))
        uye(ogrenci(class_level=9, class_section="A"))
        uye(ogrenci(class_level=10, class_section="Ç"))
        uye(ogrenci(class_level=10, class_section="C"))
        uye(personel())
        ozet = disa_aktarim.uye_ozeti()
        assert ozet["total"] == 5
        assert {s["member_type"]: s["count"] for s in ozet["by_type"]} == {
            "STUDENT": 4,
            "TEACHER": 1,
            "STAFF": 0,
        }
        # Türk alfabesi: 10/C, 10/Ç'den önce gelir.
        assert [(s["class_label"], s["count"]) for s in ozet["by_class"]] == [
            ("9/A", 2),
            ("10/C", 1),
            ("10/Ç", 1),
        ]
        satirlar = sayfa_satirlari(disa_aktarim.build_export_xlsx(), sema.MEMBER_SUMMARY_SHEET)
        assert ["Toplam", 5] in [list(s[:2]) for s in satirlar]

    def test_dosya_turu_yoksa_disa_aktarim_kipinde_okunmaz(self) -> None:
        from apps.kutuphane.tests import sentetik_katalog

        liste = sentetik_katalog.katalog_dosyasi([{"title": "Örnek", "authors": "Deneme"}])
        with pytest.raises(ParserError) as hata:
            _ayristir(liste)
        assert str(hata.value) == ei.NOT_EXPORT_FILE_MESSAGE

    def test_desteklenmeyen_surum_okunmaz(self) -> None:
        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        kitap = load_workbook(BytesIO(dosya))
        for satir in kitap[sema.INFO_SHEET].iter_rows():
            if satir[0].value == sema.INFO_VERSION:
                satir[1].value = "v9"
        cikti = BytesIO()
        kitap.save(cikti)
        with pytest.raises(ParserError, match="şema sürümü"):
            _ayristir(cikti.getvalue())

    def test_olagan_excel_listesi_yolu_dosyayi_tanir_ve_reddeder(self) -> None:
        """F10 düzeltme turu: "Excel listesi" kipinde dosya tanınmadan kabul ediliyordu —
        bütün nüshalar yeni numara alıyor, eski etiket başka kitabı açıyor, kayıttan çıkmış
        nüsha "Rafta" açılıyordu. Olağan yol dosyayı "Bilgi" sayfasından tanır ve reddeder."""
        ortak.nusha(ortak.eser(title="Örnek Olağan", authors="Deneme Yazar"))
        dosya = disa_aktarim.build_export_xlsx()
        assert ei.is_export_file(dosya)
        with pytest.raises(ParserError) as hata:
            ia.rows_from_file(dosya, source=CatalogImportSource.EXCEL)
        assert str(hata.value) == ei.EXPORT_FILE_AS_LIST_MESSAGE
        # Okulun kendi listesi olağan yolda okunur.
        from apps.kutuphane.tests import sentetik_katalog

        liste = sentetik_katalog.katalog_dosyasi([{"title": "Örnek", "authors": "Deneme"}])
        assert not ei.is_export_file(liste)
        parsed, _ozet = ia.rows_from_file(liste, source=CatalogImportSource.EXCEL)
        assert len(parsed.rows) == 1


class TestBolumler:
    def test_dosyadaki_bolumler_acilir_bilinmeyen_bolum_sorulur(self) -> None:
        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        bos_kurulum()
        # "Bölümler" sayfasında olmayan bir bölüm adı yazılırsa önizlemede sorulur.
        bozuk = _duzenle(dosya, sema.CATALOG_SHEET, {(2, "Bölüm"): "Örnek Yeni Bölüm"})
        rapor = _onizle(bozuk)
        assert [k["value"] for k in rapor.unknown_sections] == ["Örnek Yeni Bölüm"]
        with pytest.raises(ValidationError):
            _uygula(bozuk)
        rapor = _uygula(bozuk, new_sections=["Örnek Yeni Bölüm"])
        assert rapor.stats["error_rows"] == 0


class TestTmyDurdurmasi:
    def test_tmy_32_3_durdurmasi_surerken_bastan_reddedilir(self) -> None:
        from apps.kutuphane.tests.sayim_ortak import baslat, durdurma_alanlari

        katalog_kur()
        dosya = disa_aktarim.build_export_xlsx()
        bos_kurulum()
        ortak.nusha(ortak.eser(title="Örnek Sayılan", authors="Deneme Yazar"))
        baslat(**durdurma_alanlari())
        with pytest.raises(ValidationError):
            _onizle(dosya)


# ---------------------------------------------------------------------------
# docs/disa-aktarim.md ↔ şema
# ---------------------------------------------------------------------------
def _belge() -> str:
    from pathlib import Path

    yerel = Path(__file__).resolve().parents[4]
    for kok in (yerel, Path("/repo")):
        yol = kok / "docs" / "disa-aktarim.md"
        if yol.is_file():
            return yol.read_text(encoding="utf-8")
    raise AssertionError("docs/disa-aktarim.md bulunamadı.")


def test_belgedeki_sutun_tablosu_semayla_birebir() -> None:
    satirlar = [
        s for s in _belge().splitlines() if s.startswith("| ") and s.split("|")[1].strip().isdigit()
    ]
    hucreler = [[h.strip() for h in s.strip("|").split("|")] for s in satirlar]
    assert [h[1] for h in hucreler] == list(sema.HEADERS)
    assert [int(h[0]) for h in hucreler] == list(range(1, len(sema.HEADERS) + 1))
    degerler = {h[1]: h[2] for h in hucreler}
    for sutun in sema.EXTRA_COLUMNS:
        assert degerler[sutun.header] == sutun.accepted_values, sutun.header
    assert degerler["Kaynak Türü"] == " / ".join(sema.RESOURCE_TYPE_LABELS)
    belge = _belge()
    assert f"`{sema.EXPORT_SCHEMA_VERSION}`" in belge
    assert sema.EXPORT_FILE_KIND in belge
    for sayfa in sema.SHEETS:
        assert f"**{sayfa}**" in belge, sayfa
