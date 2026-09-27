"""F10-D uçları: dışa aktarım, üye özeti, dökümler özeti, E17, E11 ve İçe Aktarma'nın
"dışa aktarım dosyası" kipi (API yüzeyi).

Görevli kipinde bütün uçların 403 alması genel dolaşmadadır (`test_uc_kapilari.py`,
`apps/okul/tests/test_kip_koruma.py`); burada yanıt biçimi, dosya başlıkları ve
sözleşmeli 400'ler sınanır. Bütün veriler uydurmadır.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.kutuphane import export_schema as sema
from apps.kutuphane import tmy_dokumleri
from apps.kutuphane.models import CatalogImportSource, Copy, Work
from apps.kutuphane.tests import ortak
from apps.kutuphane.tests.disa_aktarim_ortak import bos_kurulum, katalog_kur
from apps.kutuphane.tests.dolasim_ortak import odunc_nushasi, ogrenci, uye
from apps.kutuphane.tests.sayim_ortak import baslat, onayla, tamamla
from apps.okul.kip import KIP

pytestmark = pytest.mark.django_db

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture
def istemci() -> APIClient:
    return APIClient()


def test_disa_aktarim_dosyasi_indirilir(istemci: APIClient) -> None:
    katalog_kur()
    yanit = istemci.get(reverse("library-export"))
    assert yanit.status_code == 200
    assert yanit["Content-Type"] == XLSX and yanit["Cache-Control"] == "no-store"
    assert "attachment" in yanit["Content-Disposition"]
    govde = b"".join(yanit.streaming_content)  # type: ignore[attr-defined]
    kitap = load_workbook(BytesIO(govde))
    assert kitap.sheetnames == list(sema.SHEETS)


def test_uye_ozeti_kisisiz(istemci: APIClient) -> None:
    uye(ogrenci(first_name="Denemeozet", last_name="Gizlioğlu", class_level=11, class_section="B"))
    yanit = istemci.get(reverse("library-export-member-summary"))
    assert yanit.status_code == 200
    govde = yanit.json()
    assert govde["total"] == 1 and govde["by_class"][0]["class_label"] == "11/B"
    assert "Gizlioğlu" not in yanit.content.decode("utf-8")


def test_dokumler_ozeti(istemci: APIClient) -> None:
    ortak.nusha(ortak.eser(title="Özetteki"))
    govde = istemci.get(reverse("library-report-documents")).json()
    assert govde["export"]["works"] == 1 and govde["export"]["copies"] == 1
    assert [e["value"] for e in govde["catalog_listing"]["axes"]] == ["title", "author", "subject"]
    assert govde["library_register"]["years"]
    assert govde["management_account"]["stocktakes"] == []
    assert govde["catalog_listing"]["unsectioned_works"] == 1


def test_alfabetik_katalog_pdf_xlsx_ve_gecersiz_eksen(istemci: APIClient) -> None:
    ortak.nusha(ortak.eser(title="Uçtaki"))
    adres = reverse("library-report-catalog-listing")
    pdf = istemci.get(adres, {"axis": "author"})
    assert pdf.status_code == 200 and pdf["Content-Type"] == "application/pdf"
    xlsx = istemci.get(adres, {"kind": "xlsx"})
    assert xlsx.status_code == 200 and xlsx["Content-Type"] == XLSX
    assert istemci.get(adres, {"axis": "renk"}).status_code == 400
    assert istemci.get(adres, {"kind": "docx"}).status_code == 400
    # `section=0`: bölümü yazılmamış eserler (F10 düzeltme turu); olmayan bölüm 400.
    assert istemci.get(adres, {"section": "0"}).status_code == 200
    assert istemci.get(adres, {"section": "999999"}).status_code == 400
    assert istemci.get(adres, {"section": "999999"}).status_code == 400


def test_defter_dokumu(istemci: APIClient) -> None:
    ortak.nusha(ortak.eser(title="Defterde"))
    adres = reverse("library-report-library-register")
    assert istemci.get(adres).status_code == 200
    assert istemci.get(adres, {"kind": "xlsx", "year": "2026"}).status_code == 200
    assert istemci.get(adres, {"year": "iki bin"}).status_code == 400


def test_yonetim_hesabi_yil_sonu_isaretsizse_400_gerekceli(istemci: APIClient) -> None:
    odunc_nushasi(title="Sayılan")
    sayim = tamamla(baslat())
    onayla(sayim)
    yanit = istemci.get(reverse("library-report-management-account", kwargs={"pk": sayim.pk}))
    assert yanit.status_code == 400
    assert tmy_dokumleri.NOT_YEAR_END_MESSAGE in yanit.json()["message"]
    yok = istemci.get(reverse("library-report-management-account", kwargs={"pk": 999999}))
    assert yok.status_code == 404


def test_yonetim_hesabi_yil_sonu_onayli_sayimdan(istemci: APIClient) -> None:
    odunc_nushasi(title="Sayılan")
    sayim = tamamla(baslat(is_year_end=True))
    onayla(sayim)
    adres = reverse("library-report-management-account", kwargs={"pk": sayim.pk})
    assert istemci.get(adres).status_code == 200
    assert istemci.get(adres, {"kind": "xlsx"})["Content-Type"] == XLSX


def test_sayim_yil_sonu_isareti_uctan_degisir(istemci: APIClient) -> None:
    odunc_nushasi(title="Sayılan")
    sayim = baslat()
    adres = reverse("library-stocktake-detail", kwargs={"pk": sayim.pk})
    yanit = istemci.patch(adres, {"is_year_end": True}, format="json")
    assert yanit.status_code == 200 and yanit.json()["is_year_end"] is True


def test_gorevli_kipinde_disa_aktarim_403(istemci: APIClient) -> None:
    KIP.gorevliye_gec()
    yanit = istemci.get(reverse("library-export"))
    assert yanit.status_code == 403 and yanit.json()["code"] == "kip_yetkisiz"


def test_ice_aktarma_ucu_disa_aktarim_kipinde(istemci: APIClient) -> None:
    katalog_kur()
    dosya = b"".join(istemci.get(reverse("library-export")).streaming_content)  # type: ignore[attr-defined]
    bos_kurulum()
    onizleme = istemci.post(
        reverse("library-import-preview"),
        {
            "file": SimpleUploadedFile("katalog.xlsx", dosya),
            "source": CatalogImportSource.EXPORT,
        },
        format="multipart",
    )
    assert onizleme.status_code == 200, onizleme.content
    assert onizleme.json()["dry_run"] is True and Work.objects.count() == 0
    uygulama = istemci.post(
        reverse("library-import-apply"),
        {
            "file": SimpleUploadedFile("katalog.xlsx", dosya),
            "source": CatalogImportSource.EXPORT,
        },
        format="multipart",
    )
    assert uygulama.status_code == 201, uygulama.content
    assert uygulama.json()["source"] == CatalogImportSource.EXPORT
    assert Copy.objects.count() == uygulama.json()["stats"]["copies_created"] > 0
    # Dışa aktarım dosyası olmayan bir dosya bu kipte sözleşmeli 400 alır.
    liste = istemci.post(
        reverse("library-import-preview"),
        {
            "file": SimpleUploadedFile("baska.xlsx", b"PK\x03\x04bozuk"),
            "source": CatalogImportSource.EXPORT,
        },
        format="multipart",
    )
    assert liste.status_code == 400
