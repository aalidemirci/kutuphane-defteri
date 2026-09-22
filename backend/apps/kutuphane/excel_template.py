"""Katalog Excel şablonu (openpyxl) — tasarım §8.1; OYS `excel_template.py`'den UYARLA.

Sütunların tek kaynağı `import_schema.COLUMNS`'tur; bu modül yalnız biçimi kurar.
Üç sayfa:

- **Katalog** — yalnız başlık satırı. İçe aktarım (F3) YALNIZ bu sayfayı okur;
  bu yüzden sayfada başlıktan başka hiçbir şey yoktur. OYS örnek satırı ve KVKK
  notunu veri sayfasına yazıyordu: doldurulmuş dosyada o satırlar kitap sanılırdı.
- **Sütunlar** — sözlük: başlık, zorunlu mu, ne yazılır, kabul edilen değerler,
  boş bırakılırsa ne sayılır, örnek; altında kısa notlar (kişisel veri yazılmaz,
  nüsha yazmanın iki yolu).
- **Örnek** — kamu malı klasik eserlerden satırlar; ISBN uydurulmaz (boş),
  yayınevi açıkça "Örnek Yayınevi"dir. Sözlükteki örnek değerler bu satırlardan
  gelir (test eşleşmeyi denetler).

Biçim KS şablonlarının (`apps/okul/services/templates.py`) sade düzenini izler:
varsayılan yazı tipi, kalın başlık, sabit sütun genişliği. Eklenenler şablonun
doldurulmasına hizmet eder: zorunlu sütun başlığı ayrı renkte ve açıklaması
başlık notunda; başlık satırı dondurulmuş; evet/hayır, kaynak türü ve nüsha
sayısı sütunlarında veri doğrulama; metin sütunlarında metin biçimi ("@") —
Excel ISBN'yi bilimsel gösterime, "001.4" sınıflama kodunu 1,4 sayısına
çevirmesin.

İndirilen dosyanın adı belge adı + yerel tarihtir (docs/sozluk.md §3); ön yüzün
`lib/download.ts::dosyaAdi` çıktısıyla aynı biçimdedir.
"""

from __future__ import annotations

from datetime import date
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from apps.kutuphane.import_schema import (
    CATALOG_SHEET,
    COLUMNS,
    COLUMNS_SHEET,
    EXAMPLE_SHEET,
    MAX_COPIES_PER_ROW,
    YES_NO_CHOICES,
    CatalogColumn,
    ColumnKind,
)

#: Belgenin adı: ön yüz kartının başlığı ve indirilen dosyanın adı.
TEMPLATE_DOCUMENT_NAME = "Katalog Excel Şablonu"

#: Veri doğrulama ve metin biçimi başlığın altından sayfanın sonuna kadar uygulanır
#: (Md. 7'deki 10.000 kitap eşiğini aşan okullar da tek sayfaya sığar).
_EXCEL_LAST_ROW = 1_048_576

_REQUIRED_FILL = PatternFill("solid", fgColor="F8CBAD")
_OPTIONAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_HEADER_FONT = Font(bold=True)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_HEADER_BORDER = Border(bottom=Side(style="thin", color="808080"))
_WRAP_TOP = Alignment(vertical="top", wrap_text=True)
_COMMENT_AUTHOR = "Kütüphane Defteri"

#: "Sütunlar" sayfasının sütunları: (başlık, genişlik).
_DICTIONARY_HEADERS: tuple[tuple[str, int], ...] = (
    ("Sütun", 22),
    ("Zorunlu mu?", 12),
    ("Ne yazılır?", 70),
    ("Kabul edilen değerler", 26),
    ("Boş bırakılırsa", 30),
    ("Örnek", 24),
)

#: "Sütunlar" sayfasındaki notlar (sözlüğün altında, kullanıcı dilinde).
DICTIONARY_NOTES: tuple[str, ...] = (
    f"Program yalnız “{CATALOG_SHEET}” sayfasını okur. Kitaplarınızı o sayfada başlık "
    "satırının altına yazın; başlıkları değiştirmeyin.",
    "Turuncu başlıklı sütun zorunludur; öbür sütunlar boş bırakılabilir.",
    "Aynı eserin nüshalarını iki yoldan yazabilirsiniz: tek satır açıp “Nüsha Sayısı” "
    "sütununa sayıyı yazarak ya da her nüsha için künyesi aynı ayrı bir satır açarak. "
    "Eski kayıt numarası olan nüshalar ikinci yolla, her biri ayrı satırda yazılır.",
    f"Bir satırda en çok {MAX_COPIES_PER_ROW} nüsha açılır.",
    "Bu listeye kişisel veri yazmayın: öğrenci, öğretmen ya da bağışçı adı girmez.",
    f"“{EXAMPLE_SHEET}” sayfasındaki satırlar yalnız fikir vermek içindir; o sayfa içe "
    "aktarılmaz.",
)

#: "Örnek" sayfasının satırları — kamu malı klasik eserler. ISBN uydurulmaz; yayınevi
#: ve yıllar örnektir. Satırlar sözlüğün kurallarına uyar (test denetler):
#: 1) toplu satır (3 nüsha), 2-3) aynı eserin eski kayıt numaralı iki nüshası ayrı
#: satırlarda, 4) danışma kaynağı.
EXAMPLE_ROWS: tuple[dict[str, str | int], ...] = (
    {
        "title": "Kürk Mantolu Madonna",
        "authors": "Sabahattin Ali",
        "publisher": "Örnek Yayınevi",
        "edition": "5. baskı",
        "publish_year": 2020,
        "subjects": "Türk edebiyatı, roman",
        "classification_code": "894.353",
        "language": "Türkçe",
        "resource_type": "Kitap",
        "copies": 3,
        "shelf_location": "Edebiyat",
        "is_reference": "Hayır",
    },
    {
        "title": "Telemak",
        "authors": "Fénelon",
        "translator": "Yusuf Kâmil Paşa",
        "publisher": "Örnek Yayınevi",
        "publish_year": 2018,
        "subjects": "Fransız edebiyatı, roman",
        "classification_code": "843.4",
        "language": "Türkçe",
        "resource_type": "Kitap",
        "copies": 1,
        "shelf_location": "Edebiyat",
        "old_register_no": "1452",
    },
    {
        "title": "Telemak",
        "authors": "Fénelon",
        "translator": "Yusuf Kâmil Paşa",
        "publisher": "Örnek Yayınevi",
        "publish_year": 2018,
        "subjects": "Fransız edebiyatı, roman",
        "classification_code": "843.4",
        "language": "Türkçe",
        "resource_type": "Kitap",
        "copies": 1,
        "shelf_location": "Edebiyat",
        "old_register_no": "1453",
    },
    {
        "title": "Kamus-ı Türkî",
        "authors": "Şemseddin Sami",
        "publisher": "Örnek Yayınevi",
        "publish_year": 2015,
        "subjects": "Türkçe, sözlük",
        "classification_code": "494.353",
        "language": "Türkçe",
        "resource_type": "Kitap",
        "copies": 1,
        "shelf_location": "Danışma",
        "is_reference": "Evet",
    },
)


def template_filename(today: date | None = None) -> str:
    """İndirme adı: 'Katalog-Excel-Şablonu_21.09.2026.xlsx' (yerel tarih, UTC değil)."""
    gun = today or timezone.localdate()
    return f"{TEMPLATE_DOCUMENT_NAME.replace(' ', '-')}_{gun:%d.%m.%Y}.xlsx"


def _column_width(column: CatalogColumn) -> int:
    """Başlık ve örnek sığsın; eser adı uzun yazılır, geniş tutulur."""
    if column.key == "title":
        return 40
    return max(12, len(column.header) + 4, min(30, len(column.example) + 4))


def _write_catalog_header(ws: Worksheet, *, with_comments: bool) -> None:
    for index, column in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=index, value=column.header)
        cell.font = _HEADER_FONT
        cell.fill = _REQUIRED_FILL if column.required else _OPTIONAL_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _HEADER_BORDER
        if with_comments:
            onek = "Zorunlu sütun. " if column.required else ""
            cell.comment = Comment(
                f"{onek}{column.description}", _COMMENT_AUTHOR, width=320, height=160
            )
        ws.column_dimensions[get_column_letter(index)].width = _column_width(column)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def _validation_for(column: CatalogColumn) -> DataValidation | None:
    """Sütunun Excel veri doğrulaması (yazarken yanlış değeri durdurur)."""
    values: tuple[str, ...]
    if column.kind is ColumnKind.YES_NO:
        values = YES_NO_CHOICES
    elif column.kind is ColumnKind.CHOICE:
        values = column.choices
    elif column.kind is ColumnKind.INTEGER:
        return DataValidation(
            type="whole",
            operator="between",
            formula1=str(column.min_value),
            formula2=str(column.max_value),
            allow_blank=True,
            showErrorMessage=True,
            errorTitle=column.header,
            error=f"{column.header}: {column.accepted_values} yazın.",
        )
    else:
        return None
    return DataValidation(
        type="list",
        formula1='"' + ",".join(values) + '"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle=column.header,
        error=f"{column.header}: şunlardan birini seçin: {', '.join(values)}.",
    )


def _build_catalog_sheet(ws: Worksheet) -> None:
    ws.title = CATALOG_SHEET
    _write_catalog_header(ws, with_comments=True)
    for index, column in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        if column.is_text_formatted:
            # Sütun varsayılan biçimi: kullanıcının yazacağı yeni hücreler metin kalır.
            ws.column_dimensions[letter].number_format = "@"
        validation = _validation_for(column)
        if validation is not None:
            validation.add(f"{letter}2:{letter}{_EXCEL_LAST_ROW}")
            ws.add_data_validation(validation)


def _build_dictionary_sheet(ws: Worksheet) -> None:
    for index, (header, width) in enumerate(_DICTIONARY_HEADERS, start=1):
        cell = ws.cell(row=1, column=index, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _OPTIONAL_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _HEADER_BORDER
        ws.column_dimensions[get_column_letter(index)].width = width
    for row, column in enumerate(COLUMNS, start=2):
        values = (
            column.header,
            "Evet" if column.required else "Hayır",
            column.description,
            column.accepted_values,
            column.blank_rule,
            column.example or "—",
        )
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.alignment = _WRAP_TOP
            # Örnek "894.353" ya da "1452" de metin olarak kalsın.
            cell.number_format = "@"
        ws.cell(row=row, column=1).font = _HEADER_FONT
        if column.required:
            ws.cell(row=row, column=1).fill = _REQUIRED_FILL
    ws.freeze_panes = "A2"

    notes_title_row = len(COLUMNS) + 3
    ws.cell(row=notes_title_row, column=1, value="Notlar").font = _HEADER_FONT
    last_column = get_column_letter(len(_DICTIONARY_HEADERS))
    for offset, note in enumerate(DICTIONARY_NOTES, start=1):
        row = notes_title_row + offset
        cell = ws.cell(row=row, column=1, value=note)
        cell.alignment = _WRAP_TOP
        ws.merge_cells(f"A{row}:{last_column}{row}")
        # Birleşik hücrede Excel satır yüksekliğini kendisi ayarlamaz.
        ws.row_dimensions[row].height = 15 * (1 + len(note) // 150)


def _build_example_sheet(ws: Worksheet) -> None:
    _write_catalog_header(ws, with_comments=False)
    for row, example in enumerate(EXAMPLE_ROWS, start=2):
        for index, column in enumerate(COLUMNS, start=1):
            value = example.get(column.key)
            if value is None:
                continue
            cell = ws.cell(row=row, column=index, value=value)
            if column.is_text_formatted:
                cell.number_format = "@"


def build_catalog_template() -> bytes:
    """Katalog Excel şablonunu (.xlsx) üretir: Katalog + Sütunlar + Örnek sayfaları."""
    wb = Workbook()
    wb.properties.creator = _COMMENT_AUTHOR
    wb.properties.title = TEMPLATE_DOCUMENT_NAME
    catalog = wb.active
    assert catalog is not None  # yeni çalışma kitabında etkin sayfa daima vardır
    _build_catalog_sheet(catalog)
    _build_dictionary_sheet(wb.create_sheet(COLUMNS_SHEET))
    _build_example_sheet(wb.create_sheet(EXAMPLE_SHEET))
    wb.active = 0
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
