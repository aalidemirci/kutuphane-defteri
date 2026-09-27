"""Dışa aktarım dosyasının içe aktarımı — İçe Aktarma'nın "dışa aktarım dosyası" kipi (F10).

Tasarım §8.4: "İçe aktarıcının 'dışa aktarım dosyası' kipi barkodu ve kayıt no'yu KORUR,
sayaçları ilerletir. Gidiş-dönüş testi bu kipte tanımlanır (EK-10)." Dosyanın biçimi
`export_schema`'dadır; üreticisi `disa_aktarim.py`'dir. Olağan Excel içe aktarımından
(`services.import_service`) farkları:

- **Numara korunur, asla yeniden kullanılmaz.** Barkod ve kayıt no dosyadan gelir; her
  numara yazmadan ÖNCE dört kaynağa karşı denetlenir: kayıtlı nüsha (silinmiş olanlar
  dahil — düz `unique`), boş barkod aralığında ayrılmış numara (bağlanmış, bağlanmamış ya
  da iptal edilmiş — F4), bu kurulumun numara sayacının gerisinde kalan numara (bir kez
  verilmiş sayılır) ve dosyanın kendi içindeki tekrar. Üye kartı numarası 8 hanedir ve
  9 ile başlar; nüsha barkodu 10 hanedir ve yılla başlar — biçim denetimi kart
  numaralarıyla (`IssuedCard`) çakışmayı yapısal olarak dışlar (kart biçimindeki kod ayrı
  iletiyle reddedilir). Yazmadan SONRA her yılın sayacı dosyadaki "Numara Sayaçları"na ve
  aktarılan en büyük numaraya ilerletilir: kaynak kurulumda silinmiş nüshaların ve
  bağlanmamış boş etiketlerin numaraları da yeni kurulumda bir daha verilmez.
- **Eser gruplaması dosyadan.** Aynı "Eser No"lu satırlar tek esere bağlanır; var olan
  eserlerle eşleştirme (kova, şüpheli satır kararı) YAPILMAZ — dosya bir kataloğun
  kendisidir, liste değildir. Eserin künyesi grubun ilk satırından alınır.
- **Edinim dosyadan.** Nüshalar dosyadaki edinim yolu, tarihi ve birim fiyatıyla açılan
  edinimlere bağlanır (aynı yol + tarih + fiyat + karar tek edinim — HER yolda). Komisyon
  kararı kararın tarihi ve sayısıyla yeniden kurulur: bağışta "Bağış değerlendirme" (Md.
  10/3), öbür yollarda "Kaynak seçimi" türüyle (F10 düzeltme turu: bağış dışı edinimin kararı
  sessizce düşüyordu). Başkan ve katılımcı adları dosyada yoktur (kişisel veri dışa
  aktarılmaz) ve kararın notu bunu söyler.
- **Durum.** Rafta, Onarımda, Kayıp ve kayıttan düşülmüş/devredilmiş durumlar korunur
  (kayıttan çıkışın tarihi dosyadaki "Kayıttan Çıkış Tarihi"dir). "Ödünçte" ve "Sınıf
  kitaplığında" nüsha "Rafta" açılır: ödünç ve teslim kaydı kişisel veri taşıdığı için
  dosyada yoktur ve kayıtsız "Ödünçte" nüsha masada iade alınamazdı. Satır uyarısı söyler.
- **Bölümler dosyadan.** "Bölümler" sayfasındaki bölümler (ad, DOS aralığı, kısa tarif,
  sıra) kurulumda yoksa açılır; sayfada olmayan bölüm değeri olağan yoldan sorulur.

**Yalnız boş kataloga** (F10 düzeltme turu): katalogda canlı eser varken önizleme de uygulama
da reddedilir (`NOT_EMPTY_MESSAGE`). Eşleştirme yapılmadığı için dolu kataloğa ya da ikinci
bir dosyayla uygulama nüshasız eserleri çoğaltır, kısmen aktarılmış eserin nüshasını ikiz bir
esere bağlardı.

**Hepsi ya da hiçbiri** (F10 düzeltme turu): aktarılamayan satır varken uygulama reddedilir
(`ROWS_NOT_IMPORTED_MESSAGE`). Sayaç dosyadaki "Numara Sayaçları"na ilerlediği için reddedilen
satırın numarası bu kurulumda bir daha kullanılamazdı: dosya düzeltilip yeniden
uygulandığında aynı satır "sayacın gerisinde" diye reddedilir, kitap yeni etiket isterdi.

**Olağan "Excel listesi" yoluna verilmez** (F10 düzeltme turu): o yol dosyayı "Bilgi"
sayfasından tanır ve reddeder (`is_export_file`, `EXPORT_FILE_AS_LIST_MESSAGE`); yoksa bütün
nüshalar yeni numara alır, eski etiket başka kitabı açar, kayıttan çıkmış nüsha "Rafta"
açılırdı.

**Önizleme = uygulama** (F3 ilkesi): tek yazma yolu (`_ingest`) önizlemede geri sarılır.
Aynı içerik ikinci kez UYGULANAMAZ (`payload_sha256`). TMY 32/3 durdurması sürerken baştan
reddedilir (edinim ve programa aktarım — F9, K4). Dış istek YOKTUR.
"""

from __future__ import annotations

import copy
import hashlib
import json
import logging
from collections import defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Final
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.kutuphane import barcode as barcode_module
from apps.kutuphane import card_numbers, import_schema, keys
from apps.kutuphane import export_schema as sema
from apps.kutuphane.models import (
    TERMINAL_COPY_STATUSES,
    Acquisition,
    AcquisitionMethod,
    CatalogImportSource,
    CatalogImportStatus,
    ClassificationSource,
    CommissionDecision,
    CommissionDecisionType,
    Copy,
    CopyCounter,
    CopyStatus,
    ReservedBarcode,
    ResourceType,
    Section,
    Work,
)
from apps.kutuphane.services import catalog as catalog_service
from apps.kutuphane.services import import_service, tmy_kapisi
from apps.kutuphane.services.import_service import (
    BUCKET_EXISTING,
    BUCKET_NEW,
    BUCKET_SKIPPED,
    CatalogImportReport,
    ImportRowReport,
    ParsedFile,
)
from apps.okul.excel_ogrenci import ParserError

logger = logging.getLogger("kutuphane_defteri.kutuphane")

# ---------------------------------------------------------------------------
# İletiler (kullanıcı metni — sözlüğe uyar; kişisel veri yok)
# ---------------------------------------------------------------------------
NOT_EXPORT_FILE_MESSAGE: Final = (
    "Bu dosya Kütüphane Defteri'nin dışa aktarım dosyası değil (“Bilgi” sayfasında dosya türü "
    "yok). Okulun kendi listesini “Excel listesi” olarak içe aktarın."
)
EXPORT_FILE_AS_LIST_MESSAGE: Final = (
    "Bu dosya Kütüphane Defteri'nin dışa aktarım dosyası. “Excel listesi” olarak içe aktarılırsa "
    "bütün nüshalar yeni numara alır ve kitapların üzerindeki etiketler başka kayıtları gösterir. "
    "“İçe aktarılacak dosya” seçicisinde “Dışa aktarım dosyası”nı seçin."
)
NOT_EMPTY_MESSAGE: Final = (
    "Katalogda kayıtlı eser var. Dışa aktarım dosyası yalnız boş bir kataloga geri yüklenir: "
    "dosyadaki eserler katalogdakilerle eşleştirilmez, dolu kataloğa aktarım eserleri çoğaltırdı."
)
ROWS_NOT_IMPORTED_MESSAGE: Final = (
    "Dosyada aktarılamayan {sayi} satır var. Aktarım bütün satırlarla yapılır: numara sayaçları "
    "dosyaya göre ilerlediği için aktarılmayan satırın barkodu bu kurulumda bir daha "
    "kullanılamazdı. Önizlemedeki nedenlere göre dosyayı düzeltin; düzeltilemeyen satırı "
    "dosyadan silin (o kitap sonra yeni etiketle kaydedilir) ve yeniden önizleyin."
)
UNSUPPORTED_VERSION_MESSAGE: Final = (
    "Dışa aktarım dosyasının şema sürümü (“{surum}”) bu programda okunamıyor. Dosyayı "
    "programın güncel sürümüyle yeniden dışa aktarın."
)
UNREADABLE_MESSAGE: Final = (
    "Dosya Excel olarak okunamadı. Dışa aktarım dosyası .xlsx biçimindedir; dosyayı "
    "değiştirmeden yükleyin."
)
MISSING_COLUMNS_MESSAGE: Final = "“Katalog” sayfasında şu sütunlar yok: {sutunlar}."
NO_CATALOG_MESSAGE: Final = "Dışa aktarım dosyasında “Katalog” sayfası yok."

WORK_NO_MESSAGE: Final = "“Eser No” boş ya da sayı değil; satır içe aktarılmadı."
TITLE_MESSAGE: Final = "Eser adı boş; satır içe aktarılmadı."
BARCODE_FORMAT_MESSAGE: Final = (
    "Barkod “{kod}” 10 haneli bir nüsha barkodu değil (yıl + altı hane); satır içe aktarılmadı."
)
CARD_SHAPED_MESSAGE: Final = (
    "“{kod}” üye kartı numarası biçiminde (8 hane); nüsha barkodu 10 hanedir. Satır içe "
    "aktarılmadı."
)
DUPLICATE_BARCODE_MESSAGE: Final = (
    "Barkod {kod} dosyada birden çok satırda; yalnız ilk satır aktarılır."
)
ACCESSION_MISMATCH_MESSAGE: Final = (
    "Kayıt no ({kayit}) barkodun sayı hâli değil ({kod}); satır içe aktarılmadı."
)
COPY_EXISTS_MESSAGE: Final = "Barkod {kod} bu kurulumda kayıtlı bir nüshada; satır aktarılmadı."
COPY_DELETED_MESSAGE: Final = (
    "Barkod {kod} bu kurulumda silinmiş bir nüshaya verilmişti; numara yeniden kullanılmaz. "
    "Satır aktarılmadı."
)
RESERVED_MESSAGE: Final = (
    "Barkod {kod} bu kurulumda bir boş barkod aralığında ayrılmış; numara yeniden "
    "kullanılmaz. Satır aktarılmadı."
)
COUNTER_MESSAGE: Final = (
    "Barkod {kod} bu kurulumun numara sayacının gerisinde: numara bu kurulumda daha önce "
    "verilmiş sayılır ve yeniden kullanılmaz. Satır aktarılmadı."
)
DONATION_DECISION_MESSAGE: Final = (
    "Bağış satırında komisyon kararının tarihi yok (Okul Kütüphaneleri Yönetmeliği Md. "
    "10/3); satır içe aktarılmadı."
)
STATUS_RESET_MESSAGE: Final = (
    "Dosyada “{durum}”: ödünç ve teslim kayıtları dışa aktarıma girmez; nüsha “Rafta” açıldı."
)
LOST_NOTE: Final = (
    "Dosyada “Kayıp”: kayıp dosyası dışa aktarıma girmez; nüsha “Kayıp” açıldı. Sayımda "
    "okutulursa rafa döner."
)
REPAIR_NOTE: Final = (
    "Dosyada “Onarımda”: onarım kaydı dışa aktarıma girmez; nüsha “Onarımda” açıldı "
    "(Eser Ayrıntısı → “Onarımdan dön”)."
)
EXIT_DATE_MISSING_MESSAGE: Final = (
    "Kayıttan çıkmış nüshanın “Kayıttan Çıkış Tarihi” boş; çıkış tarihi içe aktarma günü "
    "sayıldı."
)
METHOD_DEFAULT_MESSAGE: Final = "“Edinim Yolu” boş; “Mevcut koleksiyon (programa aktarım)” sayıldı."
DATE_DEFAULT_MESSAGE: Final = "“Edinim Tarihi” boş; içe aktarma günü yazıldı."
COPYLESS_COPIES_MESSAGE: Final = "Barkodu boş satırda nüsha açılmaz; yalnız eser kaydı açıldı."
WORK_MISMATCH_MESSAGE: Final = (
    "Bu satırın künyesi aynı “Eser No”lu ilk satırdan (satır {ilk}) farklı; eser ilk satırdaki "
    "künyeyle açıldı."
)
#: Dosyadan açılan edinimin ve yeniden kurulan kararın notu (kişisel veri yok).
EXPORT_ACQUISITION_NOTE: Final = "Dışa aktarım dosyasından kayda alındı."
EXPORT_DECISION_NOTE: Final = (
    "Dışa aktarım dosyasından yeniden kuruldu: kararın tarihi ve sayısı dosyadandır; başkan ve "
    "katılımcı adları dışa aktarıma girmez."
)

#: Ödünç ve teslim kaydı olmadan korunamayan durumlar → "Rafta".
RESET_STATUSES: Final[tuple[str, ...]] = (CopyStatus.ON_LOAN, CopyStatus.DELIVERED)

HEADER_SCAN_LIMIT: Final = 15
#: Toplu `__in` sorgularının parça boyu (eski SQLite değişken sınırına karşı).
_PARCA: Final = 500


# ---------------------------------------------------------------------------
# Ayrıştırma
# ---------------------------------------------------------------------------
@dataclass(slots=True)
class ExportFileRow:
    """Dışa aktarım dosyasının bir satırı (çözülmüş değerler, hata ve uyarılar)."""

    row_number: int
    work_no: int | None = None
    title: str = ""
    authors: str = ""
    translator: str = ""
    publisher: str = ""
    edition: str = ""
    publish_year: int | None = None
    isbn: str = ""
    subjects: str = ""
    classification_code: str = ""
    classification_source: str = ClassificationSource.MANUAL
    language: str = ""
    resource_type: str = ResourceType.BOOK
    call_number: str = ""
    work_section: str = ""
    barcode: str = ""
    accession_no: int | None = None
    section: str = ""
    old_register_no: str = ""
    external_asset_ref: str = ""
    status: str = CopyStatus.AVAILABLE
    file_status: str = CopyStatus.AVAILABLE
    method: str = AcquisitionMethod.EXISTING_STOCK
    acquisition_date: date | None = None
    unit_price: Decimal | None = None
    decision_date: date | None = None
    decision_no: str = ""
    exit_date: date | None = None
    is_bound_periodical: bool = False
    is_reference: bool = False
    is_out_of_print: bool = False
    is_rare_or_manuscript: bool = False
    label_printed: bool = False
    spine_label_printed: bool = False
    label_verified: bool = False
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def is_copy_row(self) -> bool:
        return bool(self.barcode)

    @property
    def importable(self) -> bool:
        return not self.errors

    def work_signature(self) -> tuple[Any, ...]:
        """Eser künyesi — aynı "Eser No"lu satırların tutarlılığı buna bakar."""
        return (
            self.title,
            self.authors,
            self.translator,
            self.publisher,
            self.edition,
            self.publish_year,
            self.isbn,
            self.subjects,
            self.classification_code,
            self.classification_source,
            self.language,
            self.resource_type,
            self.call_number,
            self.work_section,
        )

    def canonical(self) -> list[Any]:
        """İçerik özetine giren değerler (satır numarası ve iletiler dışarıda)."""
        return [
            self.work_no,
            *self.work_signature(),
            self.barcode,
            self.accession_no,
            self.section,
            self.old_register_no,
            self.external_asset_ref,
            self.file_status,
            self.method,
            self.acquisition_date.isoformat() if self.acquisition_date else None,
            str(self.unit_price) if self.unit_price is not None else None,
            self.decision_date.isoformat() if self.decision_date else None,
            self.decision_no,
            self.exit_date.isoformat() if self.exit_date else None,
            self.is_bound_periodical,
            self.is_reference,
            self.is_out_of_print,
            self.is_rare_or_manuscript,
            self.label_printed,
            self.spine_label_printed,
            self.label_verified,
        ]


@dataclass(slots=True)
class ExportSection:
    name: str
    dewey_from: str = ""
    dewey_to: str = ""
    description: str = ""
    sort_order: int = 0


@dataclass(slots=True)
class ExportFile:
    """Dışa aktarım dosyasının okunmuş hâli."""

    version: str
    rows: list[ExportFileRow]
    sections: list[ExportSection] = field(default_factory=list)
    counters: dict[int, int] = field(default_factory=dict)
    unknown_headers: list[str] = field(default_factory=list)


def _metin(value: object) -> str:
    return import_service._cell_text(value)


def _tamsayi(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else None
    ham = str(value).strip().replace(".", "") if isinstance(value, str) else str(value)
    return int(ham) if ham.isascii() and ham.isdigit() else None


def _tarih(value: object) -> tuple[date | None, bool]:
    """Tarih hücresi → (tarih, geçerli mi). Boş hücre (None, True)."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if isinstance(value, datetime):
        return value.date(), True
    if isinstance(value, date):
        return value, True
    try:
        return datetime.strptime(str(value).strip(), "%d.%m.%Y").date(), True
    except ValueError:
        return None, False


def _ondalik(value: object) -> tuple[Decimal | None, bool]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, True
    if isinstance(value, bool):
        return None, False
    try:
        if isinstance(value, int | float | Decimal):
            tutar = Decimal(str(value))
        else:
            tutar = Decimal(str(value).strip().replace(".", "").replace(",", "."))
    except InvalidOperation:
        return None, False
    if tutar < 0:
        return None, False
    return tutar.quantize(Decimal("0.01")), True


def _evet_hayir(satir: ExportFileRow, value: object, baslik: str) -> bool:
    try:
        return bool(import_schema.yes_no_value(value))
    except import_schema.CatalogValueError as exc:
        satir.errors.append(f"“{baslik}” sütunu: {exc}")
        return False


def _secim(
    satir: ExportFileRow,
    value: object,
    secenekler: Iterable[tuple[str, str]],
    baslik: str,
    varsayilan: str,
) -> str:
    if value is None or _metin(value) == "":
        return varsayilan
    kod = sema.label_to_code(secenekler, value)
    if kod is None:
        satir.errors.append(f"“{baslik}” sütunu: “{_metin(value)}” tanınmadı.")
        return varsayilan
    return kod


def _kaynak_turu(satir: ExportFileRow, value: object) -> str:
    """Kaynak türü: model etiketleri (E-kitap, E-veri tabanı dahil) ya da içe aktarım seçenekleri."""
    if value is None or _metin(value) == "":
        return ResourceType.BOOK
    kod = sema.label_to_code(ResourceType.choices, value)
    if kod is not None:
        return kod
    try:
        secenek = import_schema.resource_type_value(value)
    except import_schema.CatalogValueError as exc:
        satir.errors.append(f"“Kaynak Türü” sütunu: {exc}")
        return ResourceType.BOOK
    return import_schema.RESOURCE_TYPE_CODES[secenek]


def _hucre(ham: Sequence[Any], sutunlar: Mapping[str, int], anahtar: str) -> Any:
    indis = sutunlar.get(anahtar)
    if indis is None or indis >= len(ham):
        return None
    return ham[indis]


def _barkod_denetle(satir: ExportFileRow, value: object) -> None:
    ham = _metin(value)
    if not ham:
        return
    rakamlar = barcode_module.normalize_scan(ham)
    if card_numbers.has_card_shape(rakamlar) and len(rakamlar) == card_numbers.CARD_LENGTH:
        satir.errors.append(CARD_SHAPED_MESSAGE.format(kod=ham))
        return
    if not barcode_module.has_scan_year_prefix(rakamlar) or rakamlar != ham.replace("-", ""):
        satir.errors.append(BARCODE_FORMAT_MESSAGE.format(kod=ham))
        return
    satir.barcode = rakamlar


def _satir(row_number: int, ham: Sequence[Any], sutunlar: Mapping[str, int]) -> ExportFileRow:  # noqa: C901 — şemanın 35 sütunu tek yerde çözülür
    s = ExportFileRow(row_number=row_number)

    def h(anahtar: str) -> Any:
        return _hucre(ham, sutunlar, anahtar)

    s.work_no = _tamsayi(h("work_no"))
    if s.work_no is None or s.work_no < 1:
        s.errors.append(WORK_NO_MESSAGE)
    s.title = _metin(h("title"))
    if not s.title:
        s.errors.append(TITLE_MESSAGE)
    for alan in (
        "authors",
        "translator",
        "publisher",
        "edition",
        "isbn",
        "subjects",
        "classification_code",
        "language",
        "call_number",
        "work_section",
        "old_register_no",
        "external_asset_ref",
    ):
        setattr(s, alan, _metin(h(alan)))
    s.section = _metin(h("shelf_location"))
    s.decision_no = _metin(h("decision_no"))

    yil, yil_uyarisi = import_service._year_value(h("publish_year"))
    s.publish_year = yil
    if yil_uyarisi:
        s.warnings.append(yil_uyarisi)
    s.resource_type = _kaynak_turu(s, h("resource_type"))
    s.classification_source = _secim(
        s,
        h("classification_source"),
        ClassificationSource.choices,
        "Sınıflama Kaynağı",
        ClassificationSource.MANUAL,
    )

    for alan, baslik in (
        ("is_bound_periodical", "Ciltli Süreli Yayın"),
        ("is_reference", "Danışma Kaynağı"),
        ("is_out_of_print", "Piyasada Mevcudu Yok"),
        ("is_rare_or_manuscript", "El Yazması / Nadir Eser"),
        ("label_printed", "Barkod Etiketi Basıldı"),
        ("spine_label_printed", "Sırt Etiketi Basıldı"),
        ("label_verified", "Etiket Doğrulandı"),
    ):
        setattr(s, alan, _evet_hayir(s, h(alan), baslik))

    _barkod_denetle(s, h("barcode"))
    if not s.is_copy_row:
        nusha_sayisi = _tamsayi(h("copies"))
        if nusha_sayisi not in (None, 0) and not s.errors:
            s.warnings.append(COPYLESS_COPIES_MESSAGE)
        return s

    kayit = _tamsayi(h("accession_no"))
    if kayit is not None and kayit != barcode_module.accession_no_of(s.barcode):
        s.errors.append(ACCESSION_MISMATCH_MESSAGE.format(kayit=kayit, kod=s.barcode))
    s.accession_no = barcode_module.accession_no_of(s.barcode)

    s.file_status = _secim(s, h("status"), CopyStatus.choices, "Durum", CopyStatus.AVAILABLE)
    s.status = s.file_status
    ham_yol = h("acquisition_method")
    if ham_yol is None or _metin(ham_yol) == "":
        s.warnings.append(METHOD_DEFAULT_MESSAGE)
    s.method = _secim(
        s, ham_yol, AcquisitionMethod.choices, "Edinim Yolu", AcquisitionMethod.EXISTING_STOCK
    )
    for alan, baslik in (
        ("acquisition_date", "Edinim Tarihi"),
        ("decision_date", "Komisyon Kararı Tarihi"),
        ("exit_date", "Kayıttan Çıkış Tarihi"),
    ):
        gun, gecerli = _tarih(h(alan))
        if not gecerli:
            s.errors.append(f"“{baslik}” sütunu: tarih okunamadı (gg.aa.yyyy).")
        setattr(s, alan, gun)
    if s.acquisition_date is None and not s.errors:
        s.warnings.append(DATE_DEFAULT_MESSAGE)
    fiyat, gecerli = _ondalik(h("unit_price"))
    if not gecerli:
        s.errors.append("“Birim Fiyat” sütunu: sıfır ya da artı bir sayı olmalıdır.")
    s.unit_price = fiyat
    if s.method == AcquisitionMethod.DONATION and s.decision_date is None:
        s.errors.append(DONATION_DECISION_MESSAGE)
    return s


def _bilgi(kitap: Any) -> dict[str, Any]:
    if sema.INFO_SHEET not in kitap.sheetnames:
        raise ParserError(NOT_EXPORT_FILE_MESSAGE)
    bilgi: dict[str, Any] = {}
    for satir in kitap[sema.INFO_SHEET].iter_rows(values_only=True, max_col=2):
        if satir and satir[0] is not None:
            bilgi[_metin(satir[0])] = satir[1] if len(satir) > 1 else None
    if _metin(bilgi.get(sema.INFO_KIND)) != sema.EXPORT_FILE_KIND:
        raise ParserError(NOT_EXPORT_FILE_MESSAGE)
    surum = _metin(bilgi.get(sema.INFO_VERSION))
    if surum not in sema.SUPPORTED_VERSIONS:
        raise ParserError(UNSUPPORTED_VERSION_MESSAGE.format(surum=surum or "—"))
    bilgi[sema.INFO_VERSION] = surum
    return bilgi


def _baslik_bul(satirlar: Sequence[Sequence[Any]]) -> tuple[int, dict[str, int], list[str]]:
    for indis, satir in enumerate(satirlar[:HEADER_SCAN_LIMIT]):
        eslesen: dict[str, int] = {}
        taninmayan: list[str] = []
        for sutun, hucre in enumerate(satir):
            anahtar = import_schema.match_header(hucre) or sema.match_extra_header(hucre)
            if anahtar is None:
                if _metin(hucre):
                    taninmayan.append(_metin(hucre))
            elif anahtar not in eslesen:
                eslesen[anahtar] = sutun
        if "title" in eslesen and "barcode" in eslesen:
            return indis, eslesen, taninmayan
    return -1, {}, []


def _katalog(kitap: Any) -> tuple[list[ExportFileRow], list[str]]:
    if sema.CATALOG_SHEET not in kitap.sheetnames:
        raise ParserError(NO_CATALOG_MESSAGE)
    satirlar = [list(s) for s in kitap[sema.CATALOG_SHEET].iter_rows(values_only=True)]
    baslik, sutunlar, taninmayan = _baslik_bul(satirlar)
    zorunlu = ("title", *sema.REQUIRED_EXTRA_KEYS)
    eksik = [k for k in zorunlu if k not in sutunlar]
    if baslik < 0 or eksik:
        adlar = {
            **{c.key: c.header for c in sema.BASE_COLUMNS},
            **{c.key: c.header for c in sema.EXTRA_COLUMNS},
        }
        raise ParserError(
            MISSING_COLUMNS_MESSAGE.format(
                sutunlar=", ".join(f"“{adlar[k]}”" for k in (eksik or zorunlu))
            )
        )
    sonuc: list[ExportFileRow] = []
    for indis in range(baslik + 1, len(satirlar)):
        ham = satirlar[indis]
        if all(_metin(h) == "" for h in ham):
            continue
        sonuc.append(_satir(indis + 1, ham, sutunlar))
    return sonuc, list(dict.fromkeys(taninmayan))


def _bolumler(kitap: Any) -> list[ExportSection]:
    if sema.SECTIONS_SHEET not in kitap.sheetnames:
        return []
    sonuc: list[ExportSection] = []
    for i, satir in enumerate(kitap[sema.SECTIONS_SHEET].iter_rows(values_only=True)):
        if i == 0 or not satir or not _metin(satir[0]):
            continue
        degerler = [*list(satir), None, None, None, None][:5]
        sonuc.append(
            ExportSection(
                name=_metin(degerler[0])[:80],
                dewey_from=_metin(degerler[1])[:20],
                dewey_to=_metin(degerler[2])[:20],
                description=_metin(degerler[3])[:255],
                sort_order=max(0, _tamsayi(degerler[4]) or 0),
            )
        )
    return sonuc


def _sayaclar(kitap: Any) -> dict[int, int]:
    if sema.COUNTERS_SHEET not in kitap.sheetnames:
        return {}
    sonuc: dict[int, int] = {}
    for i, satir in enumerate(kitap[sema.COUNTERS_SHEET].iter_rows(values_only=True)):
        if i == 0 or not satir or len(satir) < 2:
            continue
        yil, son = _tamsayi(satir[0]), _tamsayi(satir[1])
        if yil is not None and son is not None and 1000 <= yil <= 9999 and son >= 0:
            sonuc[yil] = max(sonuc.get(yil, 0), min(son, barcode_module.MAX_SEQUENCE))
    return sonuc


def export_content_hash(dosya: ExportFile) -> str:
    """İçerik özeti (fikirdeşlik anahtarı) — satırlar, bölümler ve sayaçlar."""
    kanonik = {
        "version": dosya.version,
        "rows": [s.canonical() for s in dosya.rows],
        "sections": [
            [b.name, b.dewey_from, b.dewey_to, b.description, b.sort_order] for b in dosya.sections
        ],
        "counters": sorted(dosya.counters.items()),
    }
    metin = json.dumps(kanonik, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return hashlib.sha256(metin.encode("utf-8")).hexdigest()


def read_export_file(file_bytes: bytes) -> ExportFile:
    """Dışa aktarım dosyasını okur; tanınmayan dosya `ParserError`."""
    try:
        kitap = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, KeyError) as exc:
        raise ParserError(UNREADABLE_MESSAGE) from exc
    try:
        bilgi = _bilgi(kitap)
        satirlar, taninmayan = _katalog(kitap)
        return ExportFile(
            version=str(bilgi[sema.INFO_VERSION]),
            rows=satirlar,
            sections=_bolumler(kitap),
            counters=_sayaclar(kitap),
            unknown_headers=taninmayan,
        )
    finally:
        kitap.close()


def is_export_file(file_bytes: bytes) -> bool:
    """Dosya programın dışa aktarım dosyası mı? ("Bilgi" sayfasındaki dosya türü).

    Olağan "Excel listesi" yolu bununla reddeder (F10 düzeltme turu). Okunamayan dosya
    (ör. Excel 97-2003) dışa aktarım dosyası sayılmaz; olağan yol kendi iletisini verir.
    """
    try:
        kitap = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, KeyError, ValueError):
        return False
    try:
        if sema.INFO_SHEET not in kitap.sheetnames:
            return False
        for satir in kitap[sema.INFO_SHEET].iter_rows(values_only=True, max_col=2, max_row=40):
            if satir and _metin(satir[0]) == sema.INFO_KIND:
                return len(satir) > 1 and _metin(satir[1]) == sema.EXPORT_FILE_KIND
        return False
    finally:
        kitap.close()


def parse_export_file(file_bytes: bytes) -> tuple[ParsedFile, str]:
    """`import_service.rows_from_file`'ın dışa aktarım kipi → (ayrıştırma, içerik özeti)."""
    dosya = read_export_file(file_bytes)
    return (
        ParsedFile(rows=[], header_row=1, unknown_headers=dosya.unknown_headers, export=dosya),
        export_content_hash(dosya),
    )


# ---------------------------------------------------------------------------
# Plan: çakışma denetimi ve bölümler (yazmadan önce)
# ---------------------------------------------------------------------------
def _parcalar(degerler: Sequence[Any]) -> Iterable[Sequence[Any]]:
    for i in range(0, len(degerler), _PARCA):
        yield degerler[i : i + _PARCA]


def _cakismalari_isaretle(satirlar: Sequence[ExportFileRow]) -> None:
    """Numara çakışmaları — satır hatası olarak (asla yeniden kullanım yok)."""
    goruldu: set[str] = set()
    adaylar: list[ExportFileRow] = []
    for s in satirlar:
        if not s.is_copy_row or not s.importable:
            continue
        if s.barcode in goruldu:
            s.errors.append(DUPLICATE_BARCODE_MESSAGE.format(kod=s.barcode))
            continue
        goruldu.add(s.barcode)
        adaylar.append(s)
    kodlar = [s.barcode for s in adaylar]
    canli: set[str] = set()
    silinmis: set[str] = set()
    ayrilmis: set[str] = set()
    for parca in _parcalar(kodlar):
        for kod, silindi in Copy.all_objects.filter(barcode__in=parca).values_list(
            "barcode", "deleted_at"
        ):
            (silinmis if silindi is not None else canli).add(str(kod))
        ayrilmis.update(
            str(k)
            for k in ReservedBarcode.objects.filter(barcode__in=parca).values_list(
                "barcode", flat=True
            )
        )
    kayitlar = [barcode_module.accession_no_of(k) for k in kodlar]
    kayitli_nolar: set[int] = set()
    for parca in _parcalar(kayitlar):
        kayitli_nolar.update(
            int(n)
            for n in Copy.all_objects.filter(accession_no__in=parca).values_list(
                "accession_no", flat=True
            )
        )
    sayaclar = dict(CopyCounter.objects.values_list("year", "last_no"))
    for s in adaylar:
        yil, sira = int(s.barcode[:4]), int(s.barcode[4:])
        if s.barcode in canli:
            s.errors.append(COPY_EXISTS_MESSAGE.format(kod=s.barcode))
        elif s.barcode in silinmis or s.accession_no in kayitli_nolar:
            s.errors.append(COPY_DELETED_MESSAGE.format(kod=s.barcode))
        elif s.barcode in ayrilmis:
            s.errors.append(RESERVED_MESSAGE.format(kod=s.barcode))
        elif sira <= int(sayaclar.get(yil, 0)):
            s.errors.append(COUNTER_MESSAGE.format(kod=s.barcode))


def _bolum_anahtari(ad: str) -> str:
    return keys.fold_search(ad) or ad.strip()


@dataclass(slots=True)
class _BolumPlani:
    """Bölüm adı → hedef: var olan bölüm, dosyadan açılacak ya da kullanıcı karşılığı."""

    mevcut: dict[str, Section]
    dosyadan: dict[str, ExportSection]
    eslenen: dict[str, int]
    yeni_adlar: dict[str, str]
    bilinmeyen: dict[str, dict[str, Any]]


def _bolum_plani(
    dosya: ExportFile,
    *,
    section_map: Mapping[Any, Any],
    new_sections: Sequence[str],
) -> _BolumPlani:
    mevcut = import_service._section_index()
    dosyadan = {
        _bolum_anahtari(b.name): b for b in dosya.sections if _bolum_anahtari(b.name) not in mevcut
    }
    plan = _BolumPlani(
        mevcut=mevcut,
        dosyadan=dosyadan,
        eslenen=import_service._normalize_section_map(section_map),
        yeni_adlar={_bolum_anahtari(ad): str(ad).strip() for ad in new_sections},
        bilinmeyen={},
    )
    for s in dosya.rows:
        if not s.importable:
            continue
        for ad in (s.section, s.work_section):
            if not ad:
                continue
            anahtar = _bolum_anahtari(ad)
            if (
                anahtar in plan.mevcut
                or anahtar in plan.dosyadan
                or anahtar in plan.eslenen
                or anahtar in plan.yeni_adlar
            ):
                continue
            kayit = plan.bilinmeyen.setdefault(
                anahtar, {"value": ad.strip(), "rows": [], "count": 0}
            )
            kayit["count"] = int(kayit["count"]) + 1
            if len(kayit["rows"]) < import_service.MAX_SECTION_SAMPLE_ROWS:
                kayit["rows"].append(s.row_number)
    return plan


def _gruplar(satirlar: Sequence[ExportFileRow]) -> dict[int, list[ExportFileRow]]:
    """Eser No → satırlar (ilk görülme sırasıyla); künye farkı uyarısı."""
    gruplar: dict[int, list[ExportFileRow]] = {}
    for s in satirlar:
        if not s.importable or s.work_no is None:
            continue
        grup = gruplar.setdefault(s.work_no, [])
        if grup and s.work_signature() != grup[0].work_signature():
            s.warnings.append(WORK_MISMATCH_MESSAGE.format(ilk=grup[0].row_number))
        grup.append(s)
    return gruplar


def _ensure_open(dosya: ExportFile) -> None:
    """TMY 32/3 durdurması sürerken baştan reddedilir (edinim ve programa aktarım — F9, K4)."""
    yollar = {s.method for s in dosya.rows if s.is_copy_row and s.importable}
    for yol in sorted(yollar) or [AcquisitionMethod.EXISTING_STOCK]:
        tmy_kapisi.ensure_open(tmy_kapisi.edinim_islemi(yol))


# ---------------------------------------------------------------------------
# Yazma (önizleme de aynı yoldan geçer)
# ---------------------------------------------------------------------------
@dataclass(slots=True)
class _Durum:
    """Yazma sırasında paylaşılan önbellekler ve sayaçlar."""

    bolumler: dict[str, Section]
    edinimler: dict[tuple[Any, ...], Acquisition] = field(default_factory=dict)
    kararlar: dict[tuple[Any, ...], CommissionDecision] = field(default_factory=dict)
    sayaclar: dict[str, int] = field(default_factory=lambda: defaultdict(int))
    en_buyuk: dict[int, int] = field(default_factory=dict)


def _bolumleri_ac(dosya: ExportFile, plan: _BolumPlani) -> dict[str, Section]:
    """Dosyadaki bölümleri (yoksa) ve kullanıcının "yeni bölüm" kararlarını açar."""
    bolumler: dict[str, Section] = dict(plan.mevcut)
    for anahtar, eslenen_pk in plan.eslenen.items():
        hedef = Section.objects.filter(pk=eslenen_pk).first()
        if hedef is None:
            raise ValidationError({"section_map": "Seçilen bölüm bulunamadı."})
        bolumler.setdefault(anahtar, hedef)
    for anahtar, b in plan.dosyadan.items():
        bolumler[anahtar] = catalog_service.create_section(
            name=b.name,
            dewey_from=b.dewey_from,
            dewey_to=b.dewey_to,
            description=b.description,
            sort_order=b.sort_order,
        )
    for anahtar, ad in plan.yeni_adlar.items():
        if anahtar not in bolumler:
            bolumler[anahtar] = catalog_service.create_section(name=ad)
    return bolumler


def _bolum(durum: _Durum, ad: str) -> Section | None:
    return durum.bolumler.get(_bolum_anahtari(ad)) if ad else None


def _karar_turu(yol: str) -> str:
    """Yeniden kurulan kararın türü: bağışta "Bağış değerlendirme", öbür yollarda "Kaynak
    seçimi" (ayıklama kararı edinime bağlanamaz — `catalog.ensure_acquisition_decision`)."""
    if yol == AcquisitionMethod.DONATION:
        return str(CommissionDecisionType.DONATION_REVIEW)
    return str(CommissionDecisionType.SELECTION)


def _karar_anahtari(s: ExportFileRow) -> tuple[Any, ...] | None:
    if s.decision_date is None:
        return None
    return (_karar_turu(s.method), s.decision_date, s.decision_no)


def _karar(durum: _Durum, s: ExportFileRow) -> CommissionDecision | None:
    """Edinimin komisyon kararı: aynı tür + tarih + sayılı karar varsa o, yoksa yeniden kurulur.

    Bütün yollarda kurulur (F10 düzeltme turu): satın alma, Bakanlık gönderimi ve değişimin
    "Kaynak seçimi" kararı da dosyaya yazılıyor ama geri yüklemede sessizce düşüyordu.
    """
    anahtar = _karar_anahtari(s)
    if anahtar is None:
        return None
    if anahtar in durum.kararlar:
        return durum.kararlar[anahtar]
    tur = anahtar[0]
    karar: CommissionDecision | None = (
        CommissionDecision.objects.filter(
            decision_type=tur,
            decision_date=s.decision_date,
            decision_no=s.decision_no,
        )
        .order_by("pk")
        .first()
    )
    if karar is None:
        # Başkan adı dosyada yoktur (kişisel veri dışa aktarılmaz); alan boş kalır ve
        # kararın notu bunu söyler. Karar içe aktarmanın savepoint'i içinde açılır.
        karar = CommissionDecision.objects.create(
            decision_type=tur,
            decision_date=s.decision_date,
            decision_no=s.decision_no[:40],
            chair_name="",
            notes=EXPORT_DECISION_NOTE,
        )
        durum.sayaclar["decisions_created"] += 1
    return karar


def _edinim(durum: _Durum, s: ExportFileRow, bugun: date) -> tuple[Acquisition, tuple[Any, ...]]:
    gun = s.acquisition_date or bugun
    anahtar = (s.method, gun, s.unit_price, _karar_anahtari(s))
    if anahtar in durum.edinimler:
        return durum.edinimler[anahtar], anahtar
    edinim = catalog_service.create_acquisition(
        method=s.method,
        date=gun,
        unit_price=s.unit_price,
        commission_decision=_karar(durum, s),
        notes=EXPORT_ACQUISITION_NOTE,
    )
    return edinim, anahtar


def _eser_ac(durum: _Durum, ilk: ExportFileRow) -> Work:
    return catalog_service.create_work(
        title=ilk.title,
        authors=ilk.authors,
        translator=ilk.translator,
        publisher=ilk.publisher,
        edition=ilk.edition,
        publish_year=ilk.publish_year,
        isbn=ilk.isbn,
        subjects=ilk.subjects,
        classification_code=ilk.classification_code,
        classification_source=ilk.classification_source,
        call_number=ilk.call_number,
        resource_type=ilk.resource_type,
        language=ilk.language,
        section=_bolum(durum, ilk.work_section),
    )


def durumu_coz(s: ExportFileRow) -> tuple[str, str]:
    """Dosyadaki durum → (açılacak durum, satır uyarısı). Kayıtsız korunamayanlar "Rafta"."""
    if s.file_status in RESET_STATUSES:
        return CopyStatus.AVAILABLE, STATUS_RESET_MESSAGE.format(
            durum=CopyStatus(s.file_status).label
        )
    if s.file_status == CopyStatus.LOST:
        return s.file_status, LOST_NOTE
    if s.file_status == CopyStatus.IN_REPAIR:
        return s.file_status, REPAIR_NOTE
    if s.file_status in TERMINAL_COPY_STATUSES and s.exit_date is None:
        return s.file_status, EXIT_DATE_MISSING_MESSAGE
    return s.file_status, ""


def _nusha_ac(durum: _Durum, eser: Work, s: ExportFileRow, edinim: Acquisition) -> Copy:
    """Numarası KORUNARAK nüsha açar — kurallar `validate_new_copy`'den geçer."""
    alanlar = catalog_service.validate_new_copy(
        work=eser,
        acquisition=edinim,
        section=_bolum(durum, s.section),
        is_reference=s.is_reference,
        is_bound_periodical=s.is_bound_periodical,
        is_out_of_print=s.is_out_of_print,
        is_rare_or_manuscript=s.is_rare_or_manuscript,
        old_register_no=s.old_register_no,
        external_asset_ref=s.external_asset_ref,
    )
    simdi = timezone.now()
    nusha = Copy(
        accession_no=barcode_module.accession_no_of(s.barcode),
        barcode=s.barcode,
        status=durumu_coz(s)[0],
        label_printed_at=simdi if s.label_printed else None,
        spine_label_printed_at=simdi if s.spine_label_printed else None,
        label_verified_at=simdi if s.label_verified else None,
        **alanlar,
    )
    nusha.full_clean()
    nusha.save()
    if nusha.status in TERMINAL_COPY_STATUSES and s.exit_date is not None:
        # Kayıttan çıkış tarihi (onay tarihi) nüshanın son güncellemesine yazılır:
        # onay kaydı (ayıklama ya da sayım) dosyada yoktur ve 34/1 hesabı böyle nüshada
        # çıkış gününü son güncellemeden okur (`selectors_sayim.tmy_34_1`, "eski veri").
        an = timezone.make_aware(datetime.combine(s.exit_date, time(12, 0)))
        Copy.all_objects.filter(pk=nusha.pk).update(updated_at=an)
    return nusha


def _rapor_satiri(s: ExportFileRow) -> ImportRowReport:
    return ImportRowReport(
        row=s.row_number,
        title=s.title,
        bucket=BUCKET_SKIPPED,
        copies=1 if s.is_copy_row else 0,
        shelf_location=s.section,
        is_reference=s.is_reference,
        classification_source=s.classification_source,
        issues=[*s.errors, *s.warnings],
    )


def _ingest(
    dosya: ExportFile,
    *,
    plan: _BolumPlani,
    payload_sha256: str,
    file_name: str,
) -> CatalogImportReport:
    """Dosyayı YAZAR (önizleme de burayı koşar ve geri sarar).

    Satır hatası yazmayı durdurmaz: önizleme bütün satırların nedenini gösterir. Uygulama ise
    aktarılamayan satır varken sonunda reddedilir (`apply_export` — hepsi ya da hiçbiri).
    """
    bugun = timezone.localdate()
    durum = _Durum(bolumler=_bolumleri_ac(dosya, plan))
    raporlar = {s.row_number: _rapor_satiri(s) for s in dosya.rows}
    for satirlar in _gruplar(dosya.rows).values():
        eser: Work | None = None
        ilk_satir: int | None = None
        for s in satirlar:
            rapor = raporlar[s.row_number]
            try:
                with transaction.atomic():
                    yeni_eser = eser is None
                    hedef = eser if eser is not None else _eser_ac(durum, satirlar[0])
                    edinim_anahtari: tuple[Any, ...] | None = None
                    if s.is_copy_row:
                        edinim, edinim_anahtari = _edinim(durum, s, bugun)
                        _nusha_ac(durum, hedef, s, edinim)
            except ValidationError as exc:
                rapor.issues = [*s.errors, *s.warnings, *import_service._hata_iletileri(exc)]
                durum.sayaclar["error_rows"] += 1
                continue
            rapor.issues = [*s.errors, *s.warnings]
            if edinim_anahtari is not None:
                durum.edinimler[edinim_anahtari] = edinim
                rapor.copies_created = 1
                durum.sayaclar["copies_created"] += 1
                yil, sira = int(s.barcode[:4]), int(s.barcode[4:])
                durum.en_buyuk[yil] = max(durum.en_buyuk.get(yil, 0), sira)
                uyari = durumu_coz(s)[1]
                if uyari:
                    rapor.issues.append(uyari)
                if s.file_status in RESET_STATUSES:
                    durum.sayaclar["status_reset"] += 1
                if s.file_status in TERMINAL_COPY_STATUSES:
                    durum.sayaclar["exited_copies"] += 1
            if yeni_eser:
                eser, ilk_satir = hedef, s.row_number
                rapor.bucket = BUCKET_NEW
                durum.sayaclar["new_works"] += 1
            else:
                rapor.bucket = BUCKET_EXISTING
                rapor.work_row = ilk_satir
                durum.sayaclar["existing_matches"] += 1
            rapor.work = hedef.pk
    atlanan = sum(1 for s in dosya.rows if not s.importable)
    durum.sayaclar["counters_advanced"] = _sayaclari_ilerlet(dosya.counters, durum.en_buyuk)

    rapor_nesnesi = CatalogImportReport(
        payload_sha256=payload_sha256,
        source=CatalogImportSource.EXPORT,
        schema_version=dosya.version,
        file_name=file_name,
        rows=[raporlar[s.row_number] for s in dosya.rows],
        unknown_sections=sorted(plan.bilinmeyen.values(), key=lambda k: str(k["value"])),
        unknown_headers=dosya.unknown_headers,
        missing_columns=[],
        pending_decisions=[],
        label_batch=None,
    )
    sayac = dict.fromkeys(
        (
            "new_works",
            "existing_matches",
            "suspect",
            "skipped_rows",
            "error_rows",
            "copies_created",
            "reference_defaults",
            "estimated_codes",
            "corrections",
            "status_reset",
            "exited_copies",
            "decisions_created",
            "counters_advanced",
        ),
        0,
    )
    sayac.update(durum.sayaclar)
    sayac["skipped_rows"] = atlanan
    rapor_nesnesi.stats = {
        "total_rows": len(dosya.rows),
        "imported_rows": sayac["new_works"] + sayac["existing_matches"],
        "sections_created": len(plan.dosyadan)
        + sum(1 for a in plan.yeni_adlar if a not in plan.mevcut and a not in plan.dosyadan),
        **sayac,
    }
    return rapor_nesnesi


def not_imported_rows(rapor: CatalogImportReport) -> int:
    """Aktarılamayan satır sayısı: okunurken reddedilen + yazılırken hata veren."""
    return int(rapor.stats.get("skipped_rows", 0)) + int(rapor.stats.get("error_rows", 0))


def _sayaclari_ilerlet(dosya_sayaclari: Mapping[int, int], en_buyuk: Mapping[int, int]) -> int:
    """Her yılın sayacını dosyadaki sayaca ve aktarılan en büyük numaraya ilerletir.

    Sayaç GERİ GİTMEZ (`services.numbering`): hedef, var olanın gerisindeyse dokunulmaz.
    Dönen: ilerletilen yıl sayısı.
    """
    ilerleyen = 0
    for yil in sorted(set(dosya_sayaclari) | set(en_buyuk)):
        hedef = max(int(dosya_sayaclari.get(yil, 0)), int(en_buyuk.get(yil, 0)))
        if hedef < 1:
            continue
        sayac, _ = CopyCounter.objects.select_for_update().get_or_create(
            year=yil, defaults={"last_no": 0}
        )
        if hedef > int(sayac.last_no):
            sayac.last_no = hedef
            sayac.save(update_fields=["last_no"])
            ilerleyen += 1
    return ilerleyen


def ensure_empty_catalog(payload_sha256: str = "") -> None:
    """Dışa aktarım dosyası yalnız boş kataloga (canlı eser yokken) geri yüklenir.

    Aynı dosya daha önce uygulanmışsa ileti bunu da söyler (dolu kataloğun en olası nedeni).
    """
    if not Work.objects.exists():
        return
    ileti = NOT_EMPTY_MESSAGE
    onceki = import_service.applied_run(payload_sha256) if payload_sha256 else None
    if onceki is not None:
        gun = timezone.localtime(onceki.created_at).strftime("%d.%m.%Y")
        ileti += f" Bu dosya {gun} tarihinde zaten içe aktarıldı."
    raise ValidationError({"file": ileti})


def _hazirla(
    parsed: ParsedFile,
    *,
    section_map: Mapping[Any, Any] | None,
    new_sections: Sequence[str] | None,
    payload_sha256: str = "",
) -> tuple[ExportFile, _BolumPlani]:
    # Plan satırlara ileti ekler (çakışma, künye farkı): aynı ayrıştırma sonucu önizlemede
    # ve uygulamada yeniden kullanılırsa iletiler birikmesin diye kopya üzerinde çalışılır.
    dosya: ExportFile = copy.deepcopy(parsed.export)
    _ensure_open(dosya)
    ensure_empty_catalog(payload_sha256)
    _cakismalari_isaretle(dosya.rows)
    plan = _bolum_plani(dosya, section_map=section_map or {}, new_sections=new_sections or [])
    return dosya, plan


def preview_export(
    parsed: ParsedFile,
    *,
    payload_sha256: str,
    file_name: str = "",
    section_map: Mapping[Any, Any] | None = None,
    new_sections: Sequence[str] | None = None,
) -> CatalogImportReport:
    """Önizleme: yazma koşulur ve geri sarılır (sayılar uygulamanın sayılarıdır)."""
    dosya, plan = _hazirla(
        parsed, section_map=section_map, new_sections=new_sections, payload_sha256=payload_sha256
    )
    with transaction.atomic():
        rapor = _ingest(dosya, plan=plan, payload_sha256=payload_sha256, file_name=file_name)
        transaction.set_rollback(True)
    rapor.dry_run = True
    for satir in rapor.rows:
        if satir.bucket == BUCKET_NEW:
            satir.work = None
        elif satir.bucket == BUCKET_EXISTING:
            satir.work = None
    onceki = import_service.applied_run(payload_sha256)
    if onceki is not None:
        rapor.already_applied = True
        rapor.applied_at = timezone.localtime(onceki.created_at).strftime("%d.%m.%Y")
    rapor.run_id = import_service._record_run(
        rapor, status=CatalogImportStatus.DRY_RUN, acquisition=None
    )
    return rapor


@transaction.atomic
def apply_export(
    parsed: ParsedFile,
    *,
    payload_sha256: str,
    file_name: str = "",
    section_map: Mapping[Any, Any] | None = None,
    new_sections: Sequence[str] | None = None,
) -> CatalogImportReport:
    """Uygulama (tek işlem). Aynı içerik ikinci kez uygulanamaz; bölüm karşılıkları eksiksiz."""
    dosya, plan = _hazirla(
        parsed, section_map=section_map, new_sections=new_sections, payload_sha256=payload_sha256
    )
    onceki = import_service.applied_run(payload_sha256)
    if onceki is not None:
        gun = timezone.localtime(onceki.created_at).strftime("%d.%m.%Y")
        raise ValidationError(
            {
                "file": (
                    f"Bu dosya {gun} tarihinde zaten içe aktarıldı; aynı dosya ikinci kez "
                    "uygulanamaz."
                )
            }
        )
    if plan.bilinmeyen:
        degerler = ", ".join(f"“{k['value']}”" for k in list(plan.bilinmeyen.values())[:10])
        raise ValidationError(
            {
                "section_map": (
                    f"Bölüm listesinde bulunmayan değerler var: {degerler}. Her biri için var "
                    "olan bir bölüm seçin ya da yeni bölüm açın."
                )
            }
        )
    rapor = _ingest(dosya, plan=plan, payload_sha256=payload_sha256, file_name=file_name)
    aktarilmayan = not_imported_rows(rapor)
    if aktarilmayan:
        # Tek işlem (`transaction.atomic`): yazılanlar ve sayaç ilerletmesi geri sarılır.
        raise ValidationError({"file": ROWS_NOT_IMPORTED_MESSAGE.format(sayi=aktarilmayan)})
    rapor.run_id = import_service._record_run(
        rapor, status=CatalogImportStatus.APPLIED, acquisition=None
    )
    logger.info(
        "Dışa aktarım dosyası içe aktarıldı: %d satır, %d eser, %d nüsha.",
        rapor.stats["total_rows"],
        rapor.stats["new_works"],
        rapor.stats["copies_created"],
    )
    return rapor
